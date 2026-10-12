# -*- coding: utf-8 -*-
"""Touche Holding icin deterministik finans karar destek katmani.

LLM hesap yapmaz. Bu modul merkezi defterden kanitlari hesaplar, odeme ve
likidite onceliklerini kurallarla belirler. LLM yalnizca sonucu aciklar.
"""

from __future__ import annotations

import calendar
import hashlib
import json
import logging
import re
import sqlite3
import threading
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook

from config import LOGLAR, FIRMALAR, ana_excel_yolu
import finans_cekirdegi as _cekirdek
from finans_cekirdegi import _norm, faaliyet_anahtari_bul, yonetim_ozeti, veritabanini_hazirla
from finans_komutlari import acik_borc_vadesi_guncelle, turkce_tarih_coz
from kurumsal_finans import butce_sapma_raporu, kurumsal_durum
from ollama_client import metne_sor, son_metin_saglayici

log = logging.getLogger(__name__)
AUDIT_PATH = Path(LOGLAR) / "asistan_audit.jsonl"
_KREDI_CACHE = {"kimlik": None, "veri": []}
_KREDI_CACHE_LOCK = threading.RLock()


@dataclass(frozen=True)
class DanismanProfili:
    risk: str = "dusuk"
    rezerv_gun: int = 30
    yatirim_vadesi_ay: int = 3

    @classmethod
    def olustur(cls, veri=None):
        veri = veri or {}
        risk = str(veri.get("risk") or "dusuk").strip().casefold()
        if risk not in {"dusuk", "orta", "yuksek"}:
            risk = "dusuk"
        try:
            rezerv = max(7, min(int(veri.get("rezerv_gun", 30)), 180))
            vade = max(1, min(int(veri.get("yatirim_vadesi_ay", 3)), 120))
        except (TypeError, ValueError):
            rezerv, vade = 30, 3
        return cls(risk=risk, rezerv_gun=rezerv, yatirim_vadesi_ay=vade)


@contextmanager
def _baglanti():
    veritabanini_hazirla()
    # Tek DB yolu: finans_cekirdegi.DB_PATH (testlerde ve operasyonda yonlendirilebilir).
    con = sqlite3.connect(_cekirdek.DB_PATH, timeout=20)
    con.row_factory = sqlite3.Row
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        # ``with sqlite3.connect()`` kapatmaz; Windows'ta acik tutamac finans.db'yi kilitli birakir.
        con.close()


def _tarih_coz(deger):
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    metin = str(deger or "").strip()[:10]
    try:
        return datetime.strptime(metin, "%Y-%m-%d").date()
    except ValueError:
        return None


def _sayi(deger):
    try:
        return float(deger or 0)
    except (TypeError, ValueError):
        return 0.0


def _tl_yaz(deger):
    return f"{float(deger or 0):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _kredi_taksitleri():
    """Master Excel'deki açık kredi taksitlerini güncel kurla TL olarak döndürür."""
    excel = ana_excel_yolu()
    if not excel.exists():
        return []
    stat = excel.stat()
    kimlik = (str(excel.resolve()), stat.st_mtime_ns, stat.st_size)
    with _KREDI_CACHE_LOCK:
        if _KREDI_CACHE["kimlik"] == kimlik:
            return [dict(x) for x in _KREDI_CACHE["veri"]]
        wb = None
        try:
            wb = load_workbook(excel, read_only=True, data_only=False)
            if not {"04_Kredi & Vade", "03_Kayıt Merkezi", "99_Ayarlar"}.issubset(wb.sheetnames):
                return []

            ayarlar = wb["99_Ayarlar"]
            kurlar = {}
            for pb_degeri, kur_degeri in ayarlar.iter_rows(
                min_row=3, max_row=7, min_col=30, max_col=31, values_only=True
            ):
                pb = str(pb_degeri or "").strip().upper()
                kur = _sayi(kur_degeri)
                if pb and kur > 0:
                    kurlar[pb] = kur
            kurlar["TL"] = 1.0

            odenenler = {}
            ledger = wb["03_Kayıt Merkezi"]
            for satir in ledger.iter_rows(
                min_row=6, max_row=300, min_col=7, max_col=29, values_only=True
            ):
                islem_turu, tutar = satir[0], satir[2]
                taksit_id, durum, onay = satir[13], satir[15], satir[22]
                taksit_id = str(taksit_id or "").strip()
                if not taksit_id or _norm(islem_turu) != _norm("Kredi Taksit Ödemesi"):
                    continue
                durum = _norm(durum)
                onay = _norm(onay)
                if durum in {"iptal", "reddedildi"} or onay == "reddedildi":
                    continue
                odenenler[taksit_id] = odenenler.get(taksit_id, 0.0) + _sayi(tutar)

            plan = wb["04_Kredi & Vade"]
            sonuc = []
            for satir in plan.iter_rows(
                min_row=13, max_row=100, min_col=1, max_col=10, values_only=True
            ):
                taksit_id = str(satir[0] or "").strip()
                if not taksit_id:
                    continue
                kredi_id = str(satir[1] or "").strip()
                faaliyet_adi = str(satir[2] or "").strip()
                vade = _tarih_coz(satir[3])
                taksit_pb = _sayi(satir[6])
                pb = str(satir[9] or "TL").strip().upper()
                kalan_pb = max(0.0, taksit_pb - odenenler.get(taksit_id, 0.0))
                kur = kurlar.get(pb)
                if kalan_pb <= 0 or not vade or not kur:
                    continue
                sonuc.append({
                    "taksit_id": taksit_id,
                    "kredi_id": kredi_id,
                    "faaliyet": faaliyet_anahtari_bul(faaliyet_adi),
                    "faaliyet_adi": faaliyet_adi,
                    "vade": vade,
                    "tutar": round(kalan_pb * kur, 2),
                    "tutar_pb": round(kalan_pb, 2),
                    "pb": pb,
                })
            sonuc.sort(key=lambda x: (x["vade"], x["taksit_id"]))
            _KREDI_CACHE["kimlik"] = kimlik
            _KREDI_CACHE["veri"] = [dict(x) for x in sonuc]
            return sonuc
        except Exception as exc:
            log.warning("Excel kredi plani okunamadi: %s", exc)
            return []
        finally:
            if wb is not None:
                wb.close()


def _odeme_durumu(vade, bugun):
    kalan = (vade - bugun).days if vade else None
    if kalan is not None and kalan < 0:
        return "Gecikmiş", 400 + min(abs(kalan), 90), kalan
    if kalan is not None and kalan <= 7:
        return "7 gün içinde", 300 + (7 - kalan), kalan
    if kalan is not None and kalan <= 30:
        return "30 gün içinde", 200 + (30 - kalan), kalan
    if vade:
        return "İleri vadeli", 100, kalan
    return "Vade eksik", 150, kalan


def _faaliyet_kosulu(faaliyet, alias="f"):
    # Touche Holding ana seçimdir ve tüm bağlı faaliyetleri konsolide eder.
    if faaliyet and faaliyet not in {"tum", "touche_holding"}:
        return f" AND {alias}.faaliyet_key=?", [faaliyet]
    return "", []


def faaliyetler():
    # UI seçimi tek master/config ile aynı kaynaktan gelmelidir; DB'deki eski/hayalet
    # faaliyet etiketleri finans ekranına taşınmaz.
    return [
        {"key": key, "ad": firma["ad"]}
        for key, firma in FIRMALAR.items()
    ]


def _sorudan_faaliyet(soru, varsayilan="tum"):
    """Görünür faaliyet filtresi olmadan sorudaki şirket/faaliyet adını çözer."""
    if varsayilan and varsayilan not in {"tum", "touche_holding"}:
        return varsayilan
    metin = _norm(soru).replace("_", " ")
    for key, firma in FIRMALAR.items():
        if key == "touche_holding":
            continue
        adaylar = [key.replace("_", " "), firma.get("ad"), firma.get("faaliyet")]
        adaylar.extend(firma.get("faaliyet_aliaslari") or [])
        ad = _norm(firma.get("ad")).strip()
        if ad:
            adaylar.append(ad.split()[0])
        if any(_norm(aday).replace("_", " ") in metin for aday in adaylar if aday):
            return key
    return varsayilan


def _gider_run_rate(faaliyet, bugun):
    baslangic = bugun - timedelta(days=89)
    kosul, args = _faaliyet_kosulu(faaliyet)
    with _baglanti() as con:
        row = con.execute(
            "SELECT COALESCE(SUM(s.borc_kurus-s.alacak_kurus),0) gider, "
            "MIN(f.tarih) ilk_tarih FROM fisler f "
            "JOIN fis_satirlari s ON s.fis_id=f.fis_id "
            "JOIN hesaplar h ON h.kod=s.hesap_kodu "
            "WHERE h.hesap_tipi='GIDER' AND f.tarih BETWEEN ? AND ?" + kosul,
            [baslangic.isoformat(), bugun.isoformat(), *args],
        ).fetchone()
    ilk = _tarih_coz(row["ilk_tarih"]) or bugun
    gercek_gozlem_gunu = max(1, min(90, (bugun - ilk).days + 1))
    # Kisa veri gecmisini tam aylik sureye yayarak tek seferlik giderlerin
    # finansman ihtiyacini yapay bicimde sisirmesini onleriz.
    hesap_gunu = max(30, gercek_gozlem_gunu)
    gider = max(0, row["gider"] / 100)
    return {"donem_gideri": gider, "gozlem_gunu": gercek_gozlem_gunu, "gunluk_gider": gider / hesap_gunu}


def _borc_grubu(cari, aciklama):
    """Usta odemelerini tahakkuk ettikleri maliyet grubuyla eslestirir."""
    metin = _norm(f"{cari or ''} {aciklama or ''}")
    gruplar = []
    for firma in FIRMALAR.values():
        gruplar.extend((firma.get("master_data") or {}).get("maliyet_gruplari") or [])
    for grup in dict.fromkeys(gruplar):
        if _norm(grup) in metin:
            return str(grup)
    return str(cari or "Cari belirtilmemis").strip() or "Cari belirtilmemis"


def odeme_oncelikleri(faaliyet="tum", bugun=None, limit=10, kredi_taksitleri=None):
    bugun = bugun or date.today()
    kredi_taksitleri = _kredi_taksitleri() if kredi_taksitleri is None else kredi_taksitleri
    kredi_faaliyetleri = {x["faaliyet"] for x in kredi_taksitleri}
    kosul, args = _faaliyet_kosulu(faaliyet)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.faaliyet_key,TRIM(COALESCE(f.proje,'')) proje,"
            "TRIM(COALESCE(f.cari,'')) cari,TRIM(COALESCE(f.aciklama,'')) aciklama,"
            "NULLIF(TRIM(COALESCE(f.vade,'')),'') vade,s.hesap_kodu,"
            "s.alacak_kurus-s.borc_kurus bakiye "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
            "WHERE f.durum='MUHASEBELESTI' AND s.hesap_kodu IN ('300','320','331','335')" + kosul,
            args,
        ).fetchall()
    gruplu = {}
    kredi_carileri = {}
    for r in rows:
        kredi_satiri = "kredi" in _norm(f"{r['cari']} {r['aciklama']}")
        if kredi_satiri:
            kredi_carileri.setdefault(r["faaliyet_key"], r["cari"] or "Kredi Veren Banka")
        if kredi_satiri and r["faaliyet_key"] in kredi_faaliyetleri:
            continue
        cari = _borc_grubu(r["cari"], r["aciklama"])
        key = (r["faaliyet_key"], r["proje"], cari)
        item = gruplu.setdefault(key, {"faaliyet_key": r["faaliyet_key"], "proje": r["proje"], "cari": cari, "vade": None, "bakiye": 0})
        item["bakiye"] += r["bakiye"]
        if r["vade"] and (not item["vade"] or r["vade"] < item["vade"]):
            item["vade"] = r["vade"]
    sonuc = []
    for r in gruplu.values():
        if r["bakiye"] <= 0:
            continue
        vade = _tarih_coz(r["vade"])
        durum, puan, kalan = _odeme_durumu(vade, bugun)
        sonuc.append({
            "faaliyet": r["faaliyet_key"],
            "proje": r["proje"],
            "cari": r["cari"],
            "tutar": round(r["bakiye"] / 100, 2),
            "vade": vade.isoformat() if vade else None,
            "kalan_gun": kalan,
            "durum": durum,
            "_puan": puan,
        })

    # Her kredi için yalnız sıradaki açık taksit ödeme önceliğine girer.
    siradaki = {}
    for taksit in kredi_taksitleri:
        if faaliyet not in {"tum", "touche_holding"} and taksit["faaliyet"] != faaliyet:
            continue
        key = (taksit["faaliyet"], taksit["kredi_id"])
        siradaki.setdefault(key, taksit)
    for taksit in siradaki.values():
        durum, puan, kalan = _odeme_durumu(taksit["vade"], bugun)
        sonuc.append({
            "faaliyet": taksit["faaliyet"],
            "proje": taksit["taksit_id"],
            "cari": kredi_carileri.get(taksit["faaliyet"], "Kredi taksiti"),
            "tutar": taksit["tutar"],
            "vade": taksit["vade"].isoformat(),
            "kalan_gun": kalan,
            "durum": durum,
            "kaynak": "04_Kredi & Vade",
            "_puan": puan,
        })
    sonuc.sort(key=lambda x: (x["_puan"], x["tutar"]), reverse=True)
    for x in sonuc:
        x.pop("_puan", None)
    return sonuc[: max(1, min(int(limit), 50))]


def _nakit_hareket_satirlari(faaliyet="tum", baslangic=None, bitis=None, limit=None):
    kosul, args = _faaliyet_kosulu(faaliyet)
    tarih_kosulu = ""
    if baslangic:
        tarih_kosulu += " AND f.tarih>=?"
        args.append(str(baslangic))
    if bitis:
        tarih_kosulu += " AND f.tarih<=?"
        args.append(str(bitis))
    limit_sql = ""
    if limit is not None:
        limit_sql = " LIMIT ?"
        args.append(max(1, min(int(limit), 500)))
    with _baglanti() as con:
        return con.execute(
            "SELECT f.fis_id,f.faaliyet_key,f.tarih,f.proje,f.cari,f.aciklama,"
            "SUM(CASE WHEN s.hesap_kodu IN ('100','102') "
            "THEN s.borc_kurus-s.alacak_kurus ELSE 0 END) nakit_net,"
            "SUM(CASE WHEN s.hesap_kodu='300' THEN s.alacak_kurus-s.borc_kurus ELSE 0 END) kredi_artis,"
            "SUM(CASE WHEN s.hesap_kodu IN ('120','131') THEN s.alacak_kurus-s.borc_kurus ELSE 0 END) alacak_azalis "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
            "WHERE f.durum='MUHASEBELESTI'" + kosul + tarih_kosulu +
            " GROUP BY f.fis_id HAVING nakit_net<>0 "
            "ORDER BY f.tarih DESC,f.fis_id DESC" + limit_sql,
            args,
        ).fetchall()


def _nakit_hareketini_hazirla(r):
    if r["nakit_net"] < 0:
        yon = "Ödeme"
    elif r["kredi_artis"] > 0:
        yon = "Kredi kullanımı"
    elif r["alacak_azalis"] > 0:
        yon = "Tahsilat"
    else:
        yon = "Nakit girişi"
    return {
        "fis_id": r["fis_id"], "faaliyet": r["faaliyet_key"], "tarih": r["tarih"],
        "proje": r["proje"] or "", "cari": r["cari"] or "Cari belirtilmemis",
        "aciklama": r["aciklama"] or "", "yon": yon,
        "tutar": round(abs(r["nakit_net"]) / 100, 2),
    }


def son_nakit_hareketleri(faaliyet="tum", limit=10, baslangic=None, bitis=None):
    """Gerçekleşen banka/kasa giriş ve çıkışlarını en yeniden eskiye getirir."""
    rows = _nakit_hareket_satirlari(faaliyet, baslangic, bitis, limit)
    sonuc = []
    for r in rows:
        sonuc.append(_nakit_hareketini_hazirla(r))
    return sonuc


def nakit_hareket_ozeti(faaliyet="tum", baslangic=None, bitis=None):
    """Seçilen dönemin nakit hareketlerini ödeme, tahsilat ve finansman olarak ayırır."""
    rows = _nakit_hareket_satirlari(faaliyet, baslangic, bitis)
    ozet = {
        "odeme_tutari": 0, "odeme_adedi": 0,
        "tahsilat_tutari": 0, "tahsilat_adedi": 0,
        "finansman_girisi": 0, "finansman_adedi": 0,
        "diger_nakit_girisi": 0, "diger_giris_adedi": 0,
    }
    for r in rows:
        nakit = int(r["nakit_net"] or 0)
        if nakit < 0:
            ozet["odeme_tutari"] += -nakit
            ozet["odeme_adedi"] += 1
        elif int(r["kredi_artis"] or 0) > 0:
            ozet["finansman_girisi"] += nakit
            ozet["finansman_adedi"] += 1
        elif int(r["alacak_azalis"] or 0) > 0:
            ozet["tahsilat_tutari"] += nakit
            ozet["tahsilat_adedi"] += 1
        else:
            ozet["diger_nakit_girisi"] += nakit
            ozet["diger_giris_adedi"] += 1
    for alan in ("odeme_tutari", "tahsilat_tutari", "finansman_girisi", "diger_nakit_girisi"):
        ozet[alan] = round(ozet[alan] / 100, 2)
    return ozet


def _vadeli_nakit(faaliyet, bugun, gun=30, kredi_taksitleri=None):
    bitis = bugun + timedelta(days=gun)
    kosul, args = _faaliyet_kosulu(faaliyet)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.vade,s.hesap_kodu,s.borc_kurus,s.alacak_kurus "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
            "WHERE TRIM(COALESCE(f.vade,''))<>'' "
            "AND s.hesap_kodu='120'" + kosul,
            args,
        ).fetchall()
    giris = cikis = 0
    for r in rows:
        vade = _tarih_coz(r["vade"])
        if not vade or not bugun <= vade <= bitis:
            continue
        giris += max(0, r["borc_kurus"] - r["alacak_kurus"])
    kredi_taksitleri = _kredi_taksitleri() if kredi_taksitleri is None else kredi_taksitleri
    # Tahakkukların brüt tutarını değil, yapılmış ödemeler düşüldükten sonraki
    # gerçek açık bakiyeyi kullan. Kredi taksitleri de bu listede tek kez yer alır.
    for odeme in odeme_oncelikleri(faaliyet, bugun=bugun, limit=50, kredi_taksitleri=kredi_taksitleri):
        vade = _tarih_coz(odeme.get("vade"))
        if vade and bugun <= vade <= bitis:
            cikis += round(float(odeme["tutar"]) * 100)
    return round(giris / 100, 2), round(cikis / 100, 2)


def _yatirim_sepeti(kapasite, profil):
    if kapasite <= 0:
        return []
    if profil.yatirim_vadesi_ay <= 3:
        oranlar = [("Günlük likit TL aracı", 70), ("Kısa vadeli mevduat / kira sertifikası", 30)]
    elif profil.risk == "dusuk":
        oranlar = [("Para piyasası / günlük likit araç", 50), ("Vadeli mevduat", 30), ("Kısa vadeli kamu borçlanma aracı", 20)]
    elif profil.risk == "orta":
        oranlar = [("Para piyasası / günlük likit araç", 35), ("Kısa vadeli kamu borçlanma aracı", 35), ("Çeşitlendirilmiş değişken fon", 30)]
    else:
        oranlar = [("Para piyasası / günlük likit araç", 25), ("Kısa vadeli kamu borçlanma aracı", 25), ("Çeşitlendirilmiş değişken veya hisse fonu", 50)]
    return [{"arac": ad, "oran": oran, "tutar": round(kapasite * oran / 100, 2)} for ad, oran in oranlar]


def faaliyet_performansi():
    """Holding faaliyetlerini aynı hesap planı üzerinden karşılaştırır."""
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.faaliyet_key,h.hesap_tipi,s.hesap_kodu,"
            "SUM(s.borc_kurus) borc,SUM(s.alacak_kurus) alacak "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
            "JOIN hesaplar h ON h.kod=s.hesap_kodu WHERE f.durum='MUHASEBELESTI' "
            "GROUP BY f.faaliyet_key,h.hesap_tipi,s.hesap_kodu"
        ).fetchall()
    gruplar = {}
    for row in rows:
        item = gruplar.setdefault(row["faaliyet_key"], {"gelir": 0, "gider": 0, "nakit": 0, "alacak": 0, "borc": 0})
        borc, alacak = row["borc"], row["alacak"]
        if row["hesap_tipi"] == "GELIR":
            item["gelir"] += alacak - borc
        if row["hesap_tipi"] == "GIDER":
            item["gider"] += borc - alacak
        if row["hesap_kodu"] in {"100", "102"}:
            item["nakit"] += borc - alacak
        if row["hesap_kodu"] in {"120", "131"}:
            item["alacak"] += borc - alacak
        if row["hesap_kodu"] in {"300", "320", "331", "335"}:
            item["borc"] += alacak - borc
    sonuc = []
    for key, firma in FIRMALAR.items():
        if key == "touche_holding" and len(FIRMALAR) > 1:
            continue
        item = gruplar.get(key, {"gelir": 0, "gider": 0, "nakit": 0, "alacak": 0, "borc": 0})
        veri = {k: round(v / 100, 2) for k, v in item.items()}
        veri.update({
            "key": key, "ad": firma.get("ad", key),
            "net": round(veri["gelir"] - veri["gider"], 2),
            "marj": round((veri["gelir"] - veri["gider"]) / veri["gelir"] * 100, 2) if veri["gelir"] else 0,
        })
        if any(abs(veri[x]) > 0 for x in ("gelir", "gider", "nakit", "alacak", "borc")):
            sonuc.append(veri)
    return sorted(sonuc, key=lambda x: x["gelir"], reverse=True)


def veri_kalitesi(faaliyet="tum", odemeler=None):
    kosul, args = _faaliyet_kosulu(faaliyet)
    with _baglanti() as con:
        genel = con.execute(
            "SELECT COUNT(DISTINCT f.fis_id) kayit,MAX(f.tarih) son_tarih,"
            "SUM(CASE WHEN TRIM(COALESCE(f.proje,''))='' THEN 1 ELSE 0 END) projesiz,"
            "SUM(CASE WHEN TRIM(COALESCE(f.cari,''))='' THEN 1 ELSE 0 END) carisiz "
            "FROM fisler f WHERE f.durum='MUHASEBELESTI'" + kosul, args,
        ).fetchone()
        eksik_satirlar = con.execute(
            "SELECT f.fis_id,f.faaliyet_key,f.tarih,TRIM(COALESCE(f.proje,'')) proje,"
            "TRIM(COALESCE(f.cari,'')) cari,TRIM(COALESCE(f.aciklama,'')) aciklama "
            "FROM fisler f WHERE f.durum='MUHASEBELESTI' "
            "AND (TRIM(COALESCE(f.proje,''))='' OR TRIM(COALESCE(f.cari,''))='')" + kosul +
            " ORDER BY f.tarih DESC,f.fis_id DESC LIMIT 20",
            args,
        ).fetchall()
    kayit = int(genel["kayit"] or 0)
    eksik = int(genel["projesiz"] or 0) + int(genel["carisiz"] or 0)
    odemeler = odemeler if odemeler is not None else odeme_oncelikleri(faaliyet, limit=50)
    vadesiz = round(sum(x["tutar"] for x in odemeler if not x.get("vade")), 2)
    toplam = round(sum(x["tutar"] for x in odemeler), 2)
    vade_orani = (vadesiz / toplam) if toplam else 0
    eksik_orani = min(1, eksik / max(1, kayit * 2))
    skor = max(0, round(100 - vade_orani * 55 - eksik_orani * 25 - (20 if kayit < 5 else 0)))
    sorunlar = []
    if vadesiz:
        sorunlar.append(f"{vadesiz:,.2f} TL borçta vade eksik")
    if genel["projesiz"]:
        sorunlar.append(f"{int(genel['projesiz'])} kayıtta proje eksik")
    if genel["carisiz"]:
        sorunlar.append(f"{int(genel['carisiz'])} kayıtta cari eksik")
    return {
        "skor": skor, "seviye": "İyi" if skor >= 85 else ("Geliştirilmeli" if skor >= 60 else "Kritik"),
        "kayit_sayisi": kayit, "son_kayit_tarihi": genel["son_tarih"],
        "vadesiz_borc": vadesiz, "vade_tamamlama_orani": round((1 - vade_orani) * 100, 1),
        "sorunlar": sorunlar,
        "eksik_kayitlar": [{
            "fis_id": r["fis_id"], "faaliyet": r["faaliyet_key"], "tarih": r["tarih"],
            "proje": r["proje"], "cari": r["cari"], "aciklama": r["aciklama"],
            "eksikler": [x for x, deger in (("proje", r["proje"]), ("cari", r["cari"])) if not deger],
        } for r in eksik_satirlar],
    }


def nakit_projeksiyonu(faaliyet="tum", bugun=None, hafta=13, kredi_taksitleri=None):
    """Vadesi olan açık kalemlerle 13 haftalık deterministik nakit görünümü."""
    bugun = bugun or date.today()
    hafta = max(1, min(int(hafta), 26))
    ozet = yonetim_ozeti(None if faaliyet in {"tum", "touche_holding"} else faaliyet)
    kosul, args = _faaliyet_kosulu(faaliyet)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.vade,s.hesap_kodu,s.borc_kurus,s.alacak_kurus FROM fisler f "
            "JOIN fis_satirlari s ON s.fis_id=f.fis_id WHERE f.durum='MUHASEBELESTI' "
            "AND TRIM(COALESCE(f.vade,''))<>'' AND s.hesap_kodu='120'" + kosul,
            args,
        ).fetchall()
    dilimler = []
    for i in range(hafta):
        bas = bugun + timedelta(days=i * 7)
        son = bas + timedelta(days=6)
        dilimler.append({"hafta": i + 1, "baslangic": bas.isoformat(), "bitis": son.isoformat(), "giris": 0.0, "cikis": 0.0})
    for row in rows:
        vade = _tarih_coz(row["vade"])
        if not vade or vade < bugun:
            continue
        index = (vade - bugun).days // 7
        if index >= hafta:
            continue
        dilimler[index]["giris"] += max(0, row["borc_kurus"] - row["alacak_kurus"]) / 100
    kredi_taksitleri = _kredi_taksitleri() if kredi_taksitleri is None else kredi_taksitleri
    for odeme in odeme_oncelikleri(faaliyet, bugun=bugun, limit=50, kredi_taksitleri=kredi_taksitleri):
        vade = _tarih_coz(odeme.get("vade"))
        if not vade or vade < bugun:
            continue
        index = (vade - bugun).days // 7
        if index < hafta:
            dilimler[index]["cikis"] += odeme["tutar"]
    bakiye = float(ozet["nakit"])
    for item in dilimler:
        item["giris"] = round(item["giris"], 2)
        item["cikis"] = round(item["cikis"], 2)
        item["net"] = round(item["giris"] - item["cikis"], 2)
        bakiye = round(bakiye + item["net"], 2)
        item["kapanis"] = bakiye
    return dilimler


def _risk_analizi(ozet, nakit_plani, kalite, run_rate):
    hedef = max(1, nakit_plani["rezerv"] + nakit_plani["30_gun_vadeli_odeme"])
    karsilama = max(0, ozet["nakit"] + nakit_plani["30_gun_beklenen_tahsilat"]) / hedef
    borc_varlik = max(0, ozet["borc"]) / max(1, max(0, ozet["nakit"]) + max(0, ozet["alacak"]))
    puan = 0
    puan += min(45, max(0, (1 - karsilama) * 45))
    puan += min(25, borc_varlik * 20)
    puan += min(20, (100 - kalite["vade_tamamlama_orani"]) * .2)
    if run_rate["gozlem_gunu"] < 30:
        puan += 10
    puan = round(min(100, puan))
    return {
        "skor": puan,
        "seviye": "Düşük" if puan <= 30 else ("Orta" if puan <= 60 else "Yüksek"),
        "likidite_karsilama": round(karsilama * 100, 1),
        "borc_varlik_orani": round(borc_varlik * 100, 1),
    }


def _aksiyonlar(odemeler, nakit_plani, kalite):
    sonuc = []
    if nakit_plani["finansman_acigi"] > 0:
        sonuc.append({"oncelik": "Kritik", "ufuk": "Bugün", "baslik": "Finansman açığını kapat", "tutar": nakit_plani["finansman_acigi"], "neden": "Rezerv ve yakın vadeli ödemeler mevcut likiditeyi aşıyor."})
    if kalite["vadesiz_borc"] > 0:
        sonuc.append({"oncelik": "Yüksek", "ufuk": "Bugün", "baslik": "Eksik vadeleri tamamla", "tutar": kalite["vadesiz_borc"], "neden": "Vadesiz borçlar ödeme sırasını belirsizleştiriyor."})
    for odeme in odemeler[:3]:
        sonuc.append({"oncelik": "Yüksek" if odeme["durum"] in {"Gecikmiş", "7 gün içinde"} else "Planla", "ufuk": odeme["vade"] or "Vade gerekli", "baslik": f"{odeme['cari']} ödemesini planla", "tutar": odeme["tutar"], "neden": f"{odeme['proje'] or odeme['faaliyet']} · {odeme['durum']}"})
    if not sonuc:
        sonuc.append({"oncelik": "Normal", "ufuk": "Bu hafta", "baslik": "Likiditeyi izle", "tutar": 0, "neden": "Kritik finansman veya ödeme sinyali bulunmuyor."})
    return sonuc[:5]


def senaryo_hesapla(faaliyet="tum", profil=None, degisiklikler=None):
    degisiklikler = degisiklikler or {}
    degerler = {}
    for alan in ("ek_tahsilat", "ek_odeme", "yeni_yatirim"):
        try:
            deger = float(degisiklikler.get(alan, 0) or 0)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{alan} geçerli bir tutar olmalı.") from exc
        if not 0 <= deger <= 1_000_000_000_000:
            raise ValueError(f"{alan} 0 ile 1 trilyon arasında olmalı.")
        degerler[alan] = round(deger, 2)
    pano = karar_panosu(faaliyet, profil)
    n = pano["nakit_plani"]
    yeni_nakit = round(pano["ozet"]["nakit"] + degerler["ek_tahsilat"] - degerler["ek_odeme"] - degerler["yeni_yatirim"], 2)
    hedef = round(n["rezerv"] + n["30_gun_vadeli_odeme"], 2)
    kapasite = round(max(0, yeni_nakit - hedef), 2)
    acik = round(max(0, hedef - yeni_nakit - n["30_gun_beklenen_tahsilat"]), 2)
    return {
        "girdiler": degerler, "mevcut_nakit": pano["ozet"]["nakit"], "senaryo_nakit": yeni_nakit,
        "korunmasi_gereken": hedef, "yatirim_kapasitesi": kapasite, "finansman_acigi": acik,
        "karar": "Uygulanabilir" if not acik else "Finansman açığı oluşturuyor",
    }


def karar_panosu(faaliyet="tum", profil=None, bugun=None):
    bugun = bugun or date.today()
    profil = DanismanProfili.olustur(profil)
    faaliyet_arg = None if faaliyet in {"tum", "touche_holding"} else faaliyet
    ozet = yonetim_ozeti(faaliyet_arg)
    run_rate = _gider_run_rate(faaliyet, bugun)
    kredi_taksitleri = _kredi_taksitleri()
    beklenen_giris, vadeli_cikis = _vadeli_nakit(faaliyet, bugun, 30, kredi_taksitleri)
    rezerv = round(run_rate["gunluk_gider"] * profil.rezerv_gun, 2)
    yatirim_kapasitesi = round(max(0, ozet["nakit"] - rezerv - vadeli_cikis), 2)
    finansman_acigi = round(max(0, rezerv + vadeli_cikis - ozet["nakit"] - beklenen_giris), 2)
    odemeler = odeme_oncelikleri(faaliyet, bugun=bugun, limit=10, kredi_taksitleri=kredi_taksitleri)
    son_hareketler = son_nakit_hareketleri(faaliyet, limit=10)
    hareket_ozeti = nakit_hareket_ozeti(faaliyet)
    vadesiz_borc = round(sum(x["tutar"] for x in odemeler if not x["vade"]), 2)
    kalite = veri_kalitesi(faaliyet, odemeler)
    projeksiyon = nakit_projeksiyonu(faaliyet, bugun, 13, kredi_taksitleri)
    kurumsal = kurumsal_durum(faaliyet_arg, bugun)
    ay_basi = bugun.replace(day=1)
    butce_sapma = butce_sapma_raporu(faaliyet, ay_basi, bugun)

    uyarilar = []
    if vadesiz_borc:
        uyarilar.append(f"{vadesiz_borc:,.2f} TL borcun vade bilgisi eksik; ödeme sırası kesinleştirilemez.")
    if run_rate["gozlem_gunu"] < 30:
        uyarilar.append("Gider geçmişi 30 günden kısa; nakit rezervi tahmini ihtiyatlıdır.")
    if finansman_acigi:
        ana_karar = "Önce finansman açığını kapat ve ödeme vadelerini netleştir. Yeni yatırımı ertele."
        durum = "finansman_gerekli"
    elif yatirim_kapasitesi:
        ana_karar = "Zorunlu ödemeler ve işletme rezervi ayrıldıktan sonra kalan nakit yatırıma ayrılabilir."
        durum = "yatirim_degerlendirilebilir"
    else:
        ana_karar = "Nakit işletme rezervini ancak karşılıyor. Likit kal ve yeni yatırımı ertele."
        durum = "likiditeyi_koru"

    nakit_plani = {
        "gunluk_gider": round(run_rate["gunluk_gider"], 2),
        "rezerv": rezerv,
        "30_gun_beklenen_tahsilat": beklenen_giris,
        "30_gun_vadeli_odeme": vadeli_cikis,
        "yatirim_kapasitesi": yatirim_kapasitesi,
        "finansman_acigi": finansman_acigi,
    }
    return {
        "tarih": bugun.isoformat(),
        "faaliyet": faaliyet,
        "profil": asdict(profil),
        "ozet": {k: round(float(v), 2) for k, v in ozet.items()},
        "nakit_plani": nakit_plani,
        "durum": durum,
        "ana_karar": ana_karar,
        "odemeler": odemeler,
        "son_hareketler": son_hareketler,
        "nakit_hareket_ozeti": hareket_ozeti,
        "yatirim_sepeti": _yatirim_sepeti(yatirim_kapasitesi, profil),
        "nakit_projeksiyonu": projeksiyon,
        "faaliyet_performansi": faaliyet_performansi() if faaliyet in {"tum", "touche_holding"} else [],
        "veri_kalitesi": kalite,
        "risk": _risk_analizi(ozet, nakit_plani, kalite, run_rate),
        "aksiyonlar": _aksiyonlar(odemeler, nakit_plani, kalite),
        "kurumsal": kurumsal,
        "butce_sapma": butce_sapma,
        "uyarilar": uyarilar,
        "veri_notu": "Piyasa faiz/getiri verisi bağlı değil; araçlar kategori bazında gösterilir.",
    }


_AYLAR = {
    "ocak": 1, "subat": 2, "mart": 3, "nisan": 4, "mayis": 5, "haziran": 6,
    "temmuz": 7, "agustos": 8, "eylul": 9, "ekim": 10, "kasim": 11, "aralik": 12,
}
_AY_ETIKETLERI = {
    1: "Ocak", 2: "Şubat", 3: "Mart", 4: "Nisan", 5: "Mayıs", 6: "Haziran",
    7: "Temmuz", 8: "Ağustos", 9: "Eylül", 10: "Ekim", 11: "Kasım", 12: "Aralık",
}


def _donem_coz(soru, bugun=None):
    """Doğal dildeki ay/yıl ifadesini geçmiş dönem aralığına dönüştürür."""
    bugun = bugun or date.today()
    s = _norm(soru)
    if "bu ay" in s:
        ay, yil = bugun.month, bugun.year
    elif "gecen ay" in s:
        ay = 12 if bugun.month == 1 else bugun.month - 1
        yil = bugun.year - 1 if bugun.month == 1 else bugun.year
    else:
        ay = next((num for ad, num in _AYLAR.items() if re.search(rf"\b{ad}\b", s)), None)
        if not ay:
            yil_eslesme = re.search(r"\b(20\d{2})\b", s)
            if yil_eslesme:
                yil = int(yil_eslesme.group(1))
                return date(yil, 1, 1), date(yil, 12, 31), str(yil)
            return None, None, "Tüm dönem"
        yil_eslesme = re.search(r"\b(20\d{2})\b", s)
        yil = int(yil_eslesme.group(1)) if yil_eslesme else bugun.year
        # Tamamlanmış hareket sorusunda gelecekteki ay, en son geçmiş yıl demektir.
        if not yil_eslesme and ay > bugun.month:
            yil -= 1
    son_gun = calendar.monthrange(yil, ay)[1]
    return date(yil, ay, 1), date(yil, ay, son_gun), f"{_AY_ETIKETLERI[ay]} {yil}"


def _detay_istegi(soru):
    s = _norm(soru)
    return any(x in s for x in (
        "tam liste", "listeyi ver", "listele", "kalemleri", "kalemlerini",
        "kalem kalem", "tek tek", "ayrintili", "icerigi", "detay", "dokum",
        "dokumu", "hangileri",
    ))


def _gecmis_nakit_raporu(faaliyet, soru, niyet, detay_sorusu=None):
    baslangic, bitis, etiket = _donem_coz(soru)
    hareketler = son_nakit_hareketleri(faaliyet, limit=500, baslangic=baslangic, bitis=bitis)
    hedef_yon = "Ödeme" if niyet == "gecmis_odeme" else "Tahsilat"
    secilen = [x for x in hareketler if x["yon"] == hedef_yon]
    ozet = nakit_hareket_ozeti(faaliyet, baslangic, bitis)
    return {
        "etiket": etiket, "baslangic": str(baslangic or ""), "bitis": str(bitis or ""),
        "hareketler": secilen, "ozet": ozet,
        "detay": _detay_istegi(detay_sorusu if detay_sorusu is not None else soru),
    }


def _yaklasik_kelime_var(metin, hedefler, baslangic):
    """Klavye/yazım hatası içeren temel finans kelimelerini güvenli biçimde tanır."""
    kelimeler = re.findall(r"[a-z0-9]+", _norm(metin))
    for kelime in kelimeler:
        if not kelime.startswith(baslangic):
            continue
        if any(SequenceMatcher(None, kelime, hedef).ratio() >= 0.78 for hedef in hedefler):
            return True
    return False


def _niyet(soru):
    s = _norm(soru)
    odeme_var = "odeme" in s or _yaklasik_kelime_var(
        s, ("odeme", "odemeler", "odemeleri", "odemesini", "odemelerini", "odendi"), "ode"
    )
    tahsilat_var = "tahsilat" in s or _yaklasik_kelime_var(
        s, ("tahsilat", "tahsilatlar", "tahsilati", "tahsilatlari"), "tah"
    )
    tamamlanmis = any(x in s for x in (
        "yapildi", "yapilmis", "yaptik", "yaptiniz", "odendi", "tahsil edildi",
        "simdiye kadar", "bugune kadar", "gecmis",
    ))
    if tamamlanmis and tahsilat_var:
        return "gecmis_tahsilat"
    if tamamlanmis and odeme_var:
        return "gecmis_odeme"
    if odeme_var and _detay_istegi(soru):
        return "gecmis_odeme"
    if any(x in s for x in ("gecmis odeme", "son odeme")):
        return "gecmis_odeme"
    if any(x in s for x in ("yatirim", "degerlendir", "fon", "mevduat")):
        return "yatirim"
    if odeme_var or any(x in s for x in ("kime", "vade", "borc")):
        return "odeme"
    if any(x in s for x in ("finansman", "kredi", "nakit acigi")):
        return "finansman"
    if tahsilat_var or any(x in s for x in ("alacak", "kimden para", "kimden tahsil")):
        return "tahsilat"
    if any(x in s for x in (
        "veri kalitesi", "eksik veri", "eksik finans", "eksik kayit",
        "proje eksik", "cari eksik", "hangi kayit", "guvensiz finans",
    )):
        return "veri_kalitesi"
    if any(x in s for x in ("aksiyon", "ne yapmali", "yapmam gereken")):
        return "aksiyon"
    if any(x in s for x in ("risk", "tehlike", "saglik", "guvenli")):
        return "risk"
    if any(x in s for x in ("kar", "karlilik", "marj", "faaliyet")):
        return "karlilik"
    if any(x in s for x in ("nakit", "kasa", "banka", "13 hafta", "projeksiyon")):
        return "nakit"
    return "genel"


def _vade_komutu_coz(soru, pano, bugun=None):
    """Yalnızca açıkça verilen vade yazma komutlarını izinli araca dönüştürür."""
    sade = _norm(soru)
    if "vade" not in sade:
        return None
    eylemler = ("kaydet", "guncelle", "duzelt", "ata", "gir", "ekle", "olsun")
    if not any(eylem in sade for eylem in eylemler):
        return None
    yeni_vade = turkce_tarih_coz(soru, bugun=bugun)
    if not yeni_vade:
        return {"hata": "Vade komutunu anladım ancak tarihi net okuyamadım. Örnek: 13 Ekim 2026."}

    hedefler = []
    for odeme in pano.get("odemeler", []):
        # Kredi planı kendi taksit çizelgesinden yönetilir; tahakkuk vadesi aracı ona yazmaz.
        if odeme.get("kaynak") == "04_Kredi & Vade":
            continue
        cari = str(odeme.get("cari") or "").strip()
        if cari and _norm(cari) in sade:
            hedefler.append({
                "faaliyet": odeme.get("faaliyet"), "proje": odeme.get("proje") or "", "cari": cari,
            })
    if not hedefler:
        vadesiz = [x for x in pano.get("odemeler", []) if not x.get("vade") and x.get("kaynak") != "04_Kredi & Vade"]
        if len(vadesiz) == 1 and any(x in sade for x in ("bu odeme", "bunun", "bu borc")):
            x = vadesiz[0]
            hedefler = [{"faaliyet": x.get("faaliyet"), "proje": x.get("proje") or "", "cari": x.get("cari")}]
    if not hedefler:
        return {"hata": "Hangi açık borcun vadesinin değişeceğini eşleştiremedim. Cari veya iş kalemi adını yazın."}
    return {"vade": yeni_vade, "hedefler": hedefler}


def _deterministik_cevap(pano, niyet):
    n = pano["nakit_plani"]
    o = pano["ozet"]
    if niyet == "yatirim":
        if n["yatirim_kapasitesi"] <= 0:
            return (
                f"Bugun yatirima ayrilabilir guvenli nakit görünmüyor. Nakit {o['nakit']:,.2f} TL, "
                f"hedef isletme rezervi {n['rezerv']:,.2f} TL. Once likiditeyi koruyun."
            )
        return f"Hesaplanan yatirim kapasitesi {n['yatirim_kapasitesi']:,.2f} TL. Onerilen dagilim asagida."
    if niyet == "gecmis_odeme":
        rapor = pano.get("sorgu_nakit_raporu") or {}
        hareketler = rapor.get("hareketler") or [x for x in pano.get("son_hareketler", []) if x.get("yon") == "Ödeme"]
        hareket_ozeti = rapor.get("ozet") or pano.get("nakit_hareket_ozeti") or {
            "odeme_tutari": sum(float(x.get("tutar") or 0) for x in hareketler),
            "odeme_adedi": len(hareketler), "finansman_girisi": 0,
        }
        donem = rapor.get("etiket") or "Tüm dönem"
        if not hareket_ozeti["odeme_adedi"]:
            return f"{donem} için merkezi defterde tamamlanmış nakit ödeme bulunamadı."
        if rapor.get("detay"):
            satirlar = [
                f"{i}. {datetime.strptime(x['tarih'], '%Y-%m-%d').strftime('%d.%m.%Y')} · "
                f"{x['cari']} · {x['aciklama'] or x['proje'] or 'Açıklama yok'} · {_tl_yaz(x['tutar'])} TL"
                for i, x in enumerate(hareketler, 1)
            ]
            return (
                f"{donem} ödeme listesi ({len(hareketler)} kayıt):\n" + "\n".join(satirlar) +
                f"\nToplam: {_tl_yaz(hareket_ozeti['odeme_tutari'])} TL."
            )
        satirlar = [f"{x['tarih']} {x['cari']} {_tl_yaz(x['tutar'])} TL" for x in hareketler[:3]]
        cevap = (
            f"{donem} için {hareket_ozeti['odeme_adedi']} nakit ödeme yapılmış. "
            f"Toplam ödeme {_tl_yaz(hareket_ozeti['odeme_tutari'])} TL."
        )
        if hareket_ozeti.get("finansman_girisi"):
            cevap += (
                f" {_tl_yaz(hareket_ozeti['finansman_girisi'])} TL kredi kullanımı nakit girişidir; "
                "ödeme veya tahsilat toplamına dahil edilmedi."
            )
        if satirlar:
            cevap += " Son ödemeler: " + "; ".join(satirlar) + "."
        return cevap
    if niyet == "gecmis_tahsilat":
        rapor = pano.get("sorgu_nakit_raporu") or {}
        hareket_ozeti = rapor.get("ozet") or pano.get("nakit_hareket_ozeti") or {}
        hareketler = rapor.get("hareketler") or []
        donem = rapor.get("etiket") or "Tüm dönem"
        tutar = float(hareket_ozeti.get("tahsilat_tutari") or 0)
        adet = int(hareket_ozeti.get("tahsilat_adedi") or 0)
        if not adet:
            cevap = f"{donem} için merkezi defterde tamamlanmış müşteri/cari tahsilatı bulunmuyor."
        elif rapor.get("detay"):
            satirlar = [
                f"{i}. {datetime.strptime(x['tarih'], '%Y-%m-%d').strftime('%d.%m.%Y')} · "
                f"{x['cari']} · {x['aciklama'] or x['proje'] or 'Açıklama yok'} · {_tl_yaz(x['tutar'])} TL"
                for i, x in enumerate(hareketler, 1)
            ]
            cevap = (
                f"{donem} tahsilat listesi ({len(hareketler)} kayıt):\n" + "\n".join(satirlar) +
                f"\nToplam: {_tl_yaz(tutar)} TL."
            )
        else:
            cevap = f"{donem} için {adet} tahsilat yapılmış. Toplam tahsilat {_tl_yaz(tutar)} TL."
        if hareket_ozeti.get("finansman_girisi"):
            cevap += (
                f" {_tl_yaz(hareket_ozeti['finansman_girisi'])} TL kredi kullanımı finansman girişidir; "
                "tahsilat değildir."
            )
        return cevap
    if niyet == "odeme":
        if not pano["odemeler"]:
            return "Acik borc bakiyesi bulunmadi. Vade bilgileri tamamlandikca odeme sirasi burada olusur."
        ilk = pano["odemeler"][0]
        return f"Ilk incelenecek odeme {ilk['cari']} icin {ilk['tutar']:,.2f} TL. Durum: {ilk['durum']}."
    if niyet == "finansman":
        if n["finansman_acigi"]:
            return f"Mevcut kayitlara gore tahmini finansman ihtiyaci {n['finansman_acigi']:,.2f} TL."
        return "30 gunluk vade ve rezerv hesabinda finansman acigi gorunmuyor."
    if niyet == "tahsilat":
        if o["alacak"] <= 0:
            return "Merkezi defterde açık alacak görünmüyor."
        return f"Toplam açık alacak {o['alacak']:,.2f} TL. 30 gün içinde vadeye bağlanmış beklenen tahsilat {n['30_gun_beklenen_tahsilat']:,.2f} TL."
    if niyet == "risk":
        r = pano["risk"]
        return f"Finansal risk skoru 100 üzerinden {r['skor']} ve seviye {r['seviye']}. Likidite karşılama oranı %{r['likidite_karsilama']:.1f}."
    if niyet == "karlilik":
        return f"Kayıtlı gelir {o['gelir']:,.2f} TL, gider {o['gider']:,.2f} TL ve net sonuç {o['net']:,.2f} TL."
    if niyet == "veri_kalitesi":
        k = pano["veri_kalitesi"]
        detaylar = k.get("eksik_kayitlar") or []
        if detaylar:
            satirlar = []
            for x in detaylar[:5]:
                eksik = ", ".join(x["eksikler"])
                tanim = x["cari"] or x["aciklama"] or x["fis_id"]
                satirlar.append(f"{x['fis_id']} ({tanim}): {eksik} eksik")
            return "Eksik alan bulunan gerçek kayıtlar: " + "; ".join(satirlar) + "."
        sorun = "; ".join(k["sorunlar"][:3]) or "kritik eksik bulunmadı"
        return f"Veri kalite skoru 100 üzerinden {k['skor']} ({k['seviye']}). Kontrol: {sorun}."
    if niyet == "nakit":
        son = pano["nakit_projeksiyonu"][-1]["kapanis"] if pano["nakit_projeksiyonu"] else o["nakit"]
        return f"Bugünkü nakit {o['nakit']:,.2f} TL. Kayıtlı vadeler değişmezse 13. hafta tahmini kapanış {son:,.2f} TL."
    if niyet == "aksiyon":
        aksiyonlar = pano.get("aksiyonlar", [])[:3]
        if not aksiyonlar:
            return "Bugün için kritik aksiyon bulunmuyor; likidite ve vadeleri izlemeye devam edin."
        return "Öncelikli aksiyonlar: " + "; ".join(f"{x['baslik']} ({x['neden']})" for x in aksiyonlar) + "."
    return (
        f"Nakit {o['nakit']:,.2f} TL, acik alacak {o['alacak']:,.2f} TL, borc {o['borc']:,.2f} TL. "
        + pano["ana_karar"]
    )


def asistana_sor(soru, faaliyet="tum", profil=None, llm_kullan=True, aktor="cfo_asistani",
                 onceki_niyet=None, onceki_soru=None, degisiklik_onayli=False):
    soru = re.sub(r"\s+", " ", str(soru or "")).strip()[:1500]
    if not soru:
        raise ValueError("Soru bos olamaz.")
    faaliyet_sorusu = f"{soru} {onceki_soru or ''}" if _detay_istegi(soru) else soru
    faaliyet = _sorudan_faaliyet(faaliyet_sorusu, faaliyet)
    pano = karar_panosu(faaliyet, profil)
    komut = _vade_komutu_coz(soru, pano)
    if komut:
        if komut.get("hata"):
            _audit(soru, faaliyet, "vade_guncelleme", pano["durum"], "deterministik_komut")
            return {
                "cevap": komut["hata"], "niyet": "vade_guncelleme", "llm": "deterministik_komut",
                "pano": pano, "aksiyonlar": [], "kanitlar": [], "guven": "Yüksek",
                "degisiklik_yapildi": False,
            }
        if not degisiklik_onayli:
            return {
                "cevap": "Bu komut ana Excel'deki vadeyi değiştirecek. Devam etmek için onay şifresini girin.",
                "niyet": "vade_guncelleme_onayi", "llm": "deterministik_komut",
                "pano": pano, "aksiyonlar": [], "kanitlar": [], "guven": "Yüksek",
                "parola_gerekli": True, "tekrar_sorusu": soru,
                "degisiklik_yapildi": False,
            }
        sonuc = acik_borc_vadesi_guncelle(komut["hedefler"], komut["vade"], aktor=aktor)
        pano = karar_panosu(faaliyet, profil)
        adlar = ", ".join(x["cari"] for x in sonuc["hedefler"])
        tarih_metin = datetime.strptime(sonuc["vade"], "%Y-%m-%d").strftime("%d.%m.%Y")
        cevap = (
            f"Kaydedildi. {adlar} için vade {tarih_metin} olarak güncellendi. "
            f"Merkezi defterde {len(sonuc['fisler'])} fiş ve ana Excel'de "
            f"{len(sonuc['excel_satirlari'])} satır doğrulandı."
        )
        _audit(soru, faaliyet, "vade_guncelleme", pano["durum"], "deterministik_komut")
        return {
            "cevap": cevap, "niyet": "vade_guncelleme", "llm": "deterministik_komut",
            "pano": pano, "aksiyonlar": ["Dashboard yeni vadelerle yenilendi."],
            "kanitlar": [f"Vade: {tarih_metin}", f"Fişler: {', '.join(sonuc['fisler'])}"],
            "guven": "Yüksek", "degisiklik_yapildi": True, "degisiklik": sonuc,
        }
    niyet = _niyet(soru)
    rapor_sorusu = soru
    kesin_niyetler = {"gecmis_odeme", "gecmis_tahsilat"}
    if _detay_istegi(soru) and niyet not in kesin_niyetler and onceki_niyet in kesin_niyetler:
        niyet = onceki_niyet
        rapor_sorusu = str(onceki_soru or soru)
    if niyet in kesin_niyetler:
        pano["sorgu_nakit_raporu"] = _gecmis_nakit_raporu(
            faaliyet, rapor_sorusu, niyet, detay_sorusu=soru,
        )
    cevap = _deterministik_cevap(pano, niyet)
    risk = pano.get("risk") or {"skor": 0, "seviye": "Bilinmiyor"}
    kalite = pano.get("veri_kalitesi") or {"skor": 100, "seviye": "İyi", "sorunlar": []}
    aksiyonlar = [x["baslik"] for x in pano.get("aksiyonlar", [])[:3]]
    kanitlar = [
        f"Nakit: {pano['ozet']['nakit']:,.2f} TL",
        f"Borç: {pano['ozet']['borc']:,.2f} TL",
        f"30 günlük finansman açığı: {pano['nakit_plani']['finansman_acigi']:,.2f} TL",
        f"Risk skoru: {risk['skor']}/100",
    ]
    guven = "Yüksek" if kalite["skor"] >= 85 else ("Orta" if kalite["skor"] >= 60 else "Düşük")
    llm_durumu = "kullanilmadi"
    kesin_nakit_niyeti = niyet in kesin_niyetler
    if kesin_nakit_niyeti:
        rapor = pano.get("sorgu_nakit_raporu") or {}
        h = rapor.get("ozet") or pano.get("nakit_hareket_ozeti") or {}
        aksiyonlar = []
        guven = "Yüksek"
        if niyet == "gecmis_odeme":
            kanitlar = [
                f"Nakit ödeme adedi: {int(h.get('odeme_adedi') or 0)}",
                f"Toplam nakit ödeme: {_tl_yaz(h.get('odeme_tutari'))} TL",
                f"Hariç tutulan kredi kullanımı: {_tl_yaz(h.get('finansman_girisi'))} TL",
            ]
        else:
            kanitlar = [
                f"Tahsilat adedi: {int(h.get('tahsilat_adedi') or 0)}",
                f"Toplam tahsilat: {_tl_yaz(h.get('tahsilat_tutari'))} TL",
                f"Tahsilat sayılmayan kredi kullanımı: {_tl_yaz(h.get('finansman_girisi'))} TL",
            ]
        hedef_yon = "Ödeme" if niyet == "gecmis_odeme" else "Tahsilat"
        hareketler = rapor.get("hareketler") or [
            x for x in pano.get("son_hareketler", []) if x.get("yon") == hedef_yon
        ]
        if rapor.get("detay"):
            kanitlar = [
                f"{x['fis_id']}: {_tl_yaz(x['tutar'])} TL · {x['cari']} · {x['aciklama']}"
                for x in hareketler[:6]
            ]
            toplam_alan = "odeme_tutari" if niyet == "gecmis_odeme" else "tahsilat_tutari"
            kanitlar.insert(0, f"Doğrulanan toplam: {_tl_yaz(h.get(toplam_alan))} TL")
        else:
            for x in hareketler[:2]:
                kanitlar.append(f"{x['fis_id']}: {_tl_yaz(x['tutar'])} TL {x['cari']}")
    veri_detayi = niyet == "veri_kalitesi" and any(x in _norm(soru) for x in ("hangi", "hangisi", "neler"))
    if llm_kullan and not veri_detayi and not kesin_nakit_niyeti:
        kanit = {
            "ozet": pano["ozet"], "nakit_plani": pano["nakit_plani"],
            "ana_karar": pano["ana_karar"], "ilk_odemeler": pano["odemeler"][:5],
            "son_hareketler": pano.get("son_hareketler", [])[:10],
            "yatirim_sepeti": pano["yatirim_sepeti"], "uyarilar": pano["uyarilar"],
            "risk": risk, "veri_kalitesi": kalite,
            "aksiyonlar": pano.get("aksiyonlar", []), "nakit_projeksiyonu": pano.get("nakit_projeksiyonu", []),
            "faaliyet_performansi": pano.get("faaliyet_performansi", []),
            "eksik_kayitlar": kalite.get("eksik_kayitlar", []),
            "nakit_hareket_ozeti": pano.get("nakit_hareket_ozeti", {}),
        }
        prompt = (
            "Sen Touche Holding'in yerel CFO karar destek asistanısın. Yalnızca verilen kanıtı kullan. "
            "Yeni rakam, kayıt adı, piyasa getirisi, ürün veya kesin kazanç uydurma. Para transferi yapıldığını söyleme. "
            "Eksik kayıt sorularında yalnızca eksik_kayitlar listesindeki fişleri kullan. "
            "Yanıtı yönetici özeti gibi kısa, net ve Türkçe yaz; tutarları gerekçesiyle açıkla. "
            "En fazla 3 uygulanabilir aksiyon ve en fazla 4 kanıt üret. "
            "JSON döndür: {\"cevap\":\"...\",\"aksiyonlar\":[\"...\"],\"kanitlar\":[\"...\"]}.\n"
            f"Soru: {soru}\nKanit: {json.dumps(kanit, ensure_ascii=False)}"
        )
        try:
            yanit = metne_sor(prompt)
            metin = str(yanit.get("cevap") or "").strip()
            if metin:
                cevap = metin[:3000]
                llm_durumu = "evren_api" if son_metin_saglayici() == "evren" else "yerel_ollama"
                if isinstance(yanit.get("aksiyonlar"), list):
                    aksiyonlar = [str(x)[:300] for x in yanit["aksiyonlar"][:3] if str(x).strip()]
                if isinstance(yanit.get("kanitlar"), list):
                    kanitlar = [str(x)[:300] for x in yanit["kanitlar"][:4] if str(x).strip()]
        except Exception as exc:
            log.warning("LLM aciklamasi kullanilamadi: %s", exc)
            llm_durumu = "deterministik_yedek"
    _audit(soru, faaliyet, niyet, pano["durum"], llm_durumu)
    return {
        "cevap": cevap, "niyet": niyet, "llm": llm_durumu, "pano": pano,
        "aksiyonlar": aksiyonlar, "kanitlar": kanitlar, "guven": guven,
        "degisiklik_yapildi": False,
        "baglam_sorusu": rapor_sorusu if kesin_nakit_niyeti else soru,
    }


def _audit(soru, faaliyet, niyet, durum, llm):
    AUDIT_PATH.parent.mkdir(parents=True, exist_ok=True)
    kayit = {
        "zaman": datetime.now().isoformat(timespec="seconds"),
        "soru_hash": hashlib.sha256(soru.encode("utf-8")).hexdigest(),
        "faaliyet": faaliyet, "niyet": niyet, "durum": durum, "llm": llm,
    }
    with AUDIT_PATH.open("a", encoding="utf-8") as f:
        f.write(json.dumps(kayit, ensure_ascii=False) + "\n")
