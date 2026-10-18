"""Excel + SQLite arasindaki finansal yazma protokolunun kalici gunlugu.

Excel ve SQLite tek bir ACID islemi olamayacagi icin bu modul iki kaynagi
"tek islem" gibi gostermeye calismaz. Bunun yerine once SQLite'da kalici bir
PREPARED kaydi olusturur, her gecisi journal'a yazar ve yarim kalan islemleri
uygulama yeniden basladiginda guvenli bicimde inceler.
"""

from __future__ import annotations

import hashlib
import json
import logging
import sqlite3
import uuid
from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import config
import finans_cekirdegi as cekirdek
import guvenlik
from guvenlik import audit_yaz
from openpyxl import load_workbook

log = logging.getLogger(__name__)

TERMINAL_STATES = {"COMMITTED", "ABORTED"}
RECOVERABLE_STATES = {"PREPARED", "EXCEL_WRITTEN", "DB_POSTED", "PENDING_RECOVERY"}
REVIEW_STATES = {"NEEDS_REVIEW"}
ALL_STATES = TERMINAL_STATES | RECOVERABLE_STATES | REVIEW_STATES


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _json(value: Any) -> str:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
    )


def _hash_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _canonical_value(name: str, value: Any) -> Any:
    """Normalize workbook/JSON representations to one financial value."""
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if name in {"Tarih", "Vade", "Vade_Tarihi"}:
        text = str(value).strip()
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()  # noqa: DTZ007 - date-only value
            except ValueError:
                continue
    if name in {"Tutar", "Tutar_TL", "Kur"}:
        try:
            text = str(value).strip().replace(" ", "")
            if "," in text and "." in text:
                text = text.replace(".", "").replace(",", ".")
            elif "," in text:
                text = text.replace(",", ".")
            return format(Decimal(text).normalize(), "f")
        except (InvalidOperation, TypeError, ValueError):
            pass
    return str(value).strip()


def _file_hash(path: Path) -> str | None:
    if not path.exists():
        return None
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _source_hash(hazir: dict) -> str:
    value = str(hazir.get("belge_hash") or "").strip().lower()
    if len(value) == 64 and all(c in "0123456789abcdef" for c in value):
        return value
    source = str(hazir.get("kaynak_metin") or hazir.get("belge_id") or "")
    return _hash_bytes(source.encode("utf-8"))


def _financial_hash(hazir: dict) -> str:
    records = []
    for record in hazir.get("veri", {}).get("kayitlar", []):
        operation = str(record.get("Islem_Turu") or "")
        operation = {
            "Usta ödemesi": "Usta/Tedarikçi Ödemesi",
            "Ortak cari ödemesi": "Ortak Cari Ödemesi",
            "Banka masrafı": "Banka Masrafı",
            "Diğer": "Gider Tahakkuku",
        }.get(operation, operation)
        # Excel deliberately derives display labels and descriptions. The hash
        # therefore covers persisted financial semantics, not presentation text.
        semantic = {
            "Tarih": record.get("Tarih"),
            "Faaliyet": record.get("Faaliyet") or hazir.get("firma_key"),
            "Proje": record.get("Proje"),
            "Cari": record.get("Cari") or record.get("Alici") or record.get("Tedarikci"),
            "Islem_Turu": operation,
            "Hesap_Kasa": record.get("Hesap_Kasa"),
            "Tutar": record.get("Tutar") or record.get("Tutar_TL") or record.get("Toplam"),
            "PB": str(record.get("PB") or "TL").upper(),
            "Kur": record.get("Kur") or 1,
            "Vade": record.get("Vade") or record.get("Vade_Tarihi"),
            "Banka_Masrafi": record.get("Banka_Masrafi") or 0,
        }
        records.append(
            {
                str(k): _canonical_value(str(k), v)
                for k, v in semantic.items()
            }
        )
    canonical = {
        "firma_key": str(hazir.get("firma_key") or ""),
        "tur": str(hazir.get("tur") or ""),
        "kayitlar": records,
    }
    return _hash_bytes(_json(canonical).encode("utf-8"))


def _expected_ids(document_id: str, count: int) -> tuple[list[str], list[str]]:
    excel = [f"{document_id}-{i:03d}" for i in range(1, count + 1)]
    db = [f"JRN-{document_id}-{i:03d}" for i in range(1, count + 1)]
    return excel, db


def _ledger() -> tuple[str, dict, int]:
    ledger = config.FIRMALAR["touche_holding"]["ana_ledger"]
    return ledger["sayfa"], ledger["sutunlar"], ledger["ilk_veri_satiri"]


def _read_only_connection() -> sqlite3.Connection | None:
    """Open the journal without creating files, WAL sidecars or parent folders."""
    path = cekirdek.DB_PATH
    if not path.exists():
        return None
    try:
        con = sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True, timeout=30)
    except sqlite3.Error:
        return None
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    return con


def _cell_value(row: dict, name: str) -> Any:
    return row.get(name)


def _excel_rows(document_id: str, expected_ids: list[str] | None = None) -> list[dict]:
    sheet, columns, first = _ledger()
    path = config.ana_excel_yolu()
    if not path.exists():
        return []
    wanted = set(expected_ids or [])
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            return []
        ws = workbook[sheet]
        last = max(
            (ord(col) - 64 for col in columns.values() if len(col) == 1), default=26
        )
        rows = []
        for number, values in enumerate(
            ws.iter_rows(min_row=first, max_col=last, values_only=True), first
        ):

            def get(name: str, row_values=values):
                col = columns.get(name, "")
                index = ord(col) - 65 if len(col) == 1 else None
                return (
                    row_values[index]
                    if index is not None and index < len(row_values)
                    else None
                )

            source_id = str(get("Kaynak_Belge_ID") or "").strip()
            reference = str(get("Belge_Referans") or "").strip()
            if not source_id or not (
                source_id in wanted
                or reference == document_id
                or source_id.startswith(document_id + "-")
            ):
                continue
            rows.append(
                {
                    "satir": number,
                    "Kayit_ID": get("Kayit_ID"),
                    "Kaynak_Belge_ID": source_id,
                    "Belge_Referans": reference,
                    "Tarih": get("Tarih"),
                    "Faaliyet": get("Faaliyet"),
                    "Proje": get("Proje"),
                    "Cari_Tipi": get("Cari_Tipi"),
                    "Cari": get("Cari"),
                    "Islem_Turu": get("Islem_Turu"),
                    "Hesap_Kasa": get("Hesap_Kasa"),
                    "Tutar": get("Tutar"),
                    "PB": get("PB") or "TL",
                    "Kur": get("Kur") or 1,
                    "Baglanti_ID": get("Baglanti_ID"),
                    "Vade": get("Vade"),
                    "Aciklama": get("Aciklama"),
                }
            )
        return rows
    finally:
        workbook.close()


def _row_fingerprint(rows: list[dict]) -> str:
    fields = (
        "Kaynak_Belge_ID",
        "Tarih",
        "Faaliyet",
        "Proje",
        "Cari_Tipi",
        "Cari",
        "Islem_Turu",
        "Hesap_Kasa",
        "Tutar",
        "PB",
        "Kur",
        "Baglanti_ID",
        "Vade",
        "Aciklama",
    )
    payload = [
        {field: row.get(field) for field in fields}
        for row in sorted(rows, key=lambda x: str(x.get("Kaynak_Belge_ID")))
    ]
    return _hash_bytes(_json(payload).encode("utf-8"))


def _excel_financial_hash(tx: dict, rows: list[dict]) -> str:
    """Return the approved financial hash represented by persisted Excel rows."""
    if not rows:
        return ""
    return _financial_hash(_hazir_from_rows(tx, rows))


def _expected_db_snapshot(hazir: dict) -> list[dict]:
    """Build the exact SQLite content expected for the approved payload."""
    taslaklar = cekirdek.fis_taslaklari(hazir)
    return [
        {
            "fis_id": fis["fis_id"],
            "belge_id": fis["belge_id"],
            "kayit_sira": fis["sira"],
            "faaliyet_key": fis["faaliyet"],
            "tarih": fis["tarih"],
            "vade": fis["vade"],
            "proje": fis["proje"],
            "cari": fis["cari"],
            "aciklama": fis["aciklama"],
            "para_birimi": fis["pb"],
            "kur_mikro": fis["kur_mikro"],
            "satirlar": [
                {
                    "satir_no": no,
                    "hesap_kodu": line["hesap"],
                    "borc_kurus": line["borc"],
                    "alacak_kurus": line["alacak"],
                    "aciklama": line["aciklama"],
                }
                for no, line in enumerate(fis["satirlar"], 1)
            ],
        }
        for fis in taslaklar
    ]


def _snapshot_hash(snapshot: list[dict]) -> str:
    # Descriptions are presentation metadata and may be normalized by the
    # Excel writer. Amounts, currencies, accounts and debit/credit lines are
    # the integrity-bearing SQLite content.
    financial = []
    for item in snapshot:
        financial.append(
            {
                key: item.get(key)
                for key in (
                    "fis_id", "belge_id", "kayit_sira", "faaliyet_key", "tarih",
                    "vade", "proje", "cari", "para_birimi", "kur_mikro",
                )
            }
            | {
                "satirlar": [
                    {
                        key: line.get(key)
                        for key in ("satir_no", "hesap_kodu", "borc_kurus", "alacak_kurus")
                    }
                    for line in item.get("satirlar", [])
                ]
            }
        )
    return _hash_bytes(_json(financial).encode("utf-8"))


def _db_content_matches(tx: dict, hazir: dict, *, read_only: bool = False) -> bool:
    expected_hash = str(tx.get("expected_db_fingerprint") or "")
    actual = _db_snapshot(tx["document_id"], read_only=read_only)
    if not actual:
        return False
    if expected_hash:
        return _snapshot_hash(actual) == expected_hash
    if read_only:
        # A legacy journal without a persisted expected DB fingerprint cannot be
        # proven safe without reconstructing schema state. Dry-run stays read-only.
        return False
    # Backward-compatible journal rows are verified against the approved payload
    # even when they predate the persisted DB fingerprint column.
    return _snapshot_hash(actual) == _snapshot_hash(_expected_db_snapshot(hazir))


def _db_ids(document_id: str, *, read_only: bool = False) -> list[str]:
    con = _read_only_connection() if read_only else None
    if read_only and con is None:
        return []
    context = con if con is not None else cekirdek._baglanti()
    if con is not None:
        try:
            return [
                str(row[0])
                for row in con.execute(
                    "SELECT fis_id FROM fisler WHERE belge_id=? ORDER BY kayit_sira",
                    (document_id,),
                )
            ]
        finally:
            con.close()
    with context as connection:
        return [
            str(row[0])
            for row in connection.execute(
                "SELECT fis_id FROM fisler WHERE belge_id=? ORDER BY kayit_sira",
                (document_id,),
            )
        ]


def _db_snapshot(document_id: str, *, read_only: bool = False) -> list[dict]:
    con = _read_only_connection() if read_only else None
    if read_only and con is None:
        return []
    context = con if con is not None else cekirdek._baglanti()
    def read(connection):
        fisler = connection.execute(
            "SELECT fis_id,belge_id,kayit_sira,faaliyet_key,tarih,vade,proje,cari,"
            "aciklama,para_birimi,kur_mikro FROM fisler WHERE belge_id=? ORDER BY kayit_sira",
            (document_id,),
        ).fetchall()
        snapshot = []
        for fis in fisler:
            lines = connection.execute(
                "SELECT satir_no,hesap_kodu,borc_kurus,alacak_kurus,aciklama "
                "FROM fis_satirlari WHERE fis_id=? ORDER BY satir_no",
                (fis["fis_id"],),
            ).fetchall()
            snapshot.append(
                {
                    key: fis[key]
                    for key in (
                        "fis_id", "belge_id", "kayit_sira", "faaliyet_key", "tarih",
                        "vade", "proje", "cari", "aciklama", "para_birimi", "kur_mikro",
                    )
                }
                | {
                    "satirlar": [dict(line) for line in lines],
                }
            )
        return snapshot
    if con is not None:
        try:
            return read(con)
        finally:
            con.close()
    with context as connection:
        return read(connection)


def _event(tx_id: str, state: str, event: str, detail: Any = None) -> None:
    with cekirdek._baglanti() as con:
        con.execute(
            "INSERT INTO finans_islem_olaylari(transaction_id,state,event,created_at,detail) VALUES(?,?,?,?,?)",
            (
                tx_id,
                state,
                event,
                _now(),
                _json(detail) if detail is not None else None,
            ),
        )
    try:
        audit_yaz(event, transaction_id=tx_id, status=state, detail=detail or {})
    except (OSError, sqlite3.Error, RuntimeError) as exc:
        log.error("Transaction audit yazilamadi: %s", exc)
        with cekirdek._baglanti() as con:
            con.execute(
                "UPDATE finans_islemleri SET audit_repair_required=1 WHERE transaction_id=?",
                (tx_id,),
            )


def _transition(
    tx_id: str,
    state: str,
    *,
    error: str | None = None,
    excel_fingerprint: str | None = None,
    workbook_after: str | None = None,
    attempts: int | None = None,
) -> None:
    if state not in ALL_STATES:
        raise ValueError(f"Geçersiz işlem durumu: {state}")
    with cekirdek._baglanti() as con:
        current = con.execute(
            "SELECT state FROM finans_islemleri WHERE transaction_id=?", (tx_id,)
        ).fetchone()
        if not current:
            raise ValueError(f"İşlem günlüğünde bulunamadı: {tx_id}")
        sets = ["state=?", "updated_at=?"]
        values: list[Any] = [state, _now()]
        if error is not None:
            sets.append("last_error=?")
            values.append(str(error)[:2000])
        if excel_fingerprint is not None:
            sets.append("expected_excel_fingerprint=?")
            values.append(excel_fingerprint)
        if workbook_after is not None:
            sets.append("workbook_hash_after=?")
            values.append(workbook_after)
        if attempts is not None:
            sets.append("recovery_attempts=?")
            values.append(int(attempts))
        values.append(tx_id)
        con.execute(
            f"UPDATE finans_islemleri SET {', '.join(sets)} WHERE transaction_id=?",
            values,
        )
    _event(
        tx_id,
        state,
        f"transaction_{state.casefold()}",
        {"error": str(error)[:500] if error else None},
    )


def _find_transaction(document_id: str, financial_hash: str) -> dict | None:
    with cekirdek._baglanti() as con:
        row = con.execute(
            "SELECT * FROM finans_islemleri WHERE document_id=? ORDER BY created_at DESC LIMIT 1",
            (document_id,),
        ).fetchone()
    if not row:
        return None
    item = dict(row)
    if item["financial_hash"] != financial_hash:
        raise ValueError(
            "Aynı belge kimliği farklı finansal içerikle yeniden kullanılamaz."
        )
    return item


def prepare_transaction(hazir: dict, approval_identity: str) -> dict:
    """PREPARED kaydını Excel'e dokunmadan ve idempotent biçimde oluşturur."""
    cekirdek.veritabanini_hazirla()
    document_id = str(hazir.get("belge_id") or "").strip()
    if not document_id:
        raise ValueError("Belge ID olmadan işlem günlüğü oluşturulamaz.")
    records = hazir.get("veri", {}).get("kayitlar", [])
    if not records:
        raise ValueError("İşlem günlüğü için en az bir finansal kayıt gerekli.")
    financial_hash = _financial_hash(hazir)
    existing = _find_transaction(document_id, financial_hash)
    if existing:
        return existing
    tx_id = (
        "TX-"
        + uuid.uuid5(uuid.NAMESPACE_URL, f"touche:{document_id}:{financial_hash}").hex
    )
    excel_ids, db_ids = _expected_ids(document_id, len(records))
    expected_db_fingerprint = _snapshot_hash(_expected_db_snapshot(hazir))
    now = _now()
    workbook_before = _file_hash(config.ana_excel_yolu())
    try:
        with cekirdek._baglanti() as con:
            con.execute(
                "INSERT INTO finans_islemleri(transaction_id,document_id,source_hash,financial_hash,approval_identity,"
                "firma_key,belge_turu,state,created_at,updated_at,expected_excel_ids,expected_db_ids,"
                "expected_db_fingerprint,workbook_hash_before) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    tx_id,
                    document_id,
                    _source_hash(hazir),
                    financial_hash,
                    str(approval_identity),
                    str(hazir.get("firma_key") or ""),
                    str(hazir.get("tur") or ""),
                    "PREPARED",
                    now,
                    now,
                    _json(excel_ids),
                    _json(db_ids),
                    expected_db_fingerprint,
                    workbook_before,
                ),
            )
    except sqlite3.IntegrityError:
        existing = _find_transaction(document_id, financial_hash)
        if existing:
            return existing
        raise
    _event(tx_id, "PREPARED", "transaction_prepared", {"document_id": document_id})
    return {
        "transaction_id": tx_id,
        "document_id": document_id,
        "financial_hash": financial_hash,
        "source_hash": _source_hash(hazir),
        "firma_key": str(hazir.get("firma_key") or ""),
        "belge_turu": str(hazir.get("tur") or ""),
        "state": "PREPARED",
        "expected_excel_ids": _json(excel_ids),
        "expected_db_ids": _json(db_ids),
        "expected_db_fingerprint": expected_db_fingerprint,
    }


def _load_tx(tx_id: str, *, read_only: bool = False) -> dict:
    con = _read_only_connection() if read_only else None
    if read_only and con is None:
        raise ValueError("İşlem günlüğü mevcut değil.")
    context = con if con is not None else cekirdek._baglanti()
    if con is not None:
        try:
            row = con.execute(
                "SELECT * FROM finans_islemleri WHERE transaction_id=?", (tx_id,)
            ).fetchone()
        finally:
            con.close()
    else:
        with context as connection:
            row = connection.execute(
                "SELECT * FROM finans_islemleri WHERE transaction_id=?", (tx_id,)
            ).fetchone()
    if not row:
        raise ValueError(f"İşlem günlüğünde bulunamadı: {tx_id}")
    return dict(row)


def mark_audit_repair(tx_id: str, error: str) -> None:
    """Tamamlanan finansal kaydın audit yazımı başarısızsa onarım kanıtı bırakır."""
    with cekirdek._baglanti() as con:
        con.execute(
            "UPDATE finans_islemleri SET audit_repair_required=1,last_error=?,updated_at=? "
            "WHERE transaction_id=?",
            (str(error)[:2000], _now(), tx_id),
        )


def _parse_list(value: Any) -> list[str]:
    try:
        parsed = json.loads(value or "[]")
        if not isinstance(parsed, list) or not all(isinstance(x, str) for x in parsed):
            raise ValueError
        return parsed
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("İşlem günlüğü kimlik listesi bozuk.") from exc


def _hazir_from_rows(tx: dict, rows: list[dict]) -> dict:
    records = []
    for row in sorted(rows, key=lambda x: str(x.get("Kaynak_Belge_ID"))):
        record = {
            key: row.get(key)
            for key in (
                "Tarih",
                "Faaliyet",
                "Proje",
                "Cari_Tipi",
                "Cari",
                "Islem_Turu",
                "Hesap_Kasa",
                "Tutar",
                "PB",
                "Kur",
                "Baglanti_ID",
                "Vade",
                "Aciklama",
            )
        }
        for key in ("Tarih", "Vade"):
            value = record.get(key)
            if isinstance(value, (date, datetime)):
                record[key] = (
                    value.date().isoformat()
                    if isinstance(value, datetime)
                    else value.isoformat()
                )
        records.append(record)
    return {
        "basari": True,
        "belge_id": tx["document_id"],
        "belge_hash": tx.get("source_hash") or "",
        "firma_key": tx.get("firma_key") or "",
        "tur": tx.get("belge_turu") or "",
        "veri": {"kayitlar": records},
        "dogrulama": {"durum": "PASS", "skor": 100},
        "gecici_dosyalar": [],
    }


def execute_transaction(
    hazir: dict,
    approval_identity: str,
    *,
    excel_writer: Callable[..., dict],
    db_writer: Callable[..., dict] | None = None,
    excel_context: dict | None = None,
) -> dict:
    """Onayli kaydi PREPARED -> ... -> COMMITTED akisinda uygular."""
    tx = prepare_transaction(hazir, approval_identity)
    tx_id = tx["transaction_id"]
    state = tx.get("state", "PREPARED")
    if state == "COMMITTED":
        expected_excel = _parse_list(tx.get("expected_excel_ids"))
        rows = _excel_rows(tx["document_id"], expected_excel)
        valid = set(expected_excel).issubset(
            {str(row.get("Kaynak_Belge_ID")) for row in rows}
        ) and _excel_financial_hash(tx, rows) == tx.get("financial_hash")
        try:
            valid = valid and _db_content_matches(tx, hazir)
        except (KeyError, OSError, RuntimeError, sqlite3.Error, TypeError, ValueError):
            valid = False
        if not valid:
            _transition(
                tx_id,
                "NEEDS_REVIEW",
                error="Tamamlanmış işlemin Excel/SQLite finansal içeriği değişmiş.",
            )
            return {
                "basari": False,
                "committed": False,
                "status": "NEEDS_REVIEW",
                "recovery_required": True,
                "transaction_id": tx_id,
                "hata": "Tamamlanmış işlem kanıtı değişmiş; otomatik tekrar yapılmadı.",
            }
        return {
            "basari": True,
            "committed": True,
            "status": "COMMITTED",
            "transaction_id": tx_id,
            "tekrar": True,
            "fis_sayisi": len(_db_ids(tx["document_id"])),
        }
    if state in {"NEEDS_REVIEW", "ABORTED"}:
        return {
            "basari": False,
            "committed": False,
            "status": state,
            "recovery_required": state == "NEEDS_REVIEW",
            "transaction_id": tx_id,
            "hata": tx.get("last_error") or "İşlem otomatik olarak devam ettirilemez.",
        }

    expected_excel = _parse_list(tx.get("expected_excel_ids"))
    expected_db = _parse_list(tx.get("expected_db_ids"))
    rows = _excel_rows(tx["document_id"], expected_excel)
    found = {str(row.get("Kaynak_Belge_ID")) for row in rows}
    try:
        if state in {"PREPARED", "PENDING_RECOVERY"}:
            if not set(expected_excel).issubset(found):
                result = excel_writer(
                    hazir["firma_key"],
                    hazir["veri"],
                    hazir["belge_id"],
                    aktor=approval_identity,
                    baglam={
                        "onay_kaynagi": approval_identity,
                        "islem_id": tx_id,
                        **(excel_context or {}),
                    },
                )
                if not result.get("basari"):
                    code = str(result.get("hata_kodu") or "")
                    if code == "ANA_KILITLI":
                        _transition(
                            tx_id, "PENDING_RECOVERY", error=result.get("hata") or code
                        )
                        return {
                            "basari": False,
                            "committed": False,
                            "status": "PENDING_RECOVERY",
                            "recovery_required": True,
                            "transaction_id": tx_id,
                            "hata": result.get("hata") or "Excel kilitli.",
                        }
                    if code != "MUKERRER_BELGE":
                        status = (
                            "NEEDS_REVIEW"
                            if code in {"ANA_DEGISTI", "VERI_KORUMA", "KAPASITE_DOLU"}
                            else "ABORTED"
                        )
                        _transition(tx_id, status, error=result.get("hata") or code)
                        return {
                            "basari": False,
                            "committed": False,
                            "status": status,
                            "recovery_required": status == "NEEDS_REVIEW",
                            "transaction_id": tx_id,
                            "hata": result.get("hata") or "Excel kaydı başarısız.",
                        }
                declared_hash = str(result.get("ana_hash_sonra") or "").strip()
                if (
                    declared_hash
                    and _file_hash(config.ana_excel_yolu()) != declared_hash
                ):
                    _transition(
                        tx_id,
                        "NEEDS_REVIEW",
                        error="Excel yazımı sonrası çalışma kitabı dışarıdan değişti.",
                    )
                    return {
                        "basari": False,
                        "committed": False,
                        "status": "NEEDS_REVIEW",
                        "recovery_required": True,
                        "transaction_id": tx_id,
                        "hata": "Excel dışarıdan değişmiş; otomatik işlem yapılmadı.",
                    }
            rows = _excel_rows(tx["document_id"], expected_excel)
            found = {str(row.get("Kaynak_Belge_ID")) for row in rows}
            if not set(expected_excel).issubset(found):
                raise RuntimeError("Beklenen Excel satır kimlikleri doğrulanamadı.")
            if _excel_financial_hash(tx, rows) != tx["financial_hash"]:
                _transition(
                    tx_id,
                    "NEEDS_REVIEW",
                    error="Excel finansal içeriği onaylanan parmak iziyle uyuşmuyor.",
                )
                return {
                    "basari": False,
                    "committed": False,
                    "status": "NEEDS_REVIEW",
                    "recovery_required": True,
                    "transaction_id": tx_id,
                    "hata": "Excel içeriği değiştirilmiş; otomatik kayıt yapılmadı.",
                }
            fingerprint = _row_fingerprint(rows)
            _transition(
                tx_id,
                "EXCEL_WRITTEN",
                excel_fingerprint=fingerprint,
                workbook_after=_file_hash(config.ana_excel_yolu()),
            )
            state = "EXCEL_WRITTEN"

        tx = _load_tx(tx_id)
        current_fingerprint = _row_fingerprint(
            _excel_rows(tx["document_id"], expected_excel)
        )
        if (
            tx.get("expected_excel_fingerprint")
            and current_fingerprint != tx["expected_excel_fingerprint"]
        ):
            _transition(
                tx_id,
                "NEEDS_REVIEW",
                error="Excel işlem günlüğündeki satır parmak iziyle uyuşmuyor.",
            )
            return {
                "basari": False,
                "committed": False,
                "status": "NEEDS_REVIEW",
                "recovery_required": True,
                "transaction_id": tx_id,
                "hata": "Excel dışarıdan değişmiş; otomatik işlem yapılmadı.",
            }

        existing_db = _db_ids(tx["document_id"])
        if not set(expected_db).issubset(existing_db):
            writer = db_writer or __import__("finans_cekirdegi").finans_kaydi_yaz
            db_result = writer(hazir, approval_identity)
            if not db_result or db_result.get("basari") is False:
                raise RuntimeError(
                    (db_result or {}).get("hata") or "SQLite kaydı başarısız."
                )
        existing_db = _db_ids(tx["document_id"])
        if not set(expected_db).issubset(existing_db):
            raise RuntimeError("Beklenen SQLite fiş kimlikleri doğrulanamadı.")
        try:
            db_content_ok = _db_content_matches(tx, hazir)
        except Exception as exc:  # noqa: BLE001 - database proof fails closed
            _transition(
                tx_id,
                "NEEDS_REVIEW",
                error=f"SQLite fiş içeriği doğrulanamadı: {exc}",
            )
            return {
                "basari": False,
                "committed": False,
                "status": "NEEDS_REVIEW",
                "recovery_required": True,
                "transaction_id": tx_id,
                "hata": "SQLite içeriği doğrulanamadı; otomatik işlem yapılmadı.",
            }
        if not db_content_ok:
            _transition(
                tx_id,
                "NEEDS_REVIEW",
                error="SQLite fiş içeriği onaylanan tutar/hesap parmak iziyle uyuşmuyor.",
            )
            return {
                "basari": False,
                "committed": False,
                "status": "NEEDS_REVIEW",
                "recovery_required": True,
                "transaction_id": tx_id,
                "hata": "SQLite içeriği değiştirilmiş; otomatik işlem yapılmadı.",
            }
        _transition(tx_id, "DB_POSTED")
        _transition(tx_id, "COMMITTED")
        return {
            "basari": True,
            "committed": True,
            "status": "COMMITTED",
            "recovery_required": False,
            "transaction_id": tx_id,
            "fis_sayisi": len(expected_db),
            "yazilan": len(expected_excel),
        }
    except Exception as exc:  # noqa: BLE001 - finansal yazıcı hatası pending'e alınır
        current = _load_tx(tx_id)
        if current["state"] in {"EXCEL_WRITTEN", "DB_POSTED", "PENDING_RECOVERY"}:
            _transition(tx_id, "PENDING_RECOVERY", error=f"{type(exc).__name__}: {exc}")
            return {
                "basari": False,
                "committed": False,
                "status": "PENDING_RECOVERY",
                "recovery_required": True,
                "transaction_id": tx_id,
                "hata": str(exc),
            }
        _transition(tx_id, "NEEDS_REVIEW", error=f"{type(exc).__name__}: {exc}")
        return {
            "basari": False,
            "committed": False,
            "status": "NEEDS_REVIEW",
            "recovery_required": True,
            "transaction_id": tx_id,
            "hata": str(exc),
        }


def _recover_one(tx: dict, dry_run: bool = False) -> dict:
    tx_id = tx["transaction_id"]
    try:
        if tx.get("state") == "NEEDS_REVIEW":
            return {
                "transaction_id": tx_id,
                "status": "NEEDS_REVIEW",
                "action": "Manual yönetici kararı bekleniyor; otomatik kurtarma yapılmadı.",
            }
        if not dry_run:
            _event(
                tx_id,
                tx["state"],
                "transaction_recovery_attempted",
                {
                    "attempt": int(tx.get("recovery_attempts") or 0) + 1,
                },
            )
        expected_excel = _parse_list(tx.get("expected_excel_ids"))
        expected_db = _parse_list(tx.get("expected_db_ids"))
        rows = _excel_rows(tx["document_id"], expected_excel)
        found = {str(row.get("Kaynak_Belge_ID")) for row in rows}
        db_ids = _db_ids(tx["document_id"], read_only=dry_run)
        if (
            tx.get("expected_excel_fingerprint")
            and set(expected_excel).issubset(found)
            and _row_fingerprint(rows) != tx["expected_excel_fingerprint"]
        ):
            action = "NEEDS_REVIEW: Excel satır parmak izi değişmiş"
            if not dry_run:
                _transition(tx_id, "NEEDS_REVIEW", error=action)
            return {
                "transaction_id": tx_id,
                "status": "NEEDS_REVIEW",
                "action": action,
            }
        if set(expected_excel).issubset(found):
            excel_hash = _excel_financial_hash(tx, rows)
            if excel_hash != tx.get("financial_hash"):
                action = "NEEDS_REVIEW: Excel finansal içeriği onaylanan hash ile uyuşmuyor"
                if not dry_run:
                    _transition(tx_id, "NEEDS_REVIEW", error=action)
                return {
                    "transaction_id": tx_id,
                    "status": "NEEDS_REVIEW",
                    "action": action,
                }

        if set(expected_excel).issubset(found) and set(expected_db).issubset(db_ids):
            try:
                hazir = _hazir_from_rows(tx, rows)
                db_match = _db_content_matches(tx, hazir, read_only=dry_run)
            except Exception as exc:  # noqa: BLE001 - integrity proof fails closed
                db_match = False
                action = f"NEEDS_REVIEW: SQLite içeriği doğrulanamadı ({exc})"
            if not db_match:
                action = locals().get(
                    "action",
                    "NEEDS_REVIEW: SQLite fiş içeriği onaylanan içerikle uyuşmuyor",
                )
                if not dry_run:
                    _transition(tx_id, "NEEDS_REVIEW", error=action)
                return {
                    "transaction_id": tx_id,
                    "status": "NEEDS_REVIEW",
                    "action": action,
                }
            action = "COMMITTED: Excel ve SQLite içerikleri doğrulandı"
            if not dry_run and tx["state"] != "COMMITTED":
                _transition(tx_id, "DB_POSTED")
                _transition(tx_id, "COMMITTED")
                _event(
                    tx_id,
                    "COMMITTED",
                    "transaction_recovered",
                    {"document_id": tx["document_id"]},
                )
            return {"transaction_id": tx_id, "status": "COMMITTED", "action": action}
        if set(expected_excel).issubset(found) and not set(expected_db).intersection(
            db_ids
        ):
            action = "DB_POST: Excel satırları tam, SQLite fişleri eksik"
            if dry_run:
                return {
                    "transaction_id": tx_id,
                    "status": "PENDING_RECOVERY",
                    "action": action,
                }
            attempts = int(tx.get("recovery_attempts") or 0) + 1
            _transition(tx_id, "PENDING_RECOVERY", attempts=attempts)
            hazir = _hazir_from_rows(tx, rows)
            try:
                cekirdek.finans_kaydi_yaz(
                    hazir, tx.get("approval_identity") or "recovery"
                )
                if not set(expected_db).issubset(_db_ids(tx["document_id"])):
                    raise RuntimeError("Kurtarma sonrası SQLite fişleri doğrulanamadı.")
                if not _db_content_matches(tx, hazir):
                    raise RuntimeError("Kurtarma sonrası SQLite içeriği doğrulanamadı.")
                _transition(tx_id, "DB_POSTED")
                _transition(tx_id, "COMMITTED")
                _event(
                    tx_id,
                    "COMMITTED",
                    "transaction_recovered",
                    {"document_id": tx["document_id"]},
                )
                return {
                    "transaction_id": tx_id,
                    "status": "COMMITTED",
                    "action": "SQLite güvenle tamamlandı",
                }
            except Exception as exc:  # noqa: BLE001 - recovery tekrar denenebilir pending'dir
                _transition(
                    tx_id, "PENDING_RECOVERY", error=str(exc), attempts=attempts
                )
                return {
                    "transaction_id": tx_id,
                    "status": "PENDING_RECOVERY",
                    "action": str(exc),
                }
        if set(expected_db).intersection(db_ids) and not set(expected_excel).issubset(
            found
        ):
            action = "NEEDS_REVIEW: SQLite var, Excel satırları eksik; otomatik Excel yazılmadı"
        elif not found and not db_ids:
            action = "ABORTED: işlem izi bulunamadı"
        else:
            action = "NEEDS_REVIEW: eksik veya çakışan işlem izi"
        status = action.split(":", 1)[0]
        if not dry_run:
            _transition(tx_id, status, error=action)
        return {"transaction_id": tx_id, "status": status, "action": action}
    except Exception as exc:  # noqa: BLE001 - bozuk journal fail-closed NEEDS_REVIEW'dir
        action = f"NEEDS_REVIEW: bozuk journal ({exc})"
        if not dry_run:
            try:
                _transition(tx_id, "NEEDS_REVIEW", error=action)
            except Exception:
                log.exception("Bozuk journal durumu yazılamadı: %s", tx_id)
        return {"transaction_id": tx_id, "status": "NEEDS_REVIEW", "action": action}


def _journal_rows(*, read_only: bool = False) -> list[dict]:
    if read_only:
        con = _read_only_connection()
        if con is None:
            return []
        try:
            return [
                dict(row)
                for row in con.execute(
                    "SELECT * FROM finans_islemleri "
                    "WHERE state NOT IN ('COMMITTED','ABORTED') ORDER BY created_at"
                )
            ]
        finally:
            con.close()
    with cekirdek._baglanti() as con:
        return [
            dict(row)
            for row in con.execute(
                "SELECT * FROM finans_islemleri "
                "WHERE state NOT IN ('COMMITTED','ABORTED') ORDER BY created_at"
            )
        ]


def recover_pending_transactions(*, dry_run: bool = False) -> dict:
    """Terminal olmayan işlemleri güvenli ve tekrarlanabilir şekilde tarar."""
    if not dry_run:
        cekirdek.veritabanini_hazirla()
    try:
        rows = _journal_rows(read_only=dry_run)
    except sqlite3.Error as exc:
        if not dry_run:
            raise
        return {
            "dry_run": True,
            "islem_sayisi": 0,
            "sonuclar": [],
            "needs_review": 1,
            "committed": 0,
            "pending_recovery": 0,
            "hata": f"Journal salt okunur olarak doğrulanamadı: {exc}",
        }
    results = [_recover_one(row, dry_run=dry_run) for row in rows]
    return {
        "dry_run": dry_run,
        "islem_sayisi": len(results),
        "sonuclar": results,
        "needs_review": sum(x["status"] == "NEEDS_REVIEW" for x in results),
        "committed": sum(x["status"] == "COMMITTED" for x in results),
        "pending_recovery": sum(x["status"] == "PENDING_RECOVERY" for x in results),
    }


def recovery_dry_run() -> dict:
    return recover_pending_transactions(dry_run=True)


def authorize_needs_review(tx_id: str, admin_identity: str, reason: str) -> dict:
    """Require an explicit, audited decision before retrying NEEDS_REVIEW."""
    if not str(admin_identity).strip() or not str(reason).strip():
        raise ValueError("Yönetici kimliği ve karar gerekçesi zorunludur.")
    tx = _load_tx(tx_id)
    if tx["state"] != "NEEDS_REVIEW":
        raise ValueError("Yalnızca NEEDS_REVIEW işlemleri manuel kurtarmaya açılabilir.")
    _event(
        tx_id,
        "NEEDS_REVIEW",
        "transaction_admin_recovery_approved",
        {"admin_identity": str(admin_identity), "reason": str(reason)[:500]},
    )
    _transition(tx_id, "PENDING_RECOVERY", error=f"Manuel yönetici kararı: {reason}")
    return _recover_one(_load_tx(tx_id))


def repair_audit(tx_id: str) -> dict:
    """Repair missing audit evidence without reposting any financial entry."""
    tx = _load_tx(tx_id)
    if tx["state"] != "COMMITTED":
        raise ValueError("Audit onarımı yalnızca COMMITTED işlem için yapılabilir.")
    if not int(tx.get("audit_repair_required") or 0):
        return {"transaction_id": tx_id, "repaired": False, "reason": "gerekmiyor"}
    audit_path = Path(guvenlik.AUDIT_LOG)
    before = _file_hash(audit_path)
    try:
        audit_yaz(
            "transaction_audit_repaired",
            transaction_id=tx_id,
            status="COMMITTED",
            detail={"repair": True},
        )
    except (OSError, RuntimeError) as exc:
        raise RuntimeError(f"Audit onarımı başarısız: {exc}") from exc
    after = _file_hash(audit_path)
    if before == after:
        raise RuntimeError("Audit kaydı yazılmadı; onarım bayrağı korunuyor.")
    with cekirdek._baglanti() as con:
        con.execute(
            "UPDATE finans_islemleri SET audit_repair_required=0,last_error=NULL,updated_at=? "
            "WHERE transaction_id=?",
            (_now(), tx_id),
        )
        con.execute(
            "INSERT INTO finans_islem_olaylari(transaction_id,state,event,created_at,detail) "
            "VALUES(?,?,?,?,?)",
            (tx_id, "COMMITTED", "transaction_audit_repaired", _now(), _json({"repair": True})),
        )
    return {"transaction_id": tx_id, "repaired": True, "financial_posting_unchanged": True}
