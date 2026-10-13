# -*- coding: utf-8 -*-
"""ANA Excel (03_Kayıt Merkezi) ile merkezi SQLite defteri arasinda SALT-OKUNUR mutabakat.

Proje/faaliyet/holding metrikleri merkezi defterden hesaplanir; Excel'de elle girilen
ya da degistirilen satirlar deftere kendiliginden yansimaz. Bu modul farki gorunur kilar
(HICBIR SEY yazmaz, duzeltmez):

  eksik_defterde   Excel'de aktif kayit var, defterde karsiligi yok
  tutar_farki      Ikisinde de var ama TL tutari farkli (Excel sonradan degistirilmis olabilir)
  defterde_fazla   Defterde fis var, Excel'de karsiligi yok (satir Excel'den silinmis olabilir)

Eslestirme anahtari:
  * Kaynak_Belge_ID dolu (uygulamanin yazdigi)  -> fis (belge_id, kayit_sira) = ("<belge>", NNN)
  * Bos (elle girilmis / eski)                  -> belge_id ``LEGACY-<Kayit_ID>``, sira 1

Beklenen Excel satiri: Durum ``Aktif`` / bos / ``Gerçekleşti``. ``Gerçekleşti``, Gursar belge aktariminin
(excel_yazici._yaz_gursar_holding) kendi yazdigi, Excel formullerinde de ``İptal`` degil diye sayilan durumdur.

Iki ek, belirleyici eslesme (her ikisi de TL tutarini kurusu kurusuna dogrular; tutar farkliysa ESLESMEZ):
  * Banka masrafi yoldasi: Kaynak ``<belge>-NNN-MASRAF`` -> ana fisin 780 satiri (uygulamanin yazdigi, gomulu) YA DA
    ``<belge>-MASRAF`` / NNN fisi (ayri masraf fisi).
  * Kaynak anahtari tutmayan satir: ``Belge_Referans`` = fisin belge_id'si, ayni tutar ve cari, birebir (bir fis en fazla
    bir Excel satirina eslenir; ikinci ayni satir ``eksik_defterde`` kalir).
"""
from __future__ import annotations

import re
import sqlite3
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import load_workbook

import config
import finans_cekirdegi as cekirdek

_KAYNAK_RE = re.compile(r"^(?P<belge>.+)-(?P<sira>\d{3})$")
_MASRAF_RE = re.compile(r"^(?P<belge>.+)-(?P<sira>\d{3})-MASRAF$")
# Defterde karsiligi beklenen Excel durumlari (``cekirdek._norm`` ile normalize edilir)
_BEKLENEN_DURUMLAR = {"aktif", "", "gerceklesti"}
# Defterde Excel satiri olmayan, sistemin kendi urettigi fisler (acilis bakiyesi, ters kayit, grup ici cift taraf)
_DEFTER_ONEKLERI = ("OPEN-", "REV-", "IC-", "TRF-", "KUR-", "FX-")


def _kurus(deger) -> int | None:
    try:
        return int((Decimal(str(deger).replace(",", ".")) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except Exception:
        return None


def _excel_satirlari() -> list[dict]:
    ledger = config.FIRMALAR["touche_holding"]["ana_ledger"]
    s = ledger["sutunlar"]
    yol = config.ana_excel_yolu()
    if not Path(yol).exists():
        raise FileNotFoundError(f"Ana Excel bulunamadı: {yol}")
    idx = {ad: ord(h) - ord("A") for ad, h in s.items() if len(h) == 1}     # AA/AB/AC sutunlari gerekmez
    son_sutun = max(idx.values()) + 1
    wb = load_workbook(yol, read_only=True, data_only=True)
    try:
        if ledger["sayfa"] not in wb.sheetnames:
            raise ValueError(f"Sayfa yok: {ledger['sayfa']}")
        ws = wb[ledger["sayfa"]]
        satirlar = []
        for no, v in enumerate(ws.iter_rows(min_row=ledger["ilk_veri_satiri"], max_col=son_sutun, values_only=True),
                               ledger["ilk_veri_satiri"]):
            def h(ad, v=v):
                i = idx.get(ad)
                return v[i] if i is not None and i < len(v) else None
            kayit_id = str(h("Kayit_ID") or "").strip()
            if not kayit_id:
                continue
            satirlar.append({
                "satir": no, "kayit_id": kayit_id, "durum": str(h("Durum") or "").strip(),
                "kaynak_belge": str(h("Kaynak_Belge_ID") or "").strip(),
                "tutar": h("Tutar"), "pb": str(h("PB") or "TL").strip().upper(), "kur": h("Kur"),
                "cari": str(h("Cari") or ""), "proje": str(h("Proje") or ""),
                "belge_ref": str(h("Belge_Referans") or "").strip(),
            })
        return satirlar
    finally:
        wb.close()


def _beklenen_anahtar(satir: dict) -> tuple[str, int]:
    m = _KAYNAK_RE.match(satir["kaynak_belge"]) if satir["kaynak_belge"] else None
    if m:
        return m.group("belge"), int(m.group("sira"))
    return f"LEGACY-{satir['kayit_id']}", 1


def _beklenen_tl_kurus(satir: dict) -> int | None:
    tutar = _kurus(satir["tutar"])
    if tutar is None:
        return None
    if satir["pb"] in ("TL", "TRY", ""):
        return tutar
    # Kur TAM hassasiyetle kullanilir (ornegin 55,6521): kuru 2 haneye yuvarlamak, defterle ayni olan
    # tutarlara sahte "Tutar farki" uretiyordu. Defter de tutar*kur'u kurusa ROUND_HALF_UP ile yuvarlar.
    try:
        kur = Decimal(str(satir["kur"]).replace(",", ".")) if satir["kur"] not in (None, "") else Decimal(0)
    except Exception:
        return None
    if not kur:
        return None
    # tutar(kurus) * kur -> TL kurus
    return int((Decimal(tutar) * kur).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def kontrol() -> dict:
    """Mutabakat raporu. ANA Excel'i ve SQLite'i yalnizca OKUR."""
    excel = [x for x in _excel_satirlari() if cekirdek._norm(x["durum"]) in _BEKLENEN_DURUMLAR]
    yol = Path(cekirdek.DB_PATH)
    fisler: dict[tuple[str, int], dict] = {}
    if yol.exists():
        con = sqlite3.connect(yol, timeout=15)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA query_only=ON")
        try:
            for r in con.execute(
                "SELECT f.belge_id, f.kayit_sira, f.durum, f.cari, "
                "(SELECT s.borc_kurus FROM fis_satirlari s WHERE s.fis_id=f.fis_id ORDER BY s.satir_no LIMIT 1) ana_borc, "
                "(SELECT COALESCE(SUM(s.borc_kurus), 0) FROM fis_satirlari s "
                "WHERE s.fis_id=f.fis_id AND s.hesap_kodu='780') masraf "
                "FROM fisler f"
            ):
                fisler[(r["belge_id"], r["kayit_sira"])] = {
                    "durum": r["durum"], "tutar": r["ana_borc"], "cari": r["cari"], "masraf": int(r["masraf"] or 0)}
        except sqlite3.OperationalError:
            fisler = {}
        finally:
            con.close()

    eksik, fark, beklenen, ikinci_tur = [], [], set(), []
    for x in excel:
        ozet = {"satir": x["satir"], "kayit_id": x["kayit_id"], "cari": x["cari"], "proje": x["proje"],
                "excel_tutar": x["tutar"], "pb": x["pb"]}
        beklenen_tutar = _beklenen_tl_kurus(x)

        # Banka masrafi yoldasi (``<belge>-NNN-MASRAF``): ana fisin gomulu 780 satiri ya da ayri ``<belge>-MASRAF`` fisi
        m = _MASRAF_RE.match(x["kaynak_belge"]) if x["kaynak_belge"] else None
        if m:
            ana_anahtar, kardes_anahtar = (m.group("belge"), int(m.group("sira"))), (f"{m.group('belge')}-MASRAF", int(m.group("sira")))
            ana, kardes = fisler.get(ana_anahtar), fisler.get(kardes_anahtar)
            if kardes and kardes["durum"] == "MUHASEBELESTI" and beklenen_tutar is not None \
                    and kardes["tutar"] is not None and int(kardes["tutar"]) == beklenen_tutar:
                beklenen.add(kardes_anahtar)
            elif not (ana and ana["durum"] == "MUHASEBELESTI" and beklenen_tutar is not None and ana["masraf"] == beklenen_tutar):
                eksik.append(ozet)
            continue

        anahtar = _beklenen_anahtar(x)
        beklenen.add(anahtar)
        fis = fisler.get(anahtar)
        if fis is None:
            ikinci_tur.append((x, ozet, beklenen_tutar))
            continue
        if beklenen_tutar is not None and fis["tutar"] is not None and fis["durum"] == "MUHASEBELESTI" \
                and int(fis["tutar"]) != beklenen_tutar:
            fark.append({**ozet, "defter_tl_kurus": int(fis["tutar"]), "excel_tl_kurus": beklenen_tutar})

    # Kaynak anahtari tutmayan satirlar: Belge_Referans = fis belge_id'si + ayni TL tutari + ayni cari, birebir
    for x, ozet, beklenen_tutar in ikinci_tur:
        adaylar = [k for k, v in fisler.items()
                   if x["belge_ref"] and k[0] == x["belge_ref"] and k not in beklenen and v["durum"] == "MUHASEBELESTI"
                   and beklenen_tutar is not None and v["tutar"] is not None and int(v["tutar"]) == beklenen_tutar
                   and cekirdek._norm(v["cari"]) == cekirdek._norm(x["cari"])]
        if len(adaylar) == 1:
            beklenen.add(adaylar[0])
        else:
            eksik.append(ozet)
    eksik.sort(key=lambda o: o["satir"])
    fazla = sorted(
        {f"{b}-{s:03d}" for (b, s), v in fisler.items()
         if v["durum"] == "MUHASEBELESTI" and (b, s) not in beklenen and not b.startswith(_DEFTER_ONEKLERI)}
    )
    return {"excel_aktif_satir": len(excel), "defter_fis": len(fisler),
            "eksik_defterde": eksik, "tutar_farki": fark, "defterde_fazla": fazla,
            "tutarli": not (eksik or fark or fazla)}


def uyari_metinleri(rapor: dict | None = None) -> list[str]:
    """UI/holding uyarilari icin kisa, anlasilir cumleler."""
    rapor = rapor if rapor is not None else kontrol()
    u = []
    if rapor["eksik_defterde"]:
        u.append(f"Excel'de {len(rapor['eksik_defterde'])} kayıt merkezi defterde yok; bu kayıtlar proje/faaliyet "
                 "toplamlarına yansımıyor (Excel'e elle girilmiş olabilir). Ayrıntı: excel_operasyon.py mutabakat")
    if rapor["tutar_farki"]:
        u.append(f"{len(rapor['tutar_farki'])} kaydın Excel tutarı merkezi defterdekinden farklı "
                 "(Excel sonradan elle değiştirilmiş olabilir).")
    if rapor["defterde_fazla"]:
        u.append(f"Merkezi defterde {len(rapor['defterde_fazla'])} fiş var ama Excel'de aktif karşılığı yok "
                 "(satır silinmiş ya da iptal edilmiş olabilir).")
    return u
