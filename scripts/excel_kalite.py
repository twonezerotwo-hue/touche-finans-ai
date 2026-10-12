# -*- coding: utf-8 -*-
"""Excel bütünlüğü ve formül kalite kontrolleri.

witan uzak servis kullanabildiği için varsayılan olarak kapalıdır. Yerel kontrol
her zaman çalışır; uzak lint yalnız WITAN_REMOTE_ENABLED ve WITAN_API_KEY birlikte
tanımlandığında devreye girer.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import zipfile

from openpyxl import load_workbook

from config import FIRMALAR, WITAN_REMOTE_ENABLED, ana_excel_yolu


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def yerel_excel_kontrol(firma_key):
    if firma_key not in FIRMALAR:
        raise ValueError("Faaliyet bulunamadı.")
    path = ana_excel_yolu()
    sorunlar, uyarilar = [], []
    if not path.exists():
        return {"durum": "FAIL", "sorunlar": ["Excel dosyası bulunamadı."], "uyarilar": [], "istatistik": {}}
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return {"durum": "FAIL", "sorunlar": ["Yerel kalite kontrolü yalnız xlsx/xlsm destekler."], "uyarilar": [], "istatistik": {}}
    try:
        with zipfile.ZipFile(path) as zf:
            bozuk = zf.testzip()
            if bozuk:
                sorunlar.append(f"Bozuk ZIP girdisi: {bozuk}")
    except zipfile.BadZipFile:
        return {"durum": "FAIL", "sorunlar": ["Dosya geçerli bir Excel paketi değil."], "uyarilar": [], "istatistik": {}}

    wb = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    formuller, hatali_formuller = 0, []
    try:
        gerekli = set()
        firma = FIRMALAR[firma_key]
        for alan in ("ana_ledger", "ai_inbox", "holding_ozeti"):
            if firma.get(alan, {}).get("sayfa"):
                gerekli.add(firma[alan]["sayfa"])
        md = firma.get("master_data", {})
        if md.get("sayfa"):
            gerekli.add(md["sayfa"])
        eksik = sorted(gerekli.difference(wb.sheetnames))
        sorunlar.extend(f"Eksik gerekli sayfa: {x}" for x in eksik)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        formuller += 1
                        if any(err in value.upper() for err in FORMULA_ERRORS):
                            hatali_formuller.append(f"{ws.title}!{cell.coordinate}")
        sorunlar.extend(f"Hatalı formül başvurusu: {x}" for x in hatali_formuller[:50])

        ledger = firma.get("ana_ledger", {})
        if ledger.get("sayfa") in wb.sheetnames:
            ws = wb[ledger["sayfa"]]
            id_col = ledger.get("sutunlar", {}).get("Kayit_ID")
            if id_col:
                gorulen, tekrar = set(), set()
                for row in range(int(ledger.get("ilk_veri_satiri", 2)), (ws.max_row or 0) + 1):
                    value = str(ws[f"{id_col}{row}"].value or "").strip()
                    if not value:
                        continue
                    if value in gorulen:
                        tekrar.add(value)
                    gorulen.add(value)
                uyarilar.extend(f"Tekrarlanan kayıt kimliği: {x}" for x in sorted(tekrar)[:30])
    finally:
        wb.close()
    durum = "FAIL" if sorunlar else ("REVIEW" if uyarilar else "PASS")
    return {"durum": durum, "sorunlar": sorunlar, "uyarilar": uyarilar,
            "istatistik": {"sayfa": len(wb.sheetnames), "formul": formuller, "boyut": path.stat().st_size}}


def witan_excel_kontrol(firma_key):
    """Açık kullanıcı tercihi olmadan hiçbir finans dosyasını uzak servise göndermez."""
    if not WITAN_REMOTE_ENABLED:
        return {"durum": "SKIPPED", "sebep": "WITAN_REMOTE_ENABLED kapalı; finans verisi dışarı gönderilmedi."}
    if not os.environ.get("WITAN_API_KEY"):
        return {"durum": "SKIPPED", "sebep": "WITAN_API_KEY tanımlı değil."}
    exe = shutil.which("witan")
    if not exe:
        return {"durum": "SKIPPED", "sebep": "witan komutu bulunamadı."}
    path = ana_excel_yolu()
    sonuc = subprocess.run(
        [exe, "--stateless", "xlsx", "--json", "lint", str(path)],
        capture_output=True, text=True, timeout=180, shell=False,
    )
    raw = (sonuc.stdout or sonuc.stderr or "").strip()
    try:
        veri = json.loads(raw)
    except json.JSONDecodeError:
        veri = {"raw": raw[:3000]}
    return {"durum": "PASS" if sonuc.returncode == 0 else "REVIEW", "exit_code": sonuc.returncode, "rapor": veri}


def excel_kalite_kontrol(firma_key, witan=False):
    yerel = yerel_excel_kontrol(firma_key)
    sonuc = {"firma_key": firma_key, "yerel": yerel}
    if witan:
        sonuc["witan"] = witan_excel_kontrol(firma_key)
    return sonuc
