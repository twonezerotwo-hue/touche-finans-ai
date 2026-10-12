# -*- coding: utf-8 -*-
"""Tek master Excel'deki proje iş kalemlerini güvenli biçimde yönetir."""

from __future__ import annotations

import math
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from config import ana_excel_yolu
from excel_deposu import guvenli_degistir
from finans_cekirdegi import _baglanti, faaliyet_anahtari_bul, veritabani_yedekle
from guvenlik import audit_yaz
from kayit_kilidi import kilitli


SAYFA = "06_Proje İş Kalemleri"
ILK_SATIR = 6
SON_SATIR = 500

ALANLAR = {
    "metraj": "G",
    "satis_birim_fiyati": "H",
    "maliyet_turu": "J",
    "maliyet_birim_fiyati": "K",
    "durum": "O",
    "not": "P",
}
MALIYET_TURLERI = {"Birim fiyat", "Toplam sabit"}
DURUMLAR = {"Onaylı", "Taslak", "İptal"}


def _excel_yolu() -> Path:
    return ana_excel_yolu()


def _norm(deger) -> str:
    import unicodedata

    metin = unicodedata.normalize("NFKD", str(deger or ""))
    return "".join(c for c in metin if not unicodedata.combining(c)).casefold().replace("ı", "i").strip()


def _sayi(deger, alan: str) -> float:
    if isinstance(deger, bool):
        raise ValueError(f"{alan} sayı olmalı.")
    metin = str(deger if deger is not None else "").strip().replace("₺", "").replace(" ", "")
    if "," in metin and "." in metin:
        metin = metin.replace(".", "").replace(",", ".")
    elif "," in metin:
        metin = metin.replace(",", ".")
    try:
        sonuc = float(Decimal(metin))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{alan} geçerli bir sayı olmalı.") from exc
    if not math.isfinite(sonuc) or sonuc < 0 or sonuc > 1_000_000_000_000:
        raise ValueError(f"{alan} 0 ile 1 trilyon arasında olmalı.")
    return sonuc


def _tutar_kurus(deger: float) -> int:
    return int((Decimal(str(deger)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _satir_dict(v) -> dict:
    """A..R sütun değerlerinden (v[0]=A) iş kalemi sözlüğü üretir."""
    def h(i):                      # i: 0-tabanli sutun
        return v[i] if i < len(v) else None

    metraj = float(h(6) or 0)
    satis_bf = float(h(7) or 0)
    maliyet_turu = str(h(9) or "Birim fiyat").strip()
    maliyet_bf = float(h(10) or 0)
    ciro = metraj * satis_bf
    maliyet = maliyet_bf if _norm(maliyet_turu) == "toplam sabit" else metraj * maliyet_bf
    fark = ciro - maliyet
    return {
        "kalem_id": str(h(0) or "").strip(),
        "faaliyet": str(h(1) or "").strip(),
        "faaliyet_key": faaliyet_anahtari_bul(h(1)),
        "proje": str(h(2) or "").strip(),
        "maliyet_grubu": str(h(3) or "").strip(),
        "is_kalemi": str(h(4) or "").strip(),
        "birim": str(h(5) or "").strip(),
        "metraj": metraj,
        "satis_birim_fiyati": satis_bf,
        "ciro": round(ciro, 2),
        "maliyet_turu": maliyet_turu,
        "maliyet_birim_fiyati": maliyet_bf,
        "maliyet": round(maliyet, 2),
        "brut_fark": round(fark, 2),
        "brut_marj": round(fark / ciro * 100, 2) if ciro else 0,
        "durum": str(h(14) or "").strip(),
        "not": str(h(15) or "").strip(),
        "hakedis_kayit_id": str(h(16) or "").strip(),
        "gider_kayit_id": str(h(17) or "").strip(),
    }


def _satir_verisi(ws, satir: int) -> dict:
    """Tek satir (normal, yazilabilir calisma kitabi icin)."""
    satirlar = list(ws.iter_rows(min_row=satir, max_row=satir, min_col=1, max_col=18, values_only=True))
    return _satir_dict(satirlar[0] if satirlar else ())


def _kalemler(ws) -> list[dict]:
    """Tum satirlari TEK gecişte okur (salt-okunur sayfada hucre adresiyle erisim O(n^2)'dir)."""
    sonuc = []
    for v in ws.iter_rows(min_row=ILK_SATIR, max_row=SON_SATIR, min_col=1, max_col=18, values_only=True):
        if str(v[0] or "").strip():
            sonuc.append(_satir_dict(v))
    return sonuc


def _ozet(kalemler: list[dict]) -> dict:
    onayli = [x for x in kalemler if _norm(x["durum"]) == "onayli"]
    ciro = round(sum(x["ciro"] for x in onayli), 2)
    maliyet = round(sum(x["maliyet"] for x in onayli), 2)
    fark = round(ciro - maliyet, 2)
    return {
        "kalem_sayisi": len(kalemler),
        "onayli_kalem_sayisi": len(onayli),
        "ciro": ciro,
        "maliyet": maliyet,
        "brut_fark": fark,
        "brut_marj": round(fark / ciro * 100, 2) if ciro else 0,
    }


def listele(faaliyet: str = "tum") -> dict:
    yol = _excel_yolu()
    if not yol.exists():
        raise FileNotFoundError(f"Ana Excel bulunamadı: {yol}")
    wb = load_workbook(yol, read_only=True, data_only=False)
    try:
        if SAYFA not in wb.sheetnames:
            raise ValueError(f"Ana Excel'de {SAYFA} sayfası bulunamadı.")
        kalemler = _kalemler(wb[SAYFA])
    finally:
        wb.close()
    if faaliyet not in {"tum", "touche_holding"}:
        kalemler = [x for x in kalemler if x["faaliyet_key"] == faaliyet]
    return {"kalemler": kalemler, "ozet": _ozet(kalemler), "kaynak": str(yol)}


def _bagli_toplamlar(ws) -> dict[str, float]:
    toplamlar: dict[str, float] = {}
    for kalem in _kalemler(ws):
        if kalem["hakedis_kayit_id"]:
            toplamlar.setdefault(kalem["hakedis_kayit_id"], 0)
        if kalem["gider_kayit_id"]:
            toplamlar.setdefault(kalem["gider_kayit_id"], 0)
        if _norm(kalem["durum"]) != "onayli":
            continue
        if kalem["hakedis_kayit_id"]:
            toplamlar[kalem["hakedis_kayit_id"]] = toplamlar.get(kalem["hakedis_kayit_id"], 0) + kalem["ciro"]
        if kalem["gider_kayit_id"]:
            toplamlar[kalem["gider_kayit_id"]] = toplamlar.get(kalem["gider_kayit_id"], 0) + kalem["maliyet"]
    return {key: round(value, 2) for key, value in toplamlar.items()}


def _defteri_senkronla(toplamlar: dict[str, float]) -> None:
    with _baglanti() as con:
        for kayit_id, tutar in toplamlar.items():
            belge_id = f"LEGACY-{kayit_id}"
            satirlar = con.execute(
                "SELECT s.fis_id,s.satir_no,s.hesap_kodu,s.borc_kurus,s.alacak_kurus,h.hesap_tipi "
                "FROM fis_satirlari s JOIN fisler f ON f.fis_id=s.fis_id "
                "JOIN hesaplar h ON h.kod=s.hesap_kodu WHERE f.belge_id=? ORDER BY s.fis_id,s.satir_no",
                (belge_id,),
            ).fetchall()
            if len(satirlar) != 2 or len({x["fis_id"] for x in satirlar}) != 1:
                raise ValueError(f"{kayit_id} için merkez defter eşleşmesi eksik veya belirsiz.")
            kurus = _tutar_kurus(tutar)
            for satir in satirlar:
                borc_tarafi = satir["borc_kurus"] > 0 or (
                    satir["borc_kurus"] == satir["alacak_kurus"] == 0
                    and satir["hesap_tipi"] in {"VARLIK", "GIDER"}
                )
                con.execute(
                    "UPDATE fis_satirlari SET borc_kurus=?,alacak_kurus=? WHERE fis_id=? AND satir_no=?",
                    (kurus if borc_tarafi else 0, 0 if borc_tarafi else kurus, satir["fis_id"], satir["satir_no"]),
                )
            kontrol = con.execute(
                "SELECT SUM(borc_kurus) b,SUM(alacak_kurus) a FROM fis_satirlari WHERE fis_id=?",
                (satirlar[0]["fis_id"],),
            ).fetchone()
            if kontrol["b"] != kontrol["a"]:
                raise RuntimeError(f"{kayit_id} güncellemesi merkez defteri dengesiz bıraktı.")


@kilitli
def guncelle(kalem_id: str, degisiklikler: dict, aktor: str) -> dict:
    kalem_id = str(kalem_id or "").strip()
    if not kalem_id or not isinstance(degisiklikler, dict) or not degisiklikler:
        raise ValueError("Kalem kimliği ve en az bir değişiklik gerekli.")
    bilinmeyen = set(degisiklikler) - set(ALANLAR)
    if bilinmeyen:
        raise ValueError(f"Değiştirilemeyen alan: {', '.join(sorted(bilinmeyen))}")

    temiz = {}
    for alan, deger in degisiklikler.items():
        if alan in {"metraj", "satis_birim_fiyati", "maliyet_birim_fiyati"}:
            temiz[alan] = _sayi(deger, alan)
        elif alan == "maliyet_turu":
            if deger not in MALIYET_TURLERI:
                raise ValueError("Maliyet türü Birim fiyat veya Toplam sabit olmalı.")
            temiz[alan] = deger
        elif alan == "durum":
            if deger not in DURUMLAR:
                raise ValueError("Durum Onaylı, Taslak veya İptal olmalı.")
            temiz[alan] = deger
        else:
            temiz[alan] = str(deger or "").strip()[:500]

    durum = {}

    def degistir(wb):
        if SAYFA not in wb.sheetnames:
            raise ValueError(f"Ana Excel'de {SAYFA} sayfası bulunamadı.")
        ws = wb[SAYFA]
        hedef_satir = next((r for r in range(ILK_SATIR, SON_SATIR + 1) if str(ws[f"A{r}"].value or "").strip() == kalem_id), None)
        if not hedef_satir:
            raise ValueError(f"İş kalemi bulunamadı: {kalem_id}")
        durum["once"] = _satir_verisi(ws, hedef_satir)
        for alan, deger in temiz.items():
            ws[f"{ALANLAR[alan]}{hedef_satir}"] = deger
        durum["sonra"] = _satir_verisi(ws, hedef_satir)
        durum["ozet"] = _ozet([x for x in _kalemler(ws) if x["faaliyet_key"] == durum["sonra"]["faaliyet_key"]])
        durum["toplamlar"] = _bagli_toplamlar(ws)
        return {"basari": True}

    def defteri_esitle(_sonuc):
        # ANA yerlestirildikten sonra calisir; hata verirse ANA otomatik geri alinir.
        veritabani_yedekle()
        _defteri_senkronla(durum["toplamlar"])

    sonuc = guvenli_degistir(
        degistir, aktor=str(aktor), kaynak="proje_is_kalemleri.guncelle",
        beklenen_sayfalar=["03_Kayıt Merkezi", SAYFA], ek_adim=defteri_esitle,
        degisim_izni={SAYFA: set(ALANLAR.values())},
        audit_ek=lambda s: {"kalem_id": kalem_id, "once": durum["once"], "sonra": durum["sonra"],
                            "bagli_toplamlar": durum["toplamlar"]},
    )
    audit_yaz(
        "proje_is_kalemi_guncellendi", aktor=str(aktor), kalem_id=kalem_id,
        once=durum["once"], sonra=durum["sonra"], islem_id=sonuc.get("islem_id"),
        bagli_toplamlar=durum["toplamlar"],
    )
    return {"basari": True, "kalem": durum["sonra"], "ozet": durum["ozet"]}
