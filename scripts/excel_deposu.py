# -*- coding: utf-8 -*-
"""Touche Holding Smart Excel deposu: TEK ANA (MAIN) + TEK YEDEK (BACKUP).

Calisma zamaninda finansal Excel'e yazan HER kod bu modulden gecer. Baska hicbir
modul ana dosyayi acip kaydetmez, ``os.replace``/kopya ile uzerine yazmaz.

Her onayli yazma icin akis (kilit altinda):

  1. Ana dosya var mi / gecerli mi? (ZIP + openpyxl)  -> degilse HICBIR sey degismez
  2. Degisiklik bellekte uygulanir ve dogrulanir       -> reddedilirse HICBIR sey degismez
  3. Gecici dosyaya kaydedilir, yeniden acilip dogrulanir
     (sayfalar, formuller, mevcut hucre degerleri, grafik/resim kaybi)
  4. YEDEK = mevcut ANA olarak yenilenir ve hash ile dogrulanir
  5. ANA hala ayni mi? (disaridan degisim kontrolu) -> degistiyse iptal
  6. Gecici dosya ANA uzerine atomik olarak yerlestirilir (os.replace)
  7. Yerlestirme hash ile dogrulanir; sorun varsa YEDEK'ten geri alinir
  8. Gecici dosyalar silinir; islem denetim kaydina yazilir

Sonuc: ANA = V11, YEDEK = V10. Tarihli yedek, islem basina kopya veya yedek klasoru
YOKTUR. Geri donus yalniz bilincli bir operasyon eylemidir ve ANA ile YEDEK'i
ters cevirir (swap); yapay zeka hicbir zaman geri donus yapamaz.
"""

from __future__ import annotations

import hashlib
import logging
import math
import os
import re
import shutil
import time
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import config
from guvenlik import audit_yaz
from kayit_kilidi import yazma_kilidi

log = logging.getLogger(__name__)

GERI_DON_ONAY_KELIMESI = "GERI-DON"

# openpyxl'in koruyamadigi, kaybolursa veri/gorunum kaybi olusturan paket parcalari.
# Bunlardan biri ANA'da varsa ve yeni surumde yoksa yazma REDDEDILIR.
_KORUNACAK_PAKET = (
    "xl/charts/", "xl/media/", "xl/pivotTables/", "xl/pivotCache/", "xl/embeddings/",
    "xl/activeX/", "xl/slicers/", "xl/timelines/", "xl/vbaProject.bin",
)
# Kaybi yalniz UYARI olarak denetim kaydina yazilan ozellikler (yaygin, bazen
# openpyxl tarafindan yeniden gruplanir; tek basina yazmayi engellemez).
_UYARI_OZELLIKLERI = ("<x14:", "<dataValidation ", "<conditionalFormatting", "<definedName ")

_GECICI_RE = re.compile(r"^\.(?P<govde>.+)\.[0-9a-f]{32}\.tmp(?:\.xlsx)?$")


class DepoHatasi(Exception):
    """Excel deposu islemi guvenli bicimde reddetti/basarisiz oldu (dosya korundu)."""

    def __init__(self, kod: str, mesaj: str, **ek):
        super().__init__(mesaj)
        self.kod = kod
        self.mesaj = mesaj
        self.ek = ek

    def __str__(self):  # pragma: no cover - basit
        return self.mesaj


# ---------------------------------------------------------------------------
# Yardimcilar
# ---------------------------------------------------------------------------

def sha256(yol) -> str:
    h = hashlib.sha256()
    with Path(yol).open("rb") as f:
        for parca in iter(lambda: f.read(1024 * 1024), b""):
            h.update(parca)
    return h.hexdigest()


def _gecici_yol(hedef: Path) -> Path:
    """Hedefle ayni klasorde (ayni disk) gecici dosya adi; os.replace atomik kalir."""
    hedef = Path(hedef)
    return hedef.parent / f".{hedef.stem}.{uuid.uuid4().hex}.tmp{hedef.suffix}"


def _sil(yol):
    try:
        if yol is not None:
            Path(yol).unlink(missing_ok=True)
    except OSError as exc:  # pragma: no cover - acik dosya vb.
        log.warning("Gecici dosya silinemedi: %s (%s)", yol, exc)


def gecici_kalintilari_temizle() -> list[str]:
    """Bu sistemin biraktigi gecici dosyalari siler (ANA ve YEDEK klasorlerinde).

    Yalnizca ``.<ana veya yedek adi>.<32 hex>.tmp.xlsx`` desenine uyanlar silinir;
    baska hicbir dosyaya dokunulmaz. Yazma kilidi altinda cagrildiginda canli bir
    gecici dosya olamayacagindan guvenlidir.
    """
    silinen = []
    adlar = {config.ana_excel_yolu().stem, config.yedek_excel_yolu().stem}
    klasorler = {config.ana_excel_yolu().parent, config.yedek_excel_yolu().parent}
    for klasor in klasorler:
        if not klasor.exists():
            continue
        for dosya in klasor.iterdir():
            eslesme = _GECICI_RE.match(dosya.name)
            if eslesme and eslesme.group("govde") in adlar and dosya.is_file():
                _sil(dosya)
                silinen.append(str(dosya))
    return silinen


def gecici_dosyalar() -> list[str]:
    """Kalinti gecici dosyalarin listesi (silmeden)."""
    sonuc = []
    adlar = {config.ana_excel_yolu().stem, config.yedek_excel_yolu().stem}
    for klasor in {config.ana_excel_yolu().parent, config.yedek_excel_yolu().parent}:
        if klasor.exists():
            for dosya in klasor.iterdir():
                eslesme = _GECICI_RE.match(dosya.name)
                if eslesme and eslesme.group("govde") in adlar:
                    sonuc.append(str(dosya))
    return sorted(sonuc)


def dogrula(yol, beklenen_sayfalar=()) -> dict:
    """XLSX butunlugunu dogrular: ZIP + workbook.xml + openpyxl ile acilis + sayfalar."""
    yol = Path(yol)
    try:
        with zipfile.ZipFile(yol) as paket:
            if paket.testzip() is not None:
                raise DepoHatasi("DOSYA_BOZUK", f"XLSX paketi bozuk: {yol.name}")
            if "xl/workbook.xml" not in paket.namelist():
                raise DepoHatasi("DOSYA_BOZUK", f"Geçerli bir Excel çalışma kitabı değil: {yol.name}")
    except zipfile.BadZipFile as exc:
        raise DepoHatasi("DOSYA_BOZUK", f"Excel dosyası okunamıyor (ZIP hatası): {yol.name}") from exc
    try:
        wb = load_workbook(yol, read_only=True, data_only=False)
    except Exception as exc:
        raise DepoHatasi("DOSYA_BOZUK", f"Excel dosyası açılamadı: {yol.name} ({exc})") from exc
    try:
        sayfalar = list(wb.sheetnames)
    finally:
        wb.close()
    eksik = [s for s in beklenen_sayfalar if s not in sayfalar]
    if eksik:
        raise DepoHatasi("SAYFA_EKSIK", f"Excel'de beklenen sayfa yok: {', '.join(eksik)}", sayfalar=sayfalar)
    return {"sayfalar": sayfalar}


def _paket_adlari(yol) -> set[str]:
    with zipfile.ZipFile(yol) as paket:
        return set(paket.namelist())


def _paket_kaybi(onceki: Path, sonraki: Path) -> list[str]:
    once, sonra = _paket_adlari(onceki), _paket_adlari(sonraki)
    kayip = []
    for onek in _KORUNACAK_PAKET:
        if any(a.startswith(onek) or a == onek for a in once) and not any(
            a.startswith(onek) or a == onek for a in sonra
        ):
            kayip.append(onek.rstrip("/"))
    return kayip


def _ozellik_sayilari(yol: Path) -> dict[str, int]:
    sayilar = {k: 0 for k in _UYARI_OZELLIKLERI}
    with zipfile.ZipFile(yol) as paket:
        for ad in paket.namelist():
            if ad.startswith("xl/worksheets/") and ad.endswith(".xml") or ad == "xl/workbook.xml":
                govde = paket.read(ad).decode("utf-8", errors="ignore")
                for k in sayilar:
                    sayilar[k] += govde.count(k)
    return sayilar


def _ozellik_kaybi_uyarilari(onceki: Path, sonraki: Path) -> list[str]:
    once, sonra = _ozellik_sayilari(onceki), _ozellik_sayilari(sonraki)
    return [
        f"{k.strip('<').strip()} sayısı {once[k]} -> {sonra[k]}"
        for k in _UYARI_OZELLIKLERI if sonra[k] < once[k]
    ]


def _esit(a, b) -> bool:
    if a == b:
        return True
    if isinstance(a, datetime) and isinstance(b, datetime):
        return abs((a - b).total_seconds()) < 1
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return math.isclose(float(a), float(b), rel_tol=1e-12, abs_tol=1e-9)
    if isinstance(a, str) and isinstance(b, str):
        return a.strip() == b.strip()
    return False


def _formul_metni(deger) -> str:
    return str(getattr(deger, "text", deger))


def anlik_goruntu(wb) -> dict:
    """Degisiklikten once tum bos olmayan hucrelerin (deger/formul) goruntusunu alir."""
    degerler, formuller = {}, {}
    for ws in wb.worksheets:
        d, f = {}, {}
        for satir in ws.iter_rows():
            for h in satir:
                if h.value is None or h.value == "":
                    continue
                if h.data_type == "f":
                    f[h.coordinate] = _formul_metni(h.value)
                else:
                    d[h.coordinate] = h.value
        degerler[ws.title] = d
        formuller[ws.title] = f
    return {"degerler": degerler, "formuller": formuller, "sayfalar": [w.title for w in wb.worksheets]}


def _sutun_harfi(koordinat: str) -> str:
    m = re.match(r"[A-Z]+", koordinat)
    return m.group(0) if m else ""


def goruntu_karsilastir(once: dict, sonra: dict, *, degisim_izni=None, formul_degisim_izni=False,
                        sayfa_silme_izni=False) -> list[str]:
    """Mevcut verinin korunup korunmadigini denetler. Ihlal listesi dondurur (bos = temiz).

    ``degisim_izni``: {sayfa: None | {sutun harfleri} | fonksiyon(koordinat)->bool}.
    Burada izin verilen sayfa/sutun/hucrelerde mevcut degerlerin degismesine izin verilir;
    digerlerinde yalnizca BOS hucreler doldurulabilir. Formuller ``formul_degisim_izni``
    olmadikca degistirilemez.
    """
    ihlal = []
    izin = degisim_izni or {}
    for sayfa in once["sayfalar"]:
        if sayfa not in sonra["sayfalar"]:
            if not sayfa_silme_izni:
                ihlal.append(f"Sayfa kayboldu: {sayfa}")
            continue
        for koordinat, formul in once["formuller"].get(sayfa, {}).items():
            yeni = sonra["formuller"].get(sayfa, {}).get(koordinat)
            if yeni != formul and not formul_degisim_izni:
                ihlal.append(f"Formül bozuldu: {sayfa}!{koordinat}")
        izinli = izin.get(sayfa, set())
        for koordinat, eski in once["degerler"].get(sayfa, {}).items():
            yeni = sonra["degerler"].get(sayfa, {}).get(koordinat)
            if yeni is not None and _esit(eski, yeni):
                continue
            if izinli is None:
                continue
            if callable(izinli):
                if izinli(koordinat):
                    continue
            elif _sutun_harfi(koordinat) in izinli:
                continue
            ihlal.append(f"Mevcut veri değişti/silindi: {sayfa}!{koordinat}")
    return ihlal


def _kullanici_mesaji(exc: Exception) -> str | None:
    if isinstance(exc, PermissionError) or getattr(exc, "winerror", None) in (5, 32):
        return (
            "Ana Excel dosyası açık veya OneDrive tarafından kilitli. Bilgisayardaki tüm Excel "
            "pencerelerini kapatın, OneDrive eşitlemesinin tamamlanmasını bekleyin ve aynı "
            "taslakta onaya tekrar basın. Kayıt yapılmadı; taslak korunuyor."
        )
    return None


# ---------------------------------------------------------------------------
# Yedek yenileme / durum
# ---------------------------------------------------------------------------

def yedegi_yenile() -> str:
    """YEDEK'i mevcut ANA'nin birebir kopyasi yapar (hash ile dogrulanir).

    ANA gecersizse yedek ASLA ezilmez. Donus: ANA/YEDEK sha256.
    """
    ana, yedek = config.ana_excel_yolu(), config.yedek_excel_yolu()
    dogrula(ana)
    ana_hash = sha256(ana)
    yedek.parent.mkdir(parents=True, exist_ok=True)
    if yedek.exists() and sha256(yedek) == ana_hash:
        return ana_hash
    gecici = _gecici_yol(yedek)
    try:
        shutil.copyfile(ana, gecici)
        if sha256(gecici) != ana_hash:
            raise DepoHatasi("YEDEK_DOGRULAMA", "Yedek kopyası doğrulanamadı; kayıt yapılmadı.")
        dogrula(gecici)
        os.replace(gecici, yedek)
    finally:
        _sil(gecici)
    if sha256(yedek) != ana_hash:
        raise DepoHatasi("YEDEK_DOGRULAMA", "Yedek dosyası yerleştirme sonrası doğrulanamadı.")
    return ana_hash


def _dosya_bilgisi(yol: Path) -> dict:
    var = yol.exists()
    bilgi = {"yol": str(yol), "var": var, "boyut": yol.stat().st_size if var else 0,
             "degisim": datetime.fromtimestamp(yol.stat().st_mtime).isoformat(timespec="seconds") if var else None,
             "sha256": sha256(yol) if var else None}
    return bilgi


def durum() -> dict:
    """ANA/YEDEK durumunu ve kalinti gecici dosyalari raporlar (yazmaz)."""
    ana, yedek = config.ana_excel_yolu(), config.yedek_excel_yolu()
    a, y = _dosya_bilgisi(ana), _dosya_bilgisi(yedek)
    return {
        "ana": a, "yedek": y,
        "ana_yedek_ayni": bool(a["sha256"] and a["sha256"] == y["sha256"]),
        "gecici_dosyalar": gecici_dosyalar(),
    }


# ---------------------------------------------------------------------------
# Ana yazma islemi
# ---------------------------------------------------------------------------

def guvenli_degistir(degistirici, *, islem_id: str | None = None, aktor: str = "sistem",
                     kaynak: str = "", beklenen_sayfalar=(), ek_adim=None,
                     degisim_izni=None, formul_degisim_izni: bool = False,
                     kuru_calistir: bool = False, audit_ek=None) -> dict:
    """ANA Excel'i guvenli bicimde degistirir.

    ``degistirici(wb)`` bellekteki calisma kitabini degistirir ve ``{"basari": bool, ...}``
    dondurur. ``basari`` False ise veya istisna olursa ANA ve YEDEK dokunulmadan kalir.
    ``ek_adim(sonuc)`` ANA yerlestirildikten sonra calisir (ornegin SQLite defteri);
    hata verirse ANA, YEDEK'ten geri alinir ve hata yeniden firlatilir.
    ``kuru_calistir`` gecici dosyaya kadar her seyi dogrular ama ANA/YEDEK'e dokunmaz.
    """
    islem_id = islem_id or f"OP-{datetime.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:8]}"
    baslangic = time.monotonic()
    ana, yedek = config.ana_excel_yolu(), config.yedek_excel_yolu()
    olay: dict[str, Any] = {"islem_id": islem_id, "aktor": str(aktor), "kaynak": kaynak,
            "ana": str(ana), "yedek": str(yedek)}
    gecici = None
    wb = None
    ana_degisti = False
    try:
        with yazma_kilidi():
            gecici_kalintilari_temizle()
            if not ana.exists():
                raise DepoHatasi("ANA_YOK", f"Ana Excel bulunamadı: {ana}")
            dogrula(ana, beklenen_sayfalar)            # 1) ANA gecerli mi?
            ana_hash_once = sha256(ana)
            olay["ana_hash_once"] = ana_hash_once

            wb = load_workbook(ana, data_only=False, keep_links=True)
            once = anlik_goruntu(wb)
            try:
                sonuc = degistirici(wb)                # 2) bellekte uygula + dogrula
            except DepoHatasi:
                raise
            except Exception as exc:
                olay.update(durum="REDDEDILDI_HATA", hata=f"{type(exc).__name__}: {exc}")
                raise
            if not isinstance(sonuc, dict) or not sonuc.get("basari"):
                olay.update(durum="REDDEDILDI", hata=str((sonuc or {}).get("hata", "")))
                sonuc = dict(sonuc or {"basari": False, "hata": "Değiştirici sonuç döndürmedi."})
                sonuc.setdefault("yazilan", 0)
                sonuc["islem_id"] = islem_id
                sonuc["degisiklik_yapildi"] = False
                return sonuc

            if hasattr(wb, "calculation"):
                wb.calculation.fullCalcOnLoad = True
            gecici = _gecici_yol(ana)
            wb.save(gecici)                            # 3) gecici dosyaya kaydet
            wb.close()
            wb = None
            dogrula(gecici, beklenen_sayfalar)

            kontrol = load_workbook(gecici, data_only=False, keep_links=True)
            try:
                sonra = anlik_goruntu(kontrol)
            finally:
                kontrol.close()
            izin = degisim_izni(sonuc) if callable(degisim_izni) else degisim_izni
            ihlaller = goruntu_karsilastir(
                once, sonra, degisim_izni=izin, formul_degisim_izni=formul_degisim_izni,
            )
            if ihlaller:
                raise DepoHatasi(
                    "VERI_KORUMA", "Kayıt, mevcut Excel verisini bozacağı için reddedildi: "
                    + "; ".join(ihlaller[:5]) + (f" (+{len(ihlaller) - 5})" if len(ihlaller) > 5 else ""),
                    ihlaller=ihlaller)
            kayip = _paket_kaybi(ana, gecici)
            if kayip and os.environ.get("TOUCHE_PAKET_KAYBI_IZNI") != "1":
                raise DepoHatasi(
                    "PAKET_KAYBI", "Excel'deki grafik/resim/özet tablo gibi öğeler kaybolacağı için "
                    "kayıt reddedildi: " + ", ".join(kayip), kayip=kayip)
            uyarilar = _ozellik_kaybi_uyarilari(ana, gecici)
            if uyarilar:
                olay["uyarilar"] = uyarilar

            if kuru_calistir:
                olay.update(durum="KURU_CALISTIRMA")
                sonuc = dict(sonuc, islem_id=islem_id, degisiklik_yapildi=False, kuru_calistirma=True)
                return sonuc

            yedek_hash = yedegi_yenile()               # 4) YEDEK = mevcut ANA
            olay["yedek_hash"] = yedek_hash
            if sha256(ana) != ana_hash_once:           # 5) disaridan degisim?
                raise DepoHatasi(
                    "ANA_DEGISTI", "Ana Excel işlem sırasında dışarıdan değişti (Excel/OneDrive). "
                    "Kayıt yapılmadı; taslağı tekrar onaylayın.")
            yeni_hash = sha256(gecici)
            try:
                os.replace(gecici, ana)                # 6) atomik yerlestirme
            except OSError as exc:
                mesaj = _kullanici_mesaji(exc)
                if mesaj:
                    raise DepoHatasi("ANA_KILITLI", mesaj) from exc
                raise
            gecici = None
            ana_degisti = True
            if sha256(ana) != yeni_hash:               # 7) yerlestirme dogrulamasi
                raise DepoHatasi("YERLESTIRME_DOGRULAMA", "Ana Excel yerleştirme sonrası doğrulanamadı.")
            olay["ana_hash_sonra"] = yeni_hash
            if ek_adim is not None:
                ek_adim(sonuc)
            sonuc = dict(sonuc, islem_id=islem_id, degisiklik_yapildi=True,
                         ana_hash_once=ana_hash_once, ana_hash_sonra=yeni_hash, yedek_hash=yedek_hash)
            ek = audit_ek(sonuc) if callable(audit_ek) else (audit_ek or {})
            olay.update({k: v for k, v in ek.items() if k not in olay})
            olay["durum"] = "BASARILI"
            return sonuc
    except Exception as exc:
        geri_alindi = False
        if ana_degisti:
            try:
                _yedekten_yerlestir(yedek, ana)
                geri_alindi = True
            except Exception as geri_hata:  # pragma: no cover - cok nadir
                log.critical("Ana Excel yedekten geri alinamadi: %s", geri_hata)
                olay["geri_alma_hatasi"] = str(geri_hata)
            olay["durum"] = "HATA_GERI_ALINDI" if geri_alindi else "HATA_GERI_ALINAMADI"
        else:
            olay.setdefault("durum", "HATA")
        olay["hata"] = olay.get("hata") or f"{type(exc).__name__}: {exc}"
        if isinstance(exc, DepoHatasi):
            olay["hata_kodu"] = exc.kod
        raise
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:  # pragma: no cover
                pass
        _sil(gecici)
        olay["sure_ms"] = int((time.monotonic() - baslangic) * 1000)
        olay.setdefault("durum", "BILINMIYOR")
        audit_yaz("excel_islem", **olay)


def _yedekten_yerlestir(yedek: Path, ana: Path):
    """YEDEK icerigini ANA uzerine atomik olarak yerlestirir (acil geri alma)."""
    gecici = _gecici_yol(ana)
    try:
        shutil.copyfile(yedek, gecici)
        dogrula(gecici)
        os.replace(gecici, ana)
    finally:
        _sil(gecici)


# ---------------------------------------------------------------------------
# Operasyon: geri donus (swap) ve harici dosyayi ana yapma
# ---------------------------------------------------------------------------

def yedekten_geri_don(*, aktor: str, neden: str, onay: str) -> dict:
    """ANA ve YEDEK'i ters cevirir: ANA=V11,YEDEK=V10 -> ANA=V10,YEDEK=V11.

    YALNIZ bilincli bir insan/operasyon eylemi icindir (``scripts/excel_operasyon.py``).
    HTTP sunucusu, Telegram botu ve yapay zeka katmani bu fonksiyonu CAGIRMAZ
    (testlerle denetlenir). ``onay`` tam olarak ``GERI-DON`` olmalidir.
    """
    if onay != GERI_DON_ONAY_KELIMESI:
        raise DepoHatasi("ONAY_YOK", f"Geri dönüş için onay kelimesi '{GERI_DON_ONAY_KELIMESI}' gerekli.")
    if not str(aktor or "").strip() or not str(neden or "").strip():
        raise DepoHatasi("ONAY_YOK", "Geri dönüş için aktör ve neden zorunludur.")
    ana, yedek = config.ana_excel_yolu(), config.yedek_excel_yolu()
    olay: dict[str, Any] = {"aktor": str(aktor), "neden": str(neden), "ana": str(ana), "yedek": str(yedek)}
    t_ana = t_yedek = None
    try:
        with yazma_kilidi():
            gecici_kalintilari_temizle()
            if not yedek.exists():
                raise DepoHatasi("YEDEK_YOK", f"Yedek Excel bulunamadı: {yedek}")
            dogrula(yedek)
            h_yedek = sha256(yedek)
            ana_var = ana.exists()
            h_ana = sha256(ana) if ana_var else None
            olay.update(ana_hash_once=h_ana, yedek_hash_once=h_yedek)
            if ana_var and h_ana == h_yedek:
                raise DepoHatasi("AYNI_SURUM", "Ana ve yedek zaten aynı sürümde; geri dönülecek fark yok.")
            ana.parent.mkdir(parents=True, exist_ok=True)
            if ana_var:
                t_ana = _gecici_yol(yedek)
                shutil.copyfile(ana, t_ana)            # YEDEK'e gidecek kopya
                if sha256(t_ana) != h_ana:
                    raise DepoHatasi("YEDEK_DOGRULAMA", "Ana kopyası doğrulanamadı; geri dönülmedi.")
            t_yedek = _gecici_yol(ana)
            shutil.copyfile(yedek, t_yedek)            # ANA'ya gidecek kopya
            if sha256(t_yedek) != h_yedek:
                raise DepoHatasi("YEDEK_DOGRULAMA", "Yedek kopyası doğrulanamadı; geri dönülmedi.")
            os.replace(t_yedek, ana)
            t_yedek = None
            if sha256(ana) != h_yedek:
                raise DepoHatasi("YERLESTIRME_DOGRULAMA", "Geri dönüş sonrası ana dosya doğrulanamadı.")
            if t_ana is not None:
                try:
                    os.replace(t_ana, yedek)
                    t_ana = None
                except OSError as exc:  # ANA artik eski surum; yeni surumu kaybetme
                    olay["kurtarma_dosyasi"] = str(t_ana)
                    t_ana = None
                    raise DepoHatasi(
                        "YEDEK_YERLESTIRILEMEDI",
                        f"Ana geri alındı ancak önceki ana sürüm yedeğe yazılamadı: {exc}. "
                        f"Önceki sürüm şu dosyada saklandı: {olay['kurtarma_dosyasi']}") from exc
            olay.update(ana_hash_sonra=sha256(ana), yedek_hash_sonra=sha256(yedek), durum="BASARILI")
            return {"basari": True, "ana_hash": olay["ana_hash_sonra"], "yedek_hash": olay["yedek_hash_sonra"]}
    except Exception as exc:
        olay.setdefault("durum", "HATA")
        olay["hata"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        _sil(t_ana)
        _sil(t_yedek)
        audit_yaz("excel_geri_donus", **olay)


def harici_dosyayi_ana_yap(kaynak, *, aktor: str, neden: str = "") -> dict:
    """Onaylanmis, disarida duzenlenmis bir XLSX'i yeni ANA yapar (YEDEK = onceki ANA).

    Eski ``master_excel_entegre`` aracinin tek-yedekli karsiligidir.
    """
    kaynak = Path(kaynak).resolve(strict=True)
    ana = config.ana_excel_yolu()
    if kaynak == ana.resolve() or kaynak == config.yedek_excel_yolu().resolve():
        raise DepoHatasi("KAYNAK_GECERSIZ", "Kaynak dosya zaten ana veya yedek dosyadır.")
    if kaynak.suffix.lower() != ".xlsx":
        raise DepoHatasi("KAYNAK_GECERSIZ", "Yalnız .xlsx dosyaları ana dosya yapılabilir.")
    olay: dict[str, Any] = {"aktor": str(aktor), "neden": str(neden), "kaynak_dosya": str(kaynak)}
    gecici = None
    try:
        with yazma_kilidi():
            gecici_kalintilari_temizle()
            dogrula(kaynak)
            kaynak_hash = sha256(kaynak)
            olay["kaynak_hash"] = kaynak_hash
            if ana.exists():
                olay["ana_hash_once"] = sha256(ana)
                olay["yedek_hash"] = yedegi_yenile()
            ana.parent.mkdir(parents=True, exist_ok=True)
            gecici = _gecici_yol(ana)
            shutil.copyfile(kaynak, gecici)
            if sha256(gecici) != kaynak_hash:
                raise DepoHatasi("YERLESTIRME_DOGRULAMA", "Geçici dosya doğrulanamadı; ana dosya değişmedi.")
            dogrula(gecici)
            os.replace(gecici, ana)
            gecici = None
            if sha256(ana) != kaynak_hash:
                if config.yedek_excel_yolu().exists():
                    _yedekten_yerlestir(config.yedek_excel_yolu(), ana)
                raise DepoHatasi("YERLESTIRME_DOGRULAMA", "Değişim sonrası doğrulama başarısız; yedek geri yüklendi.")
            olay.update(ana_hash_sonra=kaynak_hash, durum="BASARILI")
            return {"basari": True, "ana": str(ana), "yedek": str(config.yedek_excel_yolu()), "sha256": kaynak_hash}
    except Exception as exc:
        olay.setdefault("durum", "HATA")
        olay["hata"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        _sil(gecici)
        audit_yaz("excel_harici_dosya_ana_yapildi", **olay)
