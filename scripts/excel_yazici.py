# -*- coding: utf-8 -*-
"""Firma Excel'lerine AI kaydi yazma (v2)."""

import logging
import os
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

from config import FIRMALAR, TEMP_DIR, YEDEK_SAYISI
from guvenlik import audit_yaz
from kayit_kilidi import kilitli
from master_data import (
    GURSAR_SISTEM_HESAPLARI,
    SISTEM_ISLEM_ETKILERI,
    SISTEM_ISLEM_TURLERI,
)

log = logging.getLogger(__name__)


# Alias: LLM baska isim kullanirsa bunlara eslestir
_ALAN_ALIASLARI = {
    "Alici": ["Cari", "Alici", "Alıcı", "Kisi", "Kişi", "Firma"],
    "Tutar_TL": ["Tutar", "Tutar_TL", "TutarTL", "Mebla", "Meblağ", "Amount"],
    "Islem_Turu": ["Islem_Turu", "İşlem_Türü", "IslemTuru", "Tur"],
    "Is_Kalemi": ["Is_Kalemi", "İş_Kalemi", "IsKalemi", "Kalem", "Aciklama"],
    "Kaynak_Belge": ["Kaynak_Belge", "Belge", "Belge_No", "Referans"],
    "FAST_Sorgu_No": ["FAST_Sorgu_No", "FastNo", "SorguNo", "Islem_No"],
}


def _normalize_kayit(kayit, beklenen_alanlar):
    """LLM anahtarlarini duzelt. Alias'lar ile eslestir."""
    if not isinstance(kayit, dict):
        return kayit

    def _norm(s):
        s = str(s).lower()
        return (s.replace('_', '').replace(' ', '')
                 .replace('ı', 'i').replace('İ', 'i')
                 .replace('ş', 's').replace('Ş', 's')
                 .replace('ğ', 'g').replace('Ğ', 'g')
                 .replace('ü', 'u').replace('Ü', 'u')
                 .replace('ö', 'o').replace('Ö', 'o')
                 .replace('ç', 'c').replace('Ç', 'c'))

    # Normalize edilmis kayit: {norm_anahtar: (orijinal_anahtar, deger)}
    kayit_norm = {_norm(k): (k, v) for k, v in kayit.items()}

    yeni = {}
    eslesmis_orijinaller = set()

    for alan in beklenen_alanlar:
        # 1) Once direkt eslesme
        k_norm = _norm(alan)
        if k_norm in kayit_norm:
            orijinal_k, deger = kayit_norm[k_norm]
            yeni[alan] = deger
            eslesmis_orijinaller.add(orijinal_k)
            continue

        # 2) Sonra alias'larla eslestir (Alici <-> Cari gibi)
        bulundu = False
        if alan in _ALAN_ALIASLARI:
            for alias in _ALAN_ALIASLARI[alan]:
                alias_norm = _norm(alias)
                if alias_norm in kayit_norm:
                    orijinal_k, deger = kayit_norm[alias_norm]
                    yeni[alan] = deger
                    eslesmis_orijinaller.add(orijinal_k)
                    bulundu = True
                    break

        if not bulundu:
            yeni[alan] = ""

    # LLM'in diger alanlarini da koru
    for k, v in kayit.items():
        if k not in eslesmis_orijinaller and k not in yeni:
            yeni[k] = v

    return yeni



# ============================================================
# YEDEKLEME
# ============================================================

def _yedek_al(firma_key):
    firma = FIRMALAR[firma_key]
    excel = Path(firma["excel"])
    yedek_klasor = Path(firma["yedek_klasoru"])
    yedek_klasor.mkdir(parents=True, exist_ok=True)

    zaman = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    yedek_ad = f"{excel.stem}_{zaman}.xlsx"
    yedek_yol = yedek_klasor / yedek_ad

    shutil.copy2(excel, yedek_yol)
    log.info(f"Yedek alindi: {yedek_yol.name}")
    _eski_yedekleri_temizle(yedek_klasor, tut=YEDEK_SAYISI)
    return yedek_yol


def _eski_yedekleri_temizle(yedek_klasor, tut=10):
    yedekler = sorted(yedek_klasor.glob("*.xlsx"),
                      key=lambda p: p.stat().st_mtime, reverse=True)
    for eski in yedekler[tut:]:
        try:
            eski.unlink()
        except Exception as e:
            log.warning(f"Eski yedek silinemedi: {eski.name} - {e}")


# ============================================================
# YARDIMCILAR
# ============================================================

def _ilk_bos_satir(ws, sutun, baslangic, son=500):
    """Belirtilen sutunda ilk bos satiri bul."""
    for satir in range(baslangic, son):
        if ws[f"{sutun}{satir}"].value is None:
            return satir
    return son


def _mukerrer_mi(ws, deger, sutun, ilk_satir, son_satir=500):
    if not deger:
        return False
    for satir in range(ilk_satir, son_satir):
        h = ws[f"{sutun}{satir}"]
        if h.value and str(h.value).strip() == str(deger).strip():
            return True
    return False


def _mukerrer_kayit_mi(ws, yeni_kayit, sutunlar, ilk_satir, son_satir=500):
    """Ayni Tarih + Cari + Tutar kombinasyonu var mi?
    Donus: (mukerrer_mi, satir_no)
    """
    yeni_tarih = str(yeni_kayit.get("Tarih", "")).strip()
    yeni_cari = str(yeni_kayit.get("Cari") or yeni_kayit.get("Alici") or "").strip().lower()
    yeni_tutar_raw = yeni_kayit.get("Tutar_TL") or yeni_kayit.get("Tutar") or 0

    try:
        yeni_tutar = float(yeni_tutar_raw)
    except (ValueError, TypeError):
        yeni_tutar = 0

    if not yeni_tarih and not yeni_cari and not yeni_tutar:
        return False, None

    tarih_sutun = sutunlar.get("Tarih", "B")
    cari_sutun = sutunlar.get("Alici", "E")
    tutar_sutun = sutunlar.get("Tutar_TL", "G")

    for satir in range(ilk_satir, son_satir):
        mevcut_tarih = str(ws[f"{tarih_sutun}{satir}"].value or "").strip()
        mevcut_cari = str(ws[f"{cari_sutun}{satir}"].value or "").strip().lower()
        mevcut_tutar_raw = ws[f"{tutar_sutun}{satir}"].value or 0

        try:
            mevcut_tutar = float(mevcut_tutar_raw)
        except (ValueError, TypeError):
            mevcut_tutar = 0

        if (mevcut_tarih == yeni_tarih and
            mevcut_cari == yeni_cari and
            mevcut_tutar == yeni_tutar and
            yeni_tutar > 0):
            return True, satir

    return False, None


def _gecici_yol(excel_yolu):
    # Same directory guarantees that replace stays on the destination filesystem.
    return excel_yolu.parent / f".{excel_yolu.stem}.{uuid.uuid4().hex}.tmp{excel_yolu.suffix}"


def _gecici_exceli_dogrula(gecici_yol, beklenen_sayfalar):
    """Kaydi yeniden acar; ZIP/XLSX butunlugu ile hedef sayfalari dogrular."""
    wb = load_workbook(gecici_yol, read_only=True, data_only=False)
    try:
        eksik = [sayfa for sayfa in beklenen_sayfalar if sayfa not in wb.sheetnames]
        if eksik:
            raise ValueError(f"Gecici Excel'de sayfa eksik: {', '.join(eksik)}")
    finally:
        wb.close()


# ============================================================
# ANA YAZMA FONKSIYONU
# ============================================================

@kilitli
def yaz(firma_key, veri, belge_id, test_modu=False):
    """
    Belirtilen firmaya veri yazar.
    test_modu=True ise yazma sonrasi otomatik geri alir.
    """
    if firma_key not in FIRMALAR:
        return {"basari": False, "yazilan": 0, "hata": f"Bilinmeyen firma: {firma_key}"}

    # Tek finansal kaynak: bütün faaliyetler Touche Holding master dosyasına yazılır.
    hedef_key = "touche_holding"
    hedef_keys = [hedef_key]
    excel_yolu = Path(FIRMALAR[hedef_key]["excel"])
    if not excel_yolu.exists():
        return {"basari": False, "yazilan": 0, "hata": f"Master Excel yok: {excel_yolu}"}

    yedekler = {}
    workbookler = {}
    gecici_dosyalar = {}
    degistirilenler = []
    try:
        # Once butun hedeflerin geri donus kopyasi alinir.
        for key in hedef_keys:
            yedekler[key] = _yedek_al(key)
            workbookler[key] = load_workbook(Path(FIRMALAR[key]["excel"]))

        master_wb = workbookler[hedef_key]
        master_firma = FIRMALAR[hedef_key]
        # Sohbetten gelen standart kayıtlar, faaliyetinden bağımsız olarak tek
        # master ledger şemasına yazılır. Böylece Gürsar gelir/hakedişleri eski
        # yalnız-ödeme adaptöründe yanlışlıkla gider tahakkukuna dönüşmez.
        if veri.get("_hedef") == "ana_ledger":
            _yeni_cari_tanimlarini_yaz(master_wb, veri)
            _sistem_tanimlarini_yaz(master_wb, veri)
            sonuc = _yaz_touche_ledger(master_wb, master_firma, veri, belge_id)
        elif firma_key == "touche_holding":
            sonuc = _yaz_touche(master_wb, master_firma, veri, belge_id)
        elif firma_key == "gursar_mimarlik":
            sonuc = _yaz_gursar_holding(master_wb, master_firma, veri, belge_id)
        else:
            faaliyet = FIRMALAR[firma_key].get("faaliyet") or FIRMALAR[firma_key].get("ad")
            master_veri = {
                **veri,
                "kayitlar": [
                    {**kayit, "Faaliyet": kayit.get("Faaliyet") or faaliyet}
                    for kayit in veri.get("kayitlar", [])
                ],
            }
            sonuc = _yaz_touche_ledger(master_wb, master_firma, master_veri, belge_id)

        if firma_key != "touche_holding" and sonuc["basari"]:
            sonuc["holding_yazilan"] = sonuc["yazilan"]

        if not sonuc["basari"]:
            audit_yaz("excel_yazma_reddedildi", firma=firma_key, belge_id=belge_id, hata=sonuc["hata"])
            return sonuc

        # Asil dosyaya dokunmadan gecici kaydet ve yeniden acarak dogrula.
        for key in hedef_keys:
            excel_yolu = Path(FIRMALAR[key]["excel"])
            gecici = _gecici_yol(excel_yolu)
            workbookler[key].save(gecici)
            beklenen = [FIRMALAR[key]["ana_ledger"]["sayfa"]]
            if key == "touche_holding":
                beklenen.append(FIRMALAR[key]["ai_inbox"]["sayfa"])
            _gecici_exceli_dogrula(gecici, beklenen)
            gecici_dosyalar[key] = gecici

        for wb in workbookler.values():
            wb.close()
        workbookler.clear()

        # os.replace ayni disk uzerinde atomiktir. Coklu dosyada hata olursa yedekler geri yuklenir.
        for key in hedef_keys:
            excel_yolu = Path(FIRMALAR[key]["excel"])
            os.replace(gecici_dosyalar[key], excel_yolu)
            degistirilenler.append(key)

        if test_modu:
            for key in hedef_keys:
                shutil.copy2(yedekler[key], Path(FIRMALAR[key]["excel"]))

        audit_yaz(
            "excel_yazma_basarili",
            firma=firma_key,
            belge_id=belge_id,
            yazilan=sonuc["yazilan"],
            holding_yazilan=sonuc.get("holding_yazilan", 0),
            master_excel=str(excel_yolu),
            test_modu=test_modu,
        )
        log.info("Guvenli tek-master kaydi tamamlandi: %s", excel_yolu)
        return sonuc
    except Exception as e:
        log.exception("Guvenli Excel yazma hatasi")
        for key in degistirilenler:
            try:
                shutil.copy2(yedekler[key], Path(FIRMALAR[key]["excel"]))
            except Exception as geri_hata:
                log.critical("Yedekten donulemedi (%s): %s", key, geri_hata)
        teknik_hata = str(e)
        if isinstance(e, PermissionError) or getattr(e, "winerror", None) in (5, 32):
            kullanici_hatasi = (
                "Ana Excel dosyası açık veya OneDrive tarafından kilitli. "
                "Bilgisayardaki tüm Excel pencerelerini kapatın, OneDrive eşitlemesinin "
                "tamamlanmasını bekleyin ve aynı taslakta Şifreyle Onayla'ya tekrar basın. "
                "Kayıt yapılmadı; taslak korunuyor."
            )
        else:
            kullanici_hatasi = teknik_hata
        audit_yaz("excel_yazma_hatasi", firma=firma_key, belge_id=belge_id, hata=teknik_hata)
        return {"basari": False, "yazilan": 0, "hata": kullanici_hatasi}
    finally:
        for wb in workbookler.values():
            try:
                wb.close()
            except Exception:
                pass
        for gecici in gecici_dosyalar.values():
            try:
                if Path(gecici).exists():
                    Path(gecici).unlink()
            except Exception:
                pass


# ============================================================
# TOUCHE HOLDING YAZMA
# ============================================================

def _yaz_touche(wb, firma, veri, belge_id):
    """Touche Holding 05_AI Belge Inbox'a yazar."""
    if veri.get("_hedef") == "ana_ledger":
        return _yaz_touche_ledger(wb, firma, veri, belge_id)
    inbox = firma["ai_inbox"]
    sayfa_adi = inbox["sayfa"]

    if sayfa_adi not in wb.sheetnames:
        return {"basari": False, "yazilan": 0, "hata": f"Sayfa yok: {sayfa_adi}"}

    ws = wb[sayfa_adi]

    # SHA-256 belge ozeti daha once inbox'a yazildiysa ayni belgeyi tekrar alma.
    hash_sutunu = inbox["sutunlar"].get("Belge_Hash")
    for kayit in veri.get("kayitlar", []):
        belge_hash = kayit.get("_hash")
        if hash_sutunu and belge_hash and _mukerrer_mi(
            ws, belge_hash, hash_sutunu, inbox["ilk_veri_satiri"], 500
        ):
            return {"basari": False, "yazilan": 0, "hata": f"Mukerrer SHA-256 belge: {belge_hash}"}

    # Ilk bos satir — baslik satirindan sonra
    bos_satir = _ilk_bos_satir(ws, "A", inbox["ilk_veri_satiri"], 500)

    kayitlar = veri.get("kayitlar", [])
    yazilan = 0

    for i, kayit in enumerate(kayitlar):
        satir = bos_satir + i
        _inbox_satiri_yaz(ws, satir, kayit, belge_id, inbox["sutunlar"], i)
        yazilan += 1

    return {"basari": True, "yazilan": yazilan, "hata": ""}


def _yeni_cari_tanimlarini_yaz(wb, veri):
    """Onaylanan sohbet taslagindaki acik yeni carileri master veriye ekler.

    Ledger ve tanim ayni gecici workbook'a yazildigi icin islem atomiktir.
    Modelin urettigi her ad degil, yalnizca ``_yeni_cari`` bayrakli ve kullanici
    onayindan gecmis kayitlar burada ele alinir.
    """
    oneriler = [x for x in veri.get("kayitlar", []) if x.get("_yeni_cari")]
    if not oneriler:
        return 0
    if "99_Ayarlar" not in wb.sheetnames:
        raise ValueError("Yeni cari tanımı için 99_Ayarlar sayfası bulunamadı.")
    ws = wb["99_Ayarlar"]

    def norm(value):
        import unicodedata
        metin = unicodedata.normalize("NFKD", str(value or "").casefold().replace("ı", "i"))
        return "".join(ch for ch in metin if not unicodedata.combining(ch)).strip()

    mevcut = {norm(ws[f"AI{satir}"].value) for satir in range(13, 113) if ws[f"AI{satir}"].value}
    eklendi = 0
    for kayit in oneriler:
        cari = str(kayit.get("Cari") or "").strip()
        faaliyet = str(kayit.get("Faaliyet") or "").strip()
        cari_tipi = str(kayit.get("Cari_Tipi") or "Diğer").strip()
        if not cari or not faaliyet:
            raise ValueError("Yeni cari tanımı için faaliyet ve cari adı gerekli.")
        if norm(cari) in mevcut:
            continue
        bos = next((satir for satir in range(13, 113) if ws[f"AI{satir}"].value in (None, "")), None)
        if bos is None:
            raise ValueError("99_Ayarlar cari tanım alanı dolu; yeni cari eklenemedi.")
        ws[f"AG{bos}"] = faaliyet
        ws[f"AH{bos}"] = cari_tipi
        ws[f"AI{bos}"] = cari
        mevcut.add(norm(cari))
        eklendi += 1
    return eklendi


def _sistem_tanimlarini_yaz(wb, veri):
    """Onaylanan kaydın kontrollü işlem türü ve hesaplarını 99_Ayarlar'a ekler."""
    if "99_Ayarlar" not in wb.sheetnames:
        raise ValueError("Sistem tanımları için 99_Ayarlar sayfası bulunamadı.")
    ws = wb["99_Ayarlar"]

    def norm(value):
        import unicodedata
        metin = unicodedata.normalize("NFKD", str(value or "").casefold().replace("ı", "i"))
        return "".join(ch for ch in metin if not unicodedata.combining(ch)).strip()

    gerekli_turler = {
        str(k.get("Islem_Turu") or "").strip() for k in veri.get("kayitlar", [])
        if norm(k.get("Islem_Turu")) in {norm(x) for x in SISTEM_ISLEM_TURLERI}
    }
    mevcut_turler = {
        norm(ws[f"T{satir}"].value): satir
        for satir in range(6, 31) if ws[f"T{satir}"].value
    }
    for tur in gerekli_turler:
        satir = mevcut_turler.get(norm(tur))
        if satir is None:
            satir = next(
                (aday for aday in range(6, 31) if ws[f"T{aday}"].value in (None, "")),
                None,
            )
            if satir is None:
                raise ValueError("99_Ayarlar işlem türü alanı dolu; sistem tanımı eklenemedi.")
            ws[f"T{satir}"] = tur
            mevcut_turler[norm(tur)] = satir

        etkiler = SISTEM_ISLEM_ETKILERI.get(tur)
        if etkiler is None:
            etkiler = next(
                (deger for ad, deger in SISTEM_ISLEM_ETKILERI.items() if norm(ad) == norm(tur)),
                None,
            )
        if etkiler is not None:
            for sutun, deger in zip(("U", "V", "W", "X", "Y", "Z", "AA"), etkiler):
                ws[f"{sutun}{satir}"] = deger

    sistem_hesaplari = {norm(x) for x in GURSAR_SISTEM_HESAPLARI}
    gerekli_hesaplar = [
        k for k in veri.get("kayitlar", []) if norm(k.get("Hesap_Kasa")) in sistem_hesaplari
    ]
    mevcut_hesaplar = {
        norm(ws[f"M{satir}"].value): satir
        for satir in range(13, 63) if ws[f"M{satir}"].value
    }
    mevcut_hesap_idleri = {
        str(ws[f"K{satir}"].value).strip()
        for satir in range(13, 63) if ws[f"K{satir}"].value
    }

    def yeni_hesap_id():
        sira = 1
        while f"HSP-SYS-{sira:03d}" in mevcut_hesap_idleri:
            sira += 1
        hesap_id = f"HSP-SYS-{sira:03d}"
        mevcut_hesap_idleri.add(hesap_id)
        return hesap_id

    for kayit in gerekli_hesaplar:
        hesap = str(kayit.get("Hesap_Kasa") or "").strip()
        satir = mevcut_hesaplar.get(norm(hesap))
        if satir is None:
            satir = next(
                (aday for aday in range(13, 63) if ws[f"M{aday}"].value in (None, "")),
                None,
            )
            if satir is None:
                raise ValueError("99_Ayarlar hesap alanı dolu; sistem hesabı eklenemedi.")
            ws[f"L{satir}"] = str(kayit.get("Faaliyet") or "")
            ws[f"M{satir}"] = hesap
            mevcut_hesaplar[norm(hesap)] = satir

        # Önceki sürüm yalnız L/M sütunlarını dolduruyordu. Mevcut yarım
        # tanımları da tamamla ki açılış bakiyesi ve kur formülleri çalışsın.
        if ws[f"K{satir}"].value in (None, ""):
            ws[f"K{satir}"] = yeni_hesap_id()
        if ws[f"L{satir}"].value in (None, ""):
            ws[f"L{satir}"] = str(kayit.get("Faaliyet") or "")
        if ws[f"N{satir}"].value in (None, ""):
            para_birimi = str(kayit.get("PB") or "").strip().upper()
            if not para_birimi:
                para_birimi = "USD" if "USD" in hesap.upper() else "EUR" if "EUR" in hesap.upper() else "TL"
            ws[f"N{satir}"] = para_birimi
        if ws[f"O{satir}"].value in (None, ""):
            ws[f"O{satir}"] = 0
        ws[f"P{satir}"] = f'=IF($K{satir}="","",IFERROR(VLOOKUP($N{satir},$AD$3:$AE$7,2,FALSE),0))'
        ws[f"Q{satir}"] = f'=IF($K{satir}="","",$O{satir}*$P{satir})'


def _yaz_touche_ledger(wb, firma, veri, belge_id):
    """Telegram'da acikca girilen ve onaylanan islemi holding ledger'ina yazar."""
    ledger = firma["ana_ledger"]
    if ledger["sayfa"] not in wb.sheetnames:
        return {"basari": False, "yazilan": 0, "hata": f"Sayfa yok: {ledger['sayfa']}"}
    ws = wb[ledger["sayfa"]]
    sutunlar = ledger["sutunlar"]
    gerekli = ("Tarih", "Faaliyet", "Cari_Tipi", "Cari", "Islem_Turu", "Tutar")
    for sira, kayit in enumerate(veri.get("kayitlar", []), 1):
        eksik = [alan for alan in gerekli if kayit.get(alan) in (None, "")]
        if eksik:
            return {"basari": False, "yazilan": 0,
                    "hata": f"Kayit {sira} ana ledger icin eksik: {', '.join(eksik)}. Duzelt secenegini kullanin."}
    boslar = [r for r in range(ledger['ilk_veri_satiri'],301) if ws[f'A{r}'].value in (None,'')]
    if len(boslar) < len(veri.get('kayitlar',[])):
        return {'basari':False, 'yazilan':0, 'hata':'Holding kayıt merkezi kapasitesi dolu.'}
    for index, kayit in enumerate(veri.get("kayitlar", []), 1):
        kaynak_id = f"{belge_id}-{index:03d}"
        if _mukerrer_mi(ws, kaynak_id, sutunlar["Kaynak_Belge_ID"], ledger["ilk_veri_satiri"], 500):
            return {"basari": False, "yazilan": 0, "hata": f"Mukerrer belge: {kaynak_id}"}
        degerler = {
            "Kayit_ID": f"AI-TH-{uuid.uuid4().hex[:12].upper()}",
            "Tarih": _excel_tarihi(kayit.get("Tarih")),
            "Faaliyet": kayit.get("Faaliyet"),
            "Proje": kayit.get("Proje") or "",
            "Cari_Tipi": kayit.get("Cari_Tipi"),
            "Cari": kayit.get("Cari"),
            "Islem_Turu": kayit.get("Islem_Turu"),
            "Hesap_Kasa": kayit.get("Hesap_Kasa") or "",
            "Tutar": kayit.get("Tutar"),
            "PB": kayit.get("PB") or "TL",
            "Kur": kayit.get("Kur") or 1,
            "Baglanti_ID": kayit.get("Baglanti_ID") or "",
            "Vade": kayit.get("Vade") or "",
            "Durum": kayit.get("Durum") or "Aktif",
            "Belge_Referans": belge_id,
            "Aciklama": kayit.get("Aciklama") or "Telegram metin kaydi",
            "Kaynak_Belge_ID": kaynak_id,
            "Kayit_Kaynak": "CFO Asistanı / AI / Kullanıcı Onaylı",
            "AI_Guven": kayit.get("_guven", 0),
            "Onay_Durumu": "Onaylandı",
        }
        satir = boslar[index - 1]
        for alan, sutun in sutunlar.items():
            if alan in degerler:
                cell = ws[f"{sutun}{satir}"]
                cell.value = degerler[alan]
                if isinstance(degerler[alan], str):
                    cell.data_type = 's'
    return {"basari": True, "yazilan": len(veri.get("kayitlar", [])), "hata": ""}


def _inbox_satiri_yaz(ws, satir, kayit, belge_id, sutunlar, index):
    belge_id_benzersiz = f"{belge_id}-{index+1:03d}"

    degerler = {
        "Belge_ID":         belge_id_benzersiz,
        "Yukleme_Tarihi":   datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Kaynak":           "Telegram Bot / AI",
        "Dosya_Metin_Ref":  kayit.get("_dosya_adi", ""),
        "AI_Faaliyet":      kayit.get("Faaliyet", ""),
        "AI_Proje":         kayit.get("Proje", ""),
        "AI_Cari_Tipi":     kayit.get("Cari_Tipi", ""),
        "AI_Cari":          kayit.get("Cari", ""),
        "AI_Islem_Turu":    kayit.get("Islem_Turu", ""),
        "Tutar":            kayit.get("Tutar", 0),
        "PB":               kayit.get("PB", "TL"),
        "Belge_Tarihi":     kayit.get("Tarih", ""),
        "Vade":             kayit.get("Vade", ""),
        "AI_Guven":         kayit.get("_guven", 0.90),
        "Kullanici_Onayi":  kayit.get("_kullanici_onayi", "Onaylandı"),
        "Hedef_Kayit_ID":   "",
        "AI_Aciklama":      str(kayit.get("_dogrulama_ozeti") or kayit.get("Aciklama", ""))[:200],
        "Belge_Hash":       kayit.get("_hash", ""),
    }

    for alan, sutun in sutunlar.items():
        # Q sutunundaki Excel kontrol formulu sablonun parcasidir.
        if alan == "Durum_Kontrol":
            continue
        ws[f"{sutun}{satir}"] = degerler.get(alan, "")


# ============================================================
# GURSAR YAZMA
# ============================================================

def _yaz_gursar(wb, firma, veri, belge_id):
    """Gursar Excel'ine Odeme Gecmisi sayfasina yazar."""
    ledger = firma["ana_ledger"]
    sayfa_adi = ledger["sayfa"]

    if sayfa_adi not in wb.sheetnames:
        return {"basari": False, "yazilan": 0, "hata": f"Sayfa yok: {sayfa_adi}"}

    ws = wb[sayfa_adi]

    # Mükerrer kontrolü — Kaynak_Belge sutununda belge_id
    belge_sutunu = ledger["sutunlar"]["Kaynak_Belge"]
    if _mukerrer_mi(ws, belge_id, belge_sutunu, ledger["ilk_veri_satiri"]):
        return {"basari": False, "yazilan": 0, "hata": f"Mukerrer belge: {belge_id}"}

    # Ilk bos satir
    bos_satir = _ilk_bos_satir(ws, "A", ledger["ilk_veri_satiri"], 200)

    kayitlar = veri.get("kayitlar", [])
    beklenen = list(ledger["sutunlar"].keys())
    yazilan = 0

    for i, kayit in enumerate(kayitlar):
        kayit = _normalize_kayit(kayit, beklenen)

        # Icerek bazli mukerrer kontrolu
        mukerrer, satir_no = _mukerrer_kayit_mi(
            ws, kayit, ledger["sutunlar"],
            ledger["ilk_veri_satiri"], 200
        )
        if mukerrer:
            # Cakisan kaydin detaylarini topla
            sutunlar = ledger["sutunlar"]
            mevcut_belge = ws[f"{sutunlar.get('Kaynak_Belge', 'L')}{satir_no}"].value
            mevcut_tarih = ws[f"{sutunlar.get('Tarih', 'B')}{satir_no}"].value
            mevcut_cari = ws[f"{sutunlar.get('Alici', 'E')}{satir_no}"].value
            mevcut_tutar = ws[f"{sutunlar.get('Tutar_TL', 'G')}{satir_no}"].value
            mevcut_islem = ws[f"{sutunlar.get('Islem_Turu', 'F')}{satir_no}"].value

            detay = (
                f"MUKERRER BELGE\n"
                f"Sebep: Ayni Tarih + Cari + Tutar daha once Excel'de kayitli\n"
                f"\n"
                f"Cakisan kayit (Satir {satir_no}):\n"
                f"  - Belge ID: {mevcut_belge}\n"
                f"  - Tarih: {mevcut_tarih}\n"
                f"  - Cari: {mevcut_cari}\n"
                f"  - Islem: {mevcut_islem}\n"
                f"  - Tutar: {mevcut_tutar} TL\n"
                f"\n"
                f"Yeni belge (bu fotograf):\n"
                f"  - Tarih: {kayit.get('Tarih')}\n"
                f"  - Cari: {kayit.get('Cari') or kayit.get('Alici')}\n"
                f"  - Tutar: {kayit.get('Tutar_TL') or kayit.get('Tutar')} TL"
            )
            return {
                "basari": False,
                "yazilan": 0,
                "hata": detay
            }

        satir = bos_satir + i
        _gursar_satiri_yaz(ws, satir, kayit, belge_id, ledger["sutunlar"], i)
        yazilan += 1

    return {"basari": True, "yazilan": yazilan, "hata": ""}


def _gursar_satiri_yaz(ws, satir, kayit, belge_id, sutunlar, index):
    # Her belge ve satir icin cakismayan kayit ID.
    kayit_id = f"AI-{uuid.uuid4().hex[:12].upper()}"

    tarih_str = kayit.get("Tarih") or ""
    saat_str = kayit.get("Saat") or datetime.now().strftime("%H:%M:%S")
    tutar = kayit.get("Tutar_TL") or kayit.get("Tutar") or 0
    banka_masrafi = kayit.get("Banka_Masrafi") or 0

    degerler = {
        "Kayit_ID":       kayit_id,
        "Tarih":          tarih_str,
        "Saat":           saat_str,
        "Is_Kalemi":      kayit.get("Is_Kalemi") or "",
        "Alici":          kayit.get("Alici") or kayit.get("Cari") or "",
        "Islem_Turu":     kayit.get("Islem_Turu") or "",
        "Tutar_TL":       tutar,
        "Banka_Masrafi":  banka_masrafi,
        "Islem_Kanali":   kayit.get("Islem_Kanali") or "",
        "FAST_Sorgu_No":  kayit.get("FAST_Sorgu_No") or "",
        "Kaynak_Belge":   belge_id,
        "Durum":          kayit.get("Durum") or "Bekliyor",
        "Proje":          kayit.get("Proje") or "",
        "Maliyet_Grubu":  kayit.get("Maliyet_Grubu") or "",
    }

    for alan, sutun in sutunlar.items():
        ws[f"{sutun}{satir}"] = degerler.get(alan, "")


def _yaz_gursar_holding(wb, firma, veri, belge_id):
    """Gursar odemelerini holding ana ledger'ina ayni belge baglantisiyla yansitir."""
    ledger = firma["ana_ledger"]
    sayfa_adi = ledger["sayfa"]
    if sayfa_adi not in wb.sheetnames:
        return {"basari": False, "yazilan": 0, "hata": f"Holding sayfasi yok: {sayfa_adi}"}

    ws = wb[sayfa_adi]
    sutunlar = ledger["sutunlar"]
    kaynak_sutunu = sutunlar["Kaynak_Belge_ID"]
    ilk_satir = ledger["ilk_veri_satiri"]
    yazilacaklar = []

    for index, kayit in enumerate(veri.get("kayitlar", []), 1):
        kaynak_id = f"{belge_id}-{index:03d}"
        if _mukerrer_mi(ws, kaynak_id, kaynak_sutunu, ilk_satir, 500):
            return {"basari": False, "yazilan": 0, "hata": f"Holding'de mukerrer belge: {kaynak_id}"}
        yazilacaklar.append((kaynak_id, kayit, False))
        try:
            masraf = float(kayit.get("Banka_Masrafi") or 0)
        except (TypeError, ValueError):
            masraf = 0
        if masraf > 0:
            yazilacaklar.append((f"{kaynak_id}-MASRAF", {**kayit, "Tutar": masraf}, True))

    bos_satir = _ilk_bos_satir(ws, "A", ilk_satir, 500)
    for ofset, (kaynak_id, kayit, banka_masrafi) in enumerate(yazilacaklar):
        satir = bos_satir + ofset
        islem_turu = "Banka Masrafı" if banka_masrafi else {
            "Usta ödemesi": "Usta/Tedarikçi Ödemesi",
            "Ortak cari ödemesi": "Ortak Cari Ödemesi",
            "Banka masrafı": "Banka Masrafı",
            "Diğer": "Gider Tahakkuku",
        }.get(kayit.get("Islem_Turu"), "Gider Tahakkuku")
        cari_tipi = "Banka" if banka_masrafi else {
            "Usta ödemesi": "Usta",
            "Ortak cari ödemesi": "Ortak",
            "Banka masrafı": "Banka",
            "Diğer": "Diğer",
        }.get(kayit.get("Islem_Turu"), kayit.get("Cari_Tipi") or "Diğer")
        kayit_id = f"AI-GUR-{uuid.uuid4().hex[:12].upper()}"
        degerler = {
            "Kayit_ID": kayit_id,
            "Tarih": kayit.get("Tarih") or "",
            "Faaliyet": kayit.get("Faaliyet") or "Gürsar İnşaat",
            "Proje": kayit.get("Proje") or "",
            "Cari_Tipi": cari_tipi,
            "Cari": "Ziraat" if banka_masrafi else (kayit.get("Cari") or kayit.get("Alici") or ""),
            "Islem_Turu": islem_turu,
            "Hesap_Kasa": kayit.get("Hesap_Kasa") or "Gürsar TL Banka",
            "Tutar": kayit.get("Tutar") or kayit.get("Tutar_TL") or 0,
            "PB": kayit.get("PB") or "TL",
            "Kur": 1,
            "Durum": "Gerçekleşti",
            "Belge_Referans": belge_id,
            "Aciklama": kayit.get("Is_Kalemi") or "Gürsar otomatik aktarım",
            "Kaynak_Belge_ID": kaynak_id,
            "Kayit_Kaynak": "Gürsar / Telegram Bot / AI",
            "AI_Guven": kayit.get("_guven", 0),
            "Onay_Durumu": "Onaylandı",
        }
        for alan, sutun in sutunlar.items():
            # Formul sutunlarini ve tanimlanmayan sablon alanlarini koru.
            if alan in degerler:
                ws[f"{sutun}{satir}"] = degerler[alan]

    return {"basari": True, "yazilan": len(yazilacaklar), "hata": ""}


def _excel_tarihi(value):
    if isinstance(value, datetime):
        return value
    for fmt in ('%d.%m.%Y','%d/%m/%Y','%d-%m-%Y','%Y-%m-%d','%d.%m.%y'):
        try:
            return datetime.strptime(str(value),fmt)
        except ValueError:
            pass
    return value


def _yaz_standart(wb, firma, veri, belge_id, firma_key):
    from dogrulayici import deterministik_dogrula, FAIL
    from guvenlik import finansal_fingerprint
    check = deterministik_dogrula(veri, firma_key, 'finans_belgesi')
    if check['durum'] == FAIL:
        return {'basari':False, 'yazilan':0, 'hata':'; '.join(check['hatalar'])}
    ws = wb['İşlem Kayıtları']
    sutunlar = firma['ana_ledger']['sutunlar']
    kayitlar = veri.get('kayitlar', [])
    empty = [r for r in range(6,506) if ws[f'A{r}'].value in (None,'')]
    if len(empty) < len(kayitlar):
        return {'basari':False, 'yazilan':0, 'hata':'Faaliyet Excel kapasitesi dolu (500 kayıt).'}
    hashes = {ws[f'W{r}'].value for r in range(6,506) if ws[f'W{r}'].value}
    fingerprints = {ws[f'X{r}'].value for r in range(6,506) if ws[f'X{r}'].value}
    belgeler = {ws[f'R{r}'].value for r in range(6,506) if ws[f'R{r}'].value}
    for i,kayit in enumerate(kayitlar,1):
        fingerprint = finansal_fingerprint(kayit, firma_key, 'finans_belgesi')
        if kayit.get('_hash') in hashes or fingerprint in fingerprints or f'{belge_id}-{i:03d}' in belgeler:
            return {'basari':False, 'yazilan':0, 'hata':'Mükerrer belge veya finansal fingerprint; kayıt yapılmadı.'}
        fingerprints.add(fingerprint)
        kayit['_fingerprint'] = fingerprint
    for i,(row,kayit) in enumerate(zip(empty,kayitlar),1):
        values = {**kayit, 'Kayit_ID':'AI-FA-'+uuid.uuid4().hex[:12].upper(),
                  'Durum':'Gerçekleşti', 'Kaynak_Belge_ID':f'{belge_id}-{i:03d}',
                  'Onay_Durumu':'Onaylandı', 'Kayit_Kaynak':'Telegram / '+firma['ad'],
                  'AI_Guven':kayit.get('_guven',0), 'Belge_Hash':kayit.get('_hash',''),
                  'Fingerprint':kayit['_fingerprint']}
        for alan,col in sutunlar.items():
            if col in ('M','N','O','P'):
                continue
            val = values.get(alan,'')
            if alan == 'Tarih':
                for fmt in ('%d.%m.%Y','%d/%m/%Y','%d-%m-%Y','%Y-%m-%d','%d.%m.%y'):
                    try:
                        val = datetime.strptime(str(val),fmt)
                        break
                    except ValueError:
                        pass
            cell = ws[f'{col}{row}']
            cell.value = val
            if isinstance(val,str):
                cell.data_type = 's'
    return {'basari':True, 'yazilan':len(kayitlar), 'hata':''}


@kilitli
def belge_kaydini_iptal_et(belge_id):
    """Belgeye bagli kayitlari silmeden Iptal/Reddedildi durumuna getirir."""
    belge_id = str(belge_id or "").strip()
    if not belge_id:
        return {"basari": False, "etkilenen": 0, "hata": "Belge ID bos"}

    master_key = "touche_holding"
    master = FIRMALAR[master_key]
    yol = Path(master["excel"])
    wb = None
    yedek = None
    gecici = None
    degistirildi = False
    try:
        if not yol.exists():
            return {"basari": False, "etkilenen": 0, "hata": f"Master Excel yok: {yol}"}
        wb = load_workbook(yol)
        etkilenen = 0

        ledger = master["ana_ledger"]
        ws = wb[ledger["sayfa"]]
        s = ledger["sutunlar"]
        for satir in range(ledger["ilk_veri_satiri"], (ws.max_row or 0) + 1):
            kayit_id = str(ws[f"{s['Kayit_ID']}{satir}"].value or "")
            kaynak = str(ws[f"{s['Kaynak_Belge_ID']}{satir}"].value or "")
            referans = str(ws[f"{s['Belge_Referans']}{satir}"].value or "")
            if (kayit_id == belge_id or kaynak == belge_id or kaynak.startswith(belge_id + "-")
                    or referans == belge_id):
                ws[f"{s['Durum']}{satir}"] = "İptal"
                ws[f"{s['Onay_Durumu']}{satir}"] = "Reddedildi"
                etkilenen += 1

        inbox = master["ai_inbox"]
        ws = wb[inbox["sayfa"]]
        s = inbox["sutunlar"]
        for satir in range(inbox["ilk_veri_satiri"], (ws.max_row or 0) + 1):
            kimlik = str(ws[f"{s['Belge_ID']}{satir}"].value or "")
            if kimlik == belge_id or kimlik.startswith(belge_id + "-"):
                ws[f"{s['Kullanici_Onayi']}{satir}"] = "Reddedildi"
                etkilenen += 1

        if not etkilenen:
            return {"basari": False, "etkilenen": 0, "hata": f"Belge ID bulunamadi: {belge_id}"}

        yedek = _yedek_al(master_key)
        gecici = _gecici_yol(yol)
        wb.save(gecici)
        _gecici_exceli_dogrula(gecici, [ledger["sayfa"], inbox["sayfa"]])
        wb.close()
        wb = None
        os.replace(gecici, yol)
        degistirildi = True
        audit_yaz("belge_iptal", belge_id=belge_id, etkilenen=etkilenen, dosyalar=[master_key])
        return {"basari": True, "etkilenen": etkilenen, "hata": ""}
    except Exception as exc:
        if degistirildi and yedek:
            try:
                shutil.copy2(yedek, yol)
            except Exception as geri_hata:
                log.critical("Iptal geri alma basarisiz: %s", geri_hata)
        audit_yaz("belge_iptal_hatasi", belge_id=belge_id, hata=str(exc))
        return {"basari": False, "etkilenen": 0, "hata": str(exc)}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
        if gecici:
            try:
                Path(gecici).unlink(missing_ok=True)
            except OSError:
                pass


# ============================================================
# YEDEKTEN GERI DONUS
# ============================================================

def yedekten_geri_don(firma_key, yedek_dosya_adi):
    firma = FIRMALAR[firma_key]
    excel = Path(firma["excel"])
    yedek_klasor = Path(firma["yedek_klasoru"])
    yedek = yedek_klasor / yedek_dosya_adi

    if not yedek.exists():
        return {"basari": False, "hata": f"Yedek yok: {yedek}"}

    shutil.copy2(yedek, excel)
    log.info(f"Geri donuldu: {yedek.name} -> {excel.name}")
    return {"basari": True, "hata": ""}


def yedek_listesi(firma_key):
    firma = FIRMALAR[firma_key]
    yedek_klasor = Path(firma["yedek_klasoru"])
    if not yedek_klasor.exists():
        return []
    yedekler = sorted(yedek_klasor.glob("*.xlsx"),
                      key=lambda p: p.stat().st_mtime, reverse=True)
    return [{"ad": p.name,
             "boyut_kb": round(p.stat().st_size / 1024, 1),
             "tarih": datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")}
            for p in yedekler]
