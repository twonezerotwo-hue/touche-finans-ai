"""Yeni faaliyet, proje ve cari açılışını tek Holding master dosyasında yönetir."""
import json
import math
import os
import shutil
import uuid
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from config import FIRMALAR, FAALIYET_REGISTRY, FIRMALAR_DIR, master_firma_tanimi
from faaliyet_sema import anahtar
from kayit_kilidi import kilitli
from guvenlik import audit_yaz, belge_sha256

def metin_kontrol(deger):
    deger = str(deger or '').strip()
    if not 2 <= len(deger) <= 100 or deger[0] in '=+-@' or any(ord(c) < 32 for c in deger):
        raise ValueError('Ad 2–100 karakter olmalı; formül veya kontrol karakteri içeremez.')
    return deger

def _bos(ws, sutun, ilk, son):
    for r in range(ilk, son + 1):
        if ws[f'{sutun}{r}'].value in (None, ''):
            return r
    raise ValueError(f'{ws.title}: {sutun}{ilk}:{sutun}{son} kapasitesi dolu; mevcut kayıtlar korunuyor.')

def _registry_oku():
    if not FAALIYET_REGISTRY.exists():
        return {'surum': 1, 'faaliyetler': {}}
    data = json.loads(FAALIYET_REGISTRY.read_text(encoding='utf-8'))
    if data.get('surum') != 1:
        raise ValueError('Registry sürümü desteklenmiyor.')
    return data

def _toplu_kaydet(workbooks, registry=None):
    """Tamamlanmis temp dosyalari dogrular; hata halinde degisenleri geri alir."""
    temps, backups, before, changed = {}, {}, {}, []
    batch = FIRMALAR_DIR / 'kurulum_yedek' / (datetime.now().strftime('%Y%m%d_%H%M%S_') + uuid.uuid4().hex[:8])
    batch.mkdir(parents=True, exist_ok=True)
    targets = list(workbooks)
    if registry is not None:
        targets.append(FAALIYET_REGISTRY)
    try:
        for i, target in enumerate(targets):
            target = Path(target)
            target.parent.mkdir(parents=True, exist_ok=True)
            before[target] = belge_sha256(target) if target.exists() else None
            if target.exists():
                backup = batch / f'{i}_{target.name}'
                shutil.copy2(target, backup)
                backups[target] = backup
            temp = target.with_name('.' + uuid.uuid4().hex + target.suffix)
            temps[target] = temp
            if target in workbooks:
                workbooks[target].save(temp)
                check = load_workbook(temp, read_only=True)
                try:
                    if check.sheetnames != workbooks[target].sheetnames:
                        raise ValueError('Geçici Excel sayfa doğrulaması başarısız.')
                finally:
                    check.close()
            else:
                temp.write_text(json.dumps(registry, ensure_ascii=False, indent=2), encoding='utf-8')
                json.loads(temp.read_text(encoding='utf-8'))
        for wb in workbooks.values():
            wb.close()
        for target in temps:
            actual = belge_sha256(target) if target.exists() else None
            if actual != before[target]:
                raise RuntimeError(f'Dosya işlem sırasında değişti: {target.name}')
        for target, temp in temps.items():
            os.replace(temp, target)
            changed.append(target)
    except Exception:
        for target in reversed(changed):
            if target in backups:
                shutil.copy2(backups[target], target)
            else:
                # Only the exact new file created by this batch is removed.
                target.unlink(missing_ok=True)
        raise
    finally:
        for wb in workbooks.values():
            wb.close()
        for temp in temps.values():
            temp.unlink(missing_ok=True)

@kilitli
def faaliyet_olustur(ad, proje, hesap, pb='TL', holding_payi=100, dis_ortak=''):
    ad, proje, hesap = map(metin_kontrol, (ad, proje, hesap))
    key = anahtar(ad)
    if not key:
        raise ValueError('Faaliyet adına harf veya rakam ekleyin.')
    oran = float(str(holding_payi).replace(',', '.')) / 100
    if not math.isfinite(oran) or not 0 < oran <= 1:
        raise ValueError('Holding payı 0’dan büyük, en fazla 100 olmalı.')
    if oran < 1:
        dis_ortak = metin_kontrol(dis_ortak)
    if pb not in ('TL', 'USD', 'EUR'):
        raise ValueError('Para birimi TL, USD veya EUR olmalı.')
    registry = _registry_oku()
    if key in FIRMALAR or key in registry['faaliyetler'] or any(anahtar(v['ad']) == key for v in FIRMALAR.values()):
        raise ValueError('Bu faaliyet zaten kayıtlı.')
    holding = None
    try:
        holding = load_workbook(FIRMALAR['touche_holding']['excel'])
        master = holding['99_Ayarlar']
        if any(anahtar(master[f'A{r}'].value or '') == key for r in range(13, 33)):
            raise ValueError('Holding’de aynı faaliyet zaten var; önce mevcut kayıt eşleştirilmeli.')
        r = _bos(master, 'A', 13, 32)
        for col, val in {'A':ad, 'B':oran, 'C':dis_ortak, 'D':1-oran, 'E':'Aktif'}.items():
            master[f'{col}{r}'] = val
        r = _bos(master, 'I', 13, 112)
        master[f'H{r}'], master[f'I{r}'] = ad, proje
        r = _bos(master, 'K', 13, 62)
        for col, val in {'K':'HSP-'+uuid.uuid4().hex[:10], 'L':ad, 'M':hesap, 'N':pb, 'O':0}.items():
            master[f'{col}{r}'] = val
        kayit = {'ad':ad, 'faaliyet':ad, 'excel_mode':'master', 'pb':pb,
                 'varsayilan_hesap':hesap, 'holding_payi':oran, 'dis_ortak':dis_ortak}
        registry['faaliyetler'][key] = kayit
        _toplu_kaydet({Path(FIRMALAR['touche_holding']['excel']):holding}, registry)
        FIRMALAR[key] = master_firma_tanimi(kayit)
        audit_yaz('faaliyet_olusturuldu', faaliyet=key, excel_mode='master', holding_payi=oran)
        return key
    finally:
        if holding:
            holding.close()

@kilitli
def proje_olustur(firma_key, proje):
    proje = metin_kontrol(proje)
    firma = FIRMALAR[firma_key]
    if firma_key == 'touche_holding':
        raise ValueError('Önce projenin bağlı olduğu faaliyeti /firma ile seçin.')
    from master_data import master_data_oku
    if any(anahtar(p) == anahtar(proje) for p in master_data_oku(firma_key)['projeler']):
        raise ValueError('Bu proje zaten kayıtlı.')
    holding = None
    try:
        holding = load_workbook(FIRMALAR['touche_holding']['excel'])
        faaliyet = firma.get('faaliyet', 'Gürsar İnşaat')
        ws = holding['99_Ayarlar']
        r = _bos(ws, 'I', 13, 112)
        ws[f'H{r}'], ws[f'I{r}'] = faaliyet, proje
        _toplu_kaydet({Path(FIRMALAR['touche_holding']['excel']):holding})
        audit_yaz('proje_olusturuldu', faaliyet=firma_key, proje=proje)
    finally:
        if holding:
            holding.close()


@kilitli
def cari_olustur(firma_key, cari, cari_tipi):
    cari, cari_tipi = metin_kontrol(cari), metin_kontrol(cari_tipi)
    firma = FIRMALAR[firma_key]
    if firma_key == 'touche_holding':
        raise ValueError('Cari için önce bağlı faaliyeti seçin.')
    from master_data import master_data_oku
    md = master_data_oku(firma_key)
    if cari_tipi not in md['cari_tipler']:
        raise ValueError('Cari tipi tanımsız: '+', '.join(md['cari_tipler']))
    if any(anahtar(c) == anahtar(cari) for c in md['cariler']):
        raise ValueError('Cari zaten kayıtlı.')
    holding = None
    try:
        holding = load_workbook(FIRMALAR['touche_holding']['excel'])
        ws = holding['99_Ayarlar']
        r = _bos(ws,'AI',13,112)
        ws[f'AG{r}'], ws[f'AH{r}'], ws[f'AI{r}'] = firma['faaliyet'], cari_tipi, cari
        _toplu_kaydet({Path(FIRMALAR['touche_holding']['excel']):holding})
        audit_yaz('cari_olusturuldu',faaliyet=firma_key,cari=cari,cari_tipi=cari_tipi)
    finally:
        if holding:
            holding.close()
