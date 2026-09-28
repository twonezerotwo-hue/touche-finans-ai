# -*- coding: utf-8 -*-
"""Telegram icin Excel'lerden salt-okunur finans sorgulari."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from config import FIRMALAR
from faaliyet_sema import anahtar


def _kayitlari_oku(firma_key):
    firma = FIRMALAR[firma_key]
    master = FIRMALAR["touche_holding"]
    ledger = master["ana_ledger"]
    yol = Path(master["excel"])
    wb = load_workbook(yol, read_only=True, data_only=False)
    try:
        ws = wb[ledger["sayfa"]]
        sutunlar = ledger["sutunlar"]
        kayitlar = []
        indexes = {alan:column_index_from_string(col)-1 for alan,col in sutunlar.items()}
        rows = ws.iter_rows(min_row=ledger['ilk_veri_satiri'], max_col=max(indexes.values())+1, values_only=True)
        for satir, row in enumerate(rows,ledger['ilk_veri_satiri']):
            kayit_id = row[indexes['Kayit_ID']]
            if not kayit_id:
                continue
            kayit = {alan:row[index] for alan,index in indexes.items()}
            kayit["_satir"] = satir
            kayitlar.append(kayit)
        if firma_key == "touche_holding":
            return kayitlar
        adlar = [firma.get("faaliyet"), firma.get("ad"), *firma.get("faaliyet_aliaslari", [])]
        izinli = {anahtar(ad) for ad in adlar if ad}
        return [kayit for kayit in kayitlar if anahtar(kayit.get("Faaliyet")) in izinli]
    finally:
        wb.close()


def son_kayitlar(firma_key, adet=5):
    adet = max(1, min(int(adet), 20))
    return _kayitlari_oku(firma_key)[-adet:][::-1]


def kayit_ara(firma_key, arama, limit=10):
    aranan = str(arama).casefold().strip()
    sonuc = []
    for kayit in reversed(_kayitlari_oku(firma_key)):
        metin = " | ".join(str(v) for k, v in kayit.items() if not k.startswith("_") and v is not None)
        if aranan in metin.casefold():
            sonuc.append(kayit)
            if len(sonuc) >= limit:
                break
    return sonuc


def finans_ozeti(firma_key):
    kayitlar = _kayitlari_oku(firma_key)
    aktif = [k for k in kayitlar if anahtar(k.get('Durum')) not in {'iptal', 'reddedildi'}]
    toplam = 0.0
    for kayit in aktif:
        deger = kayit.get("Tutar")
        try:
            kur = kayit.get('Kur') or 1
            toplam += float(deger or 0) * float(kur or 1)
        except (TypeError, ValueError):
            pass
    return {"kayit_sayisi": len(kayitlar), "aktif_kayit": len(aktif), "toplam_tutar": toplam}


def kayit_formatla(firma_key, kayit):
    return (
        f"{kayit.get('Kayit_ID')} | {kayit.get('Tarih') or '-'} | "
        f"{kayit.get('Faaliyet') or '-'} / {kayit.get('Proje') or '-'} | "
        f"{kayit.get('Cari') or '-'} | {kayit.get('Islem_Turu') or '-'} | "
        f"{kayit.get('Tutar') or 0} {kayit.get('PB') or 'TL'} | {kayit.get('Durum') or '-'} | "
        f"Belge: {kayit.get('Kaynak_Belge_ID') or kayit.get('Belge_Referans') or '-'}"
    )
