# -*- coding: utf-8 -*-
"""Touche Holding merkezi finans cekirdegi.

Excel dosyalari operasyonel gorunum olarak kalir. Bu modul onaylanan finansal
hareketleri kurus hassasiyetinde, dengeli ve ters kayitli bir SQLite defterinde
tutar; donem, butce, mutabakat, yetki ve raporlamayi tek yerde toplar.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from config import ARSIV, FINANS_DB, LOGLAR, MUHASEBE_EXPORT_DIR, FIRMALAR, ana_excel_yolu

DB_PATH = Path(FINANS_DB)
SYNC_QUEUE_PATH = Path(LOGLAR) / "finans_sync_queue.jsonl"

HESAP_PLANI = (
    ("100", "Kasa", "VARLIK"),
    ("101", "Alınan Çekler", "VARLIK"),
    ("102", "Bankalar", "VARLIK"),
    ("120", "Alıcılar", "VARLIK"),
    ("131", "Ortaklardan Alacaklar", "VARLIK"),
    ("159", "Verilen Sipariş Avansları", "VARLIK"),
    ("191", "İndirilecek KDV", "VARLIK"),
    ("300", "Banka Kredileri", "YUKUMLULUK"),
    ("320", "Satıcılar", "YUKUMLULUK"),
    ("331", "Ortaklara Borçlar", "YUKUMLULUK"),
    ("335", "Personele Borçlar", "YUKUMLULUK"),
    ("336", "Diğer Çeşitli Borçlar", "YUKUMLULUK"),
    ("391", "Hesaplanan KDV", "YUKUMLULUK"),
    ("500", "Sermaye", "OZKAYNAK"),
    ("600", "Yurtiçi Satışlar", "GELIR"),
    ("646", "Kambiyo Karları", "GELIR"),
    ("649", "Diğer Olağan Gelirler", "GELIR"),
    ("656", "Kambiyo Zararları", "GIDER"),
    ("740", "Hizmet Üretim Maliyeti", "GIDER"),
    ("770", "Genel Yönetim Giderleri", "GIDER"),
    ("780", "Finansman Giderleri", "GIDER"),
    ("799", "Diğer Giderler", "GIDER"),
)


SCHEMA = """
PRAGMA foreign_keys=ON;
CREATE TABLE IF NOT EXISTS schema_meta (anahtar TEXT PRIMARY KEY, deger TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS faaliyetler (
  faaliyet_key TEXT PRIMARY KEY, ad TEXT NOT NULL, ana_faaliyet_key TEXT,
  yasal_varlik INTEGER NOT NULL DEFAULT 0, sahiplik_bps INTEGER NOT NULL DEFAULT 10000,
  aktif INTEGER NOT NULL DEFAULT 1, guncelleme TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS hesaplar (
  kod TEXT PRIMARY KEY, ad TEXT NOT NULL, hesap_tipi TEXT NOT NULL, aktif INTEGER NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS donemler (
  yil INTEGER NOT NULL, ay INTEGER NOT NULL, durum TEXT NOT NULL DEFAULT 'ACIK',
  kapatan TEXT, kapanis_zamani TEXT, PRIMARY KEY(yil, ay)
);
CREATE TABLE IF NOT EXISTS fisler (
  fis_id TEXT PRIMARY KEY, belge_id TEXT NOT NULL, kayit_sira INTEGER NOT NULL,
  faaliyet_key TEXT NOT NULL, tarih TEXT NOT NULL, vade TEXT, proje TEXT, cari TEXT,
  aciklama TEXT, para_birimi TEXT NOT NULL DEFAULT 'TL', kur_mikro INTEGER NOT NULL DEFAULT 1000000,
  durum TEXT NOT NULL DEFAULT 'MUHASEBELESTI', ters_fis_id TEXT, kaynak TEXT,
  olusturan TEXT, onaylayan TEXT, olusturma_zamani TEXT NOT NULL,
  UNIQUE(belge_id, kayit_sira), FOREIGN KEY(faaliyet_key) REFERENCES faaliyetler(faaliyet_key)
);
CREATE TABLE IF NOT EXISTS fis_satirlari (
  fis_id TEXT NOT NULL, satir_no INTEGER NOT NULL, hesap_kodu TEXT NOT NULL,
  borc_kurus INTEGER NOT NULL DEFAULT 0, alacak_kurus INTEGER NOT NULL DEFAULT 0,
  aciklama TEXT, PRIMARY KEY(fis_id, satir_no),
  FOREIGN KEY(fis_id) REFERENCES fisler(fis_id), FOREIGN KEY(hesap_kodu) REFERENCES hesaplar(kod),
  CHECK(borc_kurus >= 0 AND alacak_kurus >= 0), CHECK(NOT (borc_kurus > 0 AND alacak_kurus > 0))
);
CREATE TABLE IF NOT EXISTS butceler (
  faaliyet_key TEXT NOT NULL, proje TEXT NOT NULL DEFAULT '', hesap_kodu TEXT NOT NULL,
  yil INTEGER NOT NULL, ay INTEGER NOT NULL, tutar_kurus INTEGER NOT NULL,
  guncelleyen TEXT, guncelleme TEXT NOT NULL,
  PRIMARY KEY(faaliyet_key, proje, hesap_kodu, yil, ay)
);
CREATE TABLE IF NOT EXISTS banka_hareketleri (
  hareket_id TEXT PRIMARY KEY, hesap TEXT, tarih TEXT NOT NULL, aciklama TEXT,
  tutar_kurus INTEGER NOT NULL, referans TEXT, dosya_hash TEXT NOT NULL,
  durum TEXT NOT NULL DEFAULT 'ESLESMEDI', fis_id TEXT,
  FOREIGN KEY(fis_id) REFERENCES fisler(fis_id)
);
CREATE TABLE IF NOT EXISTS kullanicilar (
  chat_id INTEGER PRIMARY KEY, rol TEXT NOT NULL, aktif INTEGER NOT NULL DEFAULT 1,
  ekleyen TEXT, eklenme TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS onay_politikasi (
  politika_id INTEGER PRIMARY KEY CHECK(politika_id=1), limit_kurus INTEGER NOT NULL,
  yuksek_tutar_onay INTEGER NOT NULL, guncelleyen TEXT, guncelleme TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS onay_oylari (
  belge_id TEXT NOT NULL, chat_id INTEGER NOT NULL, zaman TEXT NOT NULL,
  PRIMARY KEY(belge_id, chat_id)
);
CREATE TABLE IF NOT EXISTS hesap_eslemeleri (
  faaliyet_key TEXT NOT NULL, yerel_hesap_kodu TEXT NOT NULL, grup_hesap_kodu TEXT NOT NULL,
  aciklama TEXT, guncelleyen TEXT, guncelleme TEXT NOT NULL,
  PRIMARY KEY(faaliyet_key, yerel_hesap_kodu),
  FOREIGN KEY(faaliyet_key) REFERENCES faaliyetler(faaliyet_key),
  FOREIGN KEY(grup_hesap_kodu) REFERENCES hesaplar(kod)
);
CREATE TABLE IF NOT EXISTS sirketler_arasi_islemler (
  islem_id TEXT PRIMARY KEY, kaynak_faaliyet TEXT NOT NULL, hedef_faaliyet TEXT NOT NULL,
  kaynak_fis_id TEXT NOT NULL, hedef_fis_id TEXT NOT NULL, tarih TEXT NOT NULL,
  tutar_kurus INTEGER NOT NULL, para_birimi TEXT NOT NULL DEFAULT 'TL',
  aciklama TEXT, durum TEXT NOT NULL DEFAULT 'AKTIF', olusturan TEXT, olusturma_zamani TEXT NOT NULL,
  FOREIGN KEY(kaynak_faaliyet) REFERENCES faaliyetler(faaliyet_key),
  FOREIGN KEY(hedef_faaliyet) REFERENCES faaliyetler(faaliyet_key)
);
CREATE TABLE IF NOT EXISTS kurlar (
  tarih TEXT NOT NULL, para_birimi TEXT NOT NULL, kur_mikro INTEGER NOT NULL,
  kaynak TEXT, guncelleyen TEXT, guncelleme TEXT NOT NULL,
  PRIMARY KEY(tarih, para_birimi)
);
CREATE TABLE IF NOT EXISTS onay_kurallari (
  kural_id TEXT PRIMARY KEY, faaliyet_key TEXT NOT NULL DEFAULT '*',
  min_kurus INTEGER NOT NULL DEFAULT 0, max_kurus INTEGER,
  gerekli_onay INTEGER NOT NULL DEFAULT 1, gerekli_roller TEXT NOT NULL DEFAULT '[]',
  ayri_hazirlayan INTEGER NOT NULL DEFAULT 0, aktif INTEGER NOT NULL DEFAULT 1,
  guncelleyen TEXT, guncelleme TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS onay_kararlari (
  belge_id TEXT NOT NULL, aktor TEXT NOT NULL, rol TEXT NOT NULL,
  karar TEXT NOT NULL DEFAULT 'ONAY', zaman TEXT NOT NULL,
  PRIMARY KEY(belge_id, aktor)
);
CREATE TABLE IF NOT EXISTS belge_butunlugu (
  belge_id TEXT PRIMARY KEY, sha256 TEXT NOT NULL, finansal_hash TEXT,
  kaynak TEXT, kayit_zamani TEXT NOT NULL, son_dogrulama TEXT, dogrulama_durumu TEXT
);
CREATE TABLE IF NOT EXISTS banka_eslestirme_gecmisi (
  id INTEGER PRIMARY KEY AUTOINCREMENT, hareket_id TEXT NOT NULL, fis_id TEXT,
  eylem TEXT NOT NULL, skor INTEGER, aktor TEXT, zaman TEXT NOT NULL,
  FOREIGN KEY(hareket_id) REFERENCES banka_hareketleri(hareket_id),
  FOREIGN KEY(fis_id) REFERENCES fisler(fis_id)
);
CREATE INDEX IF NOT EXISTS idx_fis_tarih ON fisler(tarih);
CREATE INDEX IF NOT EXISTS idx_fis_faaliyet ON fisler(faaliyet_key);
CREATE INDEX IF NOT EXISTS idx_fis_cari ON fisler(cari);
CREATE INDEX IF NOT EXISTS idx_sirketler_arasi_tarih ON sirketler_arasi_islemler(tarih);
CREATE INDEX IF NOT EXISTS idx_banka_durum ON banka_hareketleri(durum);
"""


def _simdi():
    return datetime.now().isoformat(timespec="seconds")


def _norm(metin):
    import unicodedata
    s = unicodedata.normalize("NFKD", str(metin or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).casefold().replace("ı", "i")


def faaliyet_anahtari_bul(faaliyet_adi, varsayilan="touche_holding"):
    """Excel'deki faaliyet adini kanonik merkezi-defter anahtarina cevirir."""
    aranan = _norm(faaliyet_adi)
    if not aranan:
        return varsayilan
    for key, bilgi in FIRMALAR.items():
        adaylar = [key, bilgi.get("ad"), bilgi.get("faaliyet")]
        adaylar.extend(bilgi.get("faaliyet_aliaslari") or [])
        if aranan in {_norm(x) for x in adaylar if x}:
            return key
    return varsayilan


def _decimal(deger):
    if deger in (None, ""):
        return Decimal("0")
    metin = str(deger).strip().replace("₺", "").replace(" ", "")
    if "," in metin and "." in metin:
        metin = metin.replace(".", "").replace(",", ".")
    elif "," in metin:
        metin = metin.replace(",", ".")
    try:
        d = Decimal(metin)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Geçersiz tutar: {deger}") from exc
    if not d.is_finite():
        raise ValueError(f"Sonlu olmayan tutar: {deger}")
    return d


def _kurus(deger):
    return int((_decimal(deger) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _tarih(deger):
    if isinstance(deger, datetime):
        return deger.date()
    if isinstance(deger, date):
        return deger
    s = str(deger or "").strip()
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Geçerli işlem tarihi gerekli: {deger}")


@contextmanager
def _baglanti():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA journal_mode=WAL")
    try:
        yield con
        con.commit()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def veritabanini_hazirla():
    with _baglanti() as con:
        con.executescript(SCHEMA)
        con.execute("INSERT OR REPLACE INTO schema_meta VALUES('surum','2')")
        con.executemany("INSERT OR IGNORE INTO hesaplar(kod,ad,hesap_tipi) VALUES(?,?,?)", HESAP_PLANI)
        for key, firma in FIRMALAR.items():
            ad = firma.get("faaliyet") or firma.get("ad") or key
            pay = firma.get("holding_payi", 1)
            try:
                bps = int(Decimal(str(pay)) * 10000)
            except Exception:
                bps = 10000
            con.execute(
                "INSERT INTO faaliyetler(faaliyet_key,ad,ana_faaliyet_key,yasal_varlik,sahiplik_bps,guncelleme) "
                "VALUES(?,?,?,?,?,?) ON CONFLICT(faaliyet_key) DO UPDATE SET ad=excluded.ad, "
                "sahiplik_bps=excluded.sahiplik_bps,guncelleme=excluded.guncelleme",
                (key, ad, None if key == "touche_holding" else "touche_holding", 1 if key == "touche_holding" else 0, bps, _simdi()),
            )
            con.executemany(
                "INSERT OR IGNORE INTO hesap_eslemeleri(faaliyet_key,yerel_hesap_kodu,grup_hesap_kodu,aciklama,guncelleyen,guncelleme) "
                "VALUES(?,?,?,?,?,?)",
                [(key, kod, kod, "Varsayılan ortak hesap planı", "sistem", _simdi()) for kod, _, _ in HESAP_PLANI],
            )
        con.execute(
            "INSERT OR IGNORE INTO onay_politikasi(politika_id,limit_kurus,yuksek_tutar_onay,guncelleme) VALUES(1,25000000,1,?)",
            (_simdi(),),
        )
        con.execute(
            "INSERT OR IGNORE INTO onay_kurallari(kural_id,faaliyet_key,min_kurus,max_kurus,gerekli_onay,gerekli_roller,ayri_hazirlayan,guncelleyen,guncelleme) "
            "VALUES('VARSAYILAN','*',0,NULL,1,'[\"ONAYCI\",\"FINANS\",\"ADMIN\"]',0,'sistem',?)",
            (_simdi(),),
        )
        con.execute(
            "INSERT OR IGNORE INTO onay_kurallari(kural_id,faaliyet_key,min_kurus,max_kurus,gerekli_onay,gerekli_roller,ayri_hazirlayan,guncelleyen,guncelleme) "
            "VALUES('YUKSEK_TUTAR','*',25000000,NULL,2,'[\"ONAYCI\",\"FINANS\",\"ADMIN\"]',0,'sistem',?)",
            (_simdi(),),
        )
        con.execute(
            "UPDATE onay_kurallari SET ayri_hazirlayan=0 WHERE kural_id='YUKSEK_TUTAR' AND guncelleyen='sistem'"
        )


def _donem_acik_mi(con, tarih):
    satir = con.execute("SELECT durum FROM donemler WHERE yil=? AND ay=?", (tarih.year, tarih.month)).fetchone()
    return not satir or satir["durum"] == "ACIK"


GRUP_ICI_ANAHTARLAR = ("intercompany", "grup ici", "grup ici transfer", "sirketler arasi", "sirket arasi")


def _hesap_eslesmesi(kayit, tur):
    """Islem turunu (borc, alacak) hesap cifti ile eslestirir.

    Eslesmesi olmayan veya grup ici transfer olan islem ASLA sessizce gidere yazilmaz;
    hata verilir ve kayit yapilmaz (grup ici transfer icin ``sirketler_arasi_transfer`` kullanilir).
    """
    islem = _norm(kayit.get("Islem_Turu") or tur)
    hesap_kasa = _norm(kayit.get("Hesap_Kasa"))
    nakit_hesabi = "100" if "kasa" in hesap_kasa else "102"
    if any(x in islem for x in GRUP_ICI_ANAHTARLAR):
        raise ValueError(
            "Grup içi transfer gider olarak yazılamaz; 'Grup içi transfer' ekranından "
            "kaynak ve hedef faaliyet fişleriyle birlikte oluşturun. Kayıt yapılmadı."
        )
    if "alacak duzeltme" in islem or "mahsup" in islem:
        return "600", "120"
    if "kredi kullanimi" in islem:
        return "102", "300"
    if "cek teminatli borclanma" in islem or "diger finansal borc girisi" in islem:
        return nakit_hesabi, "336"
    if "usta/tedarikci odemesi" in islem or "tedarikci odemesi" in islem:
        return "320", nakit_hesabi
    if "ortak cepten gider" in islem:
        return "740", "331"
    if "ortak cari odemesi" in islem:
        return "331", nakit_hesabi
    if "gider tahakkuk" in islem or tur == "fatura":
        return "740", "320"
    if "hakedis" in islem or "gelir tahakkuk" in islem:
        return "120", "600"
    if "tahsilat" in islem:
        return nakit_hesabi, "120"
    if "cek alindi" in islem or tur == "cek":
        return "101", "120"
    if "ortak cari" in islem or "ortak odeme" in islem:
        return "331", nakit_hesabi
    if "banka masraf" in islem or "finansman" in islem:
        return "780", "102"
    if "kredi" in islem:
        return "300", "102"
    if "usta" in islem or "odeme" in islem or "gider" in islem:
        return "740", nakit_hesabi
    if "gelir" in islem:
        return nakit_hesabi, "649"
    raise ValueError(
        f"'{kayit.get('Islem_Turu') or tur}' işlem türü için muhasebe hesap eşlemesi tanımlı değil; "
        "bilinmeyen tür gider olarak yazılmaz. Kayıt yapılmadı."
    )


def fis_taslaklari(hazir):
    """Belge verisinden dengeli fis taslaklari olusturur; DB'ye yazmaz."""
    veritabanini_hazirla()
    taslaklar = []
    belge_id = str(hazir.get("belge_id") or "").strip()
    if not belge_id:
        raise ValueError("Belge ID olmadan muhasebe fişi oluşturulamaz.")
    faaliyet = hazir.get("firma_key") or "touche_holding"
    tur = hazir.get("tur") or ""
    with _baglanti() as con:
        for sira, kayit in enumerate(hazir.get("veri", {}).get("kayitlar", []), 1):
            pb = str(kayit.get("PB") or "TL").upper()
            try:
                kur = _decimal(kayit.get("Kur"))
            except ValueError:
                kur = Decimal("0")
            if pb != "TL" and kur <= 0:
                raise ValueError(
                    f"{pb} tutarindaki kayit icin gecerli kur (Kur) gerekli; "
                    "eksik kur 1.0 varsayilamaz."
                )
            if pb == "TL":
                kur = Decimal("1")
            tutar_doviz = _decimal(kayit.get("Tutar") or kayit.get("Toplam"))
            tutar = _kurus(tutar_doviz * kur)
            if tutar <= 0:
                continue
            gun = _tarih(kayit.get("Tarih") or kayit.get("Vade_Tarihi") or date.today())
            if not _donem_acik_mi(con, gun):
                raise ValueError(f"{gun:%Y-%m} dönemi kapalı; kayıt yapılamaz.")
            borc_hesap, alacak_hesap = _hesap_eslesmesi(kayit, tur)
            masraf = _kurus(kayit.get("Banka_Masrafi"))
            satirlar = [
                {"hesap": borc_hesap, "borc": tutar, "alacak": 0, "aciklama": "Ana hareket"},
                {"hesap": alacak_hesap, "borc": 0, "alacak": tutar + masraf, "aciklama": "Karşı hesap"},
            ]
            if masraf:
                satirlar.insert(1, {"hesap": "780", "borc": masraf, "alacak": 0, "aciklama": "Banka masrafı"})
            if sum(x["borc"] for x in satirlar) != sum(x["alacak"] for x in satirlar):
                raise ValueError("Muhasebe fişi dengeli değil.")
            taslaklar.append({
                "fis_id": f"JRN-{belge_id}-{sira:03d}", "belge_id": belge_id, "sira": sira,
                "faaliyet": faaliyet, "tarih": gun.isoformat(),
                "vade": str(kayit.get("Vade") or kayit.get("Vade_Tarihi") or ""),
                "proje": str(kayit.get("Proje") or ""),
                "cari": str(kayit.get("Cari") or kayit.get("Alici") or kayit.get("Tedarikci") or ""),
                "aciklama": str(kayit.get("Aciklama") or kayit.get("Is_Kalemi") or kayit.get("Urun") or tur),
                "pb": pb,
                "kur_mikro": max(1, int(kur * 1_000_000)),
                "kaynak": str(kayit.get("_dosya_adi") or hazir.get("kaynak_metin") or ""),
                "is_kalemi_id": str(kayit.get("Is_Kalemi_ID") or "").strip() or None,
                "satirlar": satirlar,
            })
    return taslaklar


def _is_kalemi_kolonu_hazirla(con):
    """fisler.is_kalemi_id (nullable, ek) yoksa ekler; var olan veriye dokunmaz."""
    if not any(r[1] == "is_kalemi_id" for r in con.execute("PRAGMA table_info(fisler)").fetchall()):
        con.execute("ALTER TABLE fisler ADD COLUMN is_kalemi_id TEXT")


def finans_kaydi_yaz(hazir, onaylayan="sistem"):
    taslaklar = fis_taslaklari(hazir)
    if not taslaklar:
        return {"basari": True, "fis_sayisi": 0, "fisler": []}
    with _baglanti() as con:
        finansal_ham = json.dumps(
            hazir.get("veri", {}).get("kayitlar", []), ensure_ascii=False,
            sort_keys=True, separators=(",", ":"), default=str,
        )
        finansal_hash = hashlib.sha256(finansal_ham.encode("utf-8")).hexdigest()
        belge_hash = str(hazir.get("belge_hash") or "").strip().lower()
        if not belge_hash:
            for kayit in hazir.get("veri", {}).get("kayitlar", []):
                aday = str(kayit.get("_hash") or "").strip().lower()
                if len(aday) == 64:
                    belge_hash = aday
                    break
        if len(belge_hash) != 64:
            belge_hash = finansal_hash
        mevcut_manifest = con.execute(
            "SELECT sha256,finansal_hash FROM belge_butunlugu WHERE belge_id=?",
            (str(hazir.get("belge_id")),),
        ).fetchone()
        if mevcut_manifest and (
            mevcut_manifest["sha256"] != belge_hash
            or mevcut_manifest["finansal_hash"] != finansal_hash
        ):
            raise ValueError("Belge kimliği daha önce farklı içerikle kaydedilmiş.")
        for fis in taslaklar:
            mevcut = con.execute("SELECT fis_id FROM fisler WHERE belge_id=? AND kayit_sira=?", (fis["belge_id"], fis["sira"])).fetchone()
            if mevcut:
                continue
            con.execute(
                "INSERT INTO fisler(fis_id,belge_id,kayit_sira,faaliyet_key,tarih,vade,proje,cari,aciklama,para_birimi,kur_mikro,kaynak,olusturan,onaylayan,olusturma_zamani) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (fis["fis_id"], fis["belge_id"], fis["sira"], fis["faaliyet"], fis["tarih"], fis["vade"],
                 fis["proje"], fis["cari"], fis["aciklama"], fis["pb"], fis["kur_mikro"], fis["kaynak"],
                 onaylayan, onaylayan, _simdi()),
            )
            con.executemany(
                "INSERT INTO fis_satirlari(fis_id,satir_no,hesap_kodu,borc_kurus,alacak_kurus,aciklama) VALUES(?,?,?,?,?,?)",
                [(fis["fis_id"], no, s["hesap"], s["borc"], s["alacak"], s["aciklama"]) for no, s in enumerate(fis["satirlar"], 1)],
            )
            if fis.get("is_kalemi_id"):
                _is_kalemi_kolonu_hazirla(con)
                con.execute("UPDATE fisler SET is_kalemi_id=? WHERE fis_id=?", (fis["is_kalemi_id"], fis["fis_id"]))
        con.execute(
            "INSERT OR IGNORE INTO belge_butunlugu(belge_id,sha256,finansal_hash,kaynak,kayit_zamani,dogrulama_durumu) "
            "VALUES(?,?,?,?,?,?)",
            (str(hazir.get("belge_id")), belge_hash, finansal_hash,
             str(hazir.get("foto_yolu") or hazir.get("kaynak_metin") or ""), _simdi(),
             str((hazir.get("dogrulama") or {}).get("durum") or "KAYITLI")),
        )
        ids = [x["fis_id"] for x in taslaklar]
    return {"basari": True, "fis_sayisi": len(ids), "fisler": ids}


def finans_kaydi_kuyruga_al(hazir, onaylayan="sistem", hata=""):
    """Excel yazılıp DB geçici olarak yazılamazsa kaydı kaybetmeden saklar."""
    SYNC_QUEUE_PATH.parent.mkdir(parents=True, exist_ok=True)
    kayit = {"hazir": hazir, "onaylayan": str(onaylayan), "hata": str(hata), "zaman": _simdi()}
    with SYNC_QUEUE_PATH.open("a", encoding="utf-8") as f:
        f.write(json.dumps(kayit, ensure_ascii=False, default=str) + "\n")


def senkron_kuyrugunu_isle():
    if not SYNC_QUEUE_PATH.exists():
        return {"islenen": 0, "kalan": 0}
    satirlar = [x for x in SYNC_QUEUE_PATH.read_text(encoding="utf-8").splitlines() if x.strip()]
    kalan, islenen = [], 0
    for satir in satirlar:
        try:
            kayit = json.loads(satir)
            finans_kaydi_yaz(kayit["hazir"], kayit.get("onaylayan", "senkron"))
            islenen += 1
        except Exception:
            kalan.append(satir)
    gecici = SYNC_QUEUE_PATH.with_suffix(".tmp")
    if kalan:
        gecici.write_text("\n".join(kalan) + "\n", encoding="utf-8")
        os.replace(gecici, SYNC_QUEUE_PATH)
    else:
        gecici.unlink(missing_ok=True)
        SYNC_QUEUE_PATH.unlink(missing_ok=True)
    return {"islenen": islenen, "kalan": len(kalan)}


def belgeyi_ters_kaydet(belge_id, aktor="sistem"):
    veritabanini_hazirla()
    with _baglanti() as con:
        fisler = con.execute("SELECT * FROM fisler WHERE belge_id=? AND durum='MUHASEBELESTI'", (belge_id,)).fetchall()
        tersler = []
        for fis in fisler:
            ters_id = f"REV-{fis['fis_id']}"
            if con.execute("SELECT 1 FROM fisler WHERE fis_id=?", (ters_id,)).fetchone():
                continue
            gun = date.today()
            if not _donem_acik_mi(con, gun):
                raise ValueError("Cari dönem kapalı; ters kayıt oluşturulamadı.")
            con.execute(
                "INSERT INTO fisler(fis_id,belge_id,kayit_sira,faaliyet_key,tarih,vade,proje,cari,aciklama,para_birimi,kur_mikro,durum,ters_fis_id,kaynak,olusturan,onaylayan,olusturma_zamani) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (ters_id, f"REV-{belge_id}", fis["kayit_sira"], fis["faaliyet_key"], gun.isoformat(), None,
                 fis["proje"], fis["cari"], f"Ters kayıt: {fis['fis_id']}", fis["para_birimi"], fis["kur_mikro"],
                 "TERS", fis["fis_id"], fis["kaynak"], aktor, aktor, _simdi()),
            )
            satirlar = con.execute("SELECT * FROM fis_satirlari WHERE fis_id=? ORDER BY satir_no", (fis["fis_id"],)).fetchall()
            con.executemany(
                "INSERT INTO fis_satirlari VALUES(?,?,?,?,?,?)",
                [(ters_id, s["satir_no"], s["hesap_kodu"], s["alacak_kurus"], s["borc_kurus"], "Ters kayıt") for s in satirlar],
            )
            con.execute("UPDATE fisler SET durum='IPTAL', ters_fis_id=? WHERE fis_id=?", (ters_id, fis["fis_id"]))
            tersler.append(ters_id)
    return {"basari": True, "ters_fisler": tersler}


def donem_durum_degistir(yil, ay, kapat, aktor):
    yil, ay = int(yil), int(ay)
    if yil < 2000 or ay not in range(1, 13):
        raise ValueError("Dönem YYYY-AA biçiminde olmalı.")
    veritabanini_hazirla()
    durum = "KAPALI" if kapat else "ACIK"
    with _baglanti() as con:
        if kapat:
            fark = con.execute(
                "SELECT COALESCE(SUM(s.borc_kurus-s.alacak_kurus),0) fark FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id WHERE substr(f.tarih,1,7)=?",
                (f"{yil:04d}-{ay:02d}",),
            ).fetchone()["fark"]
            if fark != 0:
                raise ValueError(f"Dönem dengede değil: {fark/100:,.2f} TL")
        con.execute(
            "INSERT INTO donemler(yil,ay,durum,kapatan,kapanis_zamani) VALUES(?,?,?,?,?) "
            "ON CONFLICT(yil,ay) DO UPDATE SET durum=excluded.durum,kapatan=excluded.kapatan,kapanis_zamani=excluded.kapanis_zamani",
            (yil, ay, durum, str(aktor), _simdi()),
        )
    return durum


def acilis_bakiyesi_kaydet(faaliyet, hesap, tarih, tutar, aktor):
    """Bir hesap için özkaynak karşılıklı, dengeli açılış fişi oluşturur."""
    veritabanini_hazirla()
    gun = _tarih(tarih)
    kurus = _kurus(tutar)
    if kurus == 0:
        raise ValueError("Açılış bakiyesi sıfır olamaz.")
    with _baglanti() as con:
        faaliyet_var = con.execute(
            "SELECT 1 FROM faaliyetler WHERE faaliyet_key=? AND aktif=1", (faaliyet,)
        ).fetchone()
        hesap_satiri = con.execute(
            "SELECT hesap_tipi FROM hesaplar WHERE kod=? AND aktif=1", (str(hesap),)
        ).fetchone()
        if not faaliyet_var:
            raise ValueError("Faaliyet bulunamadı.")
        if not hesap_satiri:
            raise ValueError("Hesap kodu bulunamadı.")
        if str(hesap) == "500":
            raise ValueError("Açılış hesabı 500 Sermaye olamaz.")
        if not _donem_acik_mi(con, gun):
            raise ValueError(f"{gun:%Y-%m} dönemi kapalı; kayıt yapılamaz.")
        belge_id = f"OPEN-{faaliyet}-{hesap}-{gun:%Y%m%d}"
        fis_id = f"JRN-{belge_id}-001"
        if con.execute("SELECT 1 FROM fisler WHERE fis_id=?", (fis_id,)).fetchone():
            raise ValueError("Bu hesap ve tarih için açılış bakiyesi zaten var.")
        normal_borc = hesap_satiri["hesap_tipi"] in {"VARLIK", "GIDER"}
        if kurus < 0:
            normal_borc = not normal_borc
        tutar_kurus = abs(kurus)
        con.execute(
            "INSERT INTO fisler(fis_id,belge_id,kayit_sira,faaliyet_key,tarih,proje,cari,aciklama,para_birimi,kur_mikro,kaynak,olusturan,onaylayan,olusturma_zamani) "
            "VALUES(?,?,?,?,?,'','','Açılış bakiyesi','TL',1000000,'telegram',?,?,?)",
            (fis_id, belge_id, 1, faaliyet, gun.isoformat(), str(aktor), str(aktor), _simdi()),
        )
        satirlar = (
            [(fis_id, 1, str(hesap), tutar_kurus, 0, "Açılış bakiyesi"),
             (fis_id, 2, "500", 0, tutar_kurus, "Açılış karşı hesabı")]
            if normal_borc else
            [(fis_id, 1, "500", tutar_kurus, 0, "Açılış karşı hesabı"),
             (fis_id, 2, str(hesap), 0, tutar_kurus, "Açılış bakiyesi")]
        )
        con.executemany("INSERT INTO fis_satirlari VALUES(?,?,?,?,?,?)", satirlar)
    return fis_id


def butce_kaydet(faaliyet, proje, hesap, yil, ay, tutar, aktor):
    veritabanini_hazirla()
    kurus = _kurus(tutar)
    with _baglanti() as con:
        if not con.execute("SELECT 1 FROM faaliyetler WHERE faaliyet_key=?", (faaliyet,)).fetchone():
            raise ValueError("Faaliyet bulunamadı.")
        if not con.execute("SELECT 1 FROM hesaplar WHERE kod=?", (hesap,)).fetchone():
            raise ValueError("Hesap kodu bulunamadı.")
        con.execute(
            "INSERT INTO butceler VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(faaliyet_key,proje,hesap_kodu,yil,ay) "
            "DO UPDATE SET tutar_kurus=excluded.tutar_kurus,guncelleyen=excluded.guncelleyen,guncelleme=excluded.guncelleme",
            (faaliyet, proje or "", hesap, int(yil), int(ay), kurus, str(aktor), _simdi()),
        )


def yonetim_ozeti(faaliyet=None, yil=None, ay=None):
    """Geriye donuk uyumluluk sarmalayicisi: tek finans motoruna delege eder."""
    from finans_motoru import bakiye_ozeti

    o = bakiye_ozeti(faaliyet, None, yil, ay)
    return {"gelir": o["ciro_kurus"] / 100, "gider": o["maliyet_kurus"] / 100,
            "net": (o["ciro_kurus"] - o["maliyet_kurus"]) / 100,
            "nakit": o["nakit_kurus"] / 100, "alacak": o["alacak_kurus"] / 100,
            "borc": o["borc_kurus"] / 100}


def mizan(baslangic, bitis, faaliyet=None):
    """Hesap planı bazında borç, alacak ve bakiye döndürür."""
    veritabanini_hazirla()
    bas, son = _tarih(baslangic), _tarih(bitis)
    where, args = ["f.tarih BETWEEN ? AND ?"], [bas.isoformat(), son.isoformat()]
    if faaliyet:
        where.append("f.faaliyet_key=?")
        args.append(faaliyet)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT h.kod,h.ad,h.hesap_tipi,COALESCE(SUM(s.borc_kurus),0) borc,"
            "COALESCE(SUM(s.alacak_kurus),0) alacak FROM hesaplar h "
            "LEFT JOIN fis_satirlari s ON s.hesap_kodu=h.kod "
            "LEFT JOIN fisler f ON f.fis_id=s.fis_id AND " + " AND ".join(where) + " "
            "GROUP BY h.kod,h.ad,h.hesap_tipi HAVING borc<>0 OR alacak<>0 ORDER BY h.kod",
            args,
        ).fetchall()
    return [{"hesap": r["kod"], "ad": r["ad"], "borc": r["borc"]/100,
             "alacak": r["alacak"]/100, "bakiye": (r["borc"]-r["alacak"])/100}
            for r in rows]


def nakit_tahmini(hafta=13, faaliyet=None, bugun=None):
    """Vadesi girilmiş açık alacak/borç fişlerinden haftalık plan üretir."""
    veritabanini_hazirla()
    hafta = int(hafta)
    if hafta < 1 or hafta > 52:
        raise ValueError("Hafta sayısı 1-52 arasında olmalı.")
    bas = _tarih(bugun) if bugun else date.today()
    son = bas + timedelta(days=hafta * 7)
    where, args = ["TRIM(COALESCE(f.vade,''))<>''"], []
    if faaliyet:
        where.append("f.faaliyet_key=?")
        args.append(faaliyet)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.vade,s.hesap_kodu,s.borc_kurus,s.alacak_kurus FROM fisler f "
            "JOIN fis_satirlari s ON s.fis_id=f.fis_id WHERE " + " AND ".join(where) +
            " AND s.hesap_kodu IN ('120','300','320','331','336')", args,
        ).fetchall()
    gruplar = {}
    for r in rows:
        try:
            vade = _tarih(r["vade"])
        except ValueError:
            continue
        if not bas <= vade <= son:
            continue
        hafta_basi = vade - timedelta(days=vade.weekday())
        veri = gruplar.setdefault(hafta_basi.isoformat(), {"giris": 0, "cikis": 0})
        if r["hesap_kodu"] == "120":
            veri["giris"] += max(0, r["borc_kurus"] - r["alacak_kurus"])
        else:
            veri["cikis"] += max(0, r["alacak_kurus"] - r["borc_kurus"])
    return [{"hafta": k, "giris": v["giris"]/100, "cikis": v["cikis"]/100,
             "net": (v["giris"]-v["cikis"])/100} for k, v in sorted(gruplar.items())]


def cari_bakiyeleri(hesap_kodu):
    veritabanini_hazirla()
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.cari, SUM(s.borc_kurus-s.alacak_kurus) bakiye FROM fisler f "
            "JOIN fis_satirlari s ON s.fis_id=f.fis_id WHERE s.hesap_kodu=? AND f.cari<>'' "
            "GROUP BY f.cari HAVING bakiye<>0 ORDER BY ABS(bakiye) DESC", (hesap_kodu,),
        ).fetchall()
    return [{"cari": r["cari"], "tutar": abs(r["bakiye"])/100} for r in rows]


def butce_raporu(faaliyet, yil, ay):
    veritabanini_hazirla()
    with _baglanti() as con:
        rows = con.execute(
            "SELECT b.hesap_kodu,h.ad,b.proje,b.tutar_kurus,COALESCE(a.gerceklesen,0) gerceklesen "
            "FROM butceler b JOIN hesaplar h ON h.kod=b.hesap_kodu LEFT JOIN ("
            "SELECT f.faaliyet_key,f.proje,s.hesap_kodu,SUM(s.borc_kurus-s.alacak_kurus) gerceklesen "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id WHERE substr(f.tarih,1,7)=? GROUP BY f.faaliyet_key,f.proje,s.hesap_kodu"
            ") a ON a.faaliyet_key=b.faaliyet_key AND a.proje=b.proje AND a.hesap_kodu=b.hesap_kodu "
            "WHERE b.faaliyet_key=? AND b.yil=? AND b.ay=? ORDER BY b.hesap_kodu,b.proje",
            (f"{int(yil):04d}-{int(ay):02d}", faaliyet, int(yil), int(ay)),
        ).fetchall()
    return [{"hesap": r["hesap_kodu"], "ad": r["ad"], "proje": r["proje"],
             "butce": r["tutar_kurus"]/100, "gerceklesen": r["gerceklesen"]/100,
             "fark": (r["tutar_kurus"]-r["gerceklesen"])/100} for r in rows]


def banka_csv_aktar(csv_yolu, hesap="Ana Banka"):
    veritabanini_hazirla()
    yol = Path(csv_yolu)
    dosya_hash = hashlib.sha256(yol.read_bytes()).hexdigest()
    eklenen = 0
    with yol.open("r", encoding="utf-8-sig", newline="") as f, _baglanti() as con:
        for sira, row in enumerate(csv.DictReader(f), 1):
            gun = _tarih(row.get("Tarih"))
            tutar = _kurus(row.get("Tutar"))
            referans = str(row.get("Referans") or "")
            kimlik = hashlib.sha256(f"{dosya_hash}|{sira}|{gun}|{tutar}|{referans}".encode()).hexdigest()
            con.execute("INSERT OR IGNORE INTO banka_hareketleri VALUES(?,?,?,?,?,?,?,'ESLESMEDI',NULL)",
                        (kimlik, hesap, gun.isoformat(), row.get("Aciklama") or "", tutar, referans, dosya_hash))
            eklenen += con.execute("SELECT changes()").fetchone()[0]
    otomatik_mutabakat()
    return eklenen


def otomatik_mutabakat():
    veritabanini_hazirla()
    with _baglanti() as con:
        hareketler = con.execute("SELECT * FROM banka_hareketleri WHERE durum='ESLESMEDI'").fetchall()
        eslesen = 0
        for h in hareketler:
            adaylar = con.execute(
                "SELECT DISTINCT f.fis_id,f.tarih FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
                "WHERE s.hesap_kodu='102' AND ABS(s.borc_kurus-s.alacak_kurus)=ABS(?) "
                "AND ABS(julianday(f.tarih)-julianday(?))<=3 AND f.fis_id NOT IN "
                "(SELECT fis_id FROM banka_hareketleri WHERE fis_id IS NOT NULL)",
                (h["tutar_kurus"], h["tarih"]),
            ).fetchall()
            if len(adaylar) == 1:
                con.execute("UPDATE banka_hareketleri SET durum='ESLESTI',fis_id=? WHERE hareket_id=?", (adaylar[0]["fis_id"], h["hareket_id"]))
                eslesen += 1
    return eslesen


def mutabakat_ozeti():
    veritabanini_hazirla()
    with _baglanti() as con:
        rows = con.execute("SELECT durum,COUNT(*) adet,SUM(ABS(tutar_kurus)) toplam FROM banka_hareketleri GROUP BY durum").fetchall()
    return {r["durum"]: {"adet": r["adet"], "tutar": (r["toplam"] or 0)/100} for r in rows}


def rol_ata(chat_id, rol, aktor):
    rol = str(rol).upper()
    if rol not in {"ADMIN", "FINANS", "ONAYCI", "IZLEYICI"}:
        raise ValueError("Rol: ADMIN, FINANS, ONAYCI veya IZLEYICI olmalı.")
    veritabanini_hazirla()
    with _baglanti() as con:
        con.execute("INSERT INTO kullanicilar VALUES(?,?,1,?,?) ON CONFLICT(chat_id) DO UPDATE SET rol=excluded.rol,aktif=1,ekleyen=excluded.ekleyen",
                    (int(chat_id), rol, str(aktor), _simdi()))


def kullanici_rolu(chat_id):
    veritabanini_hazirla()
    with _baglanti() as con:
        row = con.execute("SELECT rol FROM kullanicilar WHERE chat_id=? AND aktif=1", (int(chat_id),)).fetchone()
    return row["rol"] if row else None


def politika_ayarla(limit_tl, onay_sayisi, aktor):
    limit_kurus = _kurus(limit_tl)
    onay = int(onay_sayisi)
    if onay < 1 or onay > 5:
        raise ValueError("Onay sayısı 1-5 arasında olmalı.")
    veritabanini_hazirla()
    with _baglanti() as con:
        con.execute("UPDATE onay_politikasi SET limit_kurus=?,yuksek_tutar_onay=?,guncelleyen=?,guncelleme=? WHERE politika_id=1",
                    (limit_kurus, onay, str(aktor), _simdi()))
        con.execute(
            "INSERT INTO onay_kurallari(kural_id,faaliyet_key,min_kurus,max_kurus,gerekli_onay,gerekli_roller,ayri_hazirlayan,aktif,guncelleyen,guncelleme) "
            "VALUES('YUKSEK_TUTAR','*',?,NULL,?,'[\"ONAYCI\",\"FINANS\",\"ADMIN\"]',0,1,?,?) "
            "ON CONFLICT(kural_id) DO UPDATE SET min_kurus=excluded.min_kurus,gerekli_onay=excluded.gerekli_onay,"
            "guncelleyen=excluded.guncelleyen,guncelleme=excluded.guncelleme",
            (limit_kurus, onay, str(aktor), _simdi()),
        )


def onay_oyu_ekle(belge_id, chat_id, toplam_tl):
    veritabanini_hazirla()
    with _baglanti() as con:
        politika = con.execute("SELECT * FROM onay_politikasi WHERE politika_id=1").fetchone()
        gerekli = politika["yuksek_tutar_onay"] if _kurus(toplam_tl) >= politika["limit_kurus"] else 1
        con.execute("INSERT OR IGNORE INTO onay_oylari VALUES(?,?,?)", (belge_id, int(chat_id), _simdi()))
        mevcut = con.execute("SELECT COUNT(*) adet FROM onay_oylari WHERE belge_id=?", (belge_id,)).fetchone()["adet"]
    return {"gerekli": gerekli, "mevcut": mevcut, "tamam": mevcut >= gerekli}


def muhasebe_csv_export(baslangic, bitis, hedef=None):
    veritabanini_hazirla()
    bas, son = _tarih(baslangic), _tarih(bitis)
    hedef = Path(hedef) if hedef else Path(MUHASEBE_EXPORT_DIR) / f"muhasebe_{bas}_{son}.csv"
    hedef.parent.mkdir(parents=True, exist_ok=True)
    with _baglanti() as con:
        rows = con.execute(
            "SELECT f.tarih,f.fis_id,f.belge_id,f.faaliyet_key,f.proje,f.cari,s.hesap_kodu,h.ad,s.borc_kurus,s.alacak_kurus,f.aciklama "
            "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id JOIN hesaplar h ON h.kod=s.hesap_kodu "
            "WHERE f.tarih BETWEEN ? AND ? ORDER BY f.tarih,f.fis_id,s.satir_no", (bas.isoformat(), son.isoformat()),
        ).fetchall()
    with hedef.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Tarih","Fiş No","Belge No","Faaliyet","Proje","Cari","Hesap Kodu","Hesap Adı","Borç","Alacak","Açıklama"])
        for r in rows:
            w.writerow([r["tarih"],r["fis_id"],r["belge_id"],r["faaliyet_key"],r["proje"],r["cari"],r["hesap_kodu"],r["ad"],r["borc_kurus"]/100,r["alacak_kurus"]/100,r["aciklama"]])
    return hedef


def veritabani_yedekle(saklama_adedi=30):
    """SQLite'ın çevrimiçi backup API'siyle tutarlı bir finans DB yedeği alır."""
    veritabanini_hazirla()
    klasor = Path(ARSIV) / "veri_yedek"
    klasor.mkdir(parents=True, exist_ok=True)
    hedef = klasor / f"finans_{datetime.now():%Y%m%d_%H%M%S_%f}.db"
    kaynak = sqlite3.connect(DB_PATH, timeout=30)
    yedek = sqlite3.connect(hedef)
    try:
        kaynak.backup(yedek)
    finally:
        yedek.close()
        kaynak.close()
    dosyalar = sorted(klasor.glob("finans_*.db"), key=lambda p: p.stat().st_mtime, reverse=True)
    for eski in dosyalar[max(1, int(saklama_adedi)):]:
        eski.unlink(missing_ok=True)
    return hedef


def sistem_kontrolu():
    veritabanini_hazirla()
    with _baglanti() as con:
        dengesiz = con.execute(
            "SELECT COUNT(*) adet FROM (SELECT fis_id,SUM(borc_kurus) b,SUM(alacak_kurus) a FROM fis_satirlari GROUP BY fis_id HAVING b<>a)"
        ).fetchone()["adet"]
        sayilar = {t: con.execute(f"SELECT COUNT(*) adet FROM {t}").fetchone()["adet"] for t in ("fisler","fis_satirlari","butceler","banka_hareketleri")}
        acik = con.execute("SELECT COUNT(*) adet FROM donemler WHERE durum='ACIK'").fetchone()["adet"]
        kapali = con.execute("SELECT COUNT(*) adet FROM donemler WHERE durum='KAPALI'").fetchone()["adet"]
    kuyruk = 0
    if SYNC_QUEUE_PATH.exists():
        kuyruk = sum(1 for x in SYNC_QUEUE_PATH.read_text(encoding="utf-8").splitlines() if x.strip())
    return {"dengesiz_fis": dengesiz, "senkron_bekleyen": kuyruk,
            "acik_donem": acik, "kapali_donem": kapali, **sayilar}


def holding_excelini_aktar():
    """Merkezi deftere yansimamis Holding Excel satirlarini deftere aktarir.

    Excel'e DOKUNMAZ ve idempotenttir (``LEGACY-<Kayit_ID>`` belge kimligi, UNIQUE(belge_id, kayit_sira)).
    Uygulamanin kendi yazdigi satirlar (Kaynak_Belge_ID dolu) deftere kendi belge kimligiyle zaten
    islendigi icin ATLANIR; aksi halde ayni tutar iki kez sayilirdi.
    """
    from openpyxl import load_workbook

    firma = FIRMALAR["touche_holding"]
    ledger = firma["ana_ledger"]
    s = ledger["sutunlar"]
    idx = {ad: ord(h) - ord("A") for ad, h in s.items() if len(h) == 1}
    son_sutun = max(idx.values()) + 1
    aktarilan, atlanan, hatalar = 0, 0, []
    wb = load_workbook(ana_excel_yolu(), data_only=False, read_only=True)
    try:
        ws = wb[ledger["sayfa"]]
        for v in ws.iter_rows(min_row=ledger["ilk_veri_satiri"], max_col=son_sutun, values_only=True):
            def h(ad, v=v):
                i = idx.get(ad)
                return v[i] if i is not None and i < len(v) else None

            kayit_id = str(h("Kayit_ID") or "").strip()
            tutar = h("Tutar")
            if not kayit_id or not tutar or _norm(h("Durum")) != "aktif" or str(h("Kaynak_Belge_ID") or "").strip():
                atlanan += 1
                continue
            faaliyet_adi = h("Faaliyet")
            faaliyet = faaliyet_anahtari_bul(faaliyet_adi)
            kayit = {
                "Tarih": h("Tarih"), "Faaliyet": faaliyet_adi, "Proje": h("Proje"),
                "Cari_Tipi": h("Cari_Tipi"), "Cari": h("Cari"), "Islem_Turu": h("Islem_Turu"),
                "Hesap_Kasa": h("Hesap_Kasa"), "Tutar": tutar, "PB": h("PB") or "TL", "Kur": h("Kur") or 1,
                "Vade": h("Vade"), "Aciklama": h("Aciklama"),
                "_dosya_adi": h("Belge_Referans") or "Mevcut Holding Excel",
            }
            hazir = {"belge_id": f"LEGACY-{kayit_id}", "firma_key": faaliyet, "tur": "finans_belgesi",
                     "veri": {"kayitlar": [kayit]}}
            try:
                once = sistem_kontrolu()["fisler"]
                finans_kaydi_yaz(hazir, "excel_migrasyonu")
                aktarilan += max(0, sistem_kontrolu()["fisler"] - once)
            except Exception as exc:
                hatalar.append(f"{kayit_id}: {exc}")
    finally:
        wb.close()
    return {"aktarilan": aktarilan, "atlanan": atlanan, "hatalar": hatalar}


def holding_faaliyetlerini_uzlastir(aktor="sistem"):
    """Holding Excel'indeki faaliyet adlarini mevcut LEGACY fislerle uzlastirir.

    Yalnizca kayit kimligiyle birebir eslesen fislerin faaliyet anahtarini
    gunceller; tutar, hesap, tarih ve Excel dosyalari degismez.
    """
    from openpyxl import load_workbook

    veritabanini_hazirla()
    firma = FIRMALAR["touche_holding"]
    ledger = firma["ana_ledger"]
    wb = load_workbook(ana_excel_yolu(), data_only=False, read_only=True)
    ws = wb[ledger["sayfa"]]
    s = ledger["sutunlar"]
    eslemeler = {}
    try:
        for row in range(ledger["ilk_veri_satiri"], (ws.max_row or 0) + 1):
            kayit_id = str(ws[f"{s['Kayit_ID']}{row}"].value or "").strip()
            durum = _norm(ws[f"{s['Durum']}{row}"].value)
            if not kayit_id or durum != "aktif":
                continue
            faaliyet_adi = ws[f"{s['Faaliyet']}{row}"].value
            eslemeler[f"LEGACY-{kayit_id}"] = faaliyet_anahtari_bul(faaliyet_adi)
    finally:
        wb.close()

    degisecek = []
    with _baglanti() as con:
        for belge_id, hedef in eslemeler.items():
            rows = con.execute(
                "SELECT fis_id,faaliyet_key FROM fisler WHERE belge_id=? AND faaliyet_key<>?",
                (belge_id, hedef),
            ).fetchall()
            degisecek.extend((r["fis_id"], r["faaliyet_key"], hedef, belge_id) for r in rows)
    if not degisecek:
        return {"guncellenen": 0, "yedek": None}

    yedek = veritabani_yedekle()
    with _baglanti() as con:
        con.executemany(
            "UPDATE fisler SET faaliyet_key=? WHERE fis_id=? AND faaliyet_key=?",
            [(hedef, fis_id, eski) for fis_id, eski, hedef, _ in degisecek],
        )
    audit = Path(LOGLAR) / "faaliyet_uzlastirma.jsonl"
    with audit.open("a", encoding="utf-8") as f:
        f.write(json.dumps({
            "zaman": _simdi(), "aktor": str(aktor), "yedek": str(yedek),
            "guncellenen": len(degisecek),
            "fisler": [{"fis_id": x[0], "eski": x[1], "yeni": x[2], "belge_id": x[3]} for x in degisecek],
        }, ensure_ascii=False) + "\n")
    return {"guncellenen": len(degisecek), "yedek": str(yedek)}
