# -*- coding: utf-8 -*-
"""Smart Excel operasyon araci (yalniz bilincli insan/operator kullanimi icin).

Komutlar (Windows'ta ``py -3.12 scripts\\excel_operasyon.py <komut>``):

  durum                    ANA ve YEDEK dosyalarinin yolu, boyutu, hash'i, kalinti gecici dosyalar
  dogrula                  ANA ve YEDEK XLSX butunlugunu ve sayfalarini denetler
  envanter [--kok YOL ...] Makinede Touche Smart Excel'e benzeyen dosyalari tarar,
                           ANA ile karsilastirir ve EXCEL_FILE_INVENTORY.md uretir (HICBIR SEY SILMEZ)
  temizle  [--kok YOL ...] Envanterde "artik uretilmis kopya" ve "ANA'da bulunmayan kayit yok"
                           olarak kanitlanan dosyalari listeler. Silmek icin --sil ve
                           onay kelimesi (--onay SIL) gerekir.
  mutabakat                ANA Excel ile merkezi defteri karsilastirir (SALT-OKUNUR): Excel'de elle girilmis /
                           degistirilmis satirlar deftere yansimiyor mu? Hicbir sey degistirmez.
  defteri-aktar            Deftere yansimamis Excel satirlarini listeler; --onay AKTAR ile merkezi SQLite
                           deftere aktarir (Excel'e DOKUNMAZ, idempotent).
  geri-don                 ANA ile YEDEK'i takas eder (acil durum). --onay GERI-DON ister.
  entegre-et DOSYA         Onaylanmis, disarida duzenlenmis bir XLSX'i yeni ANA yapar (YEDEK=onceki ANA).

Bu arac yapay zeka katmanindan veya HTTP sunucusundan ASLA cagrilmaz.
"""
from __future__ import annotations

import argparse
import fnmatch
import os
import sys
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import config  # noqa: E402
import excel_deposu as depo  # noqa: E402

# Bu sistemin gecmiste urettigi ve artik gereksiz olan kopyalarin ad/klasor kaliplari.
URETILMIS_KOPYA_DESENLERI = (
    "firmalar/*/yedek/*.xlsx",            # eski tarihli yedekler
    "firmalar/kurulum_yedek/**/*.xlsx",   # eski kurulum yedekleri
    "arsiv/excel_yedek/*.xlsx",           # eski entegrasyon yedekleri
    "*_vade_[0-9]*.xlsx",                 # eski vade guncelleme yedekleri
    "Gursar_ONCE_HoldingOzeti_*.xlsx",    # eski ozet betigi yedekleri
    "*_hesap.xlsx",                       # eski faaliyet-basina kopyalar
    ".*.tmp.xlsx",                        # artik gecici dosyalar
    "*_once_*.xlsx",
)
ILGILI_AD_ANAHTARLARI = ("touche", "gursar", "gürsar", "smart excel", "ai ready", "finans yönetim", "finans yonetim")


def _hash(yol: Path) -> str:
    return depo.sha256(yol)


def _sayfalar(yol: Path):
    try:
        with zipfile.ZipFile(yol) as z:
            if "xl/workbook.xml" not in z.namelist():
                return None
        from openpyxl import load_workbook

        wb = load_workbook(yol, read_only=True, data_only=False)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return None


def _kayit_kimlikleri(yol: Path) -> set[str]:
    """03_Kayit Merkezi'ndeki Kayit_ID degerlerini dondurur (yoksa bos)."""
    from openpyxl import load_workbook

    ledger = config.FIRMALAR["touche_holding"]["ana_ledger"]
    try:
        wb = load_workbook(yol, read_only=True, data_only=True)
    except Exception:
        return set()
    try:
        if ledger["sayfa"] not in wb.sheetnames:
            return set()
        ws = wb[ledger["sayfa"]]
        sutun = ledger["sutunlar"]["Kayit_ID"]
        idx = ord(sutun) - ord("A")
        kimlikler = set()
        for satir in ws.iter_rows(min_row=ledger["ilk_veri_satiri"], values_only=True):
            if len(satir) > idx and satir[idx]:
                kimlikler.add(str(satir[idx]).strip())
        return kimlikler
    finally:
        wb.close()


def _ilgili_mi(yol: Path) -> bool:
    """Dosya Touche Smart Excel ile ilgili mi? Yalniz DOSYA ADINA veya bu sistemin kendi
    urettigi klasorlere (firmalar/, arsiv/) bakilir; ust klasor adi yeterli sayilmaz.
    """
    ad = yol.name.lower()
    if any(k in ad for k in ILGILI_AD_ANAHTARLARI):
        return True
    for uretilen in (config.BASE_DIR / "firmalar", config.BASE_DIR / "arsiv"):
        try:
            yol.resolve().relative_to(uretilen.resolve())
            return True
        except ValueError:
            continue
    return False


def _uretilmis_kopya_mi(yol: Path, kok: Path) -> bool:
    try:
        goreli = yol.relative_to(kok).as_posix()
    except ValueError:
        goreli = yol.name
    return any(fnmatch.fnmatch(goreli, d) or fnmatch.fnmatch(yol.name, d) for d in URETILMIS_KOPYA_DESENLERI)


def varsayilan_kokler() -> list[Path]:
    kokler = [config.BASE_DIR, config.ana_excel_yolu().parent, config.yedek_excel_yolu().parent]
    if config.LEGACY_GURSAR_EXCEL.parent.exists():
        kokler.append(config.LEGACY_GURSAR_EXCEL.parent)
    return list(dict.fromkeys(k for k in kokler if k.exists()))


def tara(kokler) -> list[dict]:
    ana, yedek = config.ana_excel_yolu().resolve(), config.yedek_excel_yolu().resolve()
    ana_idleri = _kayit_kimlikleri(ana) if ana.exists() else set()
    ana_hash = _hash(ana) if ana.exists() else None
    bulunanlar, gorulen = [], set()
    for kok in kokler:
        for dizin, alt, dosyalar in os.walk(kok):
            alt[:] = [a for a in alt if a not in {".git", "node_modules", "__pycache__", ".venv", "venv"}]
            for ad in dosyalar:
                yol = Path(dizin) / ad
                if yol.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
                    continue
                gercek = yol.resolve()
                if gercek in gorulen:
                    continue
                gorulen.add(gercek)
                if not _ilgili_mi(yol) and gercek not in (ana, yedek):
                    continue
                bilgi = {"yol": gercek, "boyut": yol.stat().st_size,
                         "degisim": datetime.fromtimestamp(yol.stat().st_mtime).isoformat(timespec="seconds"),
                         "sha256": _hash(yol), "sayfalar": _sayfalar(yol) if yol.suffix.lower() != ".xls" else None}
                if gercek == ana:
                    bilgi["rol"] = "MAIN"
                elif gercek == yedek:
                    bilgi["rol"] = "BACKUP"
                else:
                    ids = _kayit_kimlikleri(yol) if yol.suffix.lower() == ".xlsx" else set()
                    bilgi["tekil_kayit"] = sorted(ids - ana_idleri)
                    bilgi["ana_ile_ayni"] = bilgi["sha256"] == ana_hash
                    bilgi["uretilmis_kopya"] = _uretilmis_kopya_mi(yol, config.BASE_DIR)
                    if bilgi["ana_ile_ayni"]:
                        bilgi["rol"] = "KOPYA(ANA ile birebir ayni)"
                    elif bilgi["tekil_kayit"]:
                        bilgi["rol"] = "TEKIL_KAYIT_VAR(silinmez)"
                    elif bilgi["uretilmis_kopya"]:
                        bilgi["rol"] = "ESKI_URETILMIS_KOPYA"
                    elif yol.suffix.lower() == ".xls" or bilgi["sayfalar"] is None:
                        bilgi["rol"] = "OKUNAMADI/ESKI_FORMAT(silinmez)"
                    else:
                        bilgi["rol"] = "BELIRSIZ(silinmez)"
                bulunanlar.append(bilgi)
    return sorted(bulunanlar, key=lambda b: (b.get("rol") != "MAIN", b.get("rol") != "BACKUP", str(b["yol"])))


def silinebilir(b: dict) -> bool:
    """Yalniz kanitlanmis, bu sistemin urettigi artik kopyalar silinebilir."""
    return bool(b.get("rol") in {"ESKI_URETILMIS_KOPYA", "KOPYA(ANA ile birebir ayni)"}
                and b.get("uretilmis_kopya") and not b.get("tekil_kayit"))


def rapor_yaz(bulunanlar, hedef: Path):
    s = ["# EXCEL_FILE_INVENTORY", "",
         f"Üretim zamanı: {datetime.now().isoformat(timespec='seconds')}  ",
         f"MAIN: `{config.ana_excel_yolu()}`  ", f"BACKUP: `{config.yedek_excel_yolu()}`", "",
         "| # | Rol | Yol | Boyut | Değişim | SHA-256 | Sayfa | ANA'da olmayan kayıt | Silinebilir |",
         "|---|-----|-----|------:|---------|---------|------:|---------------------:|:-----------:|"]
    for i, b in enumerate(bulunanlar, 1):
        s.append(f"| {i} | {b.get('rol')} | `{b['yol']}` | {b['boyut']:,} | {b['degisim']} | "
                 f"`{b['sha256'][:16]}…` | {len(b['sayfalar']) if b['sayfalar'] else '-'} | "
                 f"{len(b.get('tekil_kayit', []))} | {'EVET' if silinebilir(b) else 'hayır'} |")
    sil = [b for b in bulunanlar if silinebilir(b)]
    s += ["", f"Toplam {len(bulunanlar)} ilgili dosya; silinebilir (kanıtlı) {len(sil)}.", "",
          "Kural: yalnız bu sistemin ürettiği (yedek/kurulum/tarihli/geçici) VE ANA'da bulunmayan kaydı "
          "OLMAYAN dosyalar silinebilir. Tekil kayıt içeren veya belirsiz dosyalar asla otomatik silinmez."]
    hedef.write_text("\n".join(s) + "\n", encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    alt = ap.add_subparsers(dest="komut", required=True)
    alt.add_parser("durum")
    alt.add_parser("dogrula")
    alt.add_parser("mutabakat")
    p = alt.add_parser("defteri-aktar")
    p.add_argument("--onay", default="")
    for ad in ("envanter", "temizle"):
        p = alt.add_parser(ad)
        p.add_argument("--kok", action="append", help="Taranacak klasör (birden fazla verilebilir)")
        p.add_argument("--rapor", default="EXCEL_FILE_INVENTORY.md")
        if ad == "temizle":
            p.add_argument("--sil", action="store_true")
            p.add_argument("--onay", default="")
    p = alt.add_parser("geri-don")
    p.add_argument("--neden", required=True)
    p.add_argument("--aktor", default=os.environ.get("USERNAME") or os.environ.get("USER") or "operator")
    p.add_argument("--onay", default="")
    p = alt.add_parser("entegre-et")
    p.add_argument("dosya")
    p.add_argument("--neden", default="")
    p.add_argument("--aktor", default=os.environ.get("USERNAME") or os.environ.get("USER") or "operator")
    a = ap.parse_args(argv)

    if a.komut == "durum":
        d = depo.durum()
        for rol in ("ana", "yedek"):
            x = d[rol]
            print(f"{rol.upper():6} {x['yol']}\n       var={x['var']} boyut={x['boyut']:,} degisim={x['degisim']}\n       sha256={x['sha256']}")
        print(f"ANA = YEDEK mi: {d['ana_yedek_ayni']}")
        print(f"Kalıntı geçici dosya: {len(d['gecici_dosyalar'])}")
        for g in d["gecici_dosyalar"]:
            print("  ", g)
        return 0
    if a.komut == "dogrula":
        kod = 0
        for ad, yol in (("ANA", config.ana_excel_yolu()), ("YEDEK", config.yedek_excel_yolu())):
            try:
                print(f"{ad}: OK sayfalar={depo.dogrula(yol)['sayfalar']}")
            except Exception as exc:
                print(f"{ad}: HATA {exc}")
                kod = 1
        return kod
    if a.komut == "mutabakat":
        import excel_defter_mutabakati as mut

        r = mut.kontrol()
        print(f"Excel aktif satır: {r['excel_aktif_satir']} · Defter fişi: {r['defter_fis']}")
        for ad, baslik in (("eksik_defterde", "Excel'de var, defterde YOK"), ("tutar_farki", "Tutar farkı")):
            print(f"\n{baslik}: {len(r[ad])}")
            for x in r[ad][:50]:
                print("  -", {k: v for k, v in x.items() if v not in (None, "")})
        print(f"\nDefterde fazla (Excel'de aktif karşılığı yok): {len(r['defterde_fazla'])}")
        for x in r["defterde_fazla"][:50]:
            print("  -", x)
        print("\nSonuç:", "TUTARLI" if r["tutarli"] else "FARK VAR (hiçbir şey değiştirilmedi)")
        return 0 if r["tutarli"] else 3
    if a.komut == "defteri-aktar":
        import excel_defter_mutabakati as mut
        import finans_cekirdegi as cekirdek

        eksik = mut.kontrol()["eksik_defterde"]
        print(f"Merkezi defterde karşılığı olmayan Excel kaydı: {len(eksik)}")
        for x in eksik[:50]:
            print("  -", x["kayit_id"], x["cari"], x["proje"], x["excel_tutar"], x["pb"])
        if not eksik:
            return 0
        if a.onay != "AKTAR":
            print("Bu kayıtları merkezi deftere aktarmak için: defteri-aktar --onay AKTAR (Excel'e dokunulmaz).")
            return 0
        sonuc = cekirdek.holding_excelini_aktar()
        print("Aktarım:", sonuc)
        return 0 if not sonuc["hatalar"] else 4
    if a.komut in ("envanter", "temizle"):
        kokler = [Path(k) for k in a.kok] if a.kok else varsayilan_kokler()
        bulunanlar = tara(kokler)
        rapor_yaz(bulunanlar, Path(a.rapor))
        print(f"{len(bulunanlar)} ilgili dosya bulundu; rapor: {a.rapor}")
        for b in bulunanlar:
            print(f"  [{b.get('rol')}] {b['yol']}")
        if a.komut == "envanter":
            return 0
        silinecek = [b for b in bulunanlar if silinebilir(b)]
        print(f"\nSilinebilir (kanıtlı artık kopya): {len(silinecek)}")
        for b in silinecek:
            print("  -", b["yol"])
        if not a.sil:
            print("Silmek için: temizle --sil --onay SIL")
            return 0
        if a.onay != "SIL":
            print("Onay kelimesi 'SIL' eksik; hiçbir şey silinmedi.")
            return 2
        for b in silinecek:
            b["yol"].unlink()
            print("silindi:", b["yol"])
        return 0
    if a.komut == "geri-don":
        sonuc = depo.yedekten_geri_don(aktor=a.aktor, neden=a.neden, onay=a.onay)
        print("Geri dönüş tamamlandı:", sonuc)
        return 0
    if a.komut == "entegre-et":
        sonuc = depo.harici_dosyayi_ana_yap(a.dosya, aktor=a.aktor, neden=a.neden)
        print("ANA güncellendi:", sonuc)
        return 0
    return 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except depo.DepoHatasi as exc:
        print(f"HATA [{exc.kod}]: {exc.mesaj}", file=sys.stderr)
        sys.exit(1)
