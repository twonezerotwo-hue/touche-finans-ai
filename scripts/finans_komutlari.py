# -*- coding: utf-8 -*-
"""CFO asistaninin izinli ve doğrulanabilir finans düzeltme araçları.

LLM bu modüle serbest veri yazamaz. Yalnızca deterministik olarak çözümlenen,
mevcut açık borçlarla eşleşen vade komutları merkezi defter ve tek master Excel
üzerinde birlikte uygulanır.
"""

from __future__ import annotations

import os
import shutil
import sqlite3
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from config import FINANS_DB, FIRMALAR
from finans_cekirdegi import _norm, faaliyet_anahtari_bul, veritabani_yedekle
from guvenlik import audit_yaz
from kayit_kilidi import kilitli


AYLAR = {
    "ocak": 1, "subat": 2, "mart": 3, "nisan": 4, "mayis": 5, "haziran": 6,
    "temmuz": 7, "agustos": 8, "eylul": 9, "ekim": 10, "kasim": 11, "aralik": 12,
}


def turkce_tarih_coz(metin: str, bugun: date | None = None) -> date | None:
    """13 Ekim, 13.10.2026 ve 2026-10-13 biçimlerini gerçek tarihe çevirir."""
    import re

    bugun = bugun or date.today()
    sade = _norm(metin)
    for desen, siralama in (
        (r"\b(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})\b", "ymd"),
        (r"\b(\d{1,2})[-/.](\d{1,2})[-/.](20\d{2})\b", "dmy"),
    ):
        eslesme = re.search(desen, sade)
        if eslesme:
            a, b, c = map(int, eslesme.groups())
            try:
                return date(a, b, c) if siralama == "ymd" else date(c, b, a)
            except ValueError:
                return None

    ay_ifadesi = "|".join(AYLAR)
    eslesme = re.search(rf"\b(\d{{1,2}})\s+({ay_ifadesi})(?:\s+(20\d{{2}}))?\b", sade)
    if not eslesme:
        return None
    gun, ay_adi, yil = eslesme.groups()
    yil = int(yil) if yil else bugun.year
    try:
        sonuc = date(yil, AYLAR[ay_adi], int(gun))
    except ValueError:
        return None
    if not eslesme.group(3) and sonuc < bugun:
        sonuc = date(yil + 1, sonuc.month, sonuc.day)
    return sonuc


def _hedefleri_hazirla(hedefler):
    sonuc = []
    gorulen = set()
    for hedef in hedefler or []:
        cari = str(hedef.get("cari") or "").strip()
        if not cari:
            continue
        faaliyet = str(hedef.get("faaliyet") or "").strip() or "touche_holding"
        proje = str(hedef.get("proje") or "").strip()
        anahtar = (faaliyet, _norm(proje), _norm(cari))
        if anahtar in gorulen:
            continue
        gorulen.add(anahtar)
        sonuc.append({"faaliyet": faaliyet, "proje": proje, "cari": cari})
    if not sonuc:
        raise ValueError("Vadesi güncellenecek açık borç bulunamadı.")
    return sonuc


def _hedefle_eslesir(faaliyet, proje, cari, aciklama, hedef):
    if faaliyet != hedef["faaliyet"]:
        return False
    if hedef["proje"] and _norm(proje) != _norm(hedef["proje"]):
        return False
    aranan = _norm(hedef["cari"])
    return aranan in _norm(f"{cari or ''} {aciklama or ''}")


def _db_eslesmeleri(con, hedefler):
    rows = con.execute(
        "SELECT DISTINCT f.fis_id,f.belge_id,f.faaliyet_key,f.proje,f.cari,f.aciklama "
        "FROM fisler f JOIN fis_satirlari s ON s.fis_id=f.fis_id "
        "WHERE f.durum='MUHASEBELESTI' AND s.hesap_kodu IN ('300','320','331','335') "
        "AND s.alacak_kurus>s.borc_kurus"
    ).fetchall()
    eslesmeler = []
    for hedef in hedefler:
        adaylar = [r for r in rows if _hedefle_eslesir(
            r["faaliyet_key"], r["proje"], r["cari"], r["aciklama"], hedef
        )]
        if not adaylar:
            raise ValueError(f"Açık borç tahakkuku bulunamadı: {hedef['cari']} ({hedef['proje'] or hedef['faaliyet']})")
        eslesmeler.extend(adaylar)
    return {r["fis_id"]: r for r in eslesmeler}


def _excel_eslesmeleri(wb, hedefler):
    master = FIRMALAR["touche_holding"]
    ledger = master["ana_ledger"]
    ws = wb[ledger["sayfa"]]
    s = ledger["sutunlar"]
    eslesmeler = {}
    for hedef in hedefler:
        adaylar = []
        for satir in range(ledger["ilk_veri_satiri"], (ws.max_row or 0) + 1):
            durum = _norm(ws[f"{s['Durum']}{satir}"].value)
            onay = _norm(ws[f"{s['Onay_Durumu']}{satir}"].value)
            islem = _norm(ws[f"{s['Islem_Turu']}{satir}"].value)
            if durum in {"iptal", "reddedildi"} or onay == "reddedildi" or "tahakkuk" not in islem:
                continue
            faaliyet = faaliyet_anahtari_bul(ws[f"{s['Faaliyet']}{satir}"].value)
            proje = ws[f"{s['Proje']}{satir}"].value
            cari = ws[f"{s['Cari']}{satir}"].value
            aciklama = ws[f"{s['Aciklama']}{satir}"].value
            if _hedefle_eslesir(faaliyet, proje, cari, aciklama, hedef):
                adaylar.append(satir)
        if not adaylar:
            raise ValueError(f"Ana Excel'de tahakkuk satırı bulunamadı: {hedef['cari']}")
        eslesmeler[(hedef["faaliyet"], _norm(hedef["proje"]), _norm(hedef["cari"]))] = adaylar
    return eslesmeler


@kilitli
def acik_borc_vadesi_guncelle(hedefler, yeni_vade, aktor="cfo_asistani"):
    """Eşleşen açık borç vadelerini DB ve tek master Excel'de atomik olarak günceller."""
    hedefler = _hedefleri_hazirla(hedefler)
    if isinstance(yeni_vade, datetime):
        yeni_vade = yeni_vade.date()
    if not isinstance(yeni_vade, date):
        raise ValueError("Geçerli bir vade tarihi gerekli.")

    excel_yolu = Path(FIRMALAR["touche_holding"]["excel"])
    if not excel_yolu.exists():
        raise FileNotFoundError(f"Master Excel bulunamadı: {excel_yolu}")

    db_yedek = veritabani_yedekle()
    yedek_klasoru = Path(FIRMALAR["touche_holding"]["yedek_klasoru"])
    yedek_klasoru.mkdir(parents=True, exist_ok=True)
    excel_yedek = yedek_klasoru / f"{excel_yolu.stem}_vade_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    shutil.copy2(excel_yolu, excel_yedek)
    gecici = excel_yolu.parent / f".{excel_yolu.stem}.{uuid.uuid4().hex}.tmp{excel_yolu.suffix}"
    wb = None
    con = None
    excel_degisti = False
    try:
        con = sqlite3.connect(FINANS_DB, timeout=30)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA foreign_keys=ON")
        con.execute("BEGIN IMMEDIATE")
        db_satirlari = _db_eslesmeleri(con, hedefler)

        wb = load_workbook(excel_yolu)
        excel_satirlari = _excel_eslesmeleri(wb, hedefler)
        vade_sutunu = FIRMALAR["touche_holding"]["ana_ledger"]["sutunlar"]["Vade"]
        for satirlar in excel_satirlari.values():
            for satir in satirlar:
                hucre = wb[f"03_Kayıt Merkezi"][f"{vade_sutunu}{satir}"]
                hucre.value = datetime.combine(yeni_vade, datetime.min.time())
                # Master dosyadaki diğer tarih hücreleriyle aynı Excel biçimi.
                hucre.number_format = r"dd\.mm\.yyyy"
        wb.save(gecici)
        wb.close()
        wb = None

        kontrol = load_workbook(gecici, read_only=False, data_only=False)
        try:
            ws = kontrol["03_Kayıt Merkezi"]
            for satirlar in excel_satirlari.values():
                for satir in satirlar:
                    deger = ws[f"{vade_sutunu}{satir}"].value
                    deger = deger.date() if isinstance(deger, datetime) else deger
                    if deger != yeni_vade:
                        raise RuntimeError(f"Geçici Excel vade doğrulaması başarısız: satır {satir}")
        finally:
            kontrol.close()

        for fis_id in db_satirlari:
            con.execute("UPDATE fisler SET vade=? WHERE fis_id=?", (yeni_vade.isoformat(), fis_id))
        sayi = con.execute(
            f"SELECT COUNT(*) adet FROM fisler WHERE fis_id IN ({','.join('?' for _ in db_satirlari)}) AND vade=?",
            [*db_satirlari, yeni_vade.isoformat()],
        ).fetchone()["adet"]
        if sayi != len(db_satirlari):
            raise RuntimeError("Merkezi defter vade doğrulaması başarısız.")

        os.replace(gecici, excel_yolu)
        excel_degisti = True
        con.commit()
        sonuc = {
            "basari": True,
            "vade": yeni_vade.isoformat(),
            "hedefler": hedefler,
            "fisler": sorted(db_satirlari),
            "excel_satirlari": sorted({r for satirlar in excel_satirlari.values() for r in satirlar}),
            "db_yedek": str(db_yedek),
            "excel_yedek": str(excel_yedek),
        }
        audit_yaz("asistan_vade_guncelleme", aktor=aktor, **sonuc)
        return sonuc
    except Exception as exc:
        if con is not None:
            con.rollback()
        if excel_degisti:
            shutil.copy2(excel_yedek, excel_yolu)
        audit_yaz(
            "asistan_vade_guncelleme_hatasi", aktor=aktor, vade=str(yeni_vade),
            hedefler=hedefler, hata=str(exc), db_yedek=str(db_yedek), excel_yedek=str(excel_yedek),
        )
        raise
    finally:
        if wb is not None:
            wb.close()
        if con is not None:
            con.close()
        gecici.unlink(missing_ok=True)
