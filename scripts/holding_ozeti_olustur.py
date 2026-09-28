# -*- coding: utf-8 -*-
import sys
sys.path.insert(0, r"C:\ai-excel\scripts")
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import FIRMALAR
from datetime import datetime
import shutil
from pathlib import Path

# Once yedek al
excel_yolu = Path(FIRMALAR["gursar_mimarlik"]["excel"])
yedek_klasoru = Path(FIRMALAR["gursar_mimarlik"]["yedek_klasoru"])
yedek_klasoru.mkdir(parents=True, exist_ok=True)
zaman = datetime.now().strftime("%Y%m%d_%H%M%S")
yedek = yedek_klasoru / ("Gursar_ONCE_HoldingOzeti_" + zaman + ".xlsx")
shutil.copy2(excel_yolu, yedek)
print("Yedek alindi: " + yedek.name)

# Excel'i ac
wb = load_workbook(excel_yolu)

if "Holding_Ozeti" in wb.sheetnames:
    ws = wb["Holding_Ozeti"]
    print("Holding_Ozeti sayfasi zaten var, guncelleniyor")
    for row in ws.iter_rows():
        for cell in row:
            cell.value = None
else:
    ws = wb.create_sheet("Holding_Ozeti")
    print("Holding_Ozeti sayfasi olusturuldu")

# Basliklar
baslik_font = Font(bold=True, color="FFFFFF", size=12)
baslik_fill = PatternFill("solid", fgColor="1F4E78")
baslik_align = Alignment(horizontal="center", vertical="center")

ws["A1"] = "GURSAR MIMARLIK - HOLDING OZETI"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")

# Baslik satiri
basliklar = ["Metrik", "Deger", "Kaynak", "Son Guncelleme"]
for i, baslik in enumerate(basliklar, start=1):
    c = ws.cell(row=3, column=i, value=baslik)
    c.font = baslik_font
    c.fill = baslik_fill
    c.alignment = baslik_align

# Metrikler (aciklama, formul, kaynak)
metrikler = [
    ("Toplam Ciro", "='Proje Ciro'!F14", "Proje Ciro!F14"),
    ("Usta Maliyeti", "='Proje Ciro'!I14", "Proje Ciro!I14"),
    ("Banka Masrafi", "=SUM('Ödeme Geçmişi'!H6:H100)", "Ödeme Geçmişi!H toplamı"),
    ("Ortak Cepten Gider", "='Ortak Mutabakat'!C8", "Ortak Mutabakat!C8"),
    ("Toplam Gider", "=B5+B6+B7", "Usta + Banka + Ortak"),
    ("Net Sonuc", "=B4-B8", "Ciro - Toplam Gider"),
    ("Likit Nakit", "='Alacak ve Nakit'!C9", "Alacak ve Nakit!C9"),
    ("Acik Alacak", "='Alacak ve Nakit'!C6", "Alacak ve Nakit!C6"),
    ("Cek Portfoyu", "='Alacak ve Nakit'!C7", "Alacak ve Nakit!C7"),
    ("Toplam Tahsil Edilecek", "=B11+B12", "Acik Alacak + Cek"),
    ("Operasyonel Borc (Usta)", "='Usta Ödemeleri'!G6", "Usta Ödemeleri!G6"),
    ("Kredi Borcu", "=SUM('Kredi Planı'!L16:L27)", "Kredi Planı kalan"),
    ("Ortak Cari Odemeleri", "='Ortak Cari'!J5", "Ortak Cari!J5"),
    ("Son Guncelleme Tarihi", "=NOW()", "Otomatik"),
]

for i, (metrik, formul, kaynak) in enumerate(metrikler, start=4):
    ws.cell(row=i, column=1, value=metrik).font = Font(bold=True)
    c = ws.cell(row=i, column=2, value=formul)
    c.number_format = "#,##0.00"
    ws.cell(row=i, column=3, value=kaynak)
    ws.cell(row=i, column=4, value="")

# Sutun genislikleri
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 30
ws.column_dimensions["D"].width = 22

# Alt bilgi
ws["A20"] = "NOT: Bu sayfa otomatik olusturulmustur. Degerler formullerle hesaplanir."
ws["A20"].font = Font(italic=True, color="808080")

wb.save(excel_yolu)
wb.close()

print()
print("=== Holding_Ozeti sayfasi hazir ===")
print("Toplam " + str(len(metrikler)) + " metrik eklendi (satir 4-17)")
print("Dosya: " + excel_yolu.name)
