# -*- coding: utf-8 -*-
import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from excel_yazici import _sistem_tanimlarini_yaz


class ExcelSistemTanimlariTestleri(unittest.TestCase):
    def test_cek_teminatli_borclanma_nakit_ve_borc_etkisini_tanimlar(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "99_Ayarlar"
        ws["T23"] = "Çek Teminatlı Borçlanma"
        ws["L17"] = "Gürsar İnşaat"
        ws["M17"] = "Gürsar TL Kasa"

        veri = {"kayitlar": [{
            "Faaliyet": "Gürsar İnşaat",
            "Islem_Turu": "Çek Teminatlı Borçlanma",
            "Hesap_Kasa": "Gürsar TL Kasa",
            "PB": "TL",
        }]}
        _sistem_tanimlarini_yaz(wb, veri)

        self.assertEqual([ws[f"{sutun}23"].value for sutun in "UVWXYZ"], [0, 0, 0, 1, 0, 1])
        self.assertEqual(ws["AA23"].value, 0)
        self.assertTrue(str(ws["K17"].value).startswith("HSP-SYS-"))
        self.assertEqual(ws["N17"].value, "TL")
        self.assertEqual(ws["O17"].value, 0)
        self.assertIn("VLOOKUP", ws["P17"].value)
        self.assertEqual(ws["Q17"].value, '=IF($K17="","",$O17*$P17)')


if __name__ == "__main__":
    unittest.main()
