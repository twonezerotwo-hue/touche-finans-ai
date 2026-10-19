"""Final release-readiness contracts on disposable synthetic environments.

These tests intentionally exercise the real writer, journal and recovery-vault
modules, but every source is created below pytest's temporary root.  No
production workbook, SQLite database, OneDrive folder or credential is read.
"""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

import excel_yazici
import finans_cekirdegi as cekirdek
import finans_islem_journali as journal
import pytest
import recovery_vault
import windows_dogrulama
import yetki
from cfo_test_temeli import CfoTestTemeli


def _hazir(belge_id: str, amount: int, *, vendor: str = "Sentetik Tedarikçi") -> dict:
    record = {
        "Tarih": "2026-10-12",
        "Faaliyet": "Gürsar İnşaat",
        "Proje": "Sentetik Proje",
        "Cari_Tipi": "Usta",
        "Cari": vendor,
        "Islem_Turu": "Usta/Tedarikçi Ödemesi",
        "Hesap_Kasa": "Gürsar TL Banka",
        "Tutar": amount,
        "PB": "TL",
        "Kur": 1,
        "Aciklama": "Final release readiness synthetic transaction",
    }
    return {
        "basari": True,
        "firma_key": "gursar_mimarlik",
        "tur": "finans_belgesi",
        "veri": {"kayitlar": [record]},
        "belge_id": belge_id,
        "belge_hash": hashlib.sha256(belge_id.encode()).hexdigest(),
        "dogrulama": {"durum": "PASS", "skor": 100, "hatalar": [], "uyarilar": []},
        "gecici_dosyalar": [],
    }


class SyntheticReleaseAcceptance(CfoTestTemeli):
    """Recovery, rollback and authorization evidence with synthetic data only."""

    def _post(self, belge_id: str, amount: int) -> dict:
        hazir = _hazir(belge_id, amount)
        result = excel_yazici.yaz(
            "gursar_mimarlik",
            {"_hedef": "ana_ledger", "kayitlar": hazir["veri"]["kayitlar"]},
            belge_id,
            aktor="final-release-test",
        )
        assert result["basari"] is True, result
        posted = cekirdek.finans_kaydi_yaz(hazir, "final-release-test")
        assert posted["basari"] is True, posted
        return hazir

    def test_synthetic_windows_acceptance_uses_only_temp_root(self):
        """The Windows acceptance adapter never scans a user/production path."""
        windows_dogrulama.SONUCLAR.clear()
        with tempfile.TemporaryDirectory(prefix="touche-release-windows-") as root:
            code = windows_dogrulama.main(["--klasor", root])
            assert code == 0, [x for x in windows_dogrulama.SONUCLAR if x[1] == "FAIL"]
            assert list(Path(root).iterdir()) == []
        lock_results = [status for name, status, _ in windows_dogrulama.SONUCLAR if "Excel acik" in name]
        assert lock_results == (["SKIP"] if os.name != "nt" else ["PASS", "PASS", "PASS", "PASS"])

    def test_full_vault_restore_preserves_latest_transaction_and_measures_rpo_rto(self):
        """A complete synthetic snapshot/restore keeps Excel, SQLite and audit evidence."""
        transactions = 10
        for index in range(transactions):
            self._post(f"REL-RPO-{index:03d}", 1_000 + index)
        journal_path = self.root / "journal.jsonl"
        journal_path.write_text(
            json.dumps({"event": "synthetic_commit", "document_id": "REL-RPO-009"}) + "\n",
            encoding="utf-8",
        )
        audit_before = self.audit.read_bytes()
        source_hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in (self.ana, self.root / "finans.db", self.audit, journal_path)}
        key = bytes(range(32))
        vault = self.root / "vault" / "release.vault"
        started = time.perf_counter()
        recovery_vault.create_snapshot(
            vault,
            excel_path=self.ana,
            sqlite_path=self.root / "finans.db",
            audit_path=self.audit,
            journal_path=journal_path,
            key=key,
            label="final-release-synthetic",
        )
        snapshot_seconds = time.perf_counter() - started
        restored = self.root / "restore-1"
        started = time.perf_counter()
        result = recovery_vault.restore_dry_run(vault, restored, key=key)
        restore_seconds = time.perf_counter() - started
        assert result["verified"] is True
        assert (restored / "excel.xlsx").exists()
        assert (restored / "finans.db").exists()
        assert (restored / "audit.jsonl").read_bytes() == audit_before
        assert (restored / "journal.jsonl").read_text(encoding="utf-8").startswith('{"event"')
        con = sqlite3.connect(restored / "finans.db")
        try:
            assert con.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
            assert con.execute("SELECT COUNT(*) FROM fisler").fetchone()[0] == transactions
            assert con.execute("SELECT COUNT(*) FROM fis_satirlari").fetchone()[0] == transactions * 2
            assert con.execute("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("REL-RPO-009",)).fetchone()[0] == 1
        finally:
            con.close()
        restored_excel = restored / "excel.xlsx"
        from openpyxl import load_workbook

        workbook = load_workbook(restored_excel, read_only=True, data_only=False)
        try:
            assert "03_Kayıt Merkezi" in workbook.sheetnames
        finally:
            workbook.close()
        second = recovery_vault.restore_dry_run(vault, self.root / "restore-2", key=key)
        assert second["verified"] is True
        after_hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in (self.ana, self.root / "finans.db", self.audit, journal_path)}
        assert after_hashes == source_hashes
        assert snapshot_seconds > 0 and restore_seconds > 0
        print(
            f"SYNTHETIC_RPO_RTO transactions={transactions} workbook_bytes={self.ana.stat().st_size} "
            f"sqlite_bytes={(self.root / 'finans.db').stat().st_size} "
            f"snapshot_seconds={snapshot_seconds:.4f} restore_seconds={restore_seconds:.4f} "
            "rpo_lost_transactions=0"
        )

    def test_vault_wrong_missing_damaged_and_non_overwrite_paths_fail_closed(self):
        key = b"k" * 32
        vault = self.root / "vault" / "safe.vault"
        self._post("REL-VAULT-001", 2_000)
        recovery_vault.create_snapshot(vault, excel_path=self.ana, sqlite_path=self.root / "finans.db", key=key)
        with pytest.raises(recovery_vault.VaultError, match="doğrulanamadı"):
            recovery_vault.restore_dry_run(vault, self.root / "wrong-key", key=b"x" * 32)
        damaged = bytearray(vault.read_bytes())
        damaged[-1] ^= 1
        vault.write_bytes(damaged)
        with pytest.raises(recovery_vault.VaultError, match="doğrulanamadı"):
            recovery_vault.restore_dry_run(vault, self.root / "damaged", key=key)
        with pytest.raises(recovery_vault.VaultError, match="anahtarı"):
            recovery_vault.create_snapshot(self.root / "vault" / "missing-key.vault", excel_path=self.ana, sqlite_path=self.root / "finans.db", key=None)
        missing_target = self.root / "existing"
        missing_target.mkdir()
        (missing_target / "do-not-overwrite.txt").write_text("keep", encoding="utf-8")
        with pytest.raises(recovery_vault.VaultError, match="boş"):
            recovery_vault.restore_dry_run(vault, missing_target, key=key)
        assert (missing_target / "do-not-overwrite.txt").read_text(encoding="utf-8") == "keep"

        # A failed atomic replace leaves no half-created vault behind.
        intact = self.root / "vault" / "intact.vault"
        with patch.object(recovery_vault.os, "replace", side_effect=OSError("synthetic replace failure")), pytest.raises(
            OSError, match="replace failure"
        ):
            recovery_vault.create_snapshot(intact, excel_path=self.ana, sqlite_path=self.root / "finans.db", key=key)
        assert not intact.exists()
        assert not list(intact.parent.glob(f".{intact.name}.*.tmp"))

    def test_manual_needs_review_requires_admin_reason_and_audit_without_repost(self):
        hazir = _hazir("REL-MANUAL-001", 3_000)

        def unavailable(*_args, **_kwargs):
            raise sqlite3.OperationalError("synthetic SQLite unavailable")

        result = journal.execute_transaction(
            hazir,
            "final-release-test",
            excel_writer=excel_yazici.yaz,
            db_writer=unavailable,
        )
        assert result["status"] == "PENDING_RECOVERY"
        tx_id = result["transaction_id"]
        with cekirdek._baglanti() as con:
            con.execute("UPDATE finans_islemleri SET state='NEEDS_REVIEW' WHERE transaction_id=?", (tx_id,))
        viewer = {"yetkili": True, "kimlik": "synthetic-viewer", "rol": "VIEWER", "faaliyetler": ["*"], "projeler": ["*"]}
        with pytest.raises(PermissionError):
            journal.authorize_needs_review(
                tx_id, "synthetic-viewer", "Yetkisiz rol ile kurtarma kararı", authorization=viewer
            )
        owner = {"yetkili": True, "kimlik": "synthetic-owner", "rol": "OWNER", "faaliyetler": ["*"], "projeler": ["*"]}
        approved = journal.authorize_needs_review(
            tx_id, "synthetic-owner", "Excel ve finans journal kanıtları doğrulandı", authorization=owner
        )
        assert approved["status"] == "COMMITTED"
        with cekirdek._baglanti() as con:
            assert con.execute("SELECT COUNT(*) FROM fisler WHERE belge_id=?", (hazir["belge_id"],)).fetchone()[0] == 1
        assert any(
            event.get("olay") == "transaction_admin_recovery_approved"
            and event.get("detail", {}).get("admin_identity") == "synthetic-owner"
            for event in self.audit_olaylari()
        )
        assert journal.recovery_health()["durum"] == "ok"


def test_release_rbac_matrix_is_fail_closed():
    expected_recovery = {"OWNER": True, "CFO": True, "FINANCE_MANAGER": False, "ACCOUNTANT": False, "PROJECT_MANAGER": False, "VIEWER": False}
    for role, allowed in expected_recovery.items():
        assert yetki.izin_var(role, "recovery.authorize") is allowed
        assert yetki.izin_var(role, "finans.goruntule") is True
    assert yetki.izin_var("VIEWER", "belge.onayla") is False
    assert yetki.izin_var("PROJECT_MANAGER", "finans.duzenle") is False
    assert yetki.izin_var("ACCOUNTANT", "belge.yukle") is True
    assert yetki.izin_var("OWNER", "audit.goruntule") is True
