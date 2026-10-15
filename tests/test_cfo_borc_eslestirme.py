# -*- coding: utf-8 -*-
"""CFO asistani: acik borc/alacak eslestirme, belirsizlikte soru sorma, mukerrer korumasi.

Fixture, uretimdeki Goktürk Okulu durumunu taklit eder: borc IS ADIYLA ("Taş işleri", "Saha betonu") acilmis,
odemeler KISI ADINA ("Sebahattin Çakmak", "Medet İme") yazilmis. Hicbir test uretim dosyasina dokunmaz.
"""
import hashlib
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

import cfo_eslestirme as ES
import cfo_sohbet
import finans_cekirdegi as cekirdek
import master_data
import proje_finansi as pf
import ui_belge_servisi as ui
from cfo_test_temeli import AKTOR, BUGUN, CfoTestTemeli

PROJE = "Göktürk Okulu"
FAALIYET = "gursar_mimarlik"


class AcikBorcTemeli(CfoTestTemeli):
    def setUp(self):
        super().setUp()
        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        for r, (tip, cari) in enumerate((("Usta", "Medet İme"), ("Usta", "Taş işleri"), ("Usta", "Saha betonu")), 17):
            ws[f"AG{r}"], ws[f"AH{r}"], ws[f"AI{r}"] = self.faaliyet_adi, tip, cari
        # uretim kitabindaki islem turu (alacak silme: gelir ve alacak azalir, kasa/banka hareketi yok)
        ws["T14"] = "Alacak Düzeltme / Mahsup"
        for sutun, deger in zip(("U", "V", "W", "X", "Y", "Z", "AA"), (-1, 0, -1, 0, 0, 0, 0)):
            ws[f"{sutun}14"] = deger
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()
        self._n = 0
        # --- Goktürk Okulu: borc is adiyla, odeme kisi adiyla ---
        self._fis("TAHAKKUK-TAS", "2026-09-15", "Gider Tahakkuku", 186_240, "Taş işleri", "Basamak + Bordür + Terazzo")
        self._fis("TAHAKKUK-SAHA", "2026-09-15", "Gider Tahakkuku", 140_000, "Saha betonu", "Saha betonu")
        self._fis("ODEME-1", "2026-09-16", "Usta/Tedarikçi Ödemesi", 100_000, "Sebahattin Çakmak", "Mermer / taş işleri")
        self._fis("ODEME-2", "2026-10-09", "Usta/Tedarikçi Ödemesi", 40_000, "Sebahattin Çakmak", "Taş işleri ödemesi")
        self._fis("ODEME-3", "2026-09-23", "Usta/Tedarikçi Ödemesi", 55_000, "Medet İme", "Saha betonu")
        self._fis("ODEME-4", "2026-10-09", "Usta/Tedarikçi Ödemesi", 70_000, "Medet İme",
                  "Saha betonu ödemesi - Medet İme FAST")
        self._fis("HAKEDIS-1", "2026-09-15", "Hakediş / Gelir Tahakkuku", 427_000, "İşveren", "hakediş")
        self._fis("TAHSILAT-1", "2026-09-20", "Tahsilat", 300_000, "İşveren", "tahsilat")

    def _fis(self, belge, tarih, islem, tutar, cari, aciklama="", proje=PROJE):
        kayit = {"Tarih": tarih, "Islem_Turu": islem, "Tutar": tutar, "PB": "TL", "Kur": 1, "Cari": cari,
                 "Proje": proje, "Hesap_Kasa": "Gürsar TL Banka", "Aciklama": aciklama}
        cekirdek.finans_kaydi_yaz({"belge_id": belge, "firma_key": FAALIYET, "tur": "finans_belgesi",
                                   "veri": {"kayitlar": [kayit]}}, "test")

    def fis_sayisi(self):
        with cekirdek._baglanti() as con:
            return con.execute("SELECT COUNT(*) FROM fisler").fetchone()[0]

    def kalemler(self):
        return {(a["tur"], a["cari"]): a for a in ES.eslestir(FAALIYET, PROJE, "x", 1, ES.BORC)["acik_kalemler"]
                } | {(a["tur"], a["cari"]): a for a in ES.eslestir(FAALIYET, PROJE, "x", 1, ES.ALACAK)["acik_kalemler"]}

    def cari_satiri(self, cari):
        d = pf.proje_detay(FAALIYET, PROJE, self.katalog())
        return next((x for x in d["cari_bakiyeleri"] if x["cari"] == cari), None)

    def son_fis(self):
        with cekirdek._baglanti() as con:
            return dict(con.execute("SELECT * FROM fisler WHERE belge_id LIKE '%-CHAT-%' ORDER BY olusturma_zamani DESC").fetchone())


class EslestirmeMotoruTestleri(AcikBorcTemeli):
    def test_mevcut_veri_raporu_beklenen_borc_odeme_ve_kalanlari_verir(self):
        r = ES.rapor(FAALIYET, PROJE, [("Basamak", "Taş işleri"), ("Saha Betonu", "Saha betonu")])
        b = {x["cari"]: x for x in r["borclar"]}
        self.assertEqual(b["Taş işleri"]["toplam_kurus"], 186_240 * 100)
        self.assertEqual(b["Taş işleri"]["atfedilen_kurus"], 140_000 * 100)          # 100.000 + 40.000
        self.assertEqual(b["Taş işleri"]["kalan_kurus"], 46_240 * 100)
        self.assertEqual(b["Saha betonu"]["toplam_kurus"], 140_000 * 100)
        self.assertEqual(b["Saha betonu"]["atfedilen_kurus"], 125_000 * 100)         # 55.000 + 70.000
        self.assertEqual(b["Saha betonu"]["kalan_kurus"], 15_000 * 100)
        self.assertEqual({a["fis_id"].split("-")[1] for a in r["yanlis_cari_odemeleri"]}, {"ODEME"})
        self.assertEqual(len(r["yanlis_cari_odemeleri"]), 4)
        eksi = {e["cari"]: e["acik_borc_kurus"] for e in r["eksi_bakiyeli_cariler"]}
        self.assertEqual(eksi["Sebahattin Çakmak"], -140_000 * 100)
        self.assertEqual(eksi["Medet İme"], -125_000 * 100)
        self.assertEqual({a["cari"]: a["kalan_kurus"] for a in r["acik_alacaklar"]}, {"İşveren": 127_000 * 100})
        self.assertIn("KALAN ₺46.240,00", ES.rapor_metni(r))

    def test_aciklamada_iki_borc_adi_gecen_odeme_hicbirine_atfedilmez(self):
        self._fis("BELIRSIZ-ODEME", "2026-10-10", "Usta/Tedarikçi Ödemesi", 1_000, "Yeni Usta",
                  "taş işleri ve saha betonu için karışık ödeme")
        b = {x["cari"]: x for x in ES.rapor(FAALIYET, PROJE)["borclar"]}
        self.assertEqual(b["Taş işleri"]["kalan_kurus"], 46_240 * 100)               # degismedi
        self.assertEqual(b["Saha betonu"]["kalan_kurus"], 15_000 * 100)

    def test_odeme_kendi_acik_borcu_olan_cariye_aciklamadan_baska_borca_atfedilmez(self):
        self._fis("USTA-TAHAKKUK", "2026-09-20", "Gider Tahakkuku", 50_000, "Ofis Usta", "kendi işi")
        self._fis("USTA-ODEME", "2026-10-01", "Usta/Tedarikçi Ödemesi", 10_000, "Ofis Usta", "saha betonu yanında")
        b = {x["cari"]: x for x in ES.rapor(FAALIYET, PROJE)["borclar"]}
        self.assertEqual(b["Ofis Usta"]["kalan_kurus"], 40_000 * 100)               # kendi borcunu kapatti
        self.assertEqual(b["Saha betonu"]["kalan_kurus"], 15_000 * 100)             # baska borca sayilmadi

    def test_rapor_salt_okunurdur(self):
        once = self.fis_sayisi(), hashlib.sha256(self.ana.read_bytes()).hexdigest()
        ES.rapor(FAALIYET, PROJE)
        ES.eslestir(FAALIYET, PROJE, "Medet İme", 1_000_00, ES.BORC, "saha betonu")
        self.assertEqual((self.fis_sayisi(), hashlib.sha256(self.ana.read_bytes()).hexdigest()), once)

    def test_kisi_gecmisi_tek_aday_zayif_kanit_onay_sorar(self):
        s = ES.eslestir(FAALIYET, PROJE, "Sebahattin Çakmak", 20_000 * 100, ES.BORC, "Sebahattin'e ödeme")
        self.assertEqual(s["karar"], "ONAY")
        self.assertEqual([a["cari"] for a in s["adaylar"]], ["Taş işleri"])
        self.assertEqual(s["adaylar"][0]["skor"], 75)

    def test_metin_ve_gecmis_birlikte_otomatik_eslesir(self):
        s = ES.eslestir(FAALIYET, PROJE, "Medet İme", 10_000 * 100, ES.BORC, "Medet saha betonu ödemesi")
        self.assertEqual(s["karar"], "ESLESTI")
        self.assertEqual(s["adaylar"][0]["cari"], "Saha betonu")
        self.assertEqual(s["adaylar"][0]["kalan_sonra_kurus"], 5_000 * 100)

    def test_tutar_kalani_asiyorsa_otomatik_degil_onay_sorulur(self):
        s = ES.eslestir(FAALIYET, PROJE, "Medet İme", 50_000 * 100, ES.BORC, "saha betonu")
        self.assertEqual(s["karar"], "ONAY")
        self.assertTrue(s["adaylar"][0]["asiyor"])

    def test_birden_cok_aday_belirsizdir(self):
        s = ES.eslestir(FAALIYET, PROJE, "Yeni Usta", 1_000 * 100, ES.BORC, "taş işleri ve saha betonu ödemesi")
        self.assertEqual(s["karar"], "BELIRSIZ")
        self.assertEqual({a["cari"] for a in s["adaylar"]}, {"Taş işleri", "Saha betonu"})

    def test_acik_kalem_var_ama_kanit_yoksa_sorulur_listelenir(self):
        s = ES.eslestir(FAALIYET, PROJE, "Yeni Usta", 1_000 * 100, ES.BORC, "ödeme yapıldı")
        self.assertEqual(s["karar"], "YOK")
        self.assertEqual({a["cari"] for a in s["acik_kalemler"]}, {"Taş işleri", "Saha betonu"})

    def test_kismi_ad_eslesmez_tum_jetonlar_gerekir(self):
        s = ES.eslestir(FAALIYET, PROJE, "Yeni Usta", 1_000 * 100, ES.BORC, "taş ödemesi")
        self.assertEqual(s["karar"], "YOK")

    def test_baska_projede_acik_kalem_yoksa_acik_yok(self):
        s = ES.eslestir(FAALIYET, "Bahçeşehir Villa", "Medet İme", 1_000 * 100, ES.BORC, "saha betonu")
        self.assertEqual(s["karar"], "ACIK_YOK")

    def test_dogrudan_cari_alacagi_tahsilatta_otomatik(self):
        s = ES.eslestir(FAALIYET, PROJE, "İşveren", 103_000 * 100, ES.ALACAK, "tahsilat")
        self.assertEqual(s["karar"], "ESLESTI")
        self.assertEqual(s["adaylar"][0]["kalan_kurus"], 127_000 * 100)


class CfoBorcKapamaAkisTestleri(AcikBorcTemeli):
    def _etiketli_fis_dogrula(self, cari, etiket_cari):
        fis = self.son_fis()
        self.assertEqual(fis["cari"], cari)                                   # odemeyi alan kisi cari olarak kalir
        self.assertTrue(fis["aciklama"].startswith(f"[Borç kapama: {etiket_cari}]"), fis["aciklama"])
        return fis

    def test_medet_ime_odemesi_saha_betonu_acik_borcunu_kapatir(self):
        y = self.baslat("Göktürk okulunda Medet İme'ye saha betonu için 10 bin TL ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap", y["cevap"])          # borc sorusu YOK: otomatik, kanitli
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIn("Bu kayıt yeni gider değil; mevcut borç kapama olarak işlendi.", y["cevap"])
        self.assertIn("Saha betonu / Göktürk Okulu / kalan 15.000 TL → 5.000 TL", y["cevap"])
        self.assertEqual(self.fis_sayisi(), 8)                                   # onaydan once deftere hicbir sey yazilmadi
        self.assertEqual(self.ledger(), [])
        self.onayla(y)
        self._etiketli_fis_dogrula("Medet İme", "Saha betonu")
        self.assertEqual(self.cari_satiri("Saha betonu")["acik_borc_kurus"], 130_000 * 100)   # 140.000 - 10.000
        self.assertEqual(self.kalemler()[(ES.BORC, "Saha betonu")]["kalan_kurus"], 5_000 * 100)  # 15.000 - 10.000
        self.assertEqual(self.cari_satiri("Medet İme")["odenen_kurus"], 125_000 * 100)        # eski kayit degismedi
        self.assertTrue(self.ledger()[0]["Aciklama"].startswith("[Borç kapama: Saha betonu]"))
        self.assertEqual(self.ledger()[0]["Cari"], "Medet İme")
        self.assertEqual(self.fis_sayisi(), 9)

    def test_sebahattin_cakmak_odemesi_tas_isleri_acik_borcunu_kapatir(self):
        y = self.baslat("Göktürk okulunda Sebahattin Çakmak'a 20 bin TL ödeme yapıldı")
        s = y["cfo"]["soru"]
        self.assertEqual(s["alan"], "borc", y["cevap"])                           # tek ama zayif kanit: sorar
        self.assertIn("Emin değilim", s["metin"])
        self.assertIn("Taş işleri / Göktürk Okulu / kalan 46.240 TL", s["metin"])
        self.assertEqual({o["id"] for o in s["secenekler"]}, {"kapat:Taş işleri", "yeni"})
        y = self.sec(y, "borc", "kapat:Taş işleri")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIn("Bu ödeme şu açık borcu kapatıyor: Taş işleri / Göktürk Okulu / kalan 46.240 TL → 26.240 TL", y["cevap"])
        self.onayla(y)
        self._etiketli_fis_dogrula("Sebahattin Çakmak", "Taş işleri")
        self.assertEqual(self.cari_satiri("Taş işleri")["acik_borc_kurus"], 166_240 * 100)
        self.assertEqual(self.kalemler()[(ES.BORC, "Taş işleri")]["kalan_kurus"], 26_240 * 100)

    def test_sebahattin_metinde_tas_isleri_gecerse_otomatik_kapatir(self):
        y = self.baslat("Göktürk okulunda Sebahattin Çakmak taş işleri için 20 bin TL ödeme yapıldı")
        if y["cfo"]["soru"]["alan"] == "cari":                                   # 'Taş işleri' de bir cari: ayrim sorulur
            y = self.sec(y, "cari", "Sebahattin Çakmak")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap", y["cevap"])
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIn("mevcut borç kapama olarak işlendi", y["cevap"])
        self.assertIn("Taş işleri / Göktürk Okulu", y["cevap"])

    def test_kullanici_yeni_kayit_derse_acik_borca_baglanmaz(self):
        y = self.baslat("Göktürk okulunda Sebahattin Çakmak'a 20 bin TL ödeme yapıldı")
        y = self.sec(y, "borc", "yeni")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIn("mevcut bir borca/alacağa bağlanmadı; yeni kayıt olarak hazırlandı", y["cevap"])
        self.onayla(y)
        self.assertFalse(self.son_fis()["aciklama"].startswith("[Borç kapama"))
        self.assertEqual(self.kalemler()[(ES.BORC, "Taş işleri")]["kalan_kurus"], 46_240 * 100)   # borc degismedi

    def test_hic_eslesme_yoksa_yeni_mi_mevcut_mu_diye_sorar_kayit_acmaz(self):
        y = self.baslat("Göktürk okulunda Devrim'e 7 bin TL ödeme yapıldı")      # Devrim'in gecmisi/metin kanit yok
        s = y["cfo"]["soru"]
        self.assertEqual(s["alan"], "borc", y["cevap"])
        self.assertIn("Hiç eşleşme bulamadım", s["metin"])
        self.assertIn("mevcut bir borç/alacak/hakediş kaydını mı kapatıyor, yoksa yeni bir kayıt mı açalım?", s["metin"])
        self.assertIn("yeni", {o["id"] for o in s["secenekler"]})
        self.assertEqual(ui._BEKLEYEN, {})
        self.assertEqual(self.fis_sayisi(), 8)

    def test_belirsiz_proje_kayit_acmaz_soru_sorar(self):
        y = self.baslat("Medet İme'ye 20 bin TL ödeme yapıldı")                   # ne faaliyet ne proje
        self.assertEqual(y["cfo"]["soru"]["alan"], "faaliyet")
        self.assertEqual(ui._BEKLEYEN, {})
        y = self.sec(y, "faaliyet", FAALIYET)
        self.assertEqual(y["cfo"]["soru"]["alan"], "proje")                       # proje tahmin edilmez
        self.assertEqual(ui._BEKLEYEN, {})
        self.assertEqual((self.fis_sayisi(), self.ledger()), (8, []))

    def test_belirsiz_cari_kayit_acmaz_soru_sorar(self):
        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        ws["AG20"], ws["AH20"], ws["AI20"] = self.faaliyet_adi, "Usta", "Devrim Yıldız"
        ws["AG21"], ws["AH21"], ws["AI21"] = self.faaliyet_adi, "Usta", "Devrim Kaya"
        ws["AG13"] = ws["AH13"] = ws["AI13"] = None
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()
        y = self.baslat("Gürsar Mimarlık Göktürk Okulu Devrim'e 5 bin TL ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "cari")
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]} - {"__YENI__"}, {"Devrim Yıldız", "Devrim Kaya"})
        self.assertEqual((ui._BEKLEYEN, self.fis_sayisi(), self.ledger()), ({}, 8, []))

    def test_devrim_odemesi_proje_belirtilmeden_yanlis_yere_yazilmaz(self):
        y = self.baslat("Devrim'e 10.000 TL elden ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "faaliyet")
        y = self.sec(y, "faaliyet", FAALIYET)
        self.assertEqual(y["cfo"]["soru"]["alan"], "proje")
        self.assertEqual((ui._BEKLEYEN, self.fis_sayisi(), self.ledger()), ({}, 8, []))

    def test_devrim_goktürk_okulu_belirtilirse_dogru_proje_ve_cariye_hazirlanir(self):
        y = self.baslat("Devrim'e Göktürk Okulu için 10.000 TL elden ödeme yapıldı")
        alanlar = y["cfo"]["alanlar"]
        self.assertEqual((alanlar["proje"], alanlar["cari"]), (PROJE, "Devrim"))
        self.assertNotEqual(y["cfo"]["soru"]["alan"], "faaliyet")

    def test_mukerrer_ihtimali_taslak_acmadan_once_sorulur_evet_ile_uyari_tasir(self):
        self._fis("DEVRIM-AYNI", BUGUN.isoformat(), "Usta/Tedarikçi Ödemesi", 10_000, "Devrim", "elden")
        y = self.baslat("Göktürk okulunda Devrim'e 10 bin TL ödeme yapıldı")
        for _ in range(4):                                                       # borc/hesap vb. sorulari gec
            alan = y["cfo"]["soru"]["alan"] if y["cfo"].get("soru") else None
            if alan in (None, "mukerrer_onay"):
                break
            y = self.sec(y, alan, "yeni" if alan == "borc" else "Gürsar TL Kasa")
        self.assertEqual(y["cfo"]["soru"]["alan"], "mukerrer_onay", y["cevap"])
        self.assertIn("Bu mükerrer olabilir", y["cfo"]["soru"]["metin"])
        self.assertEqual(ui._BEKLEYEN, {})                                       # henuz taslak/onay kuyrugu yok
        y = self.sec(y, "mukerrer_onay", "evet")
        self.assertIn("Dikkat: aynı gün, aynı cari ve aynı tutarla kayıt zaten var", y["cevap"])
        self.assertEqual(len(ui._BEKLEYEN), 1)

    def test_mukerrer_sorusuna_hayir_vazgecer_hicbir_sey_yazilmaz(self):
        self._fis("DEVRIM-AYNI", BUGUN.isoformat(), "Usta/Tedarikçi Ödemesi", 10_000, "Devrim", "elden")
        y = self.baslat("Göktürk okulunda Devrim'e 10 bin TL ödeme yapıldı")
        while y["cfo"].get("soru") and y["cfo"]["soru"]["alan"] != "mukerrer_onay":
            alan = y["cfo"]["soru"]["alan"]
            y = self.sec(y, alan, "yeni" if alan == "borc" else "Gürsar TL Kasa")
        y = self.sec(y, "mukerrer_onay", "hayir")
        self.assertIn("vazgeçildi", y["cevap"])
        self.assertEqual((ui._BEKLEYEN, self.ledger(), self.fis_sayisi()), ({}, [], 9))

    def test_yine_de_dendiginde_mukerrer_tekrar_sorulmaz(self):
        self._fis("DEVRIM-AYNI", BUGUN.isoformat(), "Usta/Tedarikçi Ödemesi", 10_000, "Devrim", "elden")
        y = self.baslat("Göktürk okulunda Devrim'e 10 bin TL ödeme yapıldı", yine_de=True)
        while y["cfo"].get("soru"):
            alan = y["cfo"]["soru"]["alan"]
            self.assertNotEqual(alan, "mukerrer_onay")
            y = self.sec(y, alan, "yeni" if alan == "borc" else "Gürsar TL Kasa")
        self.assertTrue(y["kayit_taslagi"])

    def test_tutar_degisince_baglanti_yeniden_hesaplanir(self):
        y = self.baslat("Göktürk okulunda Medet İme'ye saha betonu için 10 bin TL ödeme yapıldı")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        y = cfo_sohbet.duzelt(y["cfo"]["taslak_id"], "Düzelt: tutar 50 bin TL", AKTOR, self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "borc", y["cevap"])            # 50.000 > kalan 15.000: otomatik degil
        self.assertIn("tutar kalanı aşıyor", y["cfo"]["soru"]["metin"])


class BaglantiYenidenHesapTestleri(AcikBorcTemeli):
    def test_baglandiktan_sonra_tutar_yeniden_cevaplanirsa_baglanti_yeniden_hesaplanir(self):
        y = self.baslat("Göktürk okulunda Medet İme'ye saha betonu için 10 bin TL ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap")                       # baglanti otomatik kuruldu
        y = cfo_sohbet.cevapla(y["cfo"]["taslak_id"], "tutar", "50 bin TL", AKTOR, "metin", self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "borc", y["cevap"])            # 50.000 > kalan 15.000
        self.assertIn("tutar kalanı aşıyor", y["cfo"]["soru"]["metin"])


class OrtakCariVeEksikBilgiTestleri(AcikBorcTemeli):
    def _devrim_ortak_yap(self):
        wb = load_workbook(self.ana)
        wb["99_Ayarlar"]["AH13"] = "Ortak"                                   # uretimdeki gibi: Devrim = Ortak
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()

    def test_ortak_tipli_cariye_odeme_ortak_cari_odemesi_olur_borc_eslesmesi_aranmaz(self):
        self._devrim_ortak_yap()
        y = self.baslat("Göktürk okulunda Devrim'e 10 bin TL elden ödeme yapıldı")      # metinde 'ortak' kelimesi YOK
        soru = y["cfo"].get("soru")             # 'elden' ifadesi kasayi belirler -> hesap sorusu kalmayabilir
        self.assertNotEqual((soru or {}).get("alan"), "borc", y["cevap"])               # tedarikci borcu aranmaz
        y = self.sec(y, soru["alan"], "Gürsar TL Kasa") if soru and soru["alan"] == "hesap" else y
        self.assertEqual(y["cfo"]["alanlar"]["islem_turu"], "Ortak Cari Ödemesi", y["cevap"])
        self.onayla(y)
        with cekirdek._baglanti() as con:
            hesaplar = {r["hesap_kodu"] for r in con.execute(
                "SELECT s.hesap_kodu FROM fis_satirlari s JOIN fisler f ON f.fis_id=s.fis_id "
                "WHERE f.belge_id LIKE '%-CHAT-%'")}
        self.assertIn("331", hesaplar)                                                  # ortak cari hesabi
        self.assertNotIn("320", hesaplar)                                               # tedarikci hesabina yazilmadi

    def test_usta_tipli_cariye_odeme_usta_tedarikci_odemesi_kalir(self):
        y = self.baslat("Göktürk okulunda Devrim'e 10 bin TL ödeme yapıldı")
        while y["cfo"].get("soru"):
            alan = y["cfo"]["soru"]["alan"]
            y = self.sec(y, alan, "yeni" if alan == "borc" else "Gürsar TL Kasa")
        self.assertEqual(y["cfo"]["alanlar"]["islem_turu"], "Usta/Tedarikçi Ödemesi")

    def test_eksik_bilgi_listesi_istenen_bicimde_numarali_gosterilir(self):
        y = self.baslat("10 bin ödeme yapıldı")
        self.assertIn("Bu kaydı yazmadan önce şu bilgi eksik:", y["cevap"])
        self.assertIn("1. Faaliyet adı?", y["cevap"])
        self.assertIn("Hangi faaliyet için?", y["cevap"])                               # sorulan soru da var
        y = self.sec(y, "faaliyet", FAALIYET)
        self.assertIn("1. Proje adı?", y["cevap"])                                      # liste her adimda guncellenir
        self.assertNotIn("Faaliyet adı?", y["cevap"])                                   # cevaplanan alan listeden duser

    def test_taslak_hazirken_eksik_bilgi_basligi_yoktur(self):
        y = self.baslat("Göktürk okulunda Medet İme'ye saha betonu için 10 bin TL ödeme yapıldı")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertNotIn("şu bilgi eksik", y["cevap"])


class CariYaslandirmaTestleri(AcikBorcTemeli):
    def test_etiketli_odeme_yaslandirmada_acik_borcla_netlesir(self):
        import kurumsal_finans as kf

        self._fis("USTA-BORC", "2026-09-20", "Gider Tahakkuku", 50_000, "Ofis Usta", "iş bedeli")
        self._fis("USTA-ODEME", "2026-10-01", "Usta/Tedarikçi Ödemesi", 10_000, "Yeni Kişi",
                  "[Borç kapama: Ofis Usta] elden ödeme")
        satirlar = {x["cari"]: x for x in kf.cari_yaslandirma("320", FAALIYET, BUGUN)}
        self.assertEqual(satirlar["Ofis Usta"]["toplam"], 40_000)                  # 50.000 - 10.000
        self.assertNotIn("Yeni Kişi", satirlar)                                    # odeme kisi altinda avans gorunmez

    def test_etiketsiz_odeme_eski_davranisla_kisi_altinda_kalir(self):
        import kurumsal_finans as kf

        self._fis("USTA-BORC", "2026-09-20", "Gider Tahakkuku", 50_000, "Ofis Usta", "iş bedeli")
        self._fis("USTA-ODEME", "2026-10-01", "Usta/Tedarikçi Ödemesi", 10_000, "Yeni Kişi", "elden ödeme")
        satirlar = {x["cari"]: x for x in kf.cari_yaslandirma("320", FAALIYET, BUGUN)}
        self.assertEqual(satirlar["Ofis Usta"]["toplam"], 50_000)
        self.assertEqual(satirlar["Yeni Kişi"]["avans"], 10_000)


class AlacakKapamaTestleri(AcikBorcTemeli):
    def test_127_bin_alacak_103_bin_tahsilat_ve_24_bin_silme_ile_kapanir(self):
        d0 = pf.proje_detay(FAALIYET, PROJE, self.katalog())["metrikler"]
        self.assertEqual(d0["kalan_alacak"]["deger_kurus"], 127_000 * 100)
        # 1) 103.000 TL tahsilat
        y = self.baslat("Göktürk okulundan İşveren'den 103 bin TL tahsilat yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap", y["cevap"])
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIn("Bu tahsilat şu açık alacağı kapatıyor: İşveren / Göktürk Okulu / kalan 127.000 TL → 24.000 TL", y["cevap"])
        self.onayla(y)
        self.assertEqual(self.kalemler()[(ES.ALACAK, "İşveren")]["kalan_kurus"], 24_000 * 100)
        # 2) kalan 24.000 TL silme (kasa/banka hareketi yok)
        y = self.baslat("Göktürk okulundan İşveren'in 24 bin TL alacağı silindi")
        self.assertIsNone(y["cfo"].get("soru"), y["cevap"])                       # hesap sorulmaz, eslesme otomatik
        self.assertIn("Bu işlem şu açık alacağı siliyor: İşveren / Göktürk Okulu / kalan 24.000 TL → 0 TL", y["cevap"])
        self.onayla(y)
        m = pf.proje_detay(FAALIYET, PROJE, self.katalog())["metrikler"]
        self.assertEqual(m["kalan_alacak"]["deger_kurus"], 0)
        self.assertEqual(m["tahsil_edilen"]["deger_kurus"], 403_000 * 100)       # 300.000 + 103.000
        self.assertEqual(m["kazanilan_gelir"]["deger_kurus"], 403_000 * 100)     # 427.000 - 24.000 silinen alacak
        self.assertNotIn((ES.ALACAK, "İşveren"), self.kalemler())               # acik alacak kalmadi
        self.assertEqual(self.cari_satiri("İşveren")["kalan_alacak_kurus"], 0)


@unittest.skipUnless(sys.platform == "win32", "Excel paylasim kilidi Windows'a ozgu")
class ExcelKilitliVeriKaybiTestleri(AcikBorcTemeli):
    def test_excel_kilitliyken_onay_hicbir_sey_yazmaz_taslak_korunur_sonra_basarili(self):
        import windows_dogrulama as wd

        y = self.baslat("Göktürk okulunda Medet İme'ye saha betonu için 10 bin TL ödeme yapıldı")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        token = y["kayit_taslagi"]["token"]
        fis_once, ana_once = self.fis_sayisi(), hashlib.sha256(self.ana.read_bytes()).hexdigest()
        birak = wd._excel_gibi_kilitle(Path(self.ana))
        try:
            with self.assertRaises(RuntimeError) as hata:
                ui.onayla(token, AKTOR, "test-parola")
        finally:
            birak()
        mesaj = str(hata.exception)
        self.assertIn("Ana Excel dosyası açık veya OneDrive tarafından kilitli", mesaj)
        self.assertIn("aynı taslakta onaya tekrar basın", mesaj)
        self.assertIn("Kayıt yapılmadı", mesaj)
        self.assertIn("taslak korunuyor", mesaj)
        self.assertEqual(ui._BEKLEYEN[token]["durum"], "ONAY_BEKLIYOR")           # taslak korunuyor
        self.assertEqual(self.fis_sayisi(), fis_once)                            # DB'ye yarim yazilmadi
        self.assertEqual(hashlib.sha256(self.ana.read_bytes()).hexdigest(), ana_once)   # Excel'e yarim yazilmadi
        self.assertEqual(self.ledger(), [])
        self.assertEqual([p.name for p in Path(self.root).rglob("*.tmp*")], [])
        ui.onayla(token, AKTOR, "test-parola")                                   # kilit kalkinca ayni taslak onaylanir
        self.assertEqual((self.fis_sayisi(), len(self.ledger())), (fis_once + 1, 1))


if __name__ == "__main__":
    unittest.main()
