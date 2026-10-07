# -*- coding: utf-8 -*-
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import finans_danismani as fd
from finans_cekirdegi import faaliyet_anahtari_bul


class FinansDanismaniTestleri(unittest.TestCase):
    def test_vade_degisikligi_parola_onayi_olmadan_yazilmaz(self):
        pano = {"durum": "test"}
        komut = {"hedefler": [{"fis_id": "F-1"}], "vade": "2026-10-13"}
        with patch.object(fd, "karar_panosu", return_value=pano), \
             patch.object(fd, "_vade_komutu_coz", return_value=komut), \
             patch.object(fd, "acik_borc_vadesi_guncelle") as guncelle:
            sonuc = fd.asistana_sor("Taş işleri vadesini 13 Ekim yap", llm_kullan=False)
        self.assertTrue(sonuc["parola_gerekli"])
        self.assertFalse(sonuc["degisiklik_yapildi"])
        guncelle.assert_not_called()

    def test_profil_sinirlari(self):
        p = fd.DanismanProfili.olustur({"risk": "bilinmez", "rezerv_gun": 1, "yatirim_vadesi_ay": 999})
        self.assertEqual(p.risk, "dusuk")
        self.assertEqual(p.rezerv_gun, 7)
        self.assertEqual(p.yatirim_vadesi_ay, 120)

    def test_yatirim_sepeti_toplami_kapasiteyi_gecmez(self):
        p = fd.DanismanProfili("orta", 30, 12)
        sepet = fd._yatirim_sepeti(100_000, p)
        self.assertAlmostEqual(sum(x["tutar"] for x in sepet), 100_000, places=2)
        self.assertEqual(sum(x["oran"] for x in sepet), 100)

    def test_llm_hatasi_deterministik_cevabi_engellemez(self):
        pano = {
            "durum": "likiditeyi_koru", "ana_karar": "Likit kal.", "uyarilar": [],
            "ozet": {"nakit": 10, "alacak": 20, "borc": 30, "gelir": 0, "gider": 0, "net": 0},
            "nakit_plani": {"yatirim_kapasitesi": 0, "rezerv": 100, "finansman_acigi": 90,
                             "30_gun_beklenen_tahsilat": 0, "30_gun_vadeli_odeme": 0, "gunluk_gider": 1},
            "odemeler": [], "yatirim_sepeti": [],
            "son_hareketler": [],
        }
        with tempfile.TemporaryDirectory() as d, patch.object(fd, "karar_panosu", return_value=pano), \
             patch.object(fd, "metne_sor", side_effect=RuntimeError("offline")), \
             patch.object(fd, "AUDIT_PATH", Path(d) / "audit.jsonl"):
            sonuc = fd.asistana_sor("Ne kadar yatirim yapabilirim?", llm_kullan=True)
        self.assertEqual(sonuc["llm"], "deterministik_yedek")
        self.assertIn("guvenli nakit", sonuc["cevap"])

    def test_gercek_defter_panosu_negatif_kapasite_uretmez(self):
        pano = fd.karar_panosu("tum", {"rezerv_gun": 30}, bugun=date(2026, 9, 27))
        self.assertGreaterEqual(pano["nakit_plani"]["yatirim_kapasitesi"], 0)
        self.assertGreaterEqual(pano["nakit_plani"]["finansman_acigi"], 0)
        self.assertIn("ana_karar", pano)

    def test_touche_holding_secimi_konsolide_veriyi_gosterir(self):
        tumu = fd.karar_panosu("tum", {"rezerv_gun": 30}, bugun=date(2026, 9, 27))
        holding = fd.karar_panosu("touche_holding", {"rezerv_gun": 30}, bugun=date(2026, 9, 27))
        self.assertEqual(holding["ozet"], tumu["ozet"])
        self.assertEqual(holding["nakit_plani"], tumu["nakit_plani"])
        self.assertEqual(holding["odemeler"], tumu["odemeler"])

    def test_yapilmis_odeme_niyeti_acik_borctan_ayrilir(self):
        self.assertEqual(fd._niyet("100000 TL taş işleri ödemesi yapılmış mı?"), "gecmis_odeme")

    def test_yazim_hatasi_olan_odeme_listesi_niyeti_anlasilir(self):
        soru = "eylül ayı ödeöeleri kalem kalem göster"
        self.assertTrue(fd._detay_istegi(soru))
        self.assertEqual(fd._niyet(soru), "gecmis_odeme")

    def test_yapilmis_odeme_cevabi_defter_hareketini_kullanir(self):
        pano = {
            "ozet": {}, "nakit_plani": {}, "odemeler": [],
            "son_hareketler": [{"yon": "Ödeme", "tarih": "2026-09-16", "cari": "Sebahattin Çakmak", "tutar": 100000, "aciklama": "Mermer / taş işleri"}],
        }
        cevap = fd._deterministik_cevap(pano, "gecmis_odeme")
        self.assertIn("Sebahattin Çakmak", cevap)
        self.assertIn("100.000,00", cevap)

    def test_gursar_insaat_aliasi_gursar_mimarliga_eslenir(self):
        self.assertEqual(faaliyet_anahtari_bul("Gürsar İnşaat"), "gursar_mimarlik")

    def test_soru_icindeki_faaliyet_otomatik_seçilir(self):
        self.assertEqual(
            fd._sorudan_faaliyet("Gürsar'ın eylül ödemelerini kalem kalem göster"),
            "gursar_mimarlik",
        )
        self.assertEqual(fd._sorudan_faaliyet("Holding toplam borcu nedir?"), "tum")

    def test_usta_odemesi_maliyet_grubuyla_eslenir(self):
        self.assertEqual(fd._borc_grubu("Sebahattin Çakmak", "Mermer / taş işleri"), "Taş işleri")
        self.assertEqual(fd._borc_grubu("Medet İme", "Saha betonu"), "Saha betonu")

    def test_excel_kredi_plani_odenen_tutari_duser(self):
        with tempfile.TemporaryDirectory() as d:
            yol = Path(d) / "master.xlsx"
            wb = Workbook()
            wb.remove(wb.active)
            ayarlar = wb.create_sheet("99_Ayarlar")
            ayarlar["AD3"], ayarlar["AE3"] = "TL", 1
            ayarlar["AD4"], ayarlar["AE4"] = "EUR", 55
            ledger = wb.create_sheet("03_Kayıt Merkezi")
            ledger["G6"], ledger["I6"], ledger["T6"] = "Kredi Taksit Ödemesi", 100, "KRD-001-01"
            ledger["V6"], ledger["AC6"] = "Gerçekleşti", "Onaylandı"
            plan = wb.create_sheet("04_Kredi & Vade")
            plan["A13"], plan["B13"], plan["C13"] = "KRD-001-01", "KRD-001", "Gürsar İnşaat"
            plan["D13"], plan["G13"], plan["J13"] = datetime(2026, 10, 20), 462.8, "EUR"
            wb.save(yol)
            wb.close()
            with patch.dict(fd.FIRMALAR["touche_holding"], {"excel": str(yol)}):
                taksitler = fd._kredi_taksitleri()
        self.assertEqual(len(taksitler), 1)
        self.assertEqual(taksitler[0]["taksit_id"], "KRD-001-01")
        self.assertAlmostEqual(taksitler[0]["tutar"], 19_954.00, places=2)

    def test_gercek_kredi_plani_vadeli_odeme_olarak_gorunur(self):
        pano = fd.karar_panosu("tum", {"rezerv_gun": 30}, bugun=date(2026, 9, 27))
        taksit = next(x for x in pano["odemeler"] if x.get("proje") == "KRD-001-01")
        self.assertEqual(taksit["vade"], "2026-10-20")
        self.assertAlmostEqual(taksit["tutar"], 25_755.79, places=2)
        otuz_gun = sum(x["tutar"] for x in pano["odemeler"] if x.get("kalan_gun") is not None and 0 <= x["kalan_gun"] <= 30)
        self.assertAlmostEqual(pano["nakit_plani"]["30_gun_vadeli_odeme"], otuz_gun, places=2)

    def test_gelismis_pano_risk_projeksiyon_ve_veri_kalitesi_uretir(self):
        pano = fd.karar_panosu("tum", {"rezerv_gun": 30}, bugun=date(2026, 9, 27))
        self.assertEqual(len(pano["nakit_projeksiyonu"]), 13)
        self.assertIn("skor", pano["risk"])
        self.assertIn("skor", pano["veri_kalitesi"])
        self.assertTrue(pano["aksiyonlar"])

    def test_senaryo_yatirim_finansman_etkisini_hesaplar(self):
        sonuc = fd.senaryo_hesapla("tum", {"rezerv_gun": 30}, {
            "ek_tahsilat": 100_000, "ek_odeme": 25_000, "yeni_yatirim": 50_000,
        })
        self.assertAlmostEqual(sonuc["senaryo_nakit"], sonuc["mevcut_nakit"] + 25_000, places=2)
        self.assertGreaterEqual(sonuc["finansman_acigi"], 0)
        self.assertGreaterEqual(sonuc["yatirim_kapasitesi"], 0)

    def test_yeni_asistan_niyetleri(self):
        self.assertEqual(fd._niyet("Bugün hangi aksiyonları almalıyım?"), "aksiyon")
        self.assertEqual(fd._niyet("Eksik finans verileri neler?"), "veri_kalitesi")
        self.assertEqual(fd._niyet("13 haftalık nakit projeksiyonu"), "nakit")
        self.assertEqual(fd._niyet("3 kayıtta proje eksik demişsin, hangi 3 kayıt?"), "veri_kalitesi")
        self.assertEqual(fd._niyet("Şimdiye kadar ne kadar tahsilat yapıldı?"), "gecmis_tahsilat")
        self.assertEqual(fd._niyet("Eylül ayında ne kadar ödeme yaptık?"), "gecmis_odeme")

    def test_kredi_kullanimi_tahsilat_veya_odeme_sayilmaz(self):
        ozet = fd.nakit_hareket_ozeti("gursar_mimarlik")
        hareketler = fd.son_nakit_hareketleri("gursar_mimarlik", limit=500)
        odemeler = [x for x in hareketler if x["yon"] == "Ödeme"]
        finansman = [x for x in hareketler if x["yon"] == "Kredi kullanımı"]
        self.assertAlmostEqual(ozet["odeme_tutari"], sum(x["tutar"] for x in odemeler), places=2)
        self.assertEqual(ozet["odeme_adedi"], len(odemeler))
        self.assertAlmostEqual(ozet["finansman_girisi"], sum(x["tutar"] for x in finansman), places=2)
        kredi = next(x for x in hareketler if x["fis_id"] == "JRN-LEGACY-KRD-001-001")
        self.assertEqual(kredi["yon"], "Kredi kullanımı")
        self.assertNotIn(kredi, odemeler)

    def test_kesin_odeme_toplamini_llm_degistiremez(self):
        pano = {
            "durum": "finansman_gerekli", "ana_karar": "Likiditeyi koru.", "uyarilar": [],
            "ozet": {"nakit": 10, "alacak": 20, "borc": 30, "gelir": 0, "gider": 0, "net": 0},
            "nakit_plani": {"yatirim_kapasitesi": 0, "rezerv": 100, "finansman_acigi": 90,
                             "30_gun_beklenen_tahsilat": 0, "30_gun_vadeli_odeme": 0, "gunluk_gider": 1},
            "odemeler": [], "yatirim_sepeti": [], "aksiyonlar": [], "nakit_projeksiyonu": [],
            "faaliyet_performansi": [],
            "son_hareketler": [{"fis_id": "F1", "yon": "Ödeme", "tarih": "2026-09-16",
                                  "cari": "Usta", "tutar": 100_000, "aciklama": "İşçilik"}],
            "nakit_hareket_ozeti": {"odeme_tutari": 100_000, "odeme_adedi": 1,
                                      "tahsilat_tutari": 0, "tahsilat_adedi": 0,
                                      "finansman_girisi": 278_260.50},
            "risk": {"skor": 50, "seviye": "Orta"},
            "veri_kalitesi": {"skor": 80, "seviye": "Geliştirilmeli", "sorunlar": []},
        }
        rapor = {"etiket": "Tüm dönem", "detay": False, "hareketler": pano["son_hareketler"],
                 "ozet": pano["nakit_hareket_ozeti"]}
        with tempfile.TemporaryDirectory() as d, patch.object(fd, "karar_panosu", return_value=pano), \
             patch.object(fd, "_gecmis_nakit_raporu", return_value=rapor), \
             patch.object(fd, "metne_sor") as llm, patch.object(fd, "AUDIT_PATH", Path(d) / "audit.jsonl"):
            sonuc = fd.asistana_sor("Ne kadar ödeme yapıldı şimdiye kadar?", llm_kullan=True)
        llm.assert_not_called()
        self.assertIn("100.000,00 TL", sonuc["cevap"])
        self.assertIn("kredi kullanımı", sonuc["cevap"])
        self.assertEqual(sonuc["aksiyonlar"], [])

    def test_aylik_odeme_sorusu_ve_takip_listesi_baglami_korur(self):
        pano = fd.karar_panosu("gursar_mimarlik", {"rezerv_gun": 30}, bugun=date(2026, 10, 2))
        with tempfile.TemporaryDirectory() as d, patch.object(fd, "karar_panosu", return_value=pano), \
             patch.object(fd, "AUDIT_PATH", Path(d) / "audit.jsonl"), patch.object(fd, "metne_sor") as llm:
            ilk = fd.asistana_sor("Eylül ayında ne kadar ödeme yaptık?", "gursar_mimarlik", llm_kullan=True)
            liste = fd.asistana_sor(
                "Tam listeyi ver", "gursar_mimarlik", llm_kullan=True,
                onceki_niyet=ilk["niyet"], onceki_soru=ilk["baglam_sorusu"],
            )
        llm.assert_not_called()
        self.assertIn("Eylül 2026", ilk["cevap"])
        self.assertIn("182.437,89 TL", ilk["cevap"])
        self.assertIn("Eylül 2026 ödeme listesi (6 kayıt)", liste["cevap"])
        self.assertIn("5.565,21 TL", liste["cevap"])
        self.assertNotIn("5.652,21", liste["cevap"])
        self.assertIn("Toplam: 182.437,89 TL", liste["cevap"])

    def test_eksik_kayit_cevabi_yalniz_gercek_fisleri_listeler(self):
        pano = {
            "ozet": {}, "nakit_plani": {}, "odemeler": [],
            "veri_kalitesi": {"eksik_kayitlar": [{
                "fis_id": "F-ORTAK", "cari": "Devrim", "aciklama": "Ortak cari",
                "eksikler": ["proje"],
            }]},
        }
        cevap = fd._deterministik_cevap(pano, "veri_kalitesi")
        self.assertIn("F-ORTAK", cevap)
        self.assertNotIn("Taş işleri", cevap)


if __name__ == "__main__":
    unittest.main()
