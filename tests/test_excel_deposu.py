# -*- coding: utf-8 -*-
"""Tek ANA + tek YEDEK Excel deposu: guvenlik ve butunluk testleri.

Hicbir test uretim dosyasina dokunmaz; her test kendi gecici klasorunde calisir.
"""
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import config
import excel_deposu as depo
import excel_yazici
import guvenlik
from _ortam import AnaExcelYonlendirme
from yardimci_excel import ledger_satirlari, tam_akilli_excel


def kayit(**kw):
    taban = {"Tarih": "2026-10-12", "Faaliyet": "Gürsar İnşaat", "Proje": "Göktürk Okulu",
             "Cari_Tipi": "Usta", "Cari": "Devrim", "Islem_Turu": "Usta/Tedarikçi Ödemesi",
             "Hesap_Kasa": "Gürsar TL Banka", "Tutar": 10000, "PB": "TL", "Kur": 1,
             "Aciklama": "Test ödemesi"}
    taban.update(kw)
    return taban


def veri(*kayitlar):
    return {"_hedef": "ana_ledger", "kayitlar": list(kayitlar) or [kayit()]}


class DepoTestTemeli(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.ana = tam_akilli_excel(self.root / "veri" / "ANA.xlsx")
        self.yedek = self.root / "yedek" / "ANA - YEDEK.xlsx"
        self.yonlendirme = AnaExcelYonlendirme(self.ana, self.yedek).baslat()
        self.audit = self.root / "audit.jsonl"
        self._yamalar = [patch.object(guvenlik, "AUDIT_LOG", self.audit),
                         patch.object(guvenlik, "LOGLAR", self.root)]
        for y in self._yamalar:
            y.start()

    def tearDown(self):
        for y in reversed(self._yamalar):
            y.stop()
        self.yonlendirme.bitir()
        self.temp.cleanup()

    def xlsx_dosyalari(self):
        return sorted(p.relative_to(self.root).as_posix() for p in self.root.rglob("*.xlsx"))

    def audit_olaylari(self, olay=None):
        if not self.audit.exists():
            return []
        satirlar = [json.loads(x) for x in self.audit.read_text(encoding="utf-8").splitlines() if x.strip()]
        return [x for x in satirlar if olay is None or x.get("olay") == olay]


class TekAnaTekYedekTestleri(DepoTestTemeli):
    def test_ilk_yazma_yedegi_olusturur_ve_tek_yedek_kalir(self):
        v0 = depo.sha256(self.ana)
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertTrue(sonuc["basari"], sonuc)
        self.assertNotEqual(depo.sha256(self.ana), v0)
        self.assertEqual(depo.sha256(self.yedek), v0)           # YEDEK = bir onceki surum
        self.assertEqual(self.xlsx_dosyalari(), ["veri/ANA.xlsx", "yedek/ANA - YEDEK.xlsx"])
        self.assertEqual(depo.gecici_dosyalar(), [])

    def test_her_yazma_yedegi_yeniler_yeni_kopya_uretmez(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        v1 = depo.sha256(self.ana)
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=2500)), "B-2")
        self.assertEqual(depo.sha256(self.yedek), v1)            # artik V1 yedek
        for i in range(3, 9):
            excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=100 * i)), f"B-{i}")
        self.assertEqual(self.xlsx_dosyalari(), ["veri/ANA.xlsx", "yedek/ANA - YEDEK.xlsx"])
        self.assertEqual(len(ledger_satirlari(self.ana)), 8)

    def test_reddedilen_islem_yedegi_ilerletmez(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        yedek_once, ana_once = depo.sha256(self.yedek), depo.sha256(self.ana)
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tarih=None)), "B-X")   # eksik alan
        self.assertFalse(sonuc["basari"])
        self.assertEqual(depo.sha256(self.yedek), yedek_once)
        self.assertEqual(depo.sha256(self.ana), ana_once)

    def test_ayni_belge_iki_kez_yazilmaz(self):
        self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")["basari"])
        ana_once = depo.sha256(self.ana)
        tekrar = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertFalse(tekrar["basari"])
        self.assertEqual(tekrar["hata_kodu"], "MUKERRER_BELGE")
        self.assertEqual(depo.sha256(self.ana), ana_once)
        self.assertEqual(len(ledger_satirlari(self.ana)), 1)

    def test_ayni_belge_ayni_kayit_kimligini_uretir(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-SABIT")
        ilk = ledger_satirlari(self.ana)[0]["Kayit_ID"]
        self.assertEqual(ilk, excel_yazici._kayit_id("AI-TH", "B-SABIT-001"))


class AtomiklikVeKorumaTestleri(DepoTestTemeli):
    def test_yerlestirme_hatasinda_ana_degismez_gecici_kalmaz(self):
        once = depo.sha256(self.ana)
        gercek = depo.os.replace

        def hata(kaynak, hedef):
            if Path(hedef) == self.ana:
                raise OSError("simule atomik replace hatasi")
            return gercek(kaynak, hedef)

        with patch.object(depo.os, "replace", side_effect=hata):
            sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(depo.sha256(self.ana), once)
        self.assertEqual(depo.gecici_dosyalar(), [])

    def test_excel_acikken_anlasilir_mesaj_verir(self):
        gercek = depo.os.replace

        def kilitli(kaynak, hedef):
            if Path(hedef) == self.ana:
                raise PermissionError(13, "Permission denied")
            return gercek(kaynak, hedef)

        with patch.object(depo.os, "replace", side_effect=kilitli):
            sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertFalse(sonuc["basari"])
        self.assertIn("tüm Excel pencerelerini kapatın", sonuc["hata"])
        self.assertEqual(sonuc["hata_kodu"], "ANA_KILITLI")

    def _iki_surumlu(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "L-1")
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=2222)), "L-2")
        return depo.sha256(self.ana), depo.sha256(self.yedek)

    def test_excel_acikken_yazma_baslamadan_reddedilir_yedek_ilerlemez_geri_donus_calisir(self):
        """Excel dosyayi acik tutarken (paylasim kilidi) yapilan deneme YEDEK'i MAIN ile ezmez:
        'son islemi geri al' (MAIN=V3, YEDEK=V2 -> MAIN=V2) hala mumkun kalir."""
        v3, v2 = self._iki_surumlu()
        gercek_open = open

        def kilitli_open(yol, *a, **k):
            if Path(yol) == self.ana and "+" in (a[0] if a else k.get("mode", "r")):
                raise PermissionError(13, "Permission denied (WinError 32 benzeri)")
            return gercek_open(yol, *a, **k)

        with patch.object(depo, "open", create=True, side_effect=kilitli_open):
            sonuc = excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=3333)), "L-3")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(sonuc["hata_kodu"], "ANA_KILITLI")
        self.assertIn("tüm Excel pencerelerini kapatın", sonuc["hata"])
        self.assertEqual((depo.sha256(self.ana), depo.sha256(self.yedek)), (v3, v2))      # ikisi de ayni
        self.assertEqual(depo.gecici_dosyalar(), [])
        self.assertEqual(self.audit_olaylari("excel_islem")[-1]["hata_kodu"], "ANA_KILITLI")
        depo.yedekten_geri_don(aktor="ops", neden="kilit sonrasi", onay="GERI-DON")
        self.assertEqual((depo.sha256(self.ana), depo.sha256(self.yedek)), (v2, v3))
        # Excel kapatilinca ayni islem basarili olur
        # (geri donus sonrasi ANA=V2; yeni yazma YEDEK'i V2 yapar)
        self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=3333)), "L-4")["basari"])
        self.assertEqual(depo.sha256(self.yedek), v2)

    def test_disaridan_degisen_ana_yedegi_ilerletmez(self):
        v3, v2 = self._iki_surumlu()

        def eszamanli(wb):
            baska = load_workbook(self.ana)
            baska["99_Ayarlar"]["A30"] = "kullanici"
            baska.save(self.ana)
            baska.close()
            return {"basari": True}

        with self.assertRaises(depo.DepoHatasi) as cm:
            depo.guvenli_degistir(eszamanli, kaynak="test")
        self.assertEqual(cm.exception.kod, "ANA_DEGISTI")
        self.assertEqual(depo.sha256(self.yedek), v2)                  # YEDEK onceki surumde kaldi

    def test_formul_hucresi_ezilemez(self):
        once = depo.sha256(self.ana)

        def kotu(wb):
            wb["03_Kayıt Merkezi"]["S6"] = 1            # formul yerine sabit deger
            return {"basari": True}

        with self.assertRaises(depo.DepoHatasi) as cm:
            depo.guvenli_degistir(kotu, kaynak="test")
        self.assertEqual(cm.exception.kod, "VERI_KORUMA")
        self.assertIn("Formül bozuldu", cm.exception.mesaj)
        self.assertEqual(depo.sha256(self.ana), once)
        self.assertFalse(self.yedek.exists())            # basarisiz islem yedegi de olusturmaz

    def test_mevcut_veri_satiri_ezilemez(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        once, yedek_once = depo.sha256(self.ana), depo.sha256(self.yedek)

        def kotu(wb):
            wb["03_Kayıt Merkezi"]["I6"] = 1             # mevcut tutari ez
            return {"basari": True}

        with self.assertRaises(depo.DepoHatasi) as cm:
            depo.guvenli_degistir(kotu, kaynak="test")
        self.assertEqual(cm.exception.kod, "VERI_KORUMA")
        self.assertEqual(depo.sha256(self.ana), once)
        self.assertEqual(depo.sha256(self.yedek), yedek_once)

    def test_izinli_sutun_degisebilir_baska_sutun_degisemez(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")

        def iptal(wb):
            wb["03_Kayıt Merkezi"]["V6"] = "İptal"
            return {"basari": True}

        sonuc = depo.guvenli_degistir(iptal, kaynak="test",
                                      degisim_izni={"03_Kayıt Merkezi": {"V"}})
        self.assertTrue(sonuc["basari"])

        def fazla(wb):
            wb["03_Kayıt Merkezi"]["V6"] = "Aktif"
            wb["03_Kayıt Merkezi"]["I6"] = 5
            return {"basari": True}

        with self.assertRaises(depo.DepoHatasi):
            depo.guvenli_degistir(fazla, kaynak="test", degisim_izni={"03_Kayıt Merkezi": {"V"}})

    def test_kismen_dolu_satirin_uzerine_yazilmaz(self):
        wb = load_workbook(self.ana)
        ws = wb["03_Kayıt Merkezi"]
        ws["F6"] = "Elle girilmiş not"            # A bos, ama baska sutun dolu -> kismen dolu
        wb.save(self.ana)
        wb.close()
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertTrue(sonuc["basari"], sonuc)
        wb = load_workbook(self.ana)
        ws = wb["03_Kayıt Merkezi"]
        self.assertEqual(ws["F6"].value, "Elle girilmiş not")
        self.assertIsNone(ws["A6"].value)
        self.assertEqual(ws["F7"].value, "Devrim")           # bir sonraki bos satira yazildi
        self.assertTrue(str(ws["S7"].value).startswith("=IF"))
        wb.close()

    def test_kapasite_dolunca_reddeder_ve_hicbir_sey_degismez(self):
        with patch.object(config, "LEDGER_SON_SATIR", 7):
            self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")["basari"])
            self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "B-2")["basari"])
            once = depo.sha256(self.ana)
            sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-3")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(sonuc["hata_kodu"], "KAPASITE_DOLU")
        self.assertEqual(depo.sha256(self.ana), once)

    def test_islem_sirasinda_disaridan_degisen_ana_ezilmez(self):
        once = depo.sha256(self.ana)

        def kullanici_eszamanli_kaydetti(wb):
            # Kullanici ayni anda Excel'de dosyayi degistirip kaydetti.
            baska = load_workbook(self.ana)
            baska["99_Ayarlar"]["A30"] = "Kullanıcının eşzamanlı değişikliği"
            baska.save(self.ana)
            baska.close()
            wb["99_Ayarlar"]["A31"] = "Bizim değişikliğimiz"
            return {"basari": True}

        with self.assertRaises(depo.DepoHatasi) as cm:
            depo.guvenli_degistir(kullanici_eszamanli_kaydetti, kaynak="test")
        self.assertEqual(cm.exception.kod, "ANA_DEGISTI")
        wb = load_workbook(self.ana)
        self.assertEqual(wb["99_Ayarlar"]["A30"].value, "Kullanıcının eşzamanlı değişikliği")
        self.assertIsNone(wb["99_Ayarlar"]["A31"].value)
        wb.close()
        self.assertNotEqual(depo.sha256(self.ana), once)
        self.assertEqual(depo.gecici_dosyalar(), [])

    def test_bozuk_ana_yedegi_ezmez(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        yedek_once = depo.sha256(self.yedek)
        self.ana.write_bytes(b"bozuk dosya")
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-2")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(sonuc["hata_kodu"], "DOSYA_BOZUK")
        self.assertEqual(depo.sha256(self.yedek), yedek_once)

    def test_openpyxl_in_koruyamadigi_parca_varsa_yazma_reddedilir(self):
        """Gomulu nesne/dilimleyici/VBA gibi parcalar openpyxl'de kaybolur; sessiz kayip yok."""
        import shutil
        import zipfile

        gecici = self.ana.with_name("enjekte.xlsx")
        with zipfile.ZipFile(self.ana) as kaynak, zipfile.ZipFile(gecici, "w", zipfile.ZIP_DEFLATED) as hedef:
            for oge in kaynak.infolist():
                hedef.writestr(oge, kaynak.read(oge.filename))
            hedef.writestr("xl/embeddings/oleObject1.bin", b"gomulu nesne")
        shutil.move(gecici, self.ana)
        once = depo.sha256(self.ana)
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(sonuc["hata_kodu"], "PAKET_KAYBI")
        self.assertIn("xl/embeddings", sonuc["hata"])
        self.assertEqual(depo.sha256(self.ana), once)
        self.assertFalse(self.yedek.exists())

    def test_grafik_openpyxl_ile_korunur_ve_yazma_calisir(self):
        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        for i in range(1, 5):
            ws[f"AK{i}"] = i
        grafik = BarChart()
        grafik.add_data(Reference(ws, min_col=37, min_row=1, max_row=4))
        ws.add_chart(grafik, "AM2")
        wb.save(self.ana)
        wb.close()
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertTrue(sonuc["basari"], sonuc)
        import zipfile
        with zipfile.ZipFile(self.ana) as z:
            self.assertTrue(any(n.startswith("xl/charts/") for n in z.namelist()))

    def test_kuru_calistirma_hicbir_sey_degistirmez(self):
        once = depo.sha256(self.ana)
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1", kuru_calistir=True)
        self.assertTrue(sonuc["basari"])
        self.assertFalse(sonuc["degisiklik_yapildi"])
        self.assertEqual(depo.sha256(self.ana), once)
        self.assertFalse(self.yedek.exists())

    def test_artik_gecici_dosya_sonraki_islemde_silinir_baska_dosyaya_dokunulmaz(self):
        artik = self.ana.parent / f".{self.ana.stem}.{'a' * 32}.tmp.xlsx"
        artik.write_bytes(b"x")
        baska = self.ana.parent / ".baska-dosya.tmp.xlsx"
        baska.write_bytes(b"y")
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        self.assertFalse(artik.exists())
        self.assertTrue(baska.exists())

    def test_formul_enjeksiyonu_metin_olarak_yazilir(self):
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Aciklama="=HYPERLINK(\"http://x\")")), "B-1")
        wb = load_workbook(self.ana)
        h = wb["03_Kayıt Merkezi"]["X6"]
        self.assertEqual(h.data_type, "s")
        wb.close()

    def test_kilit_ic_ice_cagrida_kilitlenmez(self):
        from kayit_kilidi import kilitli

        @kilitli
        def dis():
            return excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")

        self.assertTrue(dis()["basari"])


class GeriDonusTestleri(DepoTestTemeli):
    def _iki_surum(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=777)), "B-2")
        return depo.sha256(self.ana), depo.sha256(self.yedek)

    def test_geri_donus_ana_ile_yedegi_takas_eder(self):
        v2, v1 = self._iki_surum()
        sonuc = depo.yedekten_geri_don(aktor="ops", neden="test", onay="GERI-DON")
        self.assertTrue(sonuc["basari"])
        self.assertEqual(depo.sha256(self.ana), v1)          # ANA = V1
        self.assertEqual(depo.sha256(self.yedek), v2)        # YEDEK = V2 (uclu dosya yok)
        self.assertEqual(self.xlsx_dosyalari(), ["veri/ANA.xlsx", "yedek/ANA - YEDEK.xlsx"])
        self.assertEqual(len(ledger_satirlari(self.ana)), 1)
        olay = self.audit_olaylari("excel_geri_donus")[-1]
        self.assertEqual(olay["durum"], "BASARILI")
        self.assertEqual(olay["ana_hash_once"], v2)
        self.assertEqual(olay["ana_hash_sonra"], v1)

    def test_iki_kez_geri_donus_ilk_surume_doner(self):
        v2, v1 = self._iki_surum()
        depo.yedekten_geri_don(aktor="ops", neden="a", onay="GERI-DON")
        depo.yedekten_geri_don(aktor="ops", neden="b", onay="GERI-DON")
        self.assertEqual(depo.sha256(self.ana), v2)

    def test_onay_olmadan_geri_donulmez(self):
        v2, v1 = self._iki_surum()
        for onay in ("", "evet", "geri-don", None):
            with self.assertRaises(depo.DepoHatasi):
                depo.yedekten_geri_don(aktor="ops", neden="x", onay=onay)
        with self.assertRaises(depo.DepoHatasi):
            depo.yedekten_geri_don(aktor="", neden="x", onay="GERI-DON")
        self.assertEqual(depo.sha256(self.ana), v2)

    def test_geri_donus_yedek_yerlestirilemezse_kurtarma_dosyasi_sonraki_yazmada_silinmez(self):
        """Geri donusun son adimi basarisiz olursa onceki ANA surumu desene uymayan adla korunur."""
        v2, v1 = self._iki_surum()
        gercek = depo.os.replace

        def kismen_basarisiz(kaynak, hedef):
            if Path(hedef) == self.yedek:
                raise OSError("yedek yazilamadi (test)")
            return gercek(kaynak, hedef)

        with patch.object(depo.os, "replace", side_effect=kismen_basarisiz):
            with self.assertRaises(depo.DepoHatasi) as cm:
                depo.yedekten_geri_don(aktor="ops", neden="test", onay="GERI-DON")
        self.assertEqual(cm.exception.kod, "YEDEK_YERLESTIRILEMEDI")
        kurtarma = Path(self.audit_olaylari("excel_geri_donus")[-1]["kurtarma_dosyasi"])
        self.assertTrue(kurtarma.exists())
        self.assertEqual(depo.sha256(kurtarma), v2)                 # onceki ANA (V2) kaybolmadi
        self.assertFalse(depo._GECICI_RE.match(kurtarma.name))       # temizlik desenine uymaz
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tutar=1)), "B-3")   # sonraki yazma temizlik yapar
        self.assertTrue(kurtarma.exists())

    def test_ayni_surumde_geri_donus_reddedilir(self):
        excel_yazici.yaz("gursar_mimarlik", veri(), "B-1")
        depo.yedegi_yenile()
        with self.assertRaises(depo.DepoHatasi) as cm:
            depo.yedekten_geri_don(aktor="ops", neden="x", onay="GERI-DON")
        self.assertEqual(cm.exception.kod, "AYNI_SURUM")

    def test_yapay_zeka_http_ve_telegram_katmani_geri_donusu_cagiramaz(self):
        """Geri donus yalniz operasyon aracindadir; hicbir otonom katman onu icermez."""
        scripts = Path(__file__).resolve().parents[1] / "scripts"
        # Yalniz insan/operator araclari: operasyon komutu ve TEST kitaplari uzerinde calisan Windows dogrulama araci.
        izinli = {"excel_deposu.py", "excel_operasyon.py", "windows_dogrulama.py"}
        for dosya in scripts.glob("*.py"):
            if dosya.name in izinli:
                continue
            self.assertNotIn("yedekten_geri_don", dosya.read_text(encoding="utf-8"), dosya.name)
        ui = Path(__file__).resolve().parents[1] / "ui"
        for dosya in ui.glob("*.js"):
            self.assertNotIn("geri-don", dosya.read_text(encoding="utf-8").lower(), dosya.name)


class GeciciDosyaKonumuTestleri(DepoTestTemeli):
    """Gecici Excel dosyalari proje temp alaninda olusur ve temizlenir; OneDrive klasorune dusmez."""

    def _konumlar(self):
        konumlar = []
        gercek = depo._gecici_yol

        def kaydet(hedef):
            yol = gercek(hedef)
            konumlar.append(yol.parent)
            return yol

        return konumlar, patch.object(depo, "_gecici_yol", side_effect=kaydet)

    def test_gecici_dosya_proje_temp_alaninda_olusur_ve_silinir(self):
        if not depo._ayni_birim(config.TEMP_DIR, self.ana.parent):
            self.skipTest("test klasorleri farkli diskte")
        konumlar, yama = self._konumlar()
        with yama:
            self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "T-1")["basari"])
        self.assertTrue(konumlar)
        self.assertEqual({k.resolve() for k in konumlar}, {Path(config.TEMP_DIR).resolve()})
        self.assertEqual(depo.gecici_dosyalar(), [])
        self.assertEqual(sorted(p.name for p in self.ana.parent.iterdir()), ["ANA.xlsx"])   # ANA klasorune gecici dusmedi

    def test_temp_alani_ayri_birimdeyse_hedefin_klasoru_kullanilir(self):
        konumlar, yama = self._konumlar()
        with yama, patch.object(depo, "_ayni_birim", return_value=False):
            self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "T-2")["basari"])
        self.assertTrue({k.resolve() for k in konumlar} <= {self.ana.parent.resolve(), self.yedek.parent.resolve()})
        self.assertEqual(depo.gecici_dosyalar(), [])

    def test_temp_alanindaki_artik_dosya_silinir_ilgisiz_dosyaya_dokunulmaz(self):
        temp = Path(config.TEMP_DIR)
        temp.mkdir(parents=True, exist_ok=True)
        artik = temp / f".ANA.{'a' * 32}.tmp.xlsx"
        artik.write_bytes(b"yarim")
        ilgisiz = temp / "baska.xlsx"
        ilgisiz.write_bytes(b"x")
        try:
            self.assertIn(str(artik), depo.gecici_dosyalar())
            self.assertTrue(excel_yazici.yaz("gursar_mimarlik", veri(), "T-3")["basari"])
            self.assertFalse(artik.exists())
            self.assertTrue(ilgisiz.exists())
        finally:
            ilgisiz.unlink(missing_ok=True)
            artik.unlink(missing_ok=True)


class VarsayilanYedekYoluTestleri(unittest.TestCase):
    def test_yedek_ana_ile_ayni_klasorde_backup_son_ekiyle(self):
        ana = Path("klasor") / "Touche Holding - AI Ready Finans Yönetim Sistemi.xlsx"
        self.assertEqual(config.varsayilan_yedek_yolu(ana),
                         Path("klasor") / "Touche Holding - AI Ready Finans Yönetim Sistemi_backup.xlsx")

    def test_ana_excel_ayarla_yedek_verilmezse_ayni_kurali_kullanir(self):
        eski = (config.ana_excel_yolu(), config.yedek_excel_yolu())
        try:
            config.ana_excel_ayarla(Path("x") / "Defter.xlsx")
            self.assertEqual(config.yedek_excel_yolu(), Path("x") / "Defter_backup.xlsx")
        finally:
            config.ana_excel_ayarla(*eski)


class TekYazmaHattiStatikTestleri(unittest.TestCase):
    """Mimari degismezler: kaynak kod duzeyinde, ayri bir yazma hatti/yedek mekanizmasi geri gelemez."""

    SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"

    def test_calisma_kitabini_yalniz_excel_deposu_kaydeder(self):
        kaydet = re.compile(r"\b(?:wb|workbook|master_wb|holding|kontrol|yeni_wb)\.save\(")
        for dosya in self.SCRIPTS.glob("*.py"):
            if dosya.name == "excel_deposu.py":
                continue
            self.assertIsNone(kaydet.search(dosya.read_text(encoding="utf-8")),
                              f"{dosya.name}: çalışma kitabı excel_deposu dışında kaydedilemez")

    def test_tarihli_veya_islem_basina_yedek_mekanizmasi_yok(self):
        yasak = ("YEDEK_SAYISI", "kurulum_yedek", "excel_yedek", "_yedek_al(", "yedek_klasoru")
        for dosya in self.SCRIPTS.glob("*.py"):
            if dosya.name == "excel_operasyon.py":      # eski kopyalari tanimak/temizlemek icin kalip listesi tutar
                continue
            metin = dosya.read_text(encoding="utf-8")
            for k in yasak:
                self.assertNotIn(k, metin, f"{dosya.name}: '{k}'")

    def test_kaynak_kodda_kisisel_mutlak_yol_yok(self):
        for dosya in self.SCRIPTS.glob("*.py"):
            metin = dosya.read_text(encoding="utf-8")
            self.assertNotIn("twone", metin, dosya.name)
            self.assertNotIn("C:\\Users", metin, dosya.name)

    def test_ana_excel_yolu_yalniz_config_uzerinden_cozulur(self):
        """Kullanici masaustu yolu / sabit Excel adi yalnizca config.py'de bulunur."""
        for dosya in self.SCRIPTS.glob("*.py"):
            if dosya.name in {"config.py", "excel_operasyon.py"}:
                continue
            metin = dosya.read_text(encoding="utf-8")
            self.assertNotIn("C:\\Users", metin, dosya.name)
            self.assertNotIn("AI Ready Finans", metin, dosya.name)


class HariciDosyaVeAuditTestleri(DepoTestTemeli):
    def test_harici_dosya_ana_olur_tek_yedek_korunur(self):
        yeni = tam_akilli_excel(self.root / "disarida" / "onayli.xlsx")
        wb = load_workbook(yeni)
        wb["99_Ayarlar"]["A40"] = "Harici düzenleme"
        wb.save(yeni)
        wb.close()
        once = depo.sha256(self.ana)
        sonuc = depo.harici_dosyayi_ana_yap(yeni, aktor="ops", neden="test")
        self.assertTrue(sonuc["basari"])
        self.assertEqual(depo.sha256(self.ana), depo.sha256(yeni))
        self.assertEqual(depo.sha256(self.yedek), once)
        self.assertEqual(sorted(p.name for p in (self.root / "veri").glob("*")), ["ANA.xlsx"])

    def test_audit_kaydi_islem_izini_icerir(self):
        baglam = {"orijinal_talimat": "Göktürk okuluna Devrim'e 10 bin TL ödeme yapıldı.",
                  "normalize_komut": "ODEME|gursar|goktur|devrim|10000", "onay_kaynagi": "ui:test",
                  "taslak_id": "TSL-1"}
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri(), "B-1", aktor="owner", baglam=baglam)
        self.assertTrue(sonuc["basari"])
        olay = self.audit_olaylari("excel_islem")[-1]
        self.assertEqual(olay["durum"], "BASARILI")
        self.assertEqual(olay["aktor"], "owner")
        self.assertRegex(olay["islem_id"], r"^OP-")
        self.assertEqual(olay["ana_hash_sonra"], depo.sha256(self.ana))
        self.assertEqual(olay["yedek_hash"], olay["ana_hash_once"])
        self.assertEqual(olay["orijinal_talimat"], baglam["orijinal_talimat"])
        self.assertEqual(olay["normalize_komut"], baglam["normalize_komut"])
        kayit0 = olay["kayitlar"][0]
        self.assertEqual(kayit0["transaction_id"], ledger_satirlari(self.ana)[0]["Kayit_ID"])
        self.assertEqual(kayit0["cari"], "Devrim")
        self.assertEqual(kayit0["proje"], "Göktürk Okulu")
        self.assertEqual(kayit0["tutar"], 10000)
        self.assertTrue(re.fullmatch(r"[0-9a-f]{64}", olay["kayit_hash"]))
        self.assertTrue(guvenlik.audit_zincirini_dogrula(self.audit)["gecerli"])

    def test_reddedilen_islem_de_denetlenir(self):
        excel_yazici.yaz("gursar_mimarlik", veri(kayit(Tarih=None)), "B-RED")
        olay = self.audit_olaylari("excel_islem")[-1]
        self.assertEqual(olay["durum"], "REDDEDILDI")
        self.assertIn("eksik", olay["hata"])


if __name__ == "__main__":
    unittest.main()
