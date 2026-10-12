# -*- coding: utf-8 -*-
"""Envanter/temizlik aracı: yalnız kanıtlanmış artık kopyalar silinir."""
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import config
import excel_operasyon as op
from _ortam import AnaExcelYonlendirme
from yardimci_excel import tam_akilli_excel


def kayit_ekle(yol, kayit_id):
    wb = load_workbook(yol)
    ws = wb["03_Kayıt Merkezi"]
    ws["A6"] = kayit_id
    ws["F6"] = "Test"
    wb.save(yol)
    wb.close()


class EnvanterTemizlikTestleri(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.ana = tam_akilli_excel(self.root / "Desktop" / "Touche" / "Touche Holding Smart.xlsx")
        self.yedek = self.root / "yedek" / "Touche Holding Smart - YEDEK.xlsx"
        self.yedek.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.ana, self.yedek)
        self.yonlendirme = AnaExcelYonlendirme(self.ana, self.yedek).baslat()
        self.base = self.root / "ai-excel"
        self._base_eski = config.BASE_DIR
        config.BASE_DIR = self.base

        # Gecmisten kalan, bu sistemin urettigi kopyalar
        self.eski_yedekler = []
        for i in range(3):
            y = self.base / "firmalar" / "touche_holding" / "yedek" / f"Touche Holding Smart_2026010{i}_120000.xlsx"
            y.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.ana, y)
            self.eski_yedekler.append(y)
        kurulum = self.base / "firmalar" / "kurulum_yedek" / "20260101_000000_ab" / "0_Touche Holding Smart.xlsx"
        kurulum.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.ana, kurulum)
        self.eski_yedekler.append(kurulum)
        # Tekil kayit iceren eski kopya (ANA'da olmayan Kayit_ID) -> ASLA silinmez
        self.tekil = self.base / "firmalar" / "touche_holding" / "yedek" / "Touche Holding Smart_20250101_000000.xlsx"
        shutil.copy2(self.ana, self.tekil)
        kayit_ekle(self.tekil, "ESKI-ONEMLI-KAYIT")
        # Ilgisiz kullanici dosyasi -> envantere bile girmemeli
        self.ilgisiz = self.root / "Desktop" / "Maas hesabi.xlsx"
        wb = Workbook(); wb.save(self.ilgisiz); wb.close()
        # Adi ilgili ama uretilmis kopya kalibina uymayan -> BELIRSIZ, silinmez
        self.elle = self.root / "Desktop" / "Touche" / "Touche notlarim.xlsx"
        wb = Workbook(); wb.save(self.elle); wb.close()

    def tearDown(self):
        config.BASE_DIR = self._base_eski
        self.yonlendirme.bitir()
        self.temp.cleanup()

    def _tara(self):
        return op.tara([self.base, self.root / "Desktop", self.yedek.parent])

    def test_envanter_rolleri_dogru_siniflar(self):
        roller = {Path(b["yol"]).name: b["rol"] for b in self._tara()}
        self.assertEqual(roller["Touche Holding Smart.xlsx"], "MAIN")
        self.assertEqual(roller["Touche Holding Smart - YEDEK.xlsx"], "BACKUP")
        self.assertTrue(roller["Touche Holding Smart_20250101_000000.xlsx"].startswith("TEKIL_KAYIT_VAR"))
        self.assertTrue(roller["Touche notlarim.xlsx"].startswith(("BELIRSIZ", "KAPSAMA_KANITLANAMADI")))
        self.assertNotIn("Maas hesabi.xlsx", roller)             # ilgisiz dosya envantere girmez

    def test_yalniz_kanitli_artik_kopyalar_silinebilir(self):
        silinebilir = {Path(b["yol"]) for b in self._tara() if op.silinebilir(b)}
        self.assertEqual(silinebilir, {p.resolve() for p in self.eski_yedekler})
        self.assertNotIn(self.tekil.resolve(), silinebilir)
        self.assertNotIn(self.elle.resolve(), silinebilir)
        self.assertNotIn(self.ana.resolve(), silinebilir)
        self.assertNotIn(self.yedek.resolve(), silinebilir)

    def test_silme_onay_olmadan_calismaz_onayla_yalniz_artiklari_siler(self):
        rapor = self.root / "rapor.md"
        kok = ["--kok", str(self.base), "--kok", str(self.root / "Desktop"), "--kok", str(self.yedek.parent)]
        # onaysiz -> hicbir sey silinmez
        self.assertEqual(op.main(["temizle", *kok, "--rapor", str(rapor)]), 0)
        self.assertTrue(all(p.exists() for p in self.eski_yedekler))
        self.assertEqual(op.main(["temizle", *kok, "--rapor", str(rapor), "--sil"]), 2)
        self.assertTrue(all(p.exists() for p in self.eski_yedekler))
        # onayli
        self.assertEqual(op.main(["temizle", *kok, "--rapor", str(rapor), "--sil", "--onay", "SIL"]), 0)
        self.assertFalse(any(p.exists() for p in self.eski_yedekler))
        for korunan in (self.ana, self.yedek, self.tekil, self.elle, self.ilgisiz):
            self.assertTrue(korunan.exists(), korunan)
        self.assertIn("MAIN", rapor.read_text(encoding="utf-8"))

    # --- kanit esigi: kapsami kanitlanamayan dosya, adi kalibina uysa bile ASLA silinmez ---
    def _desenli(self, ad="Touche Holding Smart_20250202_000000.xlsx"):
        yol = self.base / "firmalar" / "touche_holding" / "yedek" / ad
        yol.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(self.ana, yol)
        return yol

    def _bul(self, yol):
        return next(b for b in self._tara() if Path(b["yol"]) == yol.resolve())

    def test_bozuk_dosya_kalibina_uysa_da_silinemez(self):
        yol = self._desenli()
        yol.write_bytes(b"bozuk xlsx")
        b = self._bul(yol)
        self.assertFalse(op.silinebilir(b))
        self.assertTrue(b["rol"].startswith("OKUNAMADI"))

    def test_ledger_sayfasi_olmayan_eski_dosya_silinemez(self):
        """Eski faaliyet-basina dosyalar farkli sema kullanabilir: kimlik karsilastirmasi yapilamaz -> silinmez."""
        yol = self._desenli("Gursar_hesap.xlsx")
        wb = Workbook()
        wb.active.title = "Ödeme Geçmişi"
        wb.active["A6"] = "ESKI ODEME"
        wb.save(yol)
        wb.close()
        b = self._bul(yol)
        self.assertFalse(op.silinebilir(b))
        self.assertTrue(b["rol"].startswith("KAPSAMA_KANITLANAMADI"))
        self.assertTrue(any("03_Kayıt Merkezi" in x for x in b["kapsama_sorunlari"]))

    def test_anada_olmayan_dolu_sayfa_silmeyi_engeller(self):
        yol = self._desenli()
        wb = load_workbook(yol)
        wb.create_sheet("Eski Notlar")["A1"] = "onemli not"
        wb.save(yol)
        wb.close()
        b = self._bul(yol)
        self.assertFalse(op.silinebilir(b))
        self.assertTrue(any("Eski Notlar" in x for x in b["kapsama_sorunlari"]))

    def test_anada_olmayan_is_kalemi_kimligi_silmeyi_engeller(self):
        yol = self._desenli()
        wb = load_workbook(yol)
        wb["06_Proje İş Kalemleri"]["A20"] = "KLM-ESKI-999"
        wb.save(yol)
        wb.close()
        b = self._bul(yol)
        self.assertFalse(op.silinebilir(b))
        self.assertEqual(b["tekil_kalem"], ["KLM-ESKI-999"])
        self.assertEqual(op.siniflandir(b), "UNIQUE_DATA_REVIEW_REQUIRED")

    def test_ana_ile_ayni_icerikli_desenli_kopya_silinebilir_ve_sinifi_dogru(self):
        yol = self._desenli()
        b = self._bul(yol)
        self.assertTrue(op.silinebilir(b))
        self.assertEqual(op.siniflandir(b), "SAFE_TO_DELETE")

    def test_sinif_adlari(self):
        siniflar = {Path(b["yol"]).name: op.siniflandir(b) for b in self._tara()}
        self.assertEqual(siniflar["Touche Holding Smart.xlsx"], "CANONICAL_MAIN")
        self.assertEqual(siniflar["Touche Holding Smart - YEDEK.xlsx"], "CANONICAL_BACKUP")
        self.assertEqual(siniflar["Touche Holding Smart_20250101_000000.xlsx"], "UNIQUE_DATA_REVIEW_REQUIRED")
        self.assertEqual(siniflar["Touche notlarim.xlsx"], "UNKNOWN")
        self.assertEqual(siniflar["Touche Holding Smart_20260100_120000.xlsx"], "SAFE_TO_DELETE")

    def test_gecici_dosya_temp_sinifinda_ve_csv_yalniz_listelenir(self):
        temp = self.root / "Desktop" / "Touche" / f".Touche Holding Smart.{'b' * 32}.tmp.xlsx"
        shutil.copy2(self.ana, temp)
        csv = self.base / "arsiv" / "Touche banka hareketleri.csv"
        csv.parent.mkdir(parents=True, exist_ok=True)
        csv.write_text("Tarih;Tutar\n2026-01-01;1\n", encoding="utf-8")
        siniflar = {Path(b["yol"]).name: (op.siniflandir(b), op.silinebilir(b)) for b in self._tara()}
        self.assertEqual(siniflar[temp.name][0], "TEMP")
        self.assertEqual(siniflar["Touche banka hareketleri.csv"], ("ARCHIVE", False))     # csv asla silinmez

    def test_durum_ve_dogrula(self):
        self.assertEqual(op.main(["durum"]), 0)
        self.assertEqual(op.main(["dogrula"]), 0)
        self.ana.write_bytes(b"bozuk")
        self.assertEqual(op.main(["dogrula"]), 1)


if __name__ == "__main__":
    unittest.main()
