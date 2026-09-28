# -*- coding: utf-8 -*-
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import finans_cekirdegi as fk
import proje_is_kalemleri as pik


class ProjeIsKalemleriTestleri(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.yol = Path(self.temp.name) / "master.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "03_Kayıt Merkezi"
        ws = wb.create_sheet(pik.SAYFA)
        ws["A6"], ws["B6"], ws["C6"] = "K-001", "Gürsar İnşaat", "Göktürk Okulu"
        ws["D6"], ws["E6"], ws["F6"] = "Taş işleri", "Basamak", "mtül"
        ws["G6"], ws["H6"] = 10, 500
        ws["I6"] = '=IF(OR($G6="",$H6=""),"",$G6*$H6)'
        ws["J6"], ws["K6"] = "Birim fiyat", 300
        ws["L6"] = '=IF($J6="Toplam sabit",$K6,$G6*$K6)'
        ws["M6"], ws["N6"] = "=I6-L6", "=IFERROR(M6/I6,0)"
        ws["O6"], ws["P6"], ws["Q6"], ws["R6"] = "Onaylı", "not", "HAK-001", "GID-001"
        wb.save(self.yol)
        wb.close()
        self.excel_patch = patch.dict(pik.FIRMALAR["touche_holding"], {"excel": str(self.yol)})
        self.excel_patch.start()

    def tearDown(self):
        self.excel_patch.stop()
        self.temp.cleanup()

    def test_listele_formulleri_python_ile_hesaplar(self):
        sonuc = pik.listele("gursar_mimarlik")
        self.assertEqual(sonuc["ozet"]["kalem_sayisi"], 1)
        self.assertEqual(sonuc["ozet"]["ciro"], 5000)
        self.assertEqual(sonuc["ozet"]["maliyet"], 3000)
        self.assertEqual(sonuc["ozet"]["brut_fark"], 2000)

    def test_guncelle_formulleri_korur_ve_bagli_toplami_senkronlar(self):
        yedek = Path(self.temp.name) / "backup.xlsx"

        def backup(_):
            shutil.copy2(self.yol, yedek)
            return yedek

        with patch.object(pik, "_yedek_al", side_effect=backup), \
             patch.object(pik, "veritabani_yedekle", return_value=Path(self.temp.name) / "db.bak"), \
             patch.object(pik, "_defteri_senkronla") as sync, \
             patch.object(pik, "audit_yaz"):
            sonuc = pik.guncelle("K-001", {"metraj": "12,5", "not": "güncel"}, "test")

        self.assertTrue(sonuc["basari"])
        sync.assert_called_once_with({"HAK-001": 6250.0, "GID-001": 3750.0})
        wb = load_workbook(self.yol, data_only=False)
        try:
            self.assertEqual(wb[pik.SAYFA]["G6"].value, 12.5)
            self.assertTrue(str(wb[pik.SAYFA]["I6"].value).startswith("="))
            self.assertEqual(wb[pik.SAYFA]["P6"].value, "güncel")
        finally:
            wb.close()

    def test_bilinmeyen_alan_reddedilir(self):
        with self.assertRaisesRegex(ValueError, "Değiştirilemeyen alan"):
            pik.guncelle("K-001", {"faaliyet": "Başka"}, "test")

    def test_iptal_edilen_kalem_bagli_toplami_sifirlar(self):
        wb = load_workbook(self.yol, data_only=False)
        ws = wb[pik.SAYFA]
        ws["O6"] = "İptal"
        self.assertEqual(pik._bagli_toplamlar(ws), {"HAK-001": 0, "GID-001": 0})
        wb.close()

    def test_merkez_defter_tutarini_dengeli_gunceller(self):
        db_yolu = Path(self.temp.name) / "finans.db"
        with patch.object(fk, "DB_PATH", db_yolu):
            fk.veritabanini_hazirla()
            with fk._baglanti() as con:
                con.execute(
                    "INSERT INTO fisler(fis_id,belge_id,kayit_sira,faaliyet_key,tarih,olusturma_zamani) VALUES(?,?,?,?,?,?)",
                    ("F1", "LEGACY-HAK-001", 1, "gursar_mimarlik", "2026-09-15", "2026-09-28T00:00:00"),
                )
                con.execute("INSERT INTO fis_satirlari VALUES(?,?,?,?,?,?)", ("F1", 1, "120", 100, 0, ""))
                con.execute("INSERT INTO fis_satirlari VALUES(?,?,?,?,?,?)", ("F1", 2, "600", 0, 100, ""))
            pik._defteri_senkronla({"HAK-001": 6250.25})
            with fk._baglanti() as con:
                rows = con.execute("SELECT borc_kurus,alacak_kurus FROM fis_satirlari WHERE fis_id='F1' ORDER BY satir_no").fetchall()
            self.assertEqual([(r[0], r[1]) for r in rows], [(625025, 0), (0, 625025)])


if __name__ == "__main__":
    unittest.main()
