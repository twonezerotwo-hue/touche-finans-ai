# -*- coding: utf-8 -*-
"""CFO asistani: dogal dil -> yapisal niyet -> yonlendirmeli taslak -> mevcut onay -> Excel."""
import unittest
from decimal import Decimal

import cfo_niyet as N
import cfo_sohbet
import excel_deposu as depo
import ui_belge_servisi as ui
from cfo_test_temeli import AKTOR, BUGUN, CfoTestTemeli


class DogalDilCozumlemeTestleri(CfoTestTemeli):
    TAM = "Göktürk okuluna Gürsar Mimarlık adına Devrim'e 10 bin TL ödeme yapıldı."

    def _beklenen(self, n):
        self.assertEqual(n.tur, "KAYIT")
        self.assertEqual(n.islem, N.COUNTERPARTY_PAYMENT)
        self.assertEqual(n.faaliyet.kimlik, "gursar_mimarlik")
        self.assertEqual(n.faaliyet.gorunen(), "Gürsar Mimarlık")          # sahibin kullandigi ad
        self.assertEqual(n.faaliyet.ad, "Gürsar İnşaat")                    # Excel'deki yazim
        self.assertEqual(n.proje.ad, "Göktürk Okulu")
        self.assertEqual(n.cari.ad, "Devrim")
        self.assertEqual(n.tutar, Decimal("10000"))
        self.assertEqual(n.belirsiz, {})

    def test_tam_cumle(self):
        n = self.niyet(self.TAM)
        self._beklenen(n)
        self.assertEqual(n.pb, "TL")
        self.assertEqual(n.tarih, BUGUN)
        self.assertEqual(n.tarih_kaynagi, "varsayilan_bugun")

    def test_diger_soyleyisler(self):
        for metin in ("Göktürk okulunda Devrim'e Gürsar'dan 10.000 ödeme yaptık",
                      "Göktürk Okulu Devrim 10 bin ödeme",
                      "Gürsar Mimarlık Göktürk projesi için Devrim'e bugün 10k TL ödendi.",
                      "göktürk okulu / devrim / 10000 TL ödeme yapıldı"):
            with self.subTest(metin=metin):
                self._beklenen(self.niyet(metin))

    def test_faaliyet_projeden_cikarilir_ve_denetlenir(self):
        n = self.niyet("Göktürk Okulu Devrim 10 bin ödeme")
        self.assertTrue(any("Faaliyet projeden çıkarıldı" in x for x in n.notlar))

    def test_tutar_bicimleri(self):
        for metin, beklenen in (("Devrim'e 10 bin TL ödeme yapıldı", 10000), ("Devrim'e 10.000 TL ödeme yapıldı", 10000),
                                ("Devrim'e 10k ödeme yapıldı", 10000), ("Devrim'e 12.500,50 TL ödeme yapıldı", Decimal("12500.50")),
                                ("Devrim'e 1,5 milyon TL ödeme yapıldı", 1500000), ("Devrim'e 10.000TL ödeme yapıldı", 10000)):
            with self.subTest(metin=metin):
                self.assertEqual(self.niyet(metin).tutar, Decimal(beklenen))

    def test_tarihler(self):
        self.assertEqual(self.niyet("Devrim'e 5 bin ödeme yapıldı 05.10.2026").tarih, N.date(2026, 10, 5))
        self.assertEqual(self.niyet("Devrim'e dün 5 bin ödeme yapıldı").tarih, N.date(2026, 10, 11))
        self.assertEqual(self.niyet("Devrim'e 5 Ekim 5 bin ödeme yapıldı").tarih, N.date(2026, 10, 5))
        gecersiz = self.niyet("Devrim'e 31.02.2026 tarihinde 5 bin ödeme yapıldı")
        self.assertIsNone(gecersiz.tarih)
        self.assertTrue(gecersiz.tarih_hatasi)

    def test_sorgular_kayit_sayilmaz(self):
        for metin in ("Göktürk okulunda ne kadar ödeme yaptık?", "Devrim'e toplam ne kadar ödeme yapıldı?",
                      "Bu ay Holding toplam ne kadar ödeme yaptı?", "Eylül ödemelerini kalem kalem göster",
                      "Göktürk okulundan ne kadar kazandık?", "Devrim'e ödeme yapıldı mı?"):
            with self.subTest(metin=metin):
                self.assertEqual(self.niyet(metin).tur, "SORGU")

    def test_sohbet_kayit_degildir(self):
        for metin in ("merhaba", "ödeme", "teşekkürler"):
            self.assertNotEqual(self.niyet(metin).tur, "KAYIT", metin)

    def test_karmasik_talimatlar_mevcut_llm_hattina_birakilir(self):
        for metin, neden in (("Devrim'e 10 bin TL ve Sebahattin'e 5 bin TL ödeme yapıldı", "birden fazla tutar"),
                             ("Devrim'e 10 bin TL ödeme yapılacak", "ileri tarihli"),
                             ("Çek karşılığında 100 bin TL nakit aldık, Excel'e işle", "çek/kredi")):
            n = self.niyet(metin)
            self.assertTrue(n.karmasik, metin)
            self.assertIn(neden, n.karmasik_nedeni)

    def test_kredi_odendi_karti_rapor_sorusu_degil_kayit_talimatidir(self):
        from openpyxl import load_workbook
        import master_data

        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        ws["T20"] = "Kredi Taksit Ödemesi"
        ws["AG20"], ws["AH20"], ws["AI20"] = "Gürsar İnşaat", "Banka", "Kredi Veren Banka"
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()

        metin = "Kredi Veren Banka\nKRD-001-01 · 2026-10-20\n₺25.755,79 ödemesi yapıldı"
        n = self.niyet(metin)
        self.assertEqual(n.tur, "KAYIT")
        self.assertFalse(n.karmasik)
        self.assertEqual(n.islem, N.COUNTERPARTY_PAYMENT)
        self.assertEqual(n.faaliyet.kimlik, "gursar_mimarlik")
        self.assertEqual(n.cari.ad, "Kredi Veren Banka")
        self.assertEqual(n.tutar, Decimal("25755.79"))

        y = cfo_sohbet.baslat(n, AKTOR, self.katalog(), bugun=BUGUN)
        self.assertEqual(y["niyet"], "cfo_sihirbaz")
        self.assertEqual(y["cfo"]["alanlar"]["islem_turu"], "Kredi Taksit Ödemesi")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap")
        self.assertIn("Hangi hesap", y["cevap"])
        self.assertNotIn("Toplam ödeme", y["cevap"])

    def test_desteklenmeyen_islemler_aciklanir(self):
        n = self.niyet("Gürsar'dan Holding'e 100 bin TL grup içi transfer yaz")
        self.assertEqual(n.islem, N.TRANSFER)
        self.assertIn("Grup içi transfer", n.desteklenmeyen)

    def test_bilinmeyen_kisi_uydurma_degil_aday_olarak_isaretlenir(self):
        n = self.niyet("Gürsar Mimarlık Göktürk Okulu Mehmet'e 5 bin TL ödeme yapıldı")
        self.assertIsNone(n.cari)
        self.assertEqual(n.yeni_cari_adi, "Mehmet")

    def test_acik_is_kalemi_gideri(self):
        n = self.niyet("Göktürk okulunda elektrik iş kalemine 25 bin TL gider yaz")
        self.assertEqual(n.islem, N.PROJECT_COST)
        self.assertEqual(n.kalem.ad, "Elektrik")
        self.assertEqual(n.tutar, Decimal("25000"))

    def test_gursar_insaat_ve_mimarlik_ayni_faaliyet(self):
        for ad in ("Gürsar Mimarlık", "Gürsar İnşaat", "Gürsar"):
            n = self.niyet(f"{ad} Devrim'e 1000 TL ödeme yapıldı")
            self.assertEqual(n.faaliyet and n.faaliyet.kimlik, "gursar_mimarlik", ad)


class YonlendirmeliSihirbazTestleri(CfoTestTemeli):
    def test_tam_cumle_tek_taslak_uretir_hesap_belirsizse_yalniz_hesabi_sorar(self):
        y = self.baslat(DogalDilCozumlemeTestleri.TAM)
        # Gürsar'ın iki TL hesabı var, geçmiş kayıt yok -> yalnız HESAP sorulur; diğer her şey bilinir.
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap")
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]}, {"Gürsar TL Banka", "Gürsar TL Kasa"})
        a = y["cfo"]["alanlar"]
        self.assertEqual((a["faaliyet"], a["proje"], a["cari"], a["tutar"]),
                         ("Gürsar Mimarlık", "Göktürk Okulu", "Devrim", "10.000 TL"))
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.assertIsNotNone(y["kayit_taslagi"])
        self.assertEqual(y["kayit_taslagi"]["durum"], "ONAY_BEKLIYOR")
        self.assertEqual(y["kayit_taslagi"]["dogrulama"]["durum"], "PASS")
        self.assertEqual(self.ledger(), [])                            # onaydan once Excel'e yazilmadi

    def test_gecmis_kayittan_hesap_otomatik_ve_gorunur_gerekceli(self):
        self.gecmis_odeme_yaz()
        y = self.baslat(DogalDilCozumlemeTestleri.TAM)
        self.assertIsNotNone(y["kayit_taslagi"], y["cevap"])
        self.assertIn("Hesap: Gürsar TL Banka (otomatik:", y["cevap"])
        kayit = y["kayit_taslagi"]["kayitlar"][0]
        self.assertEqual((kayit["hesap_kasa"], kayit["cari"], kayit["proje"], kayit["islem"]),
                         ("Gürsar TL Banka", "Devrim", "Göktürk Okulu", "Usta/Tedarikçi Ödemesi"))
        self.assertTrue(kayit["otomatik_eslesmeler"])

    def test_eksik_komut_adim_adim_hicbir_alan_tekrar_sorulmaz(self):
        """Spec CASE B: '10 bin ödeme yapıldı' -> Faaliyet -> Proje -> İş kalemi/Cari -> Kime -> taslak."""
        y = self.baslat("10 bin ödeme yapıldı.")
        sorulanlar = [y["cfo"]["soru"]["alan"]]
        self.assertEqual(y["cfo"]["alanlar"]["tutar"], "10.000 TL")
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]}, {"touche_holding", "gursar_mimarlik"})
        y = self.sec(y, "faaliyet", "gursar_mimarlik")
        sorulanlar.append(y["cfo"]["soru"]["alan"])
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]},
                         {"Göktürk Okulu", "Bahçeşehir Villa", "__GENEL__"})       # + sistem 'faaliyet geneli' tanımı
        y = self.sec(y, "proje", "Göktürk Okulu")
        sorulanlar.append(y["cfo"]["soru"]["alan"])
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]},
                         {N.COUNTERPARTY_PAYMENT, N.WORK_ITEM_PAYMENT})
        self.assertIn("iş kalemine mi", y["cfo"]["soru"]["metin"])
        y = self.sec(y, "islem", N.COUNTERPARTY_PAYMENT)
        sorulanlar.append(y["cfo"]["soru"]["alan"])
        self.assertEqual(y["cfo"]["soru"]["alan"], "cari")
        y = self.sec(y, "cari", "Devrim")
        # tutar zaten bilindiği için sorulmaz; sıradaki tek eksik hesap
        sorulanlar.append(y["cfo"]["soru"]["alan"])
        self.assertEqual(sorulanlar, ["faaliyet", "proje", "islem", "cari", "hesap"])
        self.assertEqual(y["cfo"]["alanlar"]["tutar"], "10.000 TL")      # kaybolmadı
        y = self.sec(y, "hesap", "Gürsar TL Kasa")
        self.assertEqual(y["kayit_taslagi"]["kayitlar"][0]["tutar"], 10000.0)
        self.assertEqual(y["kayit_taslagi"]["kayitlar"][0]["hesap_kasa"], "Gürsar TL Kasa")

    def test_gider_komutu_yalniz_eksikleri_sorar(self):
        y = self.baslat("5 bin gider ekle.")
        alanlar = []
        while y["cfo"].get("soru"):
            soru = y["cfo"]["soru"]
            alanlar.append(soru["alan"])
            secim = {"faaliyet": "gursar_mimarlik", "proje": "Göktürk Okulu", "kalem": "KLM-001", "cari": "Sebahattin Çakmak"}[soru["alan"]]
            y = self.sec(y, soru["alan"], secim)
        self.assertEqual(alanlar, ["faaliyet", "proje", "kalem", "cari"])      # tahakkukta hesap sorulmaz
        kayit = y["kayit_taslagi"]["kayitlar"][0]
        self.assertEqual((kayit["islem"], kayit["tutar"]), ("Gider Tahakkuku", 5000.0))
        self.assertEqual(y["kayit_taslagi"]["dogrulama"]["durum"], "PASS")

    def test_serbest_metin_cevaplar_ve_secenek_tiklamasi_ayni_sonucu_verir(self):
        y = self.baslat("10 bin ödeme yapıldı")
        tid = y["cfo"]["taslak_id"]
        y = cfo_sohbet.metinle_cevapla(tid, "gürsar", AKTOR, self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "proje")
        y = cfo_sohbet.metinle_cevapla(tid, "göktürk", AKTOR, self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "islem")
        y = cfo_sohbet.metinle_cevapla(tid, "cari ödemesi", AKTOR, self.katalog())
        y = cfo_sohbet.metinle_cevapla(tid, "devrim", AKTOR, self.katalog())
        y = cfo_sohbet.metinle_cevapla(tid, "Gürsar TL Kasa", AKTOR, self.katalog())
        self.assertEqual(y["kayit_taslagi"]["kayitlar"][0]["cari"], "Devrim")

    def test_anlasilmayan_cevap_sihirbazi_bozmaz(self):
        y = self.baslat("10 bin ödeme yapıldı")
        tid = y["cfo"]["taslak_id"]
        y = cfo_sohbet.metinle_cevapla(tid, "bilmiyorum", AKTOR, self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "faaliyet")          # ayni soru, durum korunur
        self.assertIn("bulunamadı", y["cevap"])
        self.assertEqual(y["cfo"]["alanlar"]["tutar"], "10.000 TL")

    def test_tutar_eksikse_yalniz_tutar_sorulur(self):
        y = self.baslat("Gürsar Mimarlık Göktürk Okulu Devrim'e ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "tutar")
        y = cfo_sohbet.metinle_cevapla(y["cfo"]["taslak_id"], "25 bin TL", AKTOR, self.katalog())
        self.assertEqual(y["cfo"]["alanlar"]["tutar"], "25.000 TL")
        self.assertEqual(y["cfo"]["soru"]["alan"], "hesap")

    def test_bilinmeyen_kisi_onay_olmadan_cariye_eklenmez(self):
        y = self.baslat("Gürsar Mimarlık Göktürk Okulu Mehmet'e 5 bin TL ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "cari_yeni_onay")
        y = self.sec(y, "cari_yeni_onay", "evet")
        self.assertEqual(y["cfo"]["soru"]["alan"], "cari_tipi")
        y = self.sec(y, "cari_tipi", "Usta")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        kayit = y["kayit_taslagi"]["kayitlar"][0]
        self.assertTrue(kayit["yeni_cari"])
        self.assertEqual(self.ledger(), [])
        self.assertNotIn("Mehmet", {c.ad for c in self.katalog().cariler})   # onaydan once tanıma eklenmedi
        self.onayla(y)
        self.assertEqual(self.ledger()[0]["Cari"], "Mehmet")
        self.assertIn("Mehmet", {c.ad for c in self.katalog().cariler})      # onaylı kayıtla tanımlandı

    def test_belirsiz_cari_secenekleri_adaylarla_sinirlanir(self):
        from openpyxl import load_workbook
        import master_data

        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        ws["AG20"], ws["AH20"], ws["AI20"] = "Gürsar İnşaat", "Usta", "Devrim Yıldız"
        ws["AG21"], ws["AH21"], ws["AI21"] = "Gürsar İnşaat", "Usta", "Devrim Kaya"
        wb["99_Ayarlar"]["AI13"] = None      # 'Devrim' tek basina artik yok; iki farkli Devrim kaldi
        wb["99_Ayarlar"]["AG13"] = None
        wb["99_Ayarlar"]["AH13"] = None
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()
        y = self.baslat("Gürsar Mimarlık Göktürk Okulu Devrim'e 5 bin TL ödeme yapıldı")
        self.assertEqual(y["cfo"]["soru"]["alan"], "cari")
        self.assertEqual({s["id"] for s in y["cfo"]["soru"]["secenekler"]} - {"__YENI__"}, {"Devrim Yıldız", "Devrim Kaya"})
        self.assertIn("Birden fazla olası", y["cfo"]["soru"]["metin"])

    def test_gecersiz_secim_reddedilir(self):
        y = self.baslat("10 bin ödeme yapıldı")
        y = self.sec(y, "faaliyet", "gursar_mimarlik")
        y = cfo_sohbet.cevapla(y["cfo"]["taslak_id"], "proje", "Uydurma Proje", AKTOR, "secim", self.katalog())
        self.assertEqual(y["cfo"]["soru"]["alan"], "proje")
        self.assertIn("bulunamadı", y["cevap"])

    def test_baska_kullanici_taslagi_degistiremez(self):
        y = self.baslat("10 bin ödeme yapıldı")
        with self.assertRaises(PermissionError):
            cfo_sohbet.cevapla(y["cfo"]["taslak_id"], "faaliyet", "gursar_mimarlik", "baska-kullanici", "secim", self.katalog())

    def test_iptal_excel_ye_yazmaz(self):
        y = self.baslat(DogalDilCozumlemeTestleri.TAM)
        cfo_sohbet.iptal(y["cfo"]["taslak_id"], AKTOR)
        self.assertEqual(self.ledger(), [])
        self.assertFalse(self.yedek.exists())


class OnayTekrarVeDenetimTestleri(CfoTestTemeli):
    def _tam_taslak(self, metin=DogalDilCozumlemeTestleri.TAM):
        self.gecmis_odeme_yaz()
        return self.baslat(metin)

    def test_onay_tek_kayit_yazar_denetim_orijinal_talimati_tasir(self):
        y = self._tam_taslak()
        satir_once = len(self.ledger())
        sonuc = self.onayla(y)
        self.assertTrue(sonuc["tamam"])
        satirlar = self.ledger()
        self.assertEqual(len(satirlar), satir_once + 1)
        yeni = satirlar[-1]
        self.assertEqual((yeni["Faaliyet"], yeni["Proje"], yeni["Cari"], yeni["Islem_Turu"], yeni["Hesap_Kasa"], yeni["Tutar"], yeni["PB"]),
                         ("Gürsar İnşaat", "Göktürk Okulu", "Devrim", "Usta/Tedarikçi Ödemesi", "Gürsar TL Banka", 10000, "TL"))
        self.assertEqual(yeni["Onay_Durumu"], "Onaylandı")
        olay = self.audit_olaylari("excel_islem")[-1]
        self.assertEqual(olay["orijinal_talimat"], DogalDilCozumlemeTestleri.TAM)
        self.assertEqual(olay["kayitlar"][0]["transaction_id"], yeni["Kayit_ID"])
        self.assertEqual(olay["taslak_id"], y["cfo"]["taslak_id"])
        self.assertTrue(olay["normalize_komut"].startswith("COUNTERPARTY_PAYMENT|gursar_mimarlik"))
        zincir = [o["olay"] for o in self.audit_olaylari() if o.get("taslak_id") == y["cfo"]["taslak_id"]]
        self.assertEqual(zincir[0], "cfo_taslak_olustu")
        self.assertIn("cfo_taslak_hazir", zincir)
        self.assertIn("excel_islem", zincir)
        self.assertEqual(cfo_sohbet.getir(y["cfo"]["taslak_id"])["durum"], "TAMAMLANDI")

    def test_ayni_taslak_iki_kez_onaylanamaz(self):
        y = self._tam_taslak()
        self.onayla(y)
        n = len(self.ledger())
        with self.assertRaises(ValueError):
            self.onayla(y)
        self.assertEqual(len(self.ledger()), n)

    def test_birebir_ayni_komut_tekrar_yazilmaz_yine_de_ile_ikinci_kayit_olur(self):
        y1 = self._tam_taslak()
        self.onayla(y1)
        n = len(self.ledger())
        # Ayni gun, ayni kullanici, birebir ayni cumle
        y2 = self.baslat(DogalDilCozumlemeTestleri.TAM)
        self.assertEqual(y2["niyet"], "cfo_tekrar_uyarisi")
        self.assertIsNone(y2["kayit_taslagi"])
        self.assertEqual(len(self.ledger()), n)                               # yeni kayıt/taslak yok
        self.assertEqual(self.audit_olaylari("cfo_taslak_tekrar_engellendi")[-1]["orijinal_talimat"],
                         DogalDilCozumlemeTestleri.TAM)
        # Bilinçli ikinci ödeme: açıkça "yine de ekle"
        y3 = self.baslat(DogalDilCozumlemeTestleri.TAM, yine_de=True)
        self.assertIsNotNone(y3["kayit_taslagi"])
        self.assertNotEqual(y3["cfo"]["belge_id"], y1["cfo"]["belge_id"])
        self.onayla(y3)
        satirlar = self.ledger()
        self.assertEqual(len(satirlar), n + 1)
        self.assertEqual(len({s["Kayit_ID"] for s in satirlar}), len(satirlar))      # kimlikler farklı
        self.assertEqual(len({s["Kaynak_Belge_ID"] for s in satirlar}), len(satirlar))

    def test_acik_taslak_varken_ayni_komut_ikinci_taslak_uretmez(self):
        y1 = self._tam_taslak()
        y2 = self.baslat(DogalDilCozumlemeTestleri.TAM)
        self.assertEqual(y2["cfo"]["taslak_id"], y1["cfo"]["taslak_id"])
        self.assertIn("zaten açık bir taslak", y2["cevap"])
        self.assertEqual(len(ui.listele("sohbet")), 1)

    def test_onay_sirasinda_excel_yazilamazsa_taslak_korunur_ve_tekrar_denenebilir(self):
        y = self._tam_taslak()
        gercek = depo.os.replace

        def kilitli(kaynak, hedef):
            if str(hedef) == str(self.ana):
                raise PermissionError(13, "kilitli")
            return gercek(kaynak, hedef)

        n = len(self.ledger())
        from unittest.mock import patch

        with patch.object(depo.os, "replace", side_effect=kilitli), self.assertRaises(RuntimeError) as cm:
            self.onayla(y)
        self.assertIn("Excel pencerelerini kapatın", str(cm.exception))
        self.assertEqual(len(self.ledger()), n)
        self.assertEqual(ui.ozet(y["kayit_taslagi"]["token"])["durum"], "ONAY_BEKLIYOR")
        self.onayla(y)                                                       # kilit kalkınca tekrar denenir
        self.assertEqual(len(self.ledger()), n + 1)

    def test_duzelt_deterministik_tutar_ve_hesabi_degistirir_yeni_sifre_ister(self):
        y = self._tam_taslak()
        eski_token = y["kayit_taslagi"]["token"]
        d = cfo_sohbet.duzelt(y["cfo"]["taslak_id"], "Düzelt: tutar 12 bin TL, hesap Gürsar TL Kasa", AKTOR, self.katalog())
        kayit = d["kayit_taslagi"]["kayitlar"][0]
        self.assertEqual((kayit["tutar"], kayit["hesap_kasa"]), (12000.0, "Gürsar TL Kasa"))
        self.assertNotEqual(d["kayit_taslagi"]["token"], eski_token)
        self.assertEqual(ui.ozet(eski_token)["durum"], "REDDEDILDI")         # eski taslak yazılamaz
        self.onayla(d)
        self.assertEqual(self.ledger()[-1]["Tutar"], 12000)

    def test_hicbir_adimda_llm_cagrilmaz(self):
        # cfo_test_temeli LLM ve ağ çağrılarını AssertionError ile yasaklar; tüm akış sorunsuz çalışmalı.
        y = self.baslat("10 bin ödeme yapıldı")
        for alan, deger in (("faaliyet", "gursar_mimarlik"), ("proje", "Göktürk Okulu"), ("islem", N.COUNTERPARTY_PAYMENT),
                            ("cari", "Devrim"), ("hesap", "Gürsar TL Banka")):
            y = self.sec(y, alan, deger)
        self.onayla(y)
        self.assertEqual(len(self.ledger()), 1)


if __name__ == "__main__":
    unittest.main()
