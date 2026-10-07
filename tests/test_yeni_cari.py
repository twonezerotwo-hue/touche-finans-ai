# -*- coding: utf-8 -*-
import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from excel_yazici import _yeni_cari_tanimlarini_yaz


class YeniCariTestleri(unittest.TestCase):
    def test_onayli_yeni_cari_master_tanimlara_eklenir(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "99_Ayarlar"
        veri = {"kayitlar": [{
            "Faaliyet": "Gürsar İnşaat", "Cari_Tipi": "Diğer",
            "Cari": "Cemre Şensoy", "_yeni_cari": True,
        }]}
        self.assertEqual(_yeni_cari_tanimlarini_yaz(wb, veri), 1)
        self.assertEqual(ws["AG13"].value, "Gürsar İnşaat")
        self.assertEqual(ws["AH13"].value, "Diğer")
        self.assertEqual(ws["AI13"].value, "Cemre Şensoy")
        self.assertEqual(_yeni_cari_tanimlarini_yaz(wb, veri), 0)


if __name__ == "__main__":
    unittest.main()
