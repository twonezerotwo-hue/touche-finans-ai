# -*- coding: utf-8 -*-
"""Uretim Smart Excel sozlesmesini taklit eden TEST calisma kitabi ureticisi.

Sayfa ve sutun yerlesimi ``config.FIRMALAR['touche_holding']`` ile birebir aynidir.
Gercek finansal veri icermez; yalnizca test icin uydurulmus ad ve tutarlar vardir.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

# (hesap_id, faaliyet, hesap, pb)
HESAPLAR = [
    ("HSP-T1", "Gürsar İnşaat", "Gürsar TL Banka", "TL"),
    ("HSP-T2", "Gürsar İnşaat", "Gürsar TL Kasa", "TL"),
    ("HSP-T3", "Gürsar İnşaat", "Gürsar EUR Banka", "EUR"),
    ("HSP-T4", "Touche Holding", "Holding TL Banka", "TL"),
]
ISLEM_TURLERI = [
    # ad, gelir, gider, alacak, borc, cek, nakit, ortak
    ("Hakediş / Gelir Tahakkuku", 1, 0, 1, 0, 0, 0, 0),
    ("Tahsilat", 0, 0, -1, 0, 0, 1, 0),
    ("Gider Tahakkuku", 0, 1, 0, 1, 0, 0, 0),
    ("Usta/Tedarikçi Ödemesi", 0, 0, 0, -1, 0, -1, 0),
    ("Ortak Cari Ödemesi", 0, 0, 0, 0, 0, -1, 1),
    ("Kredi Kullanımı", 0, 0, 0, 1, 0, 1, 0),
    ("Diğer Gelir", 1, 0, 0, 0, 0, 1, 0),
    ("Banka Masrafı", 0, 1, 0, 0, 0, -1, 0),
]
# (faaliyet, cari_tipi, cari)
CARILER = [
    ("Gürsar İnşaat", "Usta", "Devrim"),
    ("Gürsar İnşaat", "Usta", "Sebahattin Çakmak"),
    ("Gürsar İnşaat", "Müşteri", "İşveren"),
    ("Touche Holding", "Tedarikçi", "Ofis Market"),
]
PROJELER = [("Gürsar İnşaat", "Göktürk Okulu"), ("Gürsar İnşaat", "Bahçeşehir Villa")]
# kalem_id, faaliyet, proje, grup, ad, birim, metraj, satis_bf, maliyet_turu, maliyet_bf, durum, hakedis_id, gider_id
ISLERI = [
    ("KLM-001", "Gürsar İnşaat", "Göktürk Okulu", "Elektrik", "Elektrik", "adet", 1, 400_000, "Toplam sabit", 250_000, "Onaylı", "", ""),
    ("KLM-002", "Gürsar İnşaat", "Göktürk Okulu", "Taş işleri", "Basamak", "mtül", 100, 5_000, "Birim fiyat", 3_000, "Onaylı", "", ""),
    ("KLM-003", "Gürsar İnşaat", "Göktürk Okulu", "Saha betonu", "Saha Betonu", "m3", 50, 8_000, "Birim fiyat", 5_000, "Taslak", "", ""),
]


def tam_akilli_excel(yol, faaliyet_adi="Gürsar İnşaat", ledger_formul_satiri=40):
    """Uretim sozlesmesine uygun, bos ledger'li test Smart Excel'i yazar."""
    yol = Path(yol)
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ledger = wb.active
    ledger.title = "03_Kayıt Merkezi"
    inbox = wb.create_sheet("05_AI Belge Inbox")
    ayar = wb.create_sheet("99_Ayarlar")
    wb.create_sheet("04_Kredi & Vade")
    kalem = wb.create_sheet("06_Proje İş Kalemleri")

    # --- 03_Kayıt Merkezi: formul sablonu (veri sutunlari bos) ---
    for r in range(6, 6 + ledger_formul_satiri):
        ledger[f"S{r}"] = f'=IF($A{r}="","",1)'
        ledger[f"Y{r}"] = f'=IF($A{r}="","","OK")'
    for r in range(7, 7 + ledger_formul_satiri):
        inbox[f"Q{r}"] = f'=IF($A{r}="","","ONAY BEKLİYOR")'

    # --- 99_Ayarlar ---
    ayar["A13"], ayar["B13"], ayar["E13"] = faaliyet_adi, 1, "Aktif"
    ayar["A14"], ayar["B14"], ayar["E14"] = "Touche Holding", 1, "Aktif"
    for i, (fa, pr) in enumerate(PROJELER):
        ayar[f"H{13 + i}"] = faaliyet_adi if fa == "Gürsar İnşaat" else fa
        ayar[f"I{13 + i}"] = pr
    for i, (hid, fa, hesap, pb) in enumerate(HESAPLAR):
        r = 13 + i
        ayar[f"K{r}"], ayar[f"L{r}"], ayar[f"M{r}"], ayar[f"N{r}"], ayar[f"O{r}"] = (
            hid, faaliyet_adi if fa == "Gürsar İnşaat" else fa, hesap, pb, 0)
        ayar[f"P{r}"] = f'=IF($K{r}="","",IFERROR(VLOOKUP($N{r},$AD$3:$AE$7,2,FALSE),0))'
        ayar[f"Q{r}"] = f'=IF($K{r}="","",$O{r}*$P{r})'
    for i, satir in enumerate(ISLEM_TURLERI):
        r = 6 + i
        ayar[f"T{r}"] = satir[0]
        for sutun, deger in zip(("U", "V", "W", "X", "Y", "Z", "AA"), satir[1:]):
            ayar[f"{sutun}{r}"] = deger
    for i, (pb, kur) in enumerate((("TL", 1), ("USD", 40), ("EUR", 50), ("GBP", 60), ("USDT", 40))):
        ayar[f"AD{3 + i}"], ayar[f"AE{3 + i}"] = pb, kur
    for i, (fa, tip, cari) in enumerate(CARILER):
        r = 13 + i
        ayar[f"AG{r}"] = faaliyet_adi if fa == "Gürsar İnşaat" else fa
        ayar[f"AH{r}"], ayar[f"AI{r}"] = tip, cari

    # --- 06_Proje İş Kalemleri ---
    for i, k in enumerate(ISLERI):
        r = 6 + i
        kid, fa, proje, grup, ad, birim, metraj, satis, mturu, mbf, durum, hak, gid = k
        kalem[f"A{r}"] = kid
        kalem[f"B{r}"] = faaliyet_adi if fa == "Gürsar İnşaat" else fa
        kalem[f"C{r}"], kalem[f"D{r}"], kalem[f"E{r}"], kalem[f"F{r}"] = proje, grup, ad, birim
        kalem[f"G{r}"], kalem[f"H{r}"] = metraj, satis
        kalem[f"I{r}"] = f'=IF(OR($G{r}="",$H{r}=""),"",$G{r}*$H{r})'
        kalem[f"J{r}"], kalem[f"K{r}"] = mturu, mbf
        kalem[f"L{r}"] = f'=IF($J{r}="Toplam sabit",$K{r},$G{r}*$K{r})'
        kalem[f"M{r}"], kalem[f"N{r}"] = f"=I{r}-L{r}", f"=IFERROR(M{r}/I{r},0)"
        kalem[f"O{r}"], kalem[f"P{r}"], kalem[f"Q{r}"], kalem[f"R{r}"] = durum, "", hak, gid
    wb.save(yol)
    wb.close()
    return yol


def ledger_satirlari(yol, ilk=6, son=300):
    """Test icin ledger'daki dolu satirlari sozluk listesi olarak okur."""
    from openpyxl import load_workbook

    from config import FIRMALAR

    sut = FIRMALAR["touche_holding"]["ana_ledger"]["sutunlar"]
    wb = load_workbook(yol, data_only=False)
    try:
        ws = wb["03_Kayıt Merkezi"]
        sonuc = []
        for r in range(ilk, son + 1):
            if ws[f"{sut['Kayit_ID']}{r}"].value:
                sonuc.append({alan: ws[f"{c}{r}"].value for alan, c in sut.items()} | {"_satir": r})
        return sonuc
    finally:
        wb.close()


def bugun():
    return datetime.now().date()
