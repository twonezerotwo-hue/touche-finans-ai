# -*- coding: utf-8 -*-

import copy
import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from PIL import Image

sys.path.insert(0, r"C:\ai-excel\scripts")

import excel_yazici
import guvenlik
import main as belge_main
import telegram_bot
import invoice2data_adapter
import finans_cekirdegi
import akilli_excel
import excel_kalite
import recete_motoru
from belge_dosyasi import gecicileri_temizle, model_gorseli_hazirla
from telegram_bot import _duzeltme_parse
from telegram_sorgu import finans_ozeti, kayit_ara, son_kayitlar
from config import FIRMALAR
from dogrulayici import FAIL, PASS, REVIEW, deterministik_dogrula
from guvenlik import belge_sha256, finansal_fingerprint


def _holding_olustur(yol):
    wb = Workbook()
    ledger = wb.active
    ledger.title = "03_Kayıt Merkezi"
    inbox = wb.create_sheet("05_AI Belge Inbox")
    wb.create_sheet("99_Ayarlar")
    ledger["S6"] = '=IF($A6="","",1)'
    ledger["Y6"] = '=IF($A6="","","OK")'
    inbox["Q7"] = '=IF($A7="","","ONAY BEKLİYOR")'
    wb.save(yol)
    wb.close()


def _gursar_olustur(yol):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ödeme Geçmişi"
    ws["I6"] = "=G6+H6"
    wb.save(yol)
    wb.close()


class Phase1Testleri(unittest.TestCase):
    def setUp(self):
        self.orijinal_firmalar = copy.deepcopy(FIRMALAR)
        self.orijinal_audit_log = guvenlik.AUDIT_LOG
        self.orijinal_telegram_state = telegram_bot.TELEGRAM_STATE
        self.orijinal_bot_durum = {
            "firma": copy.deepcopy(telegram_bot.KULLANICI_FIRMA),
            "dosya": copy.deepcopy(telegram_bot.BEKLEYEN_DOSYALAR),
            "onay": copy.deepcopy(telegram_bot.BEKLEYEN_ONAYLAR),
            "duzeltme": copy.deepcopy(telegram_bot.DUZELTME_BEKLEYEN),
            "metin": copy.deepcopy(telegram_bot.BEKLEYEN_METINLER),
            "iptal": copy.deepcopy(telegram_bot.IPTAL_BEKLEYEN),
        }
        self.temp = tempfile.TemporaryDirectory()
        kok = Path(self.temp.name)
        guvenlik.AUDIT_LOG = kok / "audit.jsonl"
        telegram_bot.TELEGRAM_STATE = kok / "telegram_state.json"
        telegram_bot.KULLANICI_FIRMA.clear()
        telegram_bot.BEKLEYEN_DOSYALAR.clear()
        telegram_bot.BEKLEYEN_ONAYLAR.clear()
        telegram_bot.DUZELTME_BEKLEYEN.clear()
        telegram_bot.BEKLEYEN_METINLER.clear()
        telegram_bot.IPTAL_BEKLEYEN.clear()
        self.holding = kok / "holding.xlsx"
        self.gursar = kok / "gursar.xlsx"
        _holding_olustur(self.holding)
        _gursar_olustur(self.gursar)
        FIRMALAR["touche_holding"]["excel"] = str(self.holding)
        FIRMALAR["touche_holding"]["yedek_klasoru"] = kok / "yedek_holding"
        FIRMALAR["gursar_mimarlik"]["excel"] = str(self.gursar)
        FIRMALAR["gursar_mimarlik"]["yedek_klasoru"] = kok / "yedek_gursar"
        self.orijinal_finans_db = finans_cekirdegi.DB_PATH
        self.orijinal_sync_queue = finans_cekirdegi.SYNC_QUEUE_PATH
        self.orijinal_index_dir = akilli_excel.INDEX_DIR
        self.orijinal_recipe_dir = recete_motoru.RECIPE_DIR
        self.orijinal_recipe_log = recete_motoru.RUN_LOG
        finans_cekirdegi.DB_PATH = kok / "finans.db"
        finans_cekirdegi.SYNC_QUEUE_PATH = kok / "finans_sync_queue.jsonl"
        akilli_excel.INDEX_DIR = kok / "excel_index"
        recete_motoru.RECIPE_DIR = kok / "receteler"
        recete_motoru.RUN_LOG = kok / "recete_runs.jsonl"

    def tearDown(self):
        FIRMALAR.clear()
        FIRMALAR.update(self.orijinal_firmalar)
        guvenlik.AUDIT_LOG = self.orijinal_audit_log
        telegram_bot.TELEGRAM_STATE = self.orijinal_telegram_state
        telegram_bot.KULLANICI_FIRMA.clear()
        telegram_bot.KULLANICI_FIRMA.update(self.orijinal_bot_durum["firma"])
        telegram_bot.BEKLEYEN_DOSYALAR.clear()
        telegram_bot.BEKLEYEN_DOSYALAR.update(self.orijinal_bot_durum["dosya"])
        telegram_bot.BEKLEYEN_ONAYLAR.clear()
        telegram_bot.BEKLEYEN_ONAYLAR.update(self.orijinal_bot_durum["onay"])
        telegram_bot.DUZELTME_BEKLEYEN.clear()
        telegram_bot.DUZELTME_BEKLEYEN.update(self.orijinal_bot_durum["duzeltme"])
        telegram_bot.BEKLEYEN_METINLER.clear()
        telegram_bot.BEKLEYEN_METINLER.update(self.orijinal_bot_durum["metin"])
        telegram_bot.IPTAL_BEKLEYEN.clear()
        telegram_bot.IPTAL_BEKLEYEN.update(self.orijinal_bot_durum["iptal"])
        finans_cekirdegi.DB_PATH = self.orijinal_finans_db
        finans_cekirdegi.SYNC_QUEUE_PATH = self.orijinal_sync_queue
        akilli_excel.INDEX_DIR = self.orijinal_index_dir
        recete_motoru.RECIPE_DIR = self.orijinal_recipe_dir
        recete_motoru.RUN_LOG = self.orijinal_recipe_log
        self.temp.cleanup()

    def test_sha256_tam_ozet(self):
        belge = Path(self.temp.name) / "belge.bin"
        belge.write_bytes(b"touche-phase-1")
        self.assertEqual(belge_sha256(belge), hashlib.sha256(b"touche-phase-1").hexdigest())
        self.assertEqual(len(belge_sha256(belge)), 64)

    def test_finansal_fingerprint_kanonik(self):
        a = {"Cari": "  Örnek Firma ", "Tutar": "1.250,00", "Tarih": "25.09.2026"}
        b = {"Cari": "örnek   firma", "Tutar": 1250, "Tarih": "25.09.2026"}
        self.assertEqual(finansal_fingerprint(a, "x", "fatura"), finansal_fingerprint(b, "x", "fatura"))

    def test_dogrulama_pass_review_fail(self):
        md = {
            "faaliyetler": ["Gürsar İnşaat"], "projeler": ["Göktürk Okulu"],
            "hesaplar": [], "cariler": ["Usta A"], "cari_tipler": [],
            "islem_turleri": ["Usta ödemesi"], "kurlar": {},
            "maliyet_gruplari": ["Taş işleri"],
            "proje_hiyerarsisi": [{"faaliyet": "Gürsar İnşaat", "proje": "Göktürk Okulu"}],
            "hesap_hiyerarsisi": [], "cari_hiyerarsisi": [],
        }
        veri = {"kayitlar": [{"Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi", "Tutar": "1.000,00", "Proje": "Göktürk Okulu", "Maliyet_Grubu": "Taş işleri"}]}
        self.assertEqual(deterministik_dogrula(veri, "gursar_mimarlik", "odeme_dekontu", md)["durum"], PASS)
        veri2 = {"kayitlar": [{"Tarih": "25.09.2026", "Cari": "Bilinmeyen", "Islem_Turu": "Usta ödemesi", "Tutar": 1000}]}
        self.assertEqual(deterministik_dogrula(veri2, "gursar_mimarlik", "odeme_dekontu", md)["durum"], REVIEW)
        veri3 = {"kayitlar": [{"Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi", "Tutar": "bozuk"}]}
        self.assertEqual(deterministik_dogrula(veri3, "gursar_mimarlik", "odeme_dekontu", md)["durum"], FAIL)

    def test_gursar_holding_atomik_entegrasyon(self):
        veri = {"kayitlar": [{
            "Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 1000, "Banka_Masrafi": 10, "Proje": "Göktürk Okulu",
            "Maliyet_Grubu": "Taş işleri", "_guven": .95,
        }]}
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri, "TEST-001")
        self.assertTrue(sonuc["basari"], sonuc)
        self.assertEqual(sonuc["holding_yazilan"], 2)

        wb = load_workbook(self.gursar, data_only=False)
        self.assertIsNone(wb["Ödeme Geçmişi"]["E6"].value)
        wb.close()

        wb = load_workbook(self.holding, data_only=False)
        ws = wb["03_Kayıt Merkezi"]
        self.assertEqual(ws["C6"].value, "Gürsar İnşaat")
        self.assertEqual(ws["Z6"].value, "TEST-001-001")
        self.assertTrue(str(ws["S6"].value).startswith("=IF"))
        self.assertTrue(str(ws["Y6"].value).startswith("=IF"))
        self.assertEqual(ws["Z7"].value, "TEST-001-001-MASRAF")
        wb.close()

    def test_master_replace_hatasinda_dosyalar_degismez(self):
        once_gursar = hashlib.sha256(self.gursar.read_bytes()).hexdigest()
        once_holding = hashlib.sha256(self.holding.read_bytes()).hexdigest()
        veri = {"kayitlar": [{
            "Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 1000, "Proje": "Göktürk Okulu", "Maliyet_Grubu": "Taş işleri",
        }]}
        gercek_replace = excel_yazici.os.replace
        def master_hedefte_hata(kaynak, hedef):
            if Path(hedef) == self.holding:
                raise OSError("simule atomik replace hatasi")
            return gercek_replace(kaynak, hedef)

        with patch.object(excel_yazici.os, "replace", side_effect=master_hedefte_hata):
            sonuc = excel_yazici.yaz("gursar_mimarlik", veri, "ROLLBACK-001")
        self.assertFalse(sonuc["basari"])
        self.assertEqual(hashlib.sha256(self.gursar.read_bytes()).hexdigest(), once_gursar)
        self.assertEqual(hashlib.sha256(self.holding.read_bytes()).hexdigest(), once_holding)

    def test_touche_inbox_formulu_korunur(self):
        veri = {"kayitlar": [{
            "Tarih": "25.09.2026", "Faaliyet": "Gürsar İnşaat", "Proje": "Göktürk Okulu",
            "Cari_Tipi": "Usta", "Cari": "Usta A", "Islem_Turu": "Gider Tahakkuku",
            "Tutar": 1000, "PB": "TL", "_hash": "a" * 64, "_guven": .9,
        }]}
        sonuc = excel_yazici.yaz("touche_holding", veri, "TEST-002")
        self.assertTrue(sonuc["basari"], sonuc)
        wb = load_workbook(self.holding, data_only=False)
        ws = wb["05_AI Belge Inbox"]
        self.assertEqual(ws["A7"].value, "TEST-002-001")
        self.assertEqual(ws["S7"].value, "a" * 64)
        self.assertTrue(str(ws["Q7"].value).startswith("=IF"))
        wb.close()
        tekrar = excel_yazici.yaz("touche_holding", veri, "TEST-003")
        self.assertFalse(tekrar["basari"])
        self.assertIn("Mukerrer SHA-256", tekrar["hata"])

    def test_hazirlama_excel_yazmaz_ve_duzeltme_parse_edilir(self):
        belge = Path(self.temp.name) / "dekont.jpg"
        belge.write_bytes(b"sahte-gorsel")
        ai_ciktisi = {"kayitlar": [{
            "Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 1000, "Proje": "Göktürk Okulu", "Maliyet_Grubu": "Taş işleri",
        }]}
        with patch.object(belge_main, "modele_sor", return_value=ai_ciktisi):
            hazir = belge_main.belge_hazirla(belge, "gursar_mimarlik", "odeme_dekontu")
        self.assertTrue(hazir["basari"], hazir)
        wb = load_workbook(self.gursar, data_only=False)
        self.assertIsNone(wb["Ödeme Geçmişi"]["A6"].value)
        wb.close()
        self.assertEqual(_duzeltme_parse("Tutar=1250,50\nCari=Yeni Cari")["Cari"], "Yeni Cari")

    def test_pdf_gorsele_cevrilir(self):
        pdf = Path(self.temp.name) / "dekont.pdf"
        Image.new("RGB", (800, 1100), "white").save(pdf, "PDF")
        model_yolu, geciciler = model_gorseli_hazirla(pdf)
        try:
            self.assertTrue(model_yolu.exists())
            with Image.open(model_yolu) as img:
                self.assertGreater(img.width, 0)
                self.assertGreater(img.height, 0)
        finally:
            gecicileri_temizle(geciciler)

    def test_metin_kaydi_onaydan_once_excel_yazmaz(self):
        ai_ciktisi = {"kayitlar": [{
            "Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 2500, "Proje": "Göktürk Okulu", "Maliyet_Grubu": "Taş işleri",
        }]}
        with patch.object(belge_main, "metne_sor", return_value=ai_ciktisi):
            hazir = belge_main.metin_hazirla("Gürsar için 2500 TL usta ödemesi", "gursar_mimarlik")
        self.assertTrue(hazir["basari"], hazir)
        self.assertTrue(hazir["belge_id"].startswith("GURS-TXT-"))
        wb = load_workbook(self.gursar, data_only=False)
        self.assertIsNone(wb["Ödeme Geçmişi"]["A6"].value)
        wb.close()

    def test_invoice2data_sonucu_mevcut_fatura_semasina_donusur(self):
        ham = {
            "issuer": "Örnek Tedarik A.Ş.", "supplier_vkn": "1234567890",
            "date": "26.09.2026", "invoice_number": "ABC2026000001",
            "amount": 1250.50, "amount_tax": 208.42, "currency": "TRY",
            "template_name": "ornek.yml",
        }
        veri = invoice2data_adapter._sonucu_donustur(ham)
        self.assertIsNotNone(veri)
        self.assertEqual(veri["kayitlar"][0]["Toplam"], 1250.50)
        self.assertEqual(veri["kayitlar"][0]["Tedarikci"], "Örnek Tedarik A.Ş.")
        self.assertTrue(veri["_cikarim"]["ozet_kaydi"])

    def test_invoice2data_qwenden_once_kullanilir(self):
        belge = Path(self.temp.name) / "fatura.pdf"
        belge.write_bytes(b"sahte-pdf")
        veri = {"kayitlar": [{
            "Tarih": "26.09.2026", "Fatura_No": "ABC2026000001",
            "Tedarikci": "Örnek Tedarik A.Ş.", "VKN": "1234567890",
            "Urun": "Fatura toplamı", "Miktar": 1, "Birim_Fiyat": 1250.5,
            "Toplam": 1250.5,
        }], "_cikarim": {"motor": "invoice2data", "sablon": "ornek.yml", "ozet_kaydi": True}}
        sonuc = {"eslesti": True, "veri": veri, "sablon": "ornek.yml", "ozet_kaydi": True}
        with patch.object(belge_main, "invoice2data_oku", return_value=sonuc), \
                patch.object(belge_main, "modele_sor") as qwen:
            hazir = belge_main.belge_hazirla(belge, "touche_holding", "fatura")
        self.assertTrue(hazir["basari"], hazir)
        self.assertEqual(hazir["cikarim_motoru"], "invoice2data")
        self.assertEqual(hazir["dogrulama"]["durum"], REVIEW)
        qwen.assert_not_called()

    def test_invoice2data_eslesmezse_qwen_yedegi_calisir(self):
        belge = Path(self.temp.name) / "fatura.pdf"
        belge.write_bytes(b"sahte-pdf")
        ai_ciktisi = {"kayitlar": [{
            "Tarih": "26.09.2026", "Fatura_No": "ABC2026000002",
            "Tedarikci": "Başka Tedarikçi", "Urun": "Malzeme",
            "Miktar": 2, "Birim_Fiyat": 500, "Toplam": 1000,
        }]}
        with patch.object(belge_main, "invoice2data_oku", return_value={"eslesti": False, "sebep": "eslesme_yok"}), \
                patch.object(belge_main, "modele_sor", return_value=ai_ciktisi) as qwen:
            hazir = belge_main.belge_hazirla(belge, "touche_holding", "fatura")
        self.assertTrue(hazir["basari"], hazir)
        self.assertEqual(hazir["cikarim_motoru"], "qwen2.5vl")
        qwen.assert_called_once()

    def test_merkezi_defter_dengeli_donem_kilitli_ve_ters_kayitli(self):
        hazir = {"belge_id": "DB-001", "firma_key": "gursar_mimarlik", "tur": "odeme_dekontu", "veri": {"kayitlar": [{
            "Tarih": "26.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 1000, "Banka_Masrafi": 10, "Proje": "Göktürk Okulu",
        }]}}
        sonuc = finans_cekirdegi.finans_kaydi_yaz(hazir, "test")
        self.assertEqual(sonuc["fis_sayisi"], 1)
        kontrol = finans_cekirdegi.sistem_kontrolu()
        self.assertEqual(kontrol["dengesiz_fis"], 0)
        self.assertEqual(finans_cekirdegi.yonetim_ozeti("gursar_mimarlik")["gider"], 1010)
        finans_cekirdegi.donem_durum_degistir(2026, 9, True, "test")
        with self.assertRaises(ValueError):
            finans_cekirdegi.fis_taslaklari({**hazir, "belge_id": "DB-002"})
        finans_cekirdegi.donem_durum_degistir(2026, 9, False, "test")
        finans_cekirdegi.belgeyi_ters_kaydet("DB-001", "test")
        self.assertEqual(finans_cekirdegi.yonetim_ozeti("gursar_mimarlik")["net"], 0)

    def test_butce_onay_politikasi_ve_banka_mutabakati(self):
        finans_cekirdegi.butce_kaydet("gursar_mimarlik", "Göktürk Okulu", "740", 2026, 9, 5000, "test")
        rapor = finans_cekirdegi.butce_raporu("gursar_mimarlik", 2026, 9)
        self.assertEqual(rapor[0]["butce"], 5000)
        finans_cekirdegi.politika_ayarla(1000, 2, "test")
        self.assertFalse(finans_cekirdegi.onay_oyu_ekle("ONAY-1", 1, 1500)["tamam"])
        self.assertTrue(finans_cekirdegi.onay_oyu_ekle("ONAY-1", 2, 1500)["tamam"])
        hazir = {"belge_id": "BANK-1", "firma_key": "gursar_mimarlik", "tur": "odeme_dekontu", "veri": {"kayitlar": [{
            "Tarih": "26.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi", "Tutar": 1000,
        }]}}
        finans_cekirdegi.finans_kaydi_yaz(hazir, "test")
        csv_yolu = Path(self.temp.name) / "ekstre.csv"
        csv_yolu.write_text("Tarih,Aciklama,Tutar,Referans\n26.09.2026,Usta odemesi,-1000,BANK-1\n", encoding="utf-8")
        self.assertEqual(finans_cekirdegi.banka_csv_aktar(csv_yolu), 1)
        self.assertEqual(finans_cekirdegi.mutabakat_ozeti()["ESLESTI"]["adet"], 1)

    def test_acilis_mizan_ve_nakit_tahmini(self):
        fis = finans_cekirdegi.acilis_bakiyesi_kaydet(
            "gursar_mimarlik", "102", "01.09.2026", 10000, "test"
        )
        self.assertTrue(fis.startswith("JRN-OPEN-"))
        hazir = {"belge_id": "VADE-1", "firma_key": "gursar_mimarlik", "tur": "fatura", "veri": {"kayitlar": [{
            "Tarih": "26.09.2026", "Vade": "30.09.2026", "Tedarikci": "Tedarikçi A", "Toplam": 2500,
        }]}}
        finans_cekirdegi.finans_kaydi_yaz(hazir, "test")
        mizan = finans_cekirdegi.mizan("01.09.2026", "30.09.2026", "gursar_mimarlik")
        self.assertTrue(any(r["hesap"] == "102" and r["borc"] == 10000 for r in mizan))
        tahmin = finans_cekirdegi.nakit_tahmini(13, "gursar_mimarlik", "26.09.2026")
        self.assertEqual(sum(r["cikis"] for r in tahmin), 2500)

    def test_finans_senkron_kuyrugu_kaydi_kaybetmez(self):
        hazir = {"belge_id": "SYNC-1", "firma_key": "gursar_mimarlik", "tur": "fatura", "veri": {"kayitlar": [{
            "Tarih": "26.09.2026", "Tedarikci": "Tedarikçi A", "Toplam": 750,
        }]}}
        finans_cekirdegi.finans_kaydi_kuyruga_al(hazir, "test", "geçici hata")
        self.assertEqual(finans_cekirdegi.sistem_kontrolu()["senkron_bekleyen"], 1)
        sonuc = finans_cekirdegi.senkron_kuyrugunu_isle()
        self.assertEqual(sonuc, {"islenen": 1, "kalan": 0})
        self.assertEqual(finans_cekirdegi.sistem_kontrolu()["fisler"], 1)

    def test_excel_parser_indeksi_ve_guvenli_plan(self):
        indeks = akilli_excel.indeks_olustur("touche_holding")
        self.assertGreater(indeks["total_chunks"], 0)
        self.assertEqual(akilli_excel.indeks_olustur("touche_holding")["sha256"], indeks["sha256"])
        self.assertTrue(akilli_excel.plan_dogrula({"actions": [{"name": "search_chunks"}]}))
        with self.assertRaises(ValueError):
            akilli_excel.plan_dogrula({"actions": [{"name": "write_cells", "code": "x"}]})
        cevap = akilli_excel._finansal_ozet_cevabi(
            "touche_holding", "2026 Eylül gelir ve gider nedir?", []
        )
        self.assertEqual(cevap["guven"], 100)

    def test_excel_kalite_ref_hatasini_bulur(self):
        wb = load_workbook(self.holding)
        wb["03_Kayıt Merkezi"]["B6"] = "=#REF!"
        wb.save(self.holding)
        wb.close()
        sonuc = excel_kalite.yerel_excel_kontrol("touche_holding")
        self.assertEqual(sonuc["durum"], "FAIL")
        self.assertTrue(any("B6" in x for x in sonuc["sorunlar"]))

    def test_recete_izin_listesi_kod_calistirmayi_reddeder(self):
        recete_motoru.RECIPE_DIR.mkdir(parents=True)
        (recete_motoru.RECIPE_DIR / "zararli.json").write_text(json.dumps({
            "name": "zararli", "steps": [{"action": "finance_check", "code": "import os"}]
        }), encoding="utf-8")
        with self.assertRaises(ValueError):
            recete_motoru.recete_yukle("zararli")

    def test_holding_metin_kaydi_onay_sonrasi_ledger_yazar(self):
        ai_ciktisi = {"kayitlar": [{
            "Tarih": "25.09.2026", "Faaliyet": "Gürsar İnşaat", "Proje": "Göktürk Okulu",
            "Cari_Tipi": "Usta", "Cari": "Usta A", "Islem_Turu": "Gider Tahakkuku",
            "Hesap_Kasa": "Gürsar TL Banka", "Tutar": 3000, "PB": "TL",
        }]}
        with patch.object(belge_main, "metne_sor", return_value=ai_ciktisi):
            hazir = belge_main.metin_hazirla("Holding Gürsar gider tahakkuku 3000 TL", "touche_holding")
        self.assertTrue(hazir["basari"], hazir)
        self.assertEqual(hazir["veri"]["_hedef"], "ana_ledger")
        basarili, mesaj = belge_main.hazirlanan_belgeyi_yaz(hazir, "test")
        self.assertTrue(basarili, mesaj)
        wb = load_workbook(self.holding, data_only=False)
        ws = wb["03_Kayıt Merkezi"]
        self.assertEqual(ws["C6"].value, "Gürsar İnşaat")
        self.assertEqual(ws["I6"].value, 3000)
        self.assertTrue(str(ws["S6"].value).startswith("=IF"))
        self.assertTrue(str(ws["Y6"].value).startswith("=IF"))
        wb.close()

    def test_sorgu_ve_guvenli_iptal(self):
        veri = {"kayitlar": [{
            "Tarih": "25.09.2026", "Cari": "Usta A", "Islem_Turu": "Usta ödemesi",
            "Tutar": 1000, "Proje": "Göktürk Okulu", "Maliyet_Grubu": "Taş işleri",
        }]}
        sonuc = excel_yazici.yaz("gursar_mimarlik", veri, "CANCEL-001")
        self.assertTrue(sonuc["basari"], sonuc)
        self.assertEqual(len(son_kayitlar("gursar_mimarlik", 5)), 1)
        self.assertEqual(len(kayit_ara("gursar_mimarlik", "Usta A")), 1)
        self.assertEqual(finans_ozeti("gursar_mimarlik")["aktif_kayit"], 1)
        iptal = excel_yazici.belge_kaydini_iptal_et("CANCEL-001")
        self.assertTrue(iptal["basari"], iptal)
        wb = load_workbook(self.gursar, data_only=False)
        self.assertIsNone(wb["Ödeme Geçmişi"]["M6"].value)
        wb.close()
        wb = load_workbook(self.holding, data_only=False)
        self.assertEqual(wb["03_Kayıt Merkezi"]["V6"].value, "İptal")
        self.assertEqual(wb["03_Kayıt Merkezi"]["AC6"].value, "Reddedildi")
        wb.close()

    def test_telegram_durumu_kalici(self):
        telegram_bot.KULLANICI_FIRMA[12345] = "gursar_mimarlik"
        telegram_bot.BEKLEYEN_METINLER[12345] = "100 TL gider"
        telegram_bot._durum_kaydet()
        telegram_bot.KULLANICI_FIRMA.clear()
        telegram_bot.BEKLEYEN_METINLER.clear()
        telegram_bot._durum_yukle()
        self.assertEqual(telegram_bot.KULLANICI_FIRMA[12345], "gursar_mimarlik")
        self.assertEqual(telegram_bot.BEKLEYEN_METINLER[12345], "100 TL gider")

    def test_telegram_uygulamasi_tum_handlerlari_yukler(self):
        app = telegram_bot.uygulama_olustur()
        self.assertGreaterEqual(sum(len(grup) for grup in app.handlers.values()), 16)


if __name__ == "__main__":
    unittest.main()
