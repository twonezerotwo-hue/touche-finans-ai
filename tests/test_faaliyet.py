"""Faaliyet yaşam döngüsü ve tek-master kayıt hata senaryoları."""
import asyncio
import copy
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch, AsyncMock
from openpyxl import load_workbook

sys.path.insert(0, r'C:\ai-excel\scripts')
import faaliyet_yonetimi as yonetim
import telegram_bot as bot
import guvenlik
import config
from config import FIRMALAR
from dogrulayici import deterministik_dogrula, FAIL
from excel_yazici import yaz, belge_kaydini_iptal_et
from master_data import master_data_oku
from telegram_sorgu import finans_ozeti
from test_phase1 import _holding_olustur, _gursar_olustur


class FaaliyetTestleri(unittest.TestCase):
    def setUp(self):
        self.saved = copy.deepcopy(FIRMALAR)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.holding = self.root/'holding.xlsx'
        self.gursar = self.root/'gursar.xlsx'
        _holding_olustur(self.holding)
        _gursar_olustur(self.gursar)
        w = load_workbook(self.holding)
        w['99_Ayarlar']['AH13'] = 'Tedarikçi'
        w['99_Ayarlar']['T6'] = 'Tahsilat'
        w.save(self.holding)
        w.close()
        FIRMALAR['touche_holding']['excel'] = str(self.holding)
        FIRMALAR['touche_holding']['yedek_klasoru'] = self.root/'backup_holding'
        FIRMALAR['gursar_mimarlik']['excel'] = str(self.gursar)
        FIRMALAR['gursar_mimarlik']['yedek_klasoru'] = self.root/'backup_gursar'
        self.patches = [patch.object(yonetim,'FAALIYET_REGISTRY',self.root/'registry.json'),
                        patch.object(yonetim,'FIRMALAR_DIR',self.root/'firms'),
                        patch.object(guvenlik,'AUDIT_LOG',self.root/'audit.jsonl'),
                        patch.object(bot,'TELEGRAM_STATE',self.root/'telegram.json'),
                        patch.object(bot,'YONETIM_AKIS',{}),
                        patch.object(bot,'KULLANICI_FIRMA',{}),
                        patch.object(bot,'YETKILI_CHAT_IDS',{123})]
        for p in self.patches:
            p.start()
        self.addCleanup(self.cleanup)

    def cleanup(self):
        for p in reversed(self.patches):
            p.stop()
        FIRMALAR.clear()
        FIRMALAR.update(self.saved)
        self.temp.cleanup()

    def create(self):
        key = yonetim.faaliyet_olustur('Touche Enerji','Silivri GES','Enerji TL Banka','TL',75,'Örnek Ortak')
        yonetim.cari_olustur(key,'Panel Firması','Tedarikçi')
        return key

    def record(self, **kw):
        return {'kayitlar':[{'Tarih':'26.09.2026','Proje':'Silivri GES','Cari_Tipi':'Tedarikçi',
                            'Cari':'Panel Firması','Islem_Turu':'Diğer Gider','Tutar':1250.5,
                            '_hash':'sha256-example', **kw}]}

    def test_onboarding_registry_project_and_semantics(self):
        key = self.create()
        md = master_data_oku(key)
        self.assertEqual(md['projeler'],['Silivri GES'])
        self.assertEqual(md['hesaplar'],['Enerji TL Banka'])
        self.assertIn('Tahsilat',md['islem_turleri'])
        registry = json.loads((self.root/'registry.json').read_text(encoding='utf-8'))
        self.assertEqual(registry['faaliyetler'][key]['holding_payi'],.75)
        yonetim.proje_olustur(key,'Çorlu GES')
        self.assertIn('Çorlu GES',master_data_oku(key)['projeler'])
        with patch.object(config,'FAALIYET_REGISTRY',self.root/'registry.json'):
            del FIRMALAR[key]
            config.faaliyetleri_yukle()
            self.assertEqual(FIRMALAR[key]['adapter'],'standart_v1')
        w = load_workbook(self.holding)
        self.assertEqual(w['99_Ayarlar']['A13'].value,'Touche Enerji')
        self.assertEqual(w['99_Ayarlar']['B13'].value,.75)
        self.assertEqual(w['99_Ayarlar']['I14'].value,'Çorlu GES')
        self.assertTrue(w['03_Kayıt Merkezi']['S6'].value.startswith('='))
        w.close()

    def test_duplicate_onboarding_and_invalid_input_leave_files_unchanged(self):
        self.create()
        before = self.holding.read_bytes()
        with self.assertRaises(ValueError):
            self.create()
        for pay in ('nan',-2,101):
            with self.assertRaises(ValueError):
                yonetim.faaliyet_olustur('Yeni Faaliyet','Proje A','Banka A',holding_payi=pay)
        with self.assertRaises(ValueError):
            yonetim.faaliyet_olustur('=HYPERLINK(A1)','Proje A','Banka A')
        self.assertEqual(before,self.holding.read_bytes())

    def test_onboarding_rollback_on_registry_replace_error(self):
        before = self.holding.read_bytes()
        original = os.replace
        def failing(src,dst):
            if Path(dst) == self.root/'registry.json':
                raise PermissionError('registry locked')
            return original(src,dst)
        with patch.object(yonetim.os,'replace',side_effect=failing), self.assertRaises(PermissionError):
            self.create()
        self.assertEqual(before,self.holding.read_bytes())
        self.assertFalse((self.root/'registry.json').exists())
        self.assertEqual(list(self.root.rglob('*_hesap.xlsx')), [])
        self.assertNotIn('touche_enerji',FIRMALAR)

    def test_write_single_master_duplicates_and_cancel(self):
        key = self.create()
        veri = self.record(Aciklama='=1+1')
        result = yaz(key,veri,'DOC-ENERGY-1')
        self.assertTrue(result['basari'],result)
        self.assertEqual(result['holding_yazilan'],1)
        before = self.holding.read_bytes()
        self.assertFalse(yaz(key,self.record(),'DOC-ENERGY-1')['basari'])
        self.assertEqual(before,self.holding.read_bytes())
        w = load_workbook(self.holding)
        self.assertEqual(w['03_Kayıt Merkezi']['I6'].value,1250.5)
        self.assertEqual(w['03_Kayıt Merkezi']['C6'].value,'Touche Enerji')
        self.assertEqual(w['03_Kayıt Merkezi']['X6'].data_type,'s')
        self.assertEqual(w['03_Kayıt Merkezi']['B6'].value.year,2026)
        w.close()
        self.assertEqual(finans_ozeti(key)['toplam_tutar'],1250.5)
        self.assertEqual(belge_kaydini_iptal_et('DOC-ENERGY-1')['etkilenen'],1)
        self.assertEqual(finans_ozeti(key)['toplam_tutar'],0)

    def test_invalid_routing_fx_and_nonfinite_rejected(self):
        key = self.create()
        for fields in ({'Faaliyet':'Başka Şirket'}, {'Proje':'Başka Proje'},
                       {'Hesap_Kasa':'Başka Hesap'}, {'PB':'USD'}, {'Kur':'nan'},
                       {'Tutar':'NaN'}, {'Tarih':'bozuk'}, {'Islem_Turu':'tanımsız'}):
            with self.subTest(fields=fields):
                self.assertEqual(deterministik_dogrula(self.record(**fields),key,'finans_belgesi')['durum'],FAIL)
        result = yaz(key,self.record(PB='USD',Kur=35),'FX-1')
        self.assertTrue(result['basari'],result)
        self.assertEqual(finans_ozeti(key)['toplam_tutar'],1250.5*35)

    def test_write_rollback_if_holding_locked(self):
        key = self.create()
        before = self.holding.read_bytes()
        original = os.replace
        def failing(src,dst):
            if Path(dst) == self.holding:
                raise PermissionError('holding locked')
            return original(src,dst)
        with patch('excel_yazici.os.replace',side_effect=failing):
            result = yaz(key,self.record(),'FAIL-1')
        self.assertFalse(result['basari'])
        self.assertEqual(before,self.holding.read_bytes())

    def test_gursar_project_without_changing_revenue_schedule(self):
        yonetim.proje_olustur('gursar_mimarlik','Yeni Şantiye')
        self.assertIn('Yeni Şantiye',master_data_oku('gursar_mimarlik')['projeler'])
        with self.assertRaises(ValueError):
            yonetim.proje_olustur('gursar_mimarlik','Yeni Şantiye')

    def test_telegram_wizard_confirmation_and_restart(self):
        update = SimpleNamespace(effective_chat=SimpleNamespace(id=123),
                                 message=SimpleNamespace(reply_text=AsyncMock(),text=''),callback_query=None)
        context = SimpleNamespace(args=[],bot=SimpleNamespace(send_message=AsyncMock()))
        async def run():
            await bot.yeni_faaliyet_komutu(update,context)
            for value in ['Touche Enerji','Silivri GES','TL','Enerji TL Banka','75','Örnek Ortak']:
                update.message.text = value
                await bot.duzeltme_metni_al(update,context)
            self.assertFalse((self.root/'registry.json').exists())
            token = bot.YONETIM_AKIS[123]['token']
            bot.YONETIM_AKIS.clear()
            bot._durum_yukle()
            self.assertEqual(bot.YONETIM_AKIS[123]['token'],token)
            query = SimpleNamespace(message=SimpleNamespace(chat_id=123),data='kurulum:'+token,
                                    answer=AsyncMock(),edit_message_text=AsyncMock())
            update.callback_query = query
            await bot.buton_cevabi(update,context)
            self.assertEqual(bot.KULLANICI_FIRMA[123],'touche_enerji')
            self.assertNotIn(123,bot.YONETIM_AKIS)
            before = self.holding.read_bytes()
            await bot.buton_cevabi(update,context)
            self.assertEqual(before,self.holding.read_bytes())
        asyncio.run(run())

    def test_telegram_one_time_pairing(self):
        owner = self.root/'owner.json'
        codefile = self.root/'pairing.txt'
        with patch.object(bot,'TELEGRAM_OWNER_FILE',owner), patch.object(bot,'TELEGRAM_PAIRING_CODE',codefile):
            bot.YETKILI_CHAT_IDS.clear()
            bot._eslestirme_kodu_hazirla()
            code = codefile.read_text(encoding='utf-8')
            self.assertEqual(len(code),8)
            update = SimpleNamespace(effective_chat=SimpleNamespace(id=123),
                                     message=SimpleNamespace(reply_text=AsyncMock()))
            async def run():
                await bot.yetkilendir_komutu(update,SimpleNamespace(args=['wrong']))
                self.assertFalse(owner.exists())
                await bot.yetkilendir_komutu(update,SimpleNamespace(args=[code]))
            asyncio.run(run())
            self.assertEqual(json.loads(owner.read_text())['chat_ids'],[123])
            self.assertFalse(codefile.exists())
            self.assertTrue(bot._yetkili(123))

if __name__ == '__main__':
    unittest.main()
