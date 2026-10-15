# -*- coding: utf-8 -*-
"""CFO okuma sorgulari: dogru kapsam, deterministik sayilar, SIFIR mutasyon."""
import hashlib
import unittest

import cfo_sorgu
import excel_deposu as depo
import finans_cekirdegi as cekirdek
import finans_danismani as fd
from cfo_test_temeli import BUGUN, CfoTestTemeli


class SorguTestleri(CfoTestTemeli):
    def _fis(self, belge, tarih, islem, tutar, cari, hesap="Gürsar TL Banka", proje="Göktürk Okulu",
             faaliyet="gursar_mimarlik"):
        kayit = {"Tarih": tarih, "Islem_Turu": islem, "Tutar": tutar, "PB": "TL", "Kur": 1, "Cari": cari,
                 "Proje": proje, "Hesap_Kasa": hesap, "Aciklama": f"{islem} {cari}"}
        cekirdek.finans_kaydi_yaz({"belge_id": belge, "firma_key": faaliyet, "tur": "finans_belgesi",
                                   "veri": {"kayitlar": [kayit]}}, "test")

    def setUp(self):
        super().setUp()
        self._fis("H1", "2026-09-01", "Hakediş / Gelir Tahakkuku", 500_000, "İşveren")
        self._fis("T1", "2026-09-10", "Tahsilat", 200_000, "İşveren")
        self._fis("G1", "2026-09-15", "Gider Tahakkuku", 300_000, "Devrim")
        self._fis("O1", "2026-09-20", "Usta/Tedarikçi Ödemesi", 100_000, "Devrim")
        self._fis("O2", "2026-10-02", "Usta/Tedarikçi Ödemesi", 40_000, "Sebahattin Çakmak")
        # Baska proje ve baska faaliyet: Goktürk sorularina karismamali
        self._fis("D1", "2026-09-21", "Usta/Tedarikçi Ödemesi", 9_999, "Devrim", proje="Bahçeşehir Villa")
        self._fis("H9", "2026-09-22", "Usta/Tedarikçi Ödemesi", 7_777, "Ofis Market", proje="Genel",
                  faaliyet="touche_holding")

    def sor(self, metin):
        yanit = cfo_sorgu.cevapla(metin, self.katalog(), BUGUN)
        self.assertIsNotNone(yanit, metin)
        return yanit

    def tutar(self, yanit):
        return yanit["sorgu"]["deger_kurus"] // 100

    # --- spec sorulari ---
    def test_gokturk_odemeleri_proje_kapsaminda_holding_toplami_degil(self):
        """Eski davranis: soru faaliyetsiz oldugu icin TUM holding'in odemesini donduruyordu."""
        y = self.sor("Göktürk okulunda ne kadar ödeme yaptık?")
        self.assertEqual(self.tutar(y), 140_000)                          # 100.000 + 40.000 (Bahçeşehir/Holding hariç)
        self.assertEqual(y["sorgu"]["filtre"]["proje"], "Göktürk Okulu")
        self.assertIn("Gürsar Mimarlık / Göktürk Okulu", y["cevap"])
        self.assertIn("140.000,00", y["cevap"])
        eski = fd.asistana_sor("Göktürk okulunda ne kadar ödeme yaptık?", llm_kullan=False)
        self.assertNotIn("140.000,00", eski["cevap"])                     # eski hat kapsamı bilmiyordu (regresyon kanıtı)

    def test_kazanc_tahsilatla_karistirilmaz(self):
        y = self.sor("Göktürk okulundan ne kadar kazandık?")
        self.assertEqual(self.tutar(y), 500_000)
        self.assertEqual(y["sorgu"]["ek"]["tahsil_edilen_kurus"], 200_000 * 100)
        self.assertIn("Kazanç ile tahsilat aynı şey değildir", y["cevap"])

    def test_gokturk_tahsilat_sorusu_proje_kapsaminda_ve_sifir_mutasyon(self):
        ana_once = hashlib.sha256(self.ana.read_bytes()).hexdigest()
        y = self.sor("Göktürk okulundan ne kadar tahsil ettik?")
        self.assertEqual(self.tutar(y), 200_000)                          # yalniz T1; hakedis tahakkuku (500.000) tahsilat degildir
        self.assertEqual(y["sorgu"]["filtre"]["proje"], "Göktürk Okulu")
        self.assertIn("200.000,00", y["cevap"])
        self.assertEqual(hashlib.sha256(self.ana.read_bytes()).hexdigest(), ana_once)

    def test_kalan_alacak(self):
        y = self.sor("Göktürk okulunda ne kadar alacağımız kaldı?")
        self.assertEqual(self.tutar(y), 300_000)

    def test_faaliyet_toplam_gider(self):
        y = self.sor("Gürsar'ın toplam gideri ne?")
        self.assertEqual(self.tutar(y), 300_000)
        self.assertEqual(y["sorgu"]["filtre"]["faaliyet"], "gursar_mimarlik")

    def test_cari_odemesi(self):
        y = self.sor("Devrim'e toplam ne kadar ödeme yapıldı?")
        self.assertEqual(self.tutar(y), 100_000 + 9_999)                   # tum projeler, yalniz Devrim
        y = self.sor("Göktürk okulunda Devrim'e toplam ne kadar ödeme yaptık?")
        self.assertEqual(self.tutar(y), 100_000)

    def test_bu_ay_holding_toplam_odeme(self):
        y = self.sor("Bu ay Holding toplam ne kadar ödeme yaptı?")
        # Ekim 2026: yalniz O2 (40.000) ; Holding = konsolide
        self.assertEqual(self.tutar(y), 40_000)
        self.assertIn("Holding geneli", y["cevap"])
        self.assertEqual(y["sorgu"]["donem"], "Ekim 2026")
        y = self.sor("Eylül ayında Holding toplam ne kadar ödeme yaptı?")
        self.assertEqual(self.tutar(y), 100_000 + 9_999 + 7_777)

    def test_kar_ve_ayrinti_listesi(self):
        y = self.sor("Göktürk okulunun kârı ne kadar?")
        self.assertEqual(self.tutar(y), 500_000 - 300_000)
        self.assertIn("Kâr nakit akışı değildir", y["cevap"])
        y = self.sor("Göktürk okulu ödemelerini kalem kalem göster")
        self.assertIn("1. 20.09.2026 · Devrim", y["cevap"])
        self.assertIn("2. 02.10.2026 · Sebahattin Çakmak", y["cevap"])
        y = self.sor("Göktürk okuluna ait bütün Devrim ödemelerini yaz")
        self.assertEqual(y["niyet"], "cfo_sorgu")
        self.assertEqual(self.tutar(y), 100_000)
        self.assertIn("1. 20.09.2026 · Devrim", y["cevap"])

    def test_belirsiz_ad_tahmin_edilmez(self):
        from openpyxl import load_workbook
        import master_data

        wb = load_workbook(self.ana)
        ws = wb["99_Ayarlar"]
        ws["AG20"], ws["AH20"], ws["AI20"] = "Gürsar İnşaat", "Usta", "Devrim Yıldız"
        ws["AI13"] = None
        wb.save(self.ana)
        wb.close()
        master_data.onbellegi_temizle()
        y = cfo_sorgu.cevapla("Devrim'e toplam ne kadar ödeme yapıldı?", self.katalog(), BUGUN)
        self.assertEqual(y["niyet"], "cfo_sorgu_netlestir")
        self.assertIn("tahmin yürütmedim", y["cevap"])

    # --- gunluk dil: donem ifadeleri, yon ("bize borclu" / "alacakli"), kismi proje adi ---
    def test_donem_ifadeleri_gercek_veriyle(self):
        y = self.sor("Son 30 günde Göktürk okulunda ne kadar ödeme yaptık?")      # 13.09 - 12.10
        self.assertEqual(self.tutar(y), 140_000)
        self.assertEqual(y["sorgu"]["filtre"]["bas"], "2026-09-13")
        y = self.sor("1 Ekim ile 10 Ekim arasında ne kadar ödeme yaptık")
        self.assertEqual((self.tutar(y), y["sorgu"]["filtre"]["bas"], y["sorgu"]["filtre"]["bit"]),
                         (40_000, "2026-10-01", "2026-10-10"))
        y = self.sor("Eylülde Göktürk'e ne kadar harcadık")                       # ay adi apostrofsuz ekli
        self.assertEqual(self.tutar(y), 100_000)
        y = self.sor("Bugün ne kadar ödeme yaptık")
        self.assertEqual(self.tutar(y), 0)
        self.assertIn("kayıt bulunamadı", y["cevap"])

    def test_kismi_proje_adi_kabul_edilir_kapsam_ve_not_cevapta_yazar(self):
        y = self.sor("Göktürk'ün kârı ne durumda")
        self.assertEqual(y["sorgu"]["filtre"]["proje"], "Göktürk Okulu")
        self.assertEqual(self.tutar(y), 200_000)                                   # 500.000 gelir - 300.000 gider
        self.assertIn("Not: Sorudaki ad Göktürk Okulu olarak anlaşıldı", y["cevap"])

    def test_karsi_taraf_bize_borcluysa_kalan_alacaktir(self):
        y = self.sor("İşveren bize ne kadar borçlu")
        self.assertEqual(y["sorgu"]["metrik"], "KALAN_ALACAK")
        self.assertEqual((y["sorgu"]["filtre"]["cari"], self.tutar(y)), ("İşveren", 300_000))   # 500.000 - 200.000

    def test_alacakli_ise_biz_borcluyuz_acik_borctur(self):
        y = self.sor("Devrim ne kadar alacaklı")
        self.assertEqual((y["sorgu"]["metrik"], y["sorgu"]["filtre"]["cari"]), ("ACIK_BORC", "Devrim"))
        y2 = self.sor("Devrim'e borcumuz var mı")
        self.assertEqual(self.tutar(y2), self.tutar(y))

    def test_isveren_ne_kadar_odedi_tahsilattir_iade_ise_odemedir(self):
        y = self.sor("İşveren ne kadar ödedi")
        self.assertEqual((y["sorgu"]["metrik"], self.tutar(y)), ("TAHSIL", 200_000))      # T1: 200.000 tahsilat
        self.assertIn("tahsilat olarak anlaşıldı", y["cevap"])
        y = self.sor("İşverene ne kadar ödedik")                                            # yonelme eki: bizim odememiz
        self.assertEqual((y["sorgu"]["metrik"], self.tutar(y)), ("ODENEN", 0))

    def test_gelir_ve_ciro_kazanilan_gelirdir(self):
        self.assertEqual(self.tutar(self.sor("Gelirimiz ne kadar")), 500_000)
        self.assertEqual(self.tutar(self.sor("gokturk okulunun cirosu ne")), 500_000)

    def test_desteklenmeyen_soru_none_doner_danismana_devredilir(self):
        self.assertIsNone(cfo_sorgu.cevapla("Bugün hava nasıl?", self.katalog(), BUGUN))

    # --- sifir mutasyon ---
    def _durum(self):
        satir_sayilari = {}
        with cekirdek._baglanti() as con:
            for t in ("fisler", "fis_satirlari", "belge_butunlugu"):
                satir_sayilari[t] = con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
        ana = hashlib.sha256(self.ana.read_bytes()).hexdigest()
        yedek = self.yedek.exists()
        return satir_sayilari, ana, yedek, (hashlib.sha256(self.yedek.read_bytes()).hexdigest() if yedek else None)

    def test_sorgular_excel_ve_defteri_degistirmez(self):
        once = self._durum()
        for soru in ("Göktürk okulunda ne kadar ödeme yaptık?", "Devrim'e toplam ne kadar ödeme yapıldı?",
                     "Bu ay Holding toplam ne kadar ödeme yaptı?", "Göktürk okulundan ne kadar kazandık?",
                     "Göktürk okulunda ne kadar alacağımız kaldı?", "Gürsar'ın toplam gideri ne?"):
            self.sor(soru)
        self.assertEqual(self._durum(), once)
        self.assertFalse(self.yedek.exists())                              # yedek bile ilerlemedi
        self.assertEqual(depo.gecici_dosyalar(), [])

    def test_sorgu_kayit_taslagi_olusturmaz(self):
        from cfo_niyet import analiz_et

        n = analiz_et("Göktürk okulunda ne kadar ödeme yaptık?", self.katalog(), BUGUN)
        self.assertEqual(n.tur, "SORGU")
        n = analiz_et("Göktürk okuluna ait bütün Devrim ödemelerini yaz", self.katalog(), BUGUN)
        self.assertEqual(n.tur, "SORGU")

    def test_kayit_gecmisi_neden_burada(self):
        """'Bu ödeme neden burada?' -> denetim zincirinden hikaye."""
        y = self.baslat("Göktürk Okulu Devrim 10 bin ödeme yapıldı")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.onayla(y)
        kayit_id = self.ledger()[0]["Kayit_ID"]
        g = self.sor(f"{kayit_id} kaydı neden burada?")
        self.assertEqual(g["niyet"], "cfo_sorgu")
        self.assertIn("Göktürk Okulu Devrim 10 bin ödeme yapıldı", g["cevap"])     # orijinal talimat
        self.assertIn("Excel'e yazıldı", g["cevap"])
        self.assertIn("Talimat alındı", g["cevap"])


if __name__ == "__main__":
    unittest.main()
