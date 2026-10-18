"""PR-02: sentetik Excel/SQLite tutarlilik ve kurtarma sozlesmeleri.

Bu dosya canli Excel, OneDrive, gercek SQLite veya kimlik bilgisi kullanmaz.
Her test kendi gecici workbook/DB/audit/queue dosyalarinda calisir.
"""
from __future__ import annotations

import json
import os
import shutil
import sqlite3
import subprocess
import sys
import textwrap
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

import excel_defter_mutabakati as mutabakat
import finans_cekirdegi as cekirdek
import main
import pytest
import ui_belge_servisi as ui
from cfo_test_temeli import CfoTestTemeli


class FinanceConsistencyContracts(CfoTestTemeli):
    """Tum yollar sentetik fiksturle sinirlidir; uretim yoluna erisim yoktur."""

    def setUp(self):
        super().setUp()
        self.queue = self.root / "finans_sync_queue.jsonl"
        self._queue_patch = patch.object(cekirdek, "SYNC_QUEUE_PATH", self.queue)
        self._queue_patch.start()

    def tearDown(self):
        self._queue_patch.stop()
        super().tearDown()

    def belge(self, belge_id="SYN-DOC-001", tutar=10_000, cari="Devrim",
              islem="Usta/Tedarikçi Ödemesi", pb="TL", kur=1, tarih="2026-10-12"):
        kayit = {
            "Tarih": tarih,
            "Faaliyet": self.faaliyet_adi,
            "Proje": "Göktürk Okulu",
            "Cari_Tipi": "Usta",
            "Cari": cari,
            "Islem_Turu": islem,
            "Hesap_Kasa": "Gürsar TL Banka",
            "Tutar": tutar,
            "PB": pb,
            "Kur": kur,
            "Aciklama": "sentetik PR-02 kaydi",
        }
        return {
            "basari": True,
            "foto_yolu": None,
            "model_yolu": None,
            "kaynak_metin": "sentetik PR-02",
            "firma_key": "gursar_mimarlik",
            "tur": "finans_belgesi",
            "veri": {"kayitlar": [kayit]},
            "belge_id": belge_id,
            "belge_hash": (belge_id * 64)[:64],
            "dogrulama": {"durum": "PASS", "skor": 100, "hatalar": [], "uyarilar": []},
            "gecici_dosyalar": [],
        }

    def db_scalar(self, sql, params=()):
        with cekirdek._baglanti() as con:
            row = con.execute(sql, params).fetchone()
            return row[0] if row else None

    @pytest.mark.xfail(
        strict=True,
        reason="PR-03 atomik commit: Excel yazildi/DB kuyruga alindi durumunda basari donmemeli",
    )
    def test_a01_excel_once_db_basarisizsa_tamamlandi_denememeli(self):
        """Excel-first/DB-failure yolu PR-03 atomik sonucunu simdiden kilitler."""
        hazir = self.belge("A01-DB-FAIL")
        with patch.object(main, "finans_kaydi_yaz", side_effect=sqlite3.OperationalError("synthetic db down")), \
             patch.object(main.time, "sleep"):
            ok, _ = main.hazirlanan_belgeyi_yaz(hazir, "test")

        # Mevcut davranis Excel'i yazip kuyruga aliyor ve basari donduruyor.
        # PR-03'te bu kontrat atomik hata/geri alma olarak degisecek.
        assert not ok

    def test_a02_db_on_kontrol_basarisizsa_excel_yazilmaz(self):
        hazir = self.belge("A02-PRE-DB-FAIL")
        with patch.object(main, "fis_taslaklari", side_effect=sqlite3.OperationalError("db unavailable")):
            ok, mesaj = main.hazirlanan_belgeyi_yaz(hazir, "test")
        assert not ok
        assert "Muhasebe kontrol" in mesaj
        assert self.ledger() == []
        assert not self.queue.exists()

    def test_a03_deneme_sirasinda_sqlite_rollback_yarim_fis_birakmaz(self):
        hazir = self.belge("A03-ROLLBACK")
        cekirdek.veritabanini_hazirla()
        with cekirdek._baglanti() as con:
            con.execute(
                "CREATE TRIGGER synthetic_abort BEFORE INSERT ON fis_satirlari "
                "BEGIN SELECT RAISE(ABORT, 'synthetic interruption'); END;"
            )
        with pytest.raises(sqlite3.IntegrityError):
            cekirdek.finans_kaydi_yaz(hazir, "test")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler") == 0

    def test_a04_excel_only_durum_mutabakatta_eksik_defterdir(self):
        self.gecmis_odeme_yaz("A04-EXCEL-ONLY", tutar=12_000)
        sonuc = mutabakat.kontrol()
        assert not sonuc["tutarli"]
        assert len(sonuc["eksik_defterde"]) == 1

    def test_a05_db_only_durum_mutabakatta_defter_fazlasidir(self):
        cekirdek.finans_kaydi_yaz(self.belge("A05-DB-ONLY", tutar=13_000), "test")
        sonuc = mutabakat.kontrol()
        assert not sonuc["tutarli"]
        assert any(x.startswith("A05-DB-ONLY-") for x in sonuc["defterde_fazla"])

    def test_a06_excel_tutari_degiserse_tutar_farki_raporlanir(self):
        hazir = self.belge("A06-AMOUNT")
        self.gecmis_odeme_yaz("A06-AMOUNT", tutar=10_000)
        cekirdek.finans_kaydi_yaz(hazir, "test")
        from openpyxl import load_workbook

        wb = load_workbook(self.ana)
        try:
            ws = wb["03_Kayıt Merkezi"]
            for row in range(6, 300):
                if ws[f"W{row}"].value == "A06-AMOUNT":
                    ws[f"I{row}"] = 11_000
                    break
            wb.save(self.ana)
        finally:
            wb.close()
        sonuc = mutabakat.kontrol()
        assert len(sonuc["tutar_farki"]) == 1

    def test_a07_kuyruk_kurtarma_ve_yeniden_baslatma(self):
        hazir = self.belge("A07-QUEUE")
        cekirdek.finans_kaydi_kuyruga_al(hazir, "test", "temporary")
        assert self.queue.exists()
        assert cekirdek.senkron_kuyrugunu_isle() == {"islenen": 1, "kalan": 0}
        assert not self.queue.exists()
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("A07-QUEUE",)) == 1

    def test_a08_kuyruk_hatasi_kalan_satiri_korur(self):
        self.queue.write_text("{bozuk json}\n", encoding="utf-8")
        sonuc = cekirdek.senkron_kuyrugunu_isle()
        assert sonuc == {"islenen": 0, "kalan": 1}
        assert self.queue.read_text(encoding="utf-8").strip() == "{bozuk json}"

    def test_a09_audit_incomplete_transaction_kanitini_tutar(self):
        hazir = self.belge("A09-AUDIT")
        with patch.object(main, "finans_kaydi_yaz", side_effect=RuntimeError("db down")), \
             patch.object(main.time, "sleep"):
            main.hazirlanan_belgeyi_yaz(hazir, "test")
        olaylar = self.audit_olaylari("muhasebe_fisi_kuyrukta")
        assert olaylar and olaylar[-1]["belge_id"] == "A09-AUDIT"

    def test_a10_ayni_belge_retry_idempotenttir(self):
        hazir = self.belge("A10-IDEMP", tutar=14_000)
        cekirdek.finans_kaydi_yaz(hazir, "test")
        cekirdek.finans_kaydi_yaz(hazir, "test-retry")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("A10-IDEMP",)) == 1

    def test_a11_ayni_tutar_farkli_belge_ayri_odemedir(self):
        cekirdek.finans_kaydi_yaz(self.belge("A11-FIRST", tutar=15_000), "test")
        cekirdek.finans_kaydi_yaz(self.belge("A11-SECOND", tutar=15_000), "test")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE cari=?", ("Devrim",)) == 2

    def test_a12_kuyruk_iki_kez_islenince_cift_posting_olmaz(self):
        hazir = self.belge("A12-QUEUE-IDEMP", tutar=16_000)
        cekirdek.finans_kaydi_kuyruga_al(hazir, "test", "retry")
        cekirdek.finans_kaydi_kuyruga_al(hazir, "test", "retry")
        sonuc = cekirdek.senkron_kuyrugunu_isle()
        assert sonuc == {"islenen": 2, "kalan": 0}
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("A12-QUEUE-IDEMP",)) == 1

    def test_a13_doviz_kuru_yoksa_kayit_reddedilir(self):
        with pytest.raises(ValueError, match="gecerli kur"):
            cekirdek.finans_kaydi_yaz(self.belge("A13-FX", tutar=100, pb="USD", kur=0), "test")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler") == 0

    def test_a14_bilinmeyen_islem_turu_sessizce_gider_olmaz(self):
        with pytest.raises(ValueError, match="eşlemesi"):
            cekirdek.finans_kaydi_yaz(self.belge("A14-UNKNOWN", islem="Bilinmeyen İşlem"), "test")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler") == 0

    def test_a15_es_zamanli_onayda_tek_posting_olur(self):
        taslak = self.baslat("Göktürk Okulu Devrim 17 bin ödeme yapıldı")
        taslak = self.sec(taslak, "hesap", "Gürsar TL Banka")
        token = taslak["kayit_taslagi"]["token"]
        sonuclar = []

        def onay():
            try:
                sonuclar.append(ui.onayla(token, "thread", "test-parola"))
            except (RuntimeError, ValueError) as exc:  # ikinci onay bekleyen durumda olmamalidir
                sonuclar.append(exc)

        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(lambda _: onay(), range(2)))
        basarili = sum(1 for x in sonuclar if isinstance(x, dict) and x.get("tamam") is True)
        assert basarili == 1
        assert self.db_scalar("SELECT COUNT(*) FROM fisler") == 1

    def test_fin_p01_mevcut_excel_db_ayrisini_kanitlar(self):
        """PASS diagnostic: mevcut Excel-ilk akisinda DB kopugu sessiz ayrisma uretiyor."""
        hazir = self.belge("FIN-P01-DIAG", tutar=18_000)
        with patch.object(main, "finans_kaydi_yaz", side_effect=sqlite3.OperationalError("synthetic db down")), \
             patch.object(main.time, "sleep"):
            ok, mesaj = main.hazirlanan_belgeyi_yaz(hazir, "diagnostic")

        # Bu test mevcut davranisi kanitlar; PR-03 atomik durum makinesi gelince
        # beklenen sonuc degisecek ve A01 strict XFAIL yeniden PASS olacaktir.
        assert ok is True
        assert mesaj == ""
        assert [x for x in self.ledger() if str(x.get("Kaynak_Belge_ID", "")).startswith("FIN-P01-DIAG-")]
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("FIN-P01-DIAG",)) == 0
        assert self.queue.exists()
        kuyruk = [json.loads(x) for x in self.queue.read_text(encoding="utf-8").splitlines() if x.strip()]
        assert kuyruk and kuyruk[-1]["hazir"]["belge_id"] == "FIN-P01-DIAG"
        tamam = self.audit_olaylari("belge_tamamlandi")
        assert tamam and tamam[-1]["muhasebe_fisi"] == 0
        rapor = mutabakat.kontrol()
        assert not rapor["tutarli"]
        assert len(rapor["eksik_defterde"]) == 1
        assert rapor["eksik_defterde"][0]["excel_tutar"] == 18_000

    def _crash_process_files(self, belge_id):
        """Alt surec icin tamamen sentetik env ve script hazirlar."""
        child_root = self.root / "subprocess-crash"
        child_root.mkdir(parents=True, exist_ok=True)
        ana = child_root / "ANA.xlsx"
        yedek = child_root / "ANA_backup.xlsx"
        shutil.copy2(self.ana, ana)
        shutil.copy2(self.ana, yedek)
        hazir_path = child_root / "hazir.json"
        hazir_path.write_text(json.dumps(self.belge(belge_id, tutar=19_000), ensure_ascii=False), encoding="utf-8")
        script = child_root / "child.py"
        repo = Path(__file__).resolve().parents[1]
        script.write_text(textwrap.dedent(f"""
            import json, os, sqlite3, sys
            from pathlib import Path
            sys.path.insert(0, {str(repo / 'scripts')!r})

            mode = sys.argv[1]
            belge_id = {belge_id!r}
            if mode == "crash":
                import main
                hazir = json.loads(Path({str(hazir_path)!r}).read_text(encoding="utf-8"))
                original_yaz = main.yaz
                def crash_after_excel(*args, **kwargs):
                    result = original_yaz(*args, **kwargs)
                    if result.get("basari"):
                        os._exit(47)
                    return result
                main.yaz = crash_after_excel
                main.hazirlanan_belgeyi_yaz(hazir, "subprocess-crash")
                raise AssertionError("crash enjeksiyonu tetiklenmedi")

            import config
            from openpyxl import load_workbook
            workbook = load_workbook(config.ana_excel_yolu(), read_only=True, data_only=False)
            try:
                ws = workbook["03_Kayıt Merkezi"]
                kaynak_index = ord(config.FIRMALAR["touche_holding"]["ana_ledger"]["sutunlar"]["Kaynak_Belge_ID"]) - ord("A")
                excel_count = sum(1 for row in ws.iter_rows(min_row=6, max_col=kaynak_index + 1, values_only=True)
                                  if str(row[kaynak_index] or "").startswith(f"{{belge_id}}-"))
            finally:
                workbook.close()
            db_count = 0
            if config.FINANS_DB.exists():
                with sqlite3.connect(config.FINANS_DB) as con:
                    db_count = con.execute("SELECT COUNT(*) FROM fisler WHERE belge_id=?", (belge_id,)).fetchone()[0]
            queue = Path(config.LOGLAR) / "finans_sync_queue.jsonl"
            audit = Path(config.AUDIT_LOG)
            olaylar = []
            if audit.exists():
                olaylar = [json.loads(line).get("olay") for line in audit.read_text(encoding="utf-8").splitlines() if line.strip()]
            print(json.dumps({{"excel_count": excel_count, "db_count": db_count,
                               "queue_exists": queue.exists(), "completed": "belge_tamamlandi" in olaylar,
                               "audit_events": olaylar}}, ensure_ascii=False))
        """), encoding="utf-8")
        env = os.environ.copy()
        env.update({
            "TOUCHE_BASE_DIR": str(child_root),
            "TOUCHE_MAIN_WORKBOOK": str(ana),
            "TOUCHE_BACKUP_WORKBOOK": str(yedek),
        })
        return child_root, script, env

    def test_b01_ayri_surec_cokmesinde_excel_db_ayrisini_kanitlar(self):
        """PASS diagnostic: Excel replacement sonrasi process kesilirse ikinci surec ayrismayi gorur."""
        belge_id = "B01-SUBPROCESS-CRASH"
        child_root, _script, env = self._crash_process_files(belge_id)
        result = subprocess.run(
            [sys.executable, str(child_root / "child.py"), "crash", belge_id],
            cwd=str(self.root), env=env, capture_output=True, text=True, timeout=20, check=False,
        )
        assert result.returncode == 47, result.stderr or result.stdout
        inspect = subprocess.run(
            [sys.executable, str(child_root / "child.py"), "inspect", belge_id],
            cwd=str(self.root), env=env, capture_output=True, text=True, timeout=20, check=False,
        )
        assert inspect.returncode == 0, inspect.stderr or inspect.stdout
        durum = json.loads(inspect.stdout.strip().splitlines()[-1])
        assert durum["excel_count"] == 1
        assert durum["db_count"] == 0
        assert durum["queue_exists"] is False
        assert durum["completed"] is False

    @pytest.mark.xfail(
        strict=True,
        reason="PR-03 process crash recovery markeri ve PENDING_RECOVERY kuyrugu henuz yok",
    )
    def test_b02_ayri_surec_cokmesi_kalici_kurtarma_isareti_uretmelidir(self):
        """Gelecek kontrati: ani kesilme sonrasi sistem kaydi kurtarma icin isaretlemeli."""
        belge_id = "B02-SUBPROCESS-RECOVERY"
        child_root, _script, env = self._crash_process_files(belge_id)
        result = subprocess.run(
            [sys.executable, str(child_root / "child.py"), "crash", belge_id],
            cwd=str(self.root), env=env, capture_output=True, text=True, timeout=20, check=False,
        )
        assert result.returncode == 47
        inspect = subprocess.run(
            [sys.executable, str(child_root / "child.py"), "inspect", belge_id],
            cwd=str(self.root), env=env, capture_output=True, text=True, timeout=20, check=False,
        )
        durum = json.loads(inspect.stdout.strip().splitlines()[-1])
        assert durum["queue_exists"] is True or "PENDING_RECOVERY" in durum["audit_events"]

    def test_a16_ayni_belge_id_farkli_icerik_cakismasi_reddedilir(self):
        ilk = self.belge("A16-CONTENT-CONFLICT", tutar=20_000)
        cekirdek.finans_kaydi_yaz(ilk, "test")
        ikinci = self.belge("A16-CONTENT-CONFLICT", tutar=21_000)
        with pytest.raises(ValueError, match="farklı içerikle"):
            cekirdek.finans_kaydi_yaz(ikinci, "retry-farkli")
        assert self.db_scalar("SELECT COUNT(*) FROM fisler WHERE belge_id=?", ("A16-CONTENT-CONFLICT",)) == 1

