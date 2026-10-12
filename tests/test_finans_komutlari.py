# -*- coding: utf-8 -*-
import sqlite3
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import finans_danismani as fd
import finans_komutlari as fk
from _ortam import AnaExcelYonlendirme


class FinansKomutlariTestleri(unittest.TestCase):
    def test_turkce_tarih_yil_yoksa_siradaki_tarihi_secer(self):
        self.assertEqual(fk.turkce_tarih_coz("13 Ekim", date(2026, 9, 30)), date(2026, 10, 13))
        self.assertEqual(fk.turkce_tarih_coz("13 Ekim", date(2026, 11, 1)), date(2027, 10, 13))

    def test_vade_komutu_iki_acik_borcu_cozer(self):
        pano = {"odemeler": [
            {"faaliyet": "gursar_mimarlik", "proje": "Göktürk Okulu", "cari": "Taş işleri", "vade": None},
            {"faaliyet": "gursar_mimarlik", "proje": "Göktürk Okulu", "cari": "Saha betonu", "vade": None},
        ]}
        komut = fd._vade_komutu_coz(
            "Taş işleri ve Saha betonu vadelerini 13 Ekim olarak kaydet", pano, date(2026, 9, 30)
        )
        self.assertEqual(komut["vade"], date(2026, 10, 13))
        self.assertEqual([x["cari"] for x in komut["hedefler"]], ["Taş işleri", "Saha betonu"])

    def test_vade_db_ve_tek_master_excelde_birlikte_guncellenir(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            db = root / "finans.db"
            con = sqlite3.connect(db)
            con.executescript("""
                CREATE TABLE fisler (fis_id TEXT PRIMARY KEY, belge_id TEXT, faaliyet_key TEXT,
                  proje TEXT, cari TEXT, aciklama TEXT, vade TEXT, durum TEXT);
                CREATE TABLE fis_satirlari (fis_id TEXT, hesap_kodu TEXT, borc_kurus INTEGER, alacak_kurus INTEGER);
                INSERT INTO fisler VALUES ('F1','B1','gursar_mimarlik','Göktürk Okulu','Taş işleri','Taş işleri',NULL,'MUHASEBELESTI');
                INSERT INTO fisler VALUES ('F2','B2','gursar_mimarlik','Göktürk Okulu','Saha betonu','Saha betonu',NULL,'MUHASEBELESTI');
                INSERT INTO fis_satirlari VALUES ('F1','320',0,18624000);
                INSERT INTO fis_satirlari VALUES ('F2','320',0,14000000);
            """)
            con.commit(); con.close()

            excel = root / "master.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "03_Kayıt Merkezi"
            for satir, cari in ((6, "Taş işleri"), (7, "Saha betonu")):
                ws[f"A{satir}"] = f"GID-{satir}"
                ws[f"C{satir}"] = "Gürsar İnşaat"
                ws[f"D{satir}"] = "Göktürk Okulu"
                ws[f"F{satir}"] = cari
                ws[f"G{satir}"] = "Gider Tahakkuku"
                ws[f"V{satir}"] = "Aktif"
                ws[f"X{satir}"] = cari
                ws[f"AC{satir}"] = "Onaylandı"
            wb.save(excel); wb.close()

            with patch.object(fk._cekirdek, "DB_PATH", db), \
                 patch.object(fk, "veritabani_yedekle", return_value=root / "db-yedek.db"), \
                 patch.object(fk, "audit_yaz"), \
                 AnaExcelYonlendirme(excel, root / "yedek" / "master - YEDEK.xlsx"):
                sonuc = fk.acik_borc_vadesi_guncelle([
                    {"faaliyet": "gursar_mimarlik", "proje": "Göktürk Okulu", "cari": "Taş işleri"},
                    {"faaliyet": "gursar_mimarlik", "proje": "Göktürk Okulu", "cari": "Saha betonu"},
                ], date(2026, 10, 13), aktor="test")

            con = sqlite3.connect(db)
            self.assertEqual(con.execute("SELECT COUNT(*) FROM fisler WHERE vade='2026-10-13'").fetchone()[0], 2)
            con.close()
            wb = load_workbook(excel)
            self.assertEqual(wb["03_Kayıt Merkezi"]["U6"].value, datetime(2026, 10, 13))
            self.assertEqual(wb["03_Kayıt Merkezi"]["U7"].value, datetime(2026, 10, 13))
            wb.close()
            self.assertEqual(sonuc["excel_satirlari"], [6, 7])
            # Tek yedek: islem oncesi surum; baska Excel kopyasi yok.
            self.assertTrue((root / "yedek" / "master - YEDEK.xlsx").exists())
            self.assertEqual([p.name for p in (root / "yedek").iterdir()], ["master - YEDEK.xlsx"])


if __name__ == "__main__":
    unittest.main()
