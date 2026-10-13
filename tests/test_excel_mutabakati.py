# -*- coding: utf-8 -*-
"""ANA Excel <-> merkezi defter mutabakati: salt-okunur, fark gorunur, yanlis cift sayim yok."""
import hashlib
import unittest

from openpyxl import load_workbook

import excel_defter_mutabakati as mut
import excel_operasyon
import finans_cekirdegi as cekirdek
import proje_finansi as pf
from cfo_test_temeli import CfoTestTemeli


class MutabakatTestleri(CfoTestTemeli):
    def _uygulama_kaydi(self):
        y = self.baslat("Göktürk Okulu Devrim 10 bin ödeme yapıldı")
        y = self.sec(y, "hesap", "Gürsar TL Banka")
        self.onayla(y)

    def _elle_satir(self, kayit_id="MAN-001", tutar=5000, islem="Usta/Tedarikçi Ödemesi", kaynak=None,
                    durum="Aktif", belge_ref=None, cari="Devrim"):
        """Kullanicinin Excel'e elle satir girmesini taklit eder (uygulama disi)."""
        wb = load_workbook(self.ana)
        ws = wb["03_Kayıt Merkezi"]
        satir = next(r for r in range(6, 300) if not ws[f"A{r}"].value)
        for sutun, deger in {"A": kayit_id, "B": "2026-09-05", "C": self.faaliyet_adi, "D": "Göktürk Okulu",
                             "E": "Usta", "F": cari, "G": islem, "H": "Gürsar TL Banka", "I": tutar,
                             "J": "TL", "K": 1, "V": durum, "W": belge_ref, "Z": kaynak}.items():
            if deger is not None:
                ws[f"{sutun}{satir}"] = deger
        wb.save(self.ana)
        wb.close()
        return satir

    def test_uygulamanin_yazdigi_kayit_tutarlidir(self):
        self._uygulama_kaydi()
        r = mut.kontrol()
        self.assertEqual(r["excel_aktif_satir"], 1)
        self.assertTrue(r["tutarli"], r)

    def test_elle_girilen_satir_deftere_yansimiyor_diye_raporlanir(self):
        self._uygulama_kaydi()
        self._elle_satir()
        r = mut.kontrol()
        self.assertFalse(r["tutarli"])
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["MAN-001"])
        self.assertTrue(any("merkezi defterde yok" in u for u in mut.uyari_metinleri(r)))

    def test_holding_ozeti_farki_uyari_olarak_gosterir(self):
        self._elle_satir()
        uyarilar = pf.holding_ozeti(self.katalog())["uyarilar"]
        self.assertTrue(any("merkezi defterde yok" in u for u in uyarilar), uyarilar)

    def test_excel_tutari_sonradan_degistirilirse_fark_raporlanir(self):
        self._uygulama_kaydi()
        wb = load_workbook(self.ana)
        ws = wb["03_Kayıt Merkezi"]
        satir = next(r for r in range(6, 300) if ws[f"A{r}"].value)
        ws[f"I{satir}"] = 12_345
        wb.save(self.ana)
        wb.close()
        r = mut.kontrol()
        self.assertEqual(len(r["tutar_farki"]), 1)
        self.assertEqual(r["tutar_farki"][0]["defter_tl_kurus"], 10_000 * 100)
        self.assertEqual(r["tutar_farki"][0]["excel_tl_kurus"], 12_345 * 100)

    # --- Gerçekleşti / masraf yoldasi / belge referansi eslesmesi (gercek hatalar yine yakalanmali) ---
    def _fis_yaz(self, belge, tutar, cari="Devrim", masraf=None):
        kayit = {"Tarih": "2026-09-05", "Islem_Turu": "Usta/Tedarikçi Ödemesi", "Tutar": tutar, "PB": "TL", "Kur": 1,
                 "Cari": cari, "Proje": "Göktürk Okulu"}
        if masraf:
            kayit["Banka_Masrafi"] = masraf
        cekirdek.finans_kaydi_yaz({"belge_id": belge, "firma_key": "gursar_mimarlik", "tur": "finans_belgesi",
                                   "veri": {"kayitlar": [kayit]}}, "t")

    def _gerceklesti(self, kayit_id, tutar, kaynak, cari="Devrim", belge_ref=None):
        return self._elle_satir(kayit_id, tutar, kaynak=kaynak, durum="Gerçekleşti", cari=cari, belge_ref=belge_ref)

    def test_gerceklesti_satiri_anahtarla_eslesir_ve_tutarli_olur(self):
        self._fis_yaz("GURS-T-1", 70_000, cari="Medet İme")
        self._gerceklesti("AI-T-1", 70_000, "GURS-T-1-001", cari="Medet İme")
        r = mut.kontrol()
        self.assertEqual(r["excel_aktif_satir"], 1)
        self.assertTrue(r["tutarli"], r)

    def test_gerceklesti_satirinin_defter_fisi_yoksa_artik_eksik_raporlanir(self):
        """Eskiden Gerçekleşti satirlari hic kontrol edilmiyordu: fis eksik olsa da sessiz kaliyordu."""
        self._gerceklesti("AI-T-2", 70_000, "GURS-T-2-001")
        r = mut.kontrol()
        self.assertFalse(r["tutarli"])
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["AI-T-2"])

    def test_gerceklesti_satirinin_tutari_farkliysa_tutar_farki_raporlanir(self):
        self._fis_yaz("GURS-T-3", 70_000)
        self._gerceklesti("AI-T-3", 71_000, "GURS-T-3-001")
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["tutar_farki"]], ["AI-T-3"])
        self.assertEqual(r["tutar_farki"][0]["defter_tl_kurus"], 70_000 * 100)

    def test_bekliyor_ve_iptal_satirlari_defterde_beklenmez(self):
        self._elle_satir("B-1", 100, durum="Bekliyor")
        self._elle_satir("I-1", 200, durum="İptal")
        r = mut.kontrol()
        self.assertEqual(r["excel_aktif_satir"], 0)
        self.assertTrue(r["tutarli"], r)

    def test_masraf_yoldasi_ayri_masraf_fisiyle_eslesir(self):
        """Duzeltme ile olusan form: ana fis ``<belge>`` + ayri ``<belge>-MASRAF`` fisi."""
        self._fis_yaz("GURS-T-4", 70_000, cari="Medet İme")
        self._fis_yaz("GURS-T-4-MASRAF", 16.76, cari="Ziraat")
        self._gerceklesti("AI-T-4", 70_000, "GURS-T-4-001", cari="Medet İme")
        self._gerceklesti("AI-T-4M", 16.76, "GURS-T-4-001-MASRAF", cari="Ziraat")
        r = mut.kontrol()
        self.assertTrue(r["tutarli"], r)

    def test_masraf_yoldasi_ana_fisteki_gomulu_780_satiriyle_eslesir(self):
        """Uygulamanin normal hatti: Excel'e iki satir, deftere TEK fis (780 masraf satiri gomulu)."""
        self._fis_yaz("GURS-T-5", 70_000, cari="Medet İme", masraf=16.76)
        self._gerceklesti("AI-T-5", 70_000, "GURS-T-5-001", cari="Medet İme")
        self._gerceklesti("AI-T-5M", 16.76, "GURS-T-5-001-MASRAF", cari="Ziraat")
        r = mut.kontrol()
        self.assertTrue(r["tutarli"], r)

    def test_masraf_tutari_defterle_uyusmuyorsa_raporlanir(self):
        self._fis_yaz("GURS-T-6", 70_000, masraf=16.76)
        self._gerceklesti("AI-T-6", 70_000, "GURS-T-6-001")
        self._gerceklesti("AI-T-6M", 20, "GURS-T-6-001-MASRAF", cari="Ziraat")           # defterde 16,76
        r = mut.kontrol()
        self.assertFalse(r["tutarli"])
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["AI-T-6M"])

    def test_ayri_masraf_fisinin_tutari_excelle_uyusmuyorsa_ikisi_de_raporlanir(self):
        self._fis_yaz("GURS-T-11", 70_000)
        self._fis_yaz("GURS-T-11-MASRAF", 16.76, cari="Ziraat")                              # defterde 16,76
        self._gerceklesti("AI-T-11", 70_000, "GURS-T-11-001")
        self._gerceklesti("AI-T-11M", 20, "GURS-T-11-001-MASRAF", cari="Ziraat")           # Excel'de 20
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["AI-T-11M"])
        self.assertEqual(r["defterde_fazla"], ["GURS-T-11-MASRAF-001"])

    def test_masraf_yoldasinin_ana_fisi_yoksa_ikisi_de_raporlanir(self):
        self._gerceklesti("AI-T-7", 70_000, "GURS-T-7-001")
        self._gerceklesti("AI-T-7M", 16.76, "GURS-T-7-001-MASRAF", cari="Ziraat")
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["AI-T-7", "AI-T-7M"])

    def test_masraf_fisi_excel_yoldasi_olmadan_fazla_kalir(self):
        """Ayri masraf fisi yalniz kendi Excel masraf satiriyla eslesir; yoksa gizlenmez."""
        self._fis_yaz("GURS-T-8", 70_000)
        self._fis_yaz("GURS-T-8-MASRAF", 16.76, cari="Ziraat")
        self._gerceklesti("AI-T-8", 70_000, "GURS-T-8-001")
        r = mut.kontrol()
        self.assertEqual(r["defterde_fazla"], ["GURS-T-8-MASRAF-001"])

    def test_belge_referansi_ayni_tutar_ve_cari_ile_eslesir(self):
        """Kaynak anahtari tutmayan duzeltme satiri: Belge_Referans fis belge_id'sine esit, tutar/cari ayni."""
        self._fis_yaz("MAN-DUZ-T", 40_000, cari="Sebahattin Çakmak")
        self._gerceklesti("MAN-DUZ-T1", 40_000, "MAN-DUZ-T-SEBAH-40000", cari="Sebahattin Çakmak", belge_ref="MAN-DUZ-T")
        r = mut.kontrol()
        self.assertTrue(r["tutarli"], r)

    def test_belge_referansi_tutar_farkliysa_eslesmez(self):
        self._fis_yaz("MAN-DUZ-U", 40_000, cari="Sebahattin Çakmak")
        self._gerceklesti("MAN-DUZ-U1", 41_000, "MAN-DUZ-U-SEBAH", cari="Sebahattin Çakmak", belge_ref="MAN-DUZ-U")
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["MAN-DUZ-U1"])
        self.assertEqual(r["defterde_fazla"], ["MAN-DUZ-U-001"])

    def test_belge_referansi_cari_farkliysa_eslesmez(self):
        self._fis_yaz("MAN-DUZ-V", 40_000, cari="Sebahattin Çakmak")
        self._gerceklesti("MAN-DUZ-V1", 40_000, "MAN-DUZ-V-X", cari="Başka Usta", belge_ref="MAN-DUZ-V")
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["MAN-DUZ-V1"])
        self.assertEqual(r["defterde_fazla"], ["MAN-DUZ-V-001"])

    def test_ayni_belge_referansli_ikinci_satir_cifte_sayilmaz(self):
        """Tek fis, ayni referans/tutar/cari ile iki Excel satiri: biri eslesir, digeri eksik olarak YAKALANIR."""
        self._fis_yaz("MAN-DUZ-W", 40_000, cari="Sebahattin Çakmak")
        self._gerceklesti("MAN-DUZ-W1", 40_000, "MAN-DUZ-W-A", cari="Sebahattin Çakmak", belge_ref="MAN-DUZ-W")
        self._gerceklesti("MAN-DUZ-W2", 40_000, "MAN-DUZ-W-B", cari="Sebahattin Çakmak", belge_ref="MAN-DUZ-W")
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["MAN-DUZ-W2"])
        self.assertEqual(r["defterde_fazla"], [])

    def test_elle_girilen_satir_baska_fisin_referansiyla_cifte_sayilamaz(self):
        """Uygulamanin yazdigi satir fisini zaten tuttu; ayni referansli elle satir fisi 'calamaz' ve eksik gorunur."""
        self._fis_yaz("GURS-T-9", 70_000, cari="Medet İme")
        self._gerceklesti("AI-T-9", 70_000, "GURS-T-9-001", cari="Medet İme", belge_ref="GURS-T-9")
        self._elle_satir("MAN-DUP", 70_000, cari="Medet İme", belge_ref="GURS-T-9")        # elle tekrar girilmis
        r = mut.kontrol()
        self.assertEqual([x["kayit_id"] for x in r["eksik_defterde"]], ["MAN-DUP"])

    def test_gerceklesti_satiri_deftere_aktarilmaz_cift_kayit_olusmaz(self):
        self._fis_yaz("GURS-T-10", 70_000)
        self._gerceklesti("AI-T-10", 70_000, "GURS-T-10-001")
        self._gerceklesti("AI-T-10B", 500, None)                                            # Kaynak'siz Gerçekleşti
        once = self._fis_sayisi()
        self.assertEqual(cekirdek.holding_excelini_aktar()["aktarilan"], 0)
        self.assertEqual(self._fis_sayisi(), once)

    def test_kur_iki_haneye_yuvarlanmaz_dovizli_satir_sahte_fark_uretmez(self):
        """5.000 EUR x 55,6521 = 278.260,50 TL; kur 55,65'e yuvarlanirsa 10,50 TL'lik sahte fark cikardi."""
        satir = {"tutar": 5000, "pb": "EUR", "kur": 55.6521}
        self.assertEqual(mut._beklenen_tl_kurus(satir), 27_826_050)
        # 2.275 USD x 48,4305 = 110.179,3875 -> 110.179,39 (ROUND_HALF_UP, defterle ayni)
        self.assertEqual(mut._beklenen_tl_kurus({"tutar": 2275, "pb": "USD", "kur": 48.4305}), 11_017_939)
        self.assertEqual(mut._beklenen_tl_kurus({"tutar": 100, "pb": "USD", "kur": "34,5"}), 345_000)
        self.assertIsNone(mut._beklenen_tl_kurus({"tutar": 100, "pb": "USD", "kur": None}))

    def test_excelde_olmayan_defter_fisi_raporlanir(self):
        cekirdek.finans_kaydi_yaz({"belge_id": "SADECE-DEFTER", "firma_key": "gursar_mimarlik", "tur": "finans_belgesi",
                                   "veri": {"kayitlar": [{"Tarih": "2026-09-01", "Islem_Turu": "Gider Tahakkuku",
                                                          "Tutar": 100, "PB": "TL", "Kur": 1, "Cari": "Devrim"}]}}, "t")
        r = mut.kontrol()
        self.assertEqual(r["defterde_fazla"], ["SADECE-DEFTER-001"])

    def test_sistemin_kendi_fisleri_fazla_sayilmaz(self):
        """Ters kayit, acilis bakiyesi ve grup ici fisler Excel satiri beklemez."""
        self._uygulama_kaydi()
        with cekirdek._baglanti() as con:
            belge = con.execute("SELECT belge_id FROM fisler WHERE durum='MUHASEBELESTI'").fetchone()["belge_id"]
        cekirdek.belgeyi_ters_kaydet(belge)
        self.assertEqual(mut.kontrol()["defterde_fazla"], [])

    def test_mutabakat_salt_okunurdur(self):
        self._uygulama_kaydi()
        self._elle_satir()
        once = (hashlib.sha256(self.ana.read_bytes()).hexdigest(), hashlib.sha256(self.yedek.read_bytes()).hexdigest(),
                self._fis_sayisi())
        mut.kontrol()
        mut.uyari_metinleri()
        self.assertEqual((hashlib.sha256(self.ana.read_bytes()).hexdigest(),
                          hashlib.sha256(self.yedek.read_bytes()).hexdigest(), self._fis_sayisi()), once)

    def _fis_sayisi(self):
        with cekirdek._baglanti() as con:
            return con.execute("SELECT COUNT(*) FROM fisler").fetchone()[0]

    def test_aktarim_yalniz_eksikleri_ekler_uygulama_satirlarini_cift_saymaz(self):
        self._uygulama_kaydi()
        self._elle_satir()
        fis_once = self._fis_sayisi()
        sonuc = cekirdek.holding_excelini_aktar()
        self.assertEqual(sonuc["hatalar"], [])
        self.assertEqual(sonuc["aktarilan"], 1)                       # yalniz elle girilen satir
        self.assertEqual(self._fis_sayisi(), fis_once + 1)
        self.assertTrue(mut.kontrol()["tutarli"])
        # ikinci calistirma: tekrar eklemez (idempotent)
        self.assertEqual(cekirdek.holding_excelini_aktar()["aktarilan"], 0)
        self.assertEqual(self._fis_sayisi(), fis_once + 1)

    def test_aktarim_sonrasi_proje_metrigi_elle_girilen_kaydi_icerir(self):
        self._elle_satir(tutar=7_000)
        cekirdek.holding_excelini_aktar()
        d = pf.proje_detay("gursar_mimarlik", "Göktürk Okulu", self.katalog())
        self.assertEqual(d["metrikler"]["odenen"]["deger_kurus"], 7_000 * 100)

    def test_cli_mutabakat_ve_aktar_onay_ister(self):
        self._elle_satir()
        self.assertEqual(excel_operasyon.main(["mutabakat"]), 3)        # fark var -> cikis kodu 3
        fis_once = self._fis_sayisi()
        self.assertEqual(excel_operasyon.main(["defteri-aktar"]), 0)    # onay yok -> yazmaz
        self.assertEqual(self._fis_sayisi(), fis_once)
        self.assertEqual(excel_operasyon.main(["defteri-aktar", "--onay", "AKTAR"]), 0)
        self.assertEqual(self._fis_sayisi(), fis_once + 1)
        self.assertEqual(excel_operasyon.main(["mutabakat"]), 0)


if __name__ == "__main__":
    unittest.main()
